from __future__ import annotations

import copy
import csv
import io
import json
import shutil
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .common import ROOT, Case, dump, protected_output, sha
from .oracle import reference

BENCH = ROOT / 'DataAnalysis/data/raw/spreadsheetbench/extracted/spreadsheetbench_912_v0.1/all_data_912_v0.1'


def stable_xlsx(path):
    # Fix both XML properties and zip timestamps so preparation is reproducible.
    with zipfile.ZipFile(path) as z: parts={n:z.read(n) for n in z.namelist()}
    if 'docProps/core.xml' in parts:
        import re
        parts['docProps/core.xml']=re.sub(rb'(<dcterms:(?:created|modified)[^>]*>).*?(</dcterms:(?:created|modified)>)',rb'\g<1>2026-10-18T00:00:00Z\2',parts['docProps/core.xml'])
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for n,data in sorted(parts.items()):
            info=zipfile.ZipInfo(n,(2026,10,18,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;z.writestr(info,data)


def base_book(path, changes=None):
    w=Workbook();s=w.active;s.title='订单'
    s.append(['订单号','客户编号','地区','金额','状态','数量','单价','总额','编号文本','分类','尺寸','BMI'])
    for r in [('O1','001','东',10,'完成',2,5,None,' ００１ ','有','10*20*30 mm',10),
              ('O2','002','西',20,'完成',4,5,None,' ００２ ','无','10*20*30 mm',20),
              ('O3','004','东',20,'取消',2,10,None,' ００４ ','有','10*20*30 mm',None),
              ('O4',None,'西',5,'完成',1,5,None,None,'无','10*20*30 mm',30)]:s.append(r)
    s.freeze_panes='A2';s['A1'].font=Font(bold=True,color='224466');s['D2'].number_format='0.00'
    c=w.create_sheet('客户');c.append(['客户编号','客户等级']);c.append(['001','A']);c.append(['002','B']);c.append(['003','C'])
    if changes:changes(w)
    w.save(path);w.close();stable_xlsx(path)


def spec(kind, **kw):return {'task_type':kind,**kw,'confidence':1}
def composite(*ops):return spec('composite',operations=list(ops))
def norm(sheet='订单',column='编号文本',out='标准号'):
    return spec('normalize',sheet=sheet,column=column,operations=['trim','nfkc'],output_column=out)
def lookup(key='客户编号',duplicate='block'):
    return spec('lookup',main_sheet='订单',source_sheet='客户',main_key_column=key,source_key_column='客户编号',return_columns=['客户等级'],output_column_names=['客户等级'],duplicate_policy=duplicate)
def aggregate(sheet='订单',group='地区',column='金额',filters=None):
    return spec('aggregate',sheet=sheet,group_by=[group],metrics=[{'column':column,'function':'sum','output_column':'合计'}],filters=filters or [],output_sheet='汇总')
def sort(sheet='订单',rank=False):
    return spec('sort_rank',sheet=sheet,sort_keys=[{'column':'金额','direction':'desc'}],rank_column='金额' if rank else None,rank_output_column='排名',rank_method='dense',rank_direction='desc',output_sheet='排序')
def clean(sheet='订单',operation='remove_duplicates',keys=None,filters=None,out='清洗'):
    return spec('clean',sheet=sheet,operation=operation,key_columns=['客户编号'] if keys is None else keys,filters=filters or [],output_sheet=out,keep='first')
def formula(mode='set_formula'):
    return spec('formula',sheet='订单',column='总额',start_row=2 if mode!='copy_formula' else 3,end_row=5,mode=mode,
                formula='=F2*G2' if mode=='set_formula' else None,source_row=2 if mode=='copy_formula' else None,recalculate=False)
def mapping():return spec('transform',sheet='订单',operation='map_values',map_rules=[{'source':{'sheet':'订单','display_header':'分类'},'mapping':{'有':1,'无':0},'output_column':'分类编码'}])
def impute(method='mean'):
    return spec('transform',sheet='订单',operation='impute_numeric',impute_rules=[{'source':{'sheet':'订单','display_header':'BMI'},'method':method,'output_column':'BMI填补','indicator_column':'BMI缺失'}])


def prepare(out: Path):
    from sheetpilot.schemas import TaskEnvelope
    out=protected_output(out);directory=out/'cases';directory.mkdir(parents=True,exist_ok=True)
    definitions=[]
    def add(id,cap,op,request,**kwargs):definitions.append(dict(id=id,capability=cap,spec=op,request=request,**kwargs))
    add('L1','lookup',lookup(),'按订单的客户编号从客户表补充客户等级，追加新列客户等级；精确文本匹配，未匹配留空并报告，保留行序。')
    add('L2','lookup',lookup(duplicate='last'),'按客户编号补充客户等级为新列客户等级；来源重复编号保留最后一条，未匹配留空并报告。',change='duplicate')
    add('L3','lookup',composite(norm(),lookup('标准号')),'先对订单的编号文本去首尾空格并做NFKC，新增标准号；再用标准号精确匹配客户表的客户编号，新增客户等级，未匹配留空并报告。')
    add('N1','normalize',norm(),'订单的编号文本去首尾空格并做NFKC标准化，写入新列标准号，保留原列。')
    add('N2','normalize',spec('normalize',sheet='订单',column='状态',operations=['upper'],output_column='状态大写'),'把订单的状态转成大写，写入新列状态大写，保留原列。',change='uppercase')
    add('N3','normalize',norm('Sheet2','TRACK','标准赛道'),'把Sheet2的TRACK列去除首尾空格并做NFKC，写入新列标准赛道，保留原列。',bench=('49333',1))
    add('A1','aggregate',aggregate(),'按订单的地区分组，对金额求和，输出新工作表汇总，求和列命名合计。')
    for ident,variant in [('A2',1),('A3',2)]:
        add(ident,'aggregate',aggregate('Sheet1','PRODUCT ID','QTY',[{'column':'PRODUCT ID','operator':'eq','value':111111}]),'在Sheet1筛选PRODUCT ID等于111111的记录，按PRODUCT ID分组，对QTY求和，新建汇总表，求和列命名合计。',bench=('46167',variant))
    add('S1','sort_rank',sort(),'订单按金额降序稳定排序，完整记录写入新工作表排序。')
    add('S2','sort_rank',sort(rank=True),'订单按金额降序排序，对金额做全表降序密集排名，并列同名次，排名列叫排名；写入新工作表排序。')
    flt=[{'column':'金额','operator':'gt','value':5}]
    add('S3','sort_rank',composite(clean(operation='filter_rows',filters=flt,out='中间'),sort('中间')),'先筛选订单中金额大于5的记录到新表中间；再把中间按金额降序稳定排序到新表排序。')
    add('C1','clean',clean(),'订单按客户编号精确去重，保留第一条，空编号作为一个键，完整记录写入新表清洗。',change='duplicate_main')
    add('C2','clean',clean('Raw Data','filter_rows',[],[{'column':'Task Status','operator':'eq','value':'TSCO REQ'}],'筛选结果'),'筛选Raw Data中Task Status精确等于TSCO REQ的记录，保留全部原列，输出新工作表筛选结果，不覆盖原工作表。',bench=('19-16',1))
    add('C3','clean',composite(norm(),clean(keys=['标准号'])),'先将订单编号文本去首尾空格并做NFKC，写入新列标准号；然后按标准号精确去重保留第一条，输出新表清洗。')
    for id,mode,msg in [('F1','set_formula','在订单总额列第2到5行设置公式=F2*G2，按相对引用填充，不覆盖非空值。'),
                        ('F2','fill_gaps','修复订单总额列第2到5行公式断层，使用相邻已有公式推断，不覆盖已有公式。'),
                        ('F3','copy_formula','以订单总额列第2行为样板，向第3到5行复制公式，不覆盖非空值。')]:
        add(id,'formula',formula(mode),msg+'仅检查公式文本，不进行重算。',change=mode)
    add('V1','conversion',spec('conversion',direction='csv_to_xlsx',encoding='utf-8-sig',delimiter=',',preserve_text=True),'把CSV转换成XLSX，UTF-8 BOM编码、逗号分隔，所有字段保留文本，包括前导零和长编号。',csv_encoding='utf-8-sig')
    add('V2','conversion',spec('conversion',direction='csv_to_xlsx',encoding='gb18030',delimiter=';',preserve_text=True),'把CSV转换成XLSX，GB18030编码、分号分隔，所有字段保留文本，不推断类型。',csv_encoding='gb18030')
    add('V3','conversion',spec('conversion',direction='xlsx_to_csv',source_sheet='订单',output_name='订单.csv',encoding='utf-8-sig'),'将订单工作表导出UTF-8 BOM CSV，文件名订单.csv，保留工作簿。')
    add('T1','transform',mapping(),'订单分类列有映射为1、无映射为0，新列叫分类编码，保留原列，不输出映射状态列。')
    for id,method,zh in [('T2','mean','均值'),('T3','median','中位数')]:
        op=impute(method)
        if id=='T3':op=composite(norm(),op)
        prefix='先将订单编号文本去首尾空格并做NFKC，新增标准号；然后' if id=='T3' else ''
        add(id,'transform',op,prefix+f'用订单BMI有效数值的{zh}填补空单元格，新列BMI填补，缺失指示列BMI缺失；非空格式错误保持留空，保留原列。')
    # Missing information is explicit and paired with a bounded, predeclared reply.
    templates=[('Q1','lookup',lookup(),'return_columns','按客户编号匹配客户信息。','返回客户等级，新增列客户等级。'),
               ('Q2','normalize',norm(),'column','把订单某列去首尾空格并做NFKC，新增标准号。','处理编号文本列。'),
               ('Q3','aggregate',aggregate(),'group_by','汇总订单金额，求和列合计，新表汇总。','按地区分组。'),
               ('Q4','sort_rank',sort(),'sort_keys','给订单排序，新表排序。','按金额降序稳定排序。'),
               ('Q5','clean',clean(),'key_columns','订单去重保留第一条，输出清洗表。','按客户编号精确去重。'),
               ('Q6','formula',formula(),'formula','设置订单总额第2到5行公式，仅检查文本不重算。','公式=F2*G2，按相对引用填充。'),
               ('Q7','conversion',spec('conversion',direction='xlsx_to_csv',source_sheet='订单'),'source_sheet','导出CSV。','导出订单工作表。')]
    for id,cap,op,field,request,answer in templates:
        initial=copy.deepcopy(op);initial[field]=None if field=='formula' else [] if field in ['return_columns','group_by','sort_keys','key_columns'] else ''
        if field=='return_columns':initial['output_column_names']=[]
        initial['missing_fields']=[field]
        add(id,cap,op,request,kind='clarification',initial_spec=initial,reply={'field':field,'answer':answer},change='ambiguous_returns' if id=='Q1' else None)
    add('Q8','transform',impute('median'),'订单BMI空值先用均值填补，新列BMI填补，指示列BMI缺失，保留原列。',kind='clarification',initial_spec=impute('mean'),revision='改成中位数，其他要求不变。')
    for id,request in [('R1','写一个VBA宏处理订单。'),('R2','执行任意代码删除重复记录。'),('R3','按客户名称做模糊匹配。'),('R4','把多个工作簿的数据关联起来。')]:
        add(id,'boundary',lookup(),request,kind='refusal')
    for id,fault in [('E1','timeout'),('E2','invalid_response'),('E3','step_failure'),('E4','stale_plan')]:
        op=composite(norm(),lookup('标准号')) if fault=='step_failure' else lookup()
        req='先将订单编号文本去首尾空格并做NFKC，写入标准号；再按标准号匹配客户编号，新增客户等级，未匹配留空并报告。' if fault=='step_failure' else '按客户编号从客户表补充客户等级到订单新列，未匹配留空并报告。'
        add(id,'protection',op,req,kind='fault',fault=fault)
    assert len(definitions)==40
    for d in definitions:
        folder=directory/d['id'];folder.mkdir(exist_ok=True)
        provenance={'type':'synthetic','generator':'phase1-v1','license':'project-generated'}
        kind=d.pop('kind','normal');change=d.pop('change',None);bench=d.pop('bench',None);encoding=d.pop('csv_encoding',None)
        source=folder/('input.csv' if encoding else 'input.xlsx')
        if bench:
            ident,variant=bench;original=BENCH/f'spreadsheet/{ident}/{variant}_{ident}_input.xlsx'
            if not original.is_file():raise FileNotFoundError(original)
            shutil.copyfile(original,source)
            meta=next(x for x in json.loads((BENCH/'dataset.json').read_text(encoding='utf-8')) if str(x['id'])==ident)
            provenance={'type':'spreadsheetbench_adapted','task_id':ident,'variant':variant,
                        'original_file':str(original.relative_to(ROOT)),'original_sha256':sha(original),
                        'original_instruction':meta['instruction'],'adapted_instruction':d['request'],
                        'reason':'使用当前支持的确定性操作及新列/新Sheet输出；不评价原题固定区域或宏实现',
                        'license':'CC-BY-SA-4.0','source_url':'https://huggingface.co/datasets/KAKA22/SpreadsheetBench',
                        'benchmark_score':False}
        elif encoding:
            delimiter=';' if encoding=='gb18030' else ','
            f=io.StringIO(newline='');csv.writer(f,delimiter=delimiter).writerows([['编号','名称'],['001','甲'],['12345678901234567890','乙']]);source.write_bytes(f.getvalue().encode(encoding))
        else:
            def mutate(w):
                if change=='duplicate':w['客户'].append(['001','Z'])
                elif change=='duplicate_main':w['订单']['B4']='002'
                elif change=='uppercase':
                    for i,v in enumerate(['open','Closed','aBc','open'],2):w['订单'].cell(i,5,v)
                elif change=='ambiguous_returns':
                    w['客户']['C1']='地区'
                    for i,v in enumerate(['东','西','北'],2):w['客户'].cell(i,3,v)
                elif change in ('fill_gaps','copy_formula'):
                    w['订单']['H2']='=F2*G2'
                    if change=='fill_gaps':
                        w['订单']['H4']='=F4*G4';w['订单']['H5']='=F5*G5'
            base_book(source,mutate)
        full=TaskEnvelope(goal=d['request'],spec=d['spec']).spec.model_dump(mode='json')
        d['spec']=full
        if d.get('initial_spec'):d['initial_spec']=TaskEnvelope(goal=d['request'],spec=d['initial_spec']).spec.model_dump(mode='json')
        checks=reference(source,full) if kind not in ('fault','refusal') else {'workbook':None,'csv_artifacts':[],'tolerance':1e-9}
        c=Case(**d,kind=kind,source=provenance,group=f'bench-{bench[0]}' if bench else 'synthetic-csv-v1' if encoding else 'synthetic-orders-v1',input=source.name,input_sha256=sha(source),checks=checks)
        dump(folder/'case.json',c.model_dump(mode='json'))
    from .expansion import prepare_expansion
    prepare_expansion(directory)
    from .common import load_cases
    loaded=load_cases(directory)
    dump(out/'inventory.json',{'version':'1.0','count':len(loaded),'kinds':dict(Counter(c.kind for c in loaded)),
                              'normal_capabilities':dict(Counter(c.capability for c in loaded if c.kind=='normal')),
                              'note':'28 public-file variants including original 4; adapted tasks are not benchmark scores'})
    names={'normal':'正常处理','clarification':'需要追问','refusal':'应当拒绝','fault':'故障保护'}
    sources={'synthetic':'按规则生成','research_reconstructed':'根据调研重建表格','spreadsheetbench_adapted':'使用公开原文件'}
    lines=['# 280例覆盖表','',
           '原有40例，加上新增240例；同一案例在两种方式运行仍只算一例。','',
           '|编号|测什么|主要能力|文件来源|来源组|运行方式|用户要求|','|---|---|---|---|---|---|---|']
    for c in loaded:
        request=c.request.replace('|','/').replace('\n',' ')
        lines.append(f'|{c.id}|{names[c.kind]}|{c.capability}|{sources.get(c.source["type"],c.source["type"])}|{c.group}|{"、".join(c.modes)}|{request}|')
    (out/'case-coverage.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    return directory

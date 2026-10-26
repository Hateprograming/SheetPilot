"""Independent reference calculations. Never imports sheetpilot tools/planner."""
from __future__ import annotations

import csv
import math
import statistics
import unicodedata
import re
from collections import Counter
from copy import copy
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator


def value(v):
    if isinstance(v, (datetime, date)):
        return {'date': v.isoformat()}
    if isinstance(v,time):return {'time':v.isoformat()}
    if isinstance(v,timedelta):return {'timedelta_seconds':v.total_seconds()}
    return v


def snapshot(path: Path, encoding='utf-8-sig', delimiter=','):
    if path.suffix.lower() == '.csv':
        with path.open(encoding=encoding, newline='') as f:
            return {'CSV_Data': list(csv.reader(f, delimiter=delimiter))}
    w = load_workbook(path, data_only=False)
    try:
        return {s.title: [[value(c.value) for c in row] for row in s.iter_rows()] for s in w}
    finally:
        w.close()


def compare(expected, actual, tolerance=1e-9):
    diffs = []
    if set(expected) != set(actual):
        diffs.append({'scope': 'sheets', 'expected': list(expected), 'actual': list(actual)})
    for sheet, rows in expected.items():
        if sheet not in actual:
            continue
        got = actual[sheet]
        if len(rows) != len(got):
            diffs.append({'sheet': sheet, 'scope': 'row_count', 'expected': len(rows), 'actual': len(got)})
        for ri in range(max(len(rows), len(got))):
            er = rows[ri] if ri < len(rows) else []
            ar = got[ri] if ri < len(got) else []
            for ci in range(max(len(er), len(ar))):
                e = er[ci] if ci < len(er) else None
                a = ar[ci] if ci < len(ar) else None
                # Only absent trailing cells and empty cells are equivalent.
                numeric = type(e) in (float, int) and type(a) in (float, int)
                ok = math.isclose(e, a, rel_tol=tolerance, abs_tol=tolerance) if numeric else type(e) == type(a) and e == a
                if not ok:
                    diffs.append({'sheet': sheet, 'row': ri+1, 'column': ci+1, 'expected': e, 'actual': a})
    return diffs


def unchanged_metadata(before: Path, after: Path, expected):
    """Check original layout and style on every original cell, not a sample."""
    if before.suffix == '.csv':
        return []
    a, b = load_workbook(before), load_workbook(after)
    differences = []
    try:
        if list(b.sheetnames[:len(a.sheetnames)]) != a.sheetnames:
            differences.append({'scope': 'original_sheet_order'})
        for s in a:
            if s.title not in b:
                continue
            out = b[s.title]
            def structure(ws):
                return (ws.sheet_state, str(ws.freeze_panes), sorted(map(str, ws.merged_cells.ranges)),
                        [(k,v.hidden,v.height) for k,v in ws.row_dimensions.items()],
                        [(k,v.hidden,v.width) for k,v in ws.column_dimensions.items()])
            before_layout=structure(s);after_layout=structure(out)
            if before_layout != after_layout:
                labels=['工作表显隐','冻结位置','合并单元格','行高或隐藏状态','列宽或隐藏状态']
                changed=[name for name,a,b in zip(labels,before_layout,after_layout) if a!=b]
                differences.append({'sheet':s.title,'scope':'layout','changed':changed,
                                    'before':[before_layout[i] for i,(a,b) in enumerate(zip(before_layout,after_layout)) if a!=b],
                                    'after':[after_layout[i] for i,(a,b) in enumerate(zip(before_layout,after_layout)) if a!=b]})
            for row in s:
                for c in row:
                    d = out[c.coordinate]
                    if (c.has_style or d.has_style) and (copy(c.font) != copy(d.font) or copy(c.fill) != copy(d.fill) or c.number_format != d.number_format or copy(c.alignment) != copy(d.alignment) or copy(c.border) != copy(d.border)):
                        differences.append({'sheet': s.title, 'cell': c.coordinate, 'scope': 'style'})
    finally:
        a.close(); b.close()
    return differences


def reference(path: Path, spec: dict):
    # Work on plain rows; expected computation does not invoke product execution.
    book = snapshot(path, spec.get('encoding','utf-8-sig').replace('auto','utf-8-sig'),
                    spec.get('delimiter',',').replace('auto',','))
    artifacts = []
    for op in spec['operations'] if spec['task_type']=='composite' else [spec]:
        typ = op['task_type']
        if typ == 'conversion':
            if op['direction'] == 'xlsx_to_csv':
                artifacts.append(book[op['source_sheet']])
            continue
        sheet = op.get('sheet') or op.get('main_sheet')
        table = book[sheet]; heads = table[0]; rows = table[1:]
        # Formatting can extend a sheet beyond the last populated record.
        # Keep those cells in the original-sheet snapshot, but do not create
        # artificial blank records in a new result sheet.
        while rows and all(v is None for v in rows[-1]):rows=rows[:-1]
        def col(name):
            return heads.index(name)
        def append(name, vals):
            index = len(heads); heads.append(name)
            for row, v in zip(rows, vals):
                row.extend([None] * (index-len(row))); row.append(v)
        def keep(row, filters):
            for f in filters:
                a = row[col(f['column'])]; b = f.get('value'); k = f['operator']
                ok = {'eq': lambda:a==b, 'ne':lambda:a!=b, 'gt':lambda:a>b,
                      'gte':lambda:a>=b, 'lt':lambda:a<b, 'lte':lambda:a<=b,
                      'contains':lambda:str(b) in str(a), 'is_blank':lambda:a is None,
                      'not_blank':lambda:a is not None}[k]()
                if not ok: return False
            return True
        if typ == 'lookup':
            src = book[op['source_sheet']]; sh = src[0]; mapping = {}
            for row in src[1:]:
                key = row[sh.index(op['source_key_column'])]
                if key not in mapping or op.get('duplicate_policy') == 'last': mapping[key] = row
            exceptions = []
            keys = [r[col(op['main_key_column'])] for r in rows]
            for i,key in enumerate(keys,2):
                if key is None or key not in mapping:
                    exceptions.append([i,key,'empty_key' if key is None else 'not_found'])
            for ret,out in zip(op['return_columns'],op['output_column_names']):
                append(out,[mapping[k][sh.index(ret)] if k is not None and k in mapping else None for k in keys])
            if exceptions: book['SheetPilot_Unmatched'] = [['主表行号','关联键','失败原因'],*exceptions]
        elif typ == 'normalize':
            result=[]
            for row in rows:
                v=row[col(op['column'])]
                for action in op['operations']:
                    if v is None and action!='fill_blank': continue
                    if action=='trim': v=str(v).strip()
                    elif action=='nfkc': v=unicodedata.normalize('NFKC',str(v))
                    elif action=='upper': v=str(v).upper()
                    elif action=='lower': v=str(v).lower()
                    elif action=='to_number': v=float(v)
                    elif action=='to_text': v=str(v)
                    elif action=='fill_blank': v=op.get('fill_value') if v is None or v=='' else v
                    elif action=='strip_suffix':
                        for suffix in op.get('strip_suffixes',[]):
                            if str(v).endswith(suffix):v=str(v)[:-len(suffix)];break
                    else: raise ValueError(f'unsupported reference normalization {action}')
                result.append(v)
            append(op['output_column'],result)
        elif typ=='aggregate':
            groups={}
            for row in rows:
                if keep(row,op.get('filters',[])):
                    key=tuple(row[col(k)] for k in op['group_by']);groups.setdefault(key,[]).append(row)
            out=[[ *op['group_by'], *[m['output_column'] for m in op['metrics']]]]
            for key in sorted(groups,key=lambda item:tuple(str(v) for v in item)):
                group=groups[key]
                vals=[]
                for m in op['metrics']:
                    vs=[r[col(m['column'])] for r in group if r[col(m['column'])] is not None]
                    vals.append({'sum':lambda:sum(vs),'count':lambda:len(vs),'mean':lambda:statistics.mean(vs),'min':lambda:min(vs),'max':lambda:max(vs)}[m['function']]())
                out.append([*key,*vals])
            book[op['output_sheet']]=out
        elif typ=='sort_rank':
            def sort_value(v):
                if v is None:return (4,'')
                if isinstance(v,bool):return (0,int(v))
                if isinstance(v,(int,float)):return (1,float(v))
                if isinstance(v,(date,datetime)):return (2,v.isoformat())
                return (3,str(v).casefold())
            ordered=[list(r) for r in rows]
            for k in reversed(op['sort_keys']): ordered.sort(key=lambda r:sort_value(r[col(k['column'])]),reverse=k['direction']=='desc')
            oh=list(heads)
            if op.get('rank_column'):
                vals=sorted(set(r[col(op['rank_column'])] for r in rows),key=sort_value,reverse=op.get('rank_direction','desc')=='desc')
                counts=Counter(r[col(op['rank_column'])] for r in rows)
                for r in ordered:
                    index=vals.index(r[col(op['rank_column'])]);method=op.get('rank_method','dense')
                    r.append(index+1 if method=='dense' else 1+sum(counts[v] for v in vals[:index]) if method=='min' else sum(counts[v] for v in vals[:index+1]))
                oh.append(op['rank_output_column'])
            if op.get('top_n'):
                ordered=[r for r in ordered if r[-1]<=op['top_n']] if op.get('rank_column') else ordered[:op['top_n']]
            book[op['output_sheet']]=[oh,*ordered]
        elif typ=='clean':
            if op['operation']=='keep_columns':
                indexes=sorted({col(ref['display_header']) for ref in op['keep_columns']})
                book[op['output_sheet']]=[[r[i] for i in indexes] for r in table]
                continue
            out=[];seen=set()
            counts=Counter(tuple(row[col(k)] for k in op['key_columns']) for row in rows) if op['operation']!='filter_rows' else Counter()
            source=reversed(rows) if op.get('keep')=='last' else rows
            for row in source:
                if op['operation']=='filter_rows':
                    if keep(row,op['filters']):out.append(list(row))
                elif op['operation']=='mark_duplicates':
                    key=tuple(row[col(k)] for k in op['key_columns'])
                    out.append([*row,'重复' if counts[key]>1 else None])
                else:
                    key=tuple(row[col(k)] for k in op['key_columns'])
                    if key not in seen:out.append(list(row));seen.add(key)
            if op.get('keep')=='last':out.reverse()
            book[op['output_sheet']]=[list(heads)+([op.get('duplicate_flag_column','SheetPilot_重复标记')] if op['operation']=='mark_duplicates' else []),*out]
        elif typ=='formula':
            ci=col(op['column']);start=op['start_row'];end=op['end_row']
            from openpyxl.utils import get_column_letter
            letter=get_column_letter(ci+1)
            anchors={ri:table[ri-1][ci] for ri in range(start,end+1) if isinstance(table[ri-1][ci],str) and table[ri-1][ci].startswith('=')}
            for ri in range(start,end+1):
                if op['mode']=='fill_gaps':
                    if table[ri-1][ci] is not None:continue
                    above=max((r for r in anchors if r<ri),default=None)
                    below=min((r for r in anchors if r>ri),default=None)
                    if above is None or below is None:continue
                    from_above=Translator(anchors[above],origin=f'{letter}{above}').translate_formula(f'{letter}{ri}')
                    from_below=Translator(anchors[below],origin=f'{letter}{below}').translate_formula(f'{letter}{ri}')
                    if from_above==from_below:table[ri-1][ci]=from_above
                    continue
                if op['mode']=='copy_formula':
                    sr=op['source_row'];f=table[sr-1][ci];origin=f'{letter}{sr}'
                else:f=op.get('formula') or '=F2*G2';origin=f'{letter}{start}'
                table[ri-1][ci]=Translator(f,origin=origin).translate_formula(f'{letter}{ri}')
        elif typ=='transform':
            if op['operation']=='map_values':
                for rule in op['map_rules']:
                    ci=col(rule['source']['display_header'])
                    append(rule['output_column'],[rule['mapping'][unicodedata.normalize('NFKC',str(r[ci])).strip()] for r in rows])
            elif op['operation']=='impute_numeric':
                for rule in op['impute_rules']:
                    ci=col(rule['source']['display_header']);vs=[r[ci] for r in rows if r[ci] is not None]
                    fill=statistics.mean(vs) if rule['method']=='mean' else statistics.median(vs) if rule['method']=='median' else rule['constant_value']
                    append(rule['output_column'],[r[ci] if r[ci] is not None else fill for r in rows])
                    append(rule['indicator_column'],[int(r[ci] is None) for r in rows])
            elif op['operation']=='parse_dimensions':
                rule=op['dimension_rule'];ci=col(rule['source']['display_header']);parsed=[]
                for row in rows:
                    raw=row[ci]
                    if raw is None or str(raw).strip()=='':parsed.append((None,None,None,None,'blank'));continue
                    match=re.fullmatch(r'\s*([0-9]+(?:\.[0-9]+)?)\s*[*xX×]\s*([0-9]+(?:\.[0-9]+)?)\s*[*xX×]\s*([0-9]+(?:\.[0-9]+)?)\s*mm\s*',str(raw),re.I)
                    if not match:parsed.append((None,None,None,None,'no_numeric_measurement'));continue
                    numbers=[float(v) for v in match.groups()]
                    parsed.append((*numbers,max(numbers),'parsed_3d'))
                for key,idx in [('length_column',0),('width_column',1),('height_column',2),('max_dimension_column',3),('status_column',4)]:
                    append(rule[key],[v[idx] for v in parsed])
            elif op['operation']=='derive_metric':
                rule=op['derive_rule'];indexes=[col(ref['display_header']) for ref in rule['sources']];values=[];statuses=[]
                for row in rows:
                    nums=[row[i] for i in indexes]
                    if all(v is None or v=='' for v in nums):values.append(None);statuses.append('input_missing')
                    elif any(v is None or not isinstance(v,(int,float)) or v<=0 for v in nums):values.append(None);statuses.append('invalid_input')
                    else:values.append(math.prod(nums)/1000);statuses.append('derived')
                append(rule['output_column'],values);append(rule['status_column'],statuses)
            else: raise ValueError('unsupported reference transform')
        else:raise ValueError(f'unsupported reference task {typ}')
    return {'workbook':book,'csv_artifacts':artifacts,'tolerance':1e-9}

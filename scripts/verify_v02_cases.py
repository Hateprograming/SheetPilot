from __future__ import annotations

import tempfile
from pathlib import Path

from openpyxl import load_workbook

from sheetpilot.config import Settings
from sheetpilot.schemas import (
    AggregateFunction,
    AggregateMetric,
    AggregateTaskSpec,
    CleanOperation,
    CleanTaskSpec,
    DuplicatePolicy,
    FilterOperator,
    LookupTaskSpec,
    NormalizeOperation,
    NormalizeTaskSpec,
    RankMethod,
    RowFilter,
    SortDirection,
    SortKey,
    SortRankTaskSpec,
    TaskEnvelope,
)
from sheetpilot.service import SheetPilotService


ROOT = Path(__file__).resolve().parents[2]
CASES = ROOT / "Test" / "v0.2_cases"


def service_for(path: Path) -> SheetPilotService:
    return SheetPilotService(
        Settings(
            project_root=path,
            data_dir=path / "data",
            database_path=path / "data" / "verify.sqlite3",
            llm_provider="fake",
            api_key="",
            base_url="",
            model="fake",
            max_file_bytes=20 * 1024 * 1024,
            max_sheet_rows=100_000,
        )
    )


def execute(service: SheetPilotService, file_name: str, spec):
    source = CASES / file_name
    profile = service.create_task(file_name, source.read_bytes())
    task_id = service.task_id_from_profile(profile)
    envelope = TaskEnvelope(goal="sample verification", spec=spec)
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid, validation.errors
    service.approve(task_id, "approve")
    output, report, result = service.execute_task_plan(task_id, envelope, plan)
    assert report.status in {"passed", "passed_with_warnings"}
    return output, result


def rows(path: Path, sheet: str):
    wb = load_workbook(path, read_only=True, data_only=True)
    result = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    return result


def main() -> None:
    expected_files = {"00_v0.2测试用例总览.xlsx", *[f"TC{index:02d}_" for index in range(1, 9)]}
    actual_names = {item.name for item in CASES.glob("*.xlsx")}
    assert "00_v0.2测试用例总览.xlsx" in actual_names
    for prefix in expected_files - {"00_v0.2测试用例总览.xlsx"}:
        assert any(name.startswith(prefix) for name in actual_names), prefix

    with tempfile.TemporaryDirectory(prefix="sheetpilot-v02-") as temporary:
        service = service_for(Path(temporary))

        output, result = execute(service, "TC01_跨表匹配.xlsx", LookupTaskSpec(main_sheet="订单明细", source_sheet="客户信息", main_key_column="客户编号", source_key_column="客户编号", return_columns=["客户等级", "所属地区"], output_column_names=["客户等级", "所属地区"], confidence=1))
        assert result.counters == {"matched": 3, "unmatched": 1, "empty_key": 1, "duplicate_key": 0}
        assert rows(output, "订单明细")[1][-2:] == ("A", "华东")

        output, _ = execute(service, "TC02_文本标准化.xlsx", NormalizeTaskSpec(sheet="订单数据", column="客户编号", operations=[NormalizeOperation.TRIM, NormalizeOperation.NFKC, NormalizeOperation.UPPER], output_column="客户编号_标准化", confidence=1))
        assert [item[-1] for item in rows(output, "订单数据")[1:]] == ["A001", "B002", "C003", None]

        output, result = execute(service, "TC03_聚合汇总.xlsx", AggregateTaskSpec(sheet="销售明细", group_by=["地区"], metrics=[AggregateMetric(column="销售额", function=AggregateFunction.SUM, output_column="销售额合计"), AggregateMetric(column="销售额", function=AggregateFunction.MEAN, output_column="平均销售额"), AggregateMetric(column="订单号", function=AggregateFunction.COUNT, output_column="订单数")], filters=[RowFilter(column="状态", operator=FilterOperator.EQ, value="已完成")], output_sheet="地区汇总", confidence=1))
        assert result.counters["groups"] == 3
        assert rows(output, "地区汇总")[1:] == [("华东", 400, 200, 2), ("华北", 600, 300, 2), ("华南", 500, 500, 1)]

        output, result = execute(service, "TC04_分组排名TopN.xlsx", SortRankTaskSpec(sheet="销售业绩", sort_keys=[SortKey(column="地区", direction=SortDirection.ASC), SortKey(column="销售额", direction=SortDirection.DESC)], rank_column="销售额", rank_output_column="地区排名", rank_direction=SortDirection.DESC, rank_method=RankMethod.DENSE, group_by=["地区"], top_n=2, output_sheet="地区Top2", confidence=1))
        assert result.counters["output_rows"] == 5
        assert all(item[-1] <= 2 for item in rows(output, "地区Top2")[1:])

        output, result = execute(service, "TC05_重复标记.xlsx", CleanTaskSpec(sheet="回访记录", operation=CleanOperation.MARK_DUPLICATES, key_columns=["客户编号", "月份"], output_sheet="重复检查", duplicate_flag_column="是否重复", confidence=1))
        assert result.counters["duplicate_rows"] == 4
        assert [item[-1] for item in rows(output, "重复检查")[1:]] == ["重复", "重复", None, "重复", "重复"]

        output, result = execute(service, "TC06_移除重复.xlsx", CleanTaskSpec(sheet="客户名单", operation=CleanOperation.REMOVE_DUPLICATES, key_columns=["客户编号"], keep="last", output_sheet="客户去重", confidence=1))
        assert result.counters["duplicates_removed"] == 2
        assert len(rows(output, "客户去重")) == 4

        output, result = execute(service, "TC07_条件筛选.xlsx", CleanTaskSpec(sheet="应收明细", operation=CleanOperation.FILTER_ROWS, filters=[RowFilter(column="状态", operator=FilterOperator.EQ, value="待处理"), RowFilter(column="金额", operator=FilterOperator.GTE, value="1000")], output_sheet="重点待处理", confidence=1))
        assert result.counters["output_rows"] == 3
        assert [item[0] for item in rows(output, "重点待处理")[1:]] == ["R01", "R04", "R05"]

        source = CASES / "TC08_重复来源键安全阻止.xlsx"
        profile = service.create_task(source.name, source.read_bytes())
        task_id = service.task_id_from_profile(profile)
        envelope = TaskEnvelope(goal="duplicate boundary", spec=LookupTaskSpec(main_sheet="采购明细", source_sheet="产品主数据", main_key_column="产品编码", source_key_column="产品编码", return_columns=["标准单价"], output_column_names=["标准单价"], duplicate_policy=DuplicatePolicy.BLOCK, confidence=1))
        _, validation = service.make_task_plan(task_id, envelope)
        assert not validation.valid
        assert any("重复" in error for error in validation.errors)

    print("verified 8 v0.2 sample cases")


if __name__ == "__main__":
    main()

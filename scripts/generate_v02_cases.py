from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUTPUT_DIR = Path(__file__).resolve().parents[2] / "Test" / "v0.2_cases"

NAVY = "17324D"
BLUE = "2F75B5"
PALE_BLUE = "D9EAF7"
PALE_GREEN = "E2F0D9"
PALE_YELLOW = "FFF2CC"
WHITE = "FFFFFF"
GRID = Side(style="thin", color="D9E1F2")


def style_table(ws, freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=GRID)
    ws.row_dimensions[1].height = 24
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=GRID)
    for index, column in enumerate(ws.columns, start=1):
        width = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column)
        ws.column_dimensions[get_column_letter(index)].width = min(max(width + 3, 12), 30)


def add_table(wb: Workbook, title: str, headers: list[str], rows: list[list[object]]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style_table(ws)
    return ws


def add_instructions(
    wb: Workbook,
    *,
    case_id: str,
    capability: str,
    prompt: str,
    configuration: list[tuple[str, str]],
    expected: list[str],
) -> None:
    ws = wb.create_sheet("测试说明")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 88
    ws.merge_cells("A1:B1")
    ws["A1"] = f"{case_id} · {capability}"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(name="Arial", color=WHITE, bold=True, size=15)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    rows = [("可复制的任务描述", prompt), *configuration]
    for row_index, (label, value) in enumerate(rows, start=3):
        ws.cell(row_index, 1, label)
        ws.cell(row_index, 2, value)
        ws.cell(row_index, 1).fill = PatternFill("solid", fgColor=PALE_BLUE)
        ws.cell(row_index, 1).font = Font(name="Arial", bold=True)
        ws.cell(row_index, 2).font = Font(name="Arial")
        ws.cell(row_index, 2).alignment = Alignment(wrap_text=True, vertical="top")
    start = len(rows) + 5
    ws.cell(start, 1, "预期结果")
    ws.cell(start, 1).fill = PatternFill("solid", fgColor=PALE_GREEN)
    ws.cell(start, 1).font = Font(name="Arial", bold=True)
    for offset, item in enumerate(expected):
        ws.cell(start + offset, 2, f"• {item}")
        ws.cell(start + offset, 2).font = Font(name="Arial")
        ws.cell(start + offset, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(start + len(expected) + 2, 1, "通用流程")
    ws.cell(start + len(expected) + 2, 1).fill = PatternFill("solid", fgColor=PALE_YELLOW)
    ws.cell(start + len(expected) + 2, 1).font = Font(name="Arial", bold=True)
    ws.cell(start + len(expected) + 2, 2, "上传 → 检查结构 → 确认任务参数 → 预览计划 → 审批 → 执行 → 查看验证报告 → 下载结果；执行后可测试版本恢复。")
    ws.cell(start + len(expected) + 2, 2).font = Font(name="Arial")
    ws.cell(start + len(expected) + 2, 2).alignment = Alignment(wrap_text=True, vertical="top")


def save_case(file_name: str, build) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    build(wb)
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    wb.save(OUTPUT_DIR / file_name)
    wb.close()


def build_cases() -> list[dict[str, str]]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    manifest: list[dict[str, str]] = []

    def register(case_id, file_name, capability, prompt, build):
        save_case(file_name, build)
        manifest.append({"case": case_id, "file": file_name, "capability": capability, "prompt": prompt})

    register(
        "TC01",
        "TC01_跨表匹配.xlsx",
        "跨表精确匹配",
        "按客户编号，从客户信息表把客户等级和所属地区补充到订单明细；未匹配项留空并报告，保留原行顺序。",
        lambda wb: (
            add_table(wb, "订单明细", ["订单号", "客户编号", "订单金额"], [["SO-1001", "C001", 1200], ["SO-1002", "C002", 800], ["SO-1003", "C004", 560], ["SO-1004", None, 300], ["SO-1005", "C003", 1500]]),
            add_table(wb, "客户信息", ["客户编号", "客户等级", "所属地区"], [["C001", "A", "华东"], ["C002", "B", "华北"], ["C003", "A", "华南"]]),
            add_table(wb, "核验答案", ["订单号", "客户等级", "所属地区"], [["SO-1001", "A", "华东"], ["SO-1002", "B", "华北"], ["SO-1003", None, None], ["SO-1004", None, None], ["SO-1005", "A", "华南"]]),
            add_instructions(wb, case_id="TC01", capability="跨表精确匹配", prompt="按客户编号，从客户信息表把客户等级和所属地区补充到订单明细；未匹配项留空并报告，保留原行顺序。", configuration=[("主表 / 来源表", "订单明细 / 客户信息"), ("关联列", "客户编号 / 客户编号"), ("返回列", "客户等级、所属地区"), ("比较方式", "精确文本"), ("重复策略", "阻止并报告")], expected=["匹配 3 行，未匹配 1 行，空键 1 行。", "订单明细新增 2 列，原 5 行顺序不变。", "状态允许为“通过但有警告”，并列出未匹配与空键。"]),
        ),
    )

    register(
        "TC02",
        "TC02_文本标准化.xlsx",
        "数据标准化",
        "把订单数据的客户编号统一为半角大写并去除首尾空格，写入新列“客户编号_标准化”，不要覆盖原列。",
        lambda wb: (
            add_table(wb, "订单数据", ["订单号", "客户编号", "金额"], [["N-01", " Ａ001 ", 100], ["N-02", "ｂ002", 200], ["N-03", " c003 ", 300], ["N-04", None, 50]]),
            add_table(wb, "核验答案", ["订单号", "客户编号_标准化"], [["N-01", "A001"], ["N-02", "B002"], ["N-03", "C003"], ["N-04", None]]),
            add_instructions(wb, case_id="TC02", capability="数据标准化", prompt="把订单数据的客户编号统一为半角大写并去除首尾空格，写入新列“客户编号_标准化”，不要覆盖原列。", configuration=[("工作表 / 列", "订单数据 / 客户编号"), ("操作", "去除首尾空格、统一全角/半角字符、转大写"), ("输出", "客户编号_标准化；不覆盖")], expected=["新增“客户编号_标准化”列。", "四个结果依次为 A001、B002、C003、空值。", "原“客户编号”列内容保持不变。"]),
        ),
    )

    register(
        "TC03",
        "TC03_聚合汇总.xlsx",
        "聚合与条件过滤",
        "只统计状态为“已完成”的记录，按地区汇总销售额合计、平均销售额和订单数，结果写入“地区汇总”。",
        lambda wb: (
            add_table(wb, "销售明细", ["订单号", "地区", "产品", "销售额", "状态"], [["A01", "华东", "甲", 100, "已完成"], ["A02", "华东", "乙", 300, "已完成"], ["A03", "华东", "甲", 999, "已取消"], ["A04", "华北", "甲", 200, "已完成"], ["A05", "华北", "乙", 400, "已完成"], ["A06", "华南", "甲", 500, "已完成"]]),
            add_table(wb, "核验答案", ["地区", "销售额合计", "平均销售额", "订单数"], [["华东", 400, 200, 2], ["华北", 600, 300, 2], ["华南", 500, 500, 1]]),
            add_instructions(wb, case_id="TC03", capability="聚合与条件过滤", prompt="只统计状态为“已完成”的记录，按地区汇总销售额合计、平均销售额和订单数，结果写入“地区汇总”。", configuration=[("来源表 / 分组", "销售明细 / 地区"), ("指标 1", "销售额 · 求和 → 销售额合计"), ("指标 2", "销售额 · 平均 → 平均销售额"), ("指标 3", "订单号 · 计数 → 订单数"), ("筛选", "状态 = 已完成"), ("结果 Sheet", "地区汇总")], expected=["生成 3 行地区汇总。", "华东合计为 400，取消订单 999 不参与。", "验证报告显示输入/输出行数与分组数合理。"]),
        ),
    )

    register(
        "TC04",
        "TC04_分组排名TopN.xlsx",
        "排序、分组排名与 Top N",
        "按地区对销售额做降序密集排名，每个地区保留前 2 名；结果先按地区升序、再按销售额降序，写入“地区Top2”。",
        lambda wb: (
            add_table(wb, "销售业绩", ["地区", "销售员", "销售额"], [["华东", "小林", 800], ["华东", "小周", 1000], ["华东", "小陈", 800], ["华北", "小王", 900], ["华北", "小赵", 700], ["华北", "小孙", 500]]),
            add_table(wb, "核验答案", ["地区", "销售员", "销售额", "地区排名"], [["华东", "小周", 1000, 1], ["华东", "小林", 800, 2], ["华东", "小陈", 800, 2], ["华北", "小王", 900, 1], ["华北", "小赵", 700, 2]]),
            add_instructions(wb, case_id="TC04", capability="排序、分组排名与 Top N", prompt="按地区对销售额做降序密集排名，每个地区保留前 2 名；结果先按地区升序、再按销售额降序，写入“地区Top2”。", configuration=[("来源表", "销售业绩"), ("排序", "地区升序、销售额降序"), ("排名列 / 分组", "销售额 / 地区"), ("方向 / 并列", "数值越大越前 / 密集排名"), ("Top N", "2"), ("输出", "地区Top2；排名列名“地区排名”")], expected=["输出 5 行；华东第 2 名并列的两行都保留。", "华北小孙（排名 3）被排除。", "同分记录排名一致。"]),
        ),
    )

    register(
        "TC05",
        "TC05_重复标记.xlsx",
        "重复记录标记",
        "按“客户编号+月份”标记重复记录，写入新表“重复检查”，标记列名为“是否重复”。",
        lambda wb: (
            add_table(wb, "回访记录", ["客户编号", "月份", "回访人", "结果"], [["C001", "2026-08", "小林", "完成"], ["C001", "2026-08", "小周", "完成"], ["C001", "2026-09", "小林", "完成"], ["C002", "2026-08", "小王", "未联系"], ["C002", "2026-08", "小赵", "完成"]]),
            add_table(wb, "核验答案", ["客户编号", "月份", "是否重复"], [["C001", "2026-08", "重复"], ["C001", "2026-08", "重复"], ["C001", "2026-09", ""], ["C002", "2026-08", "重复"], ["C002", "2026-08", "重复"]]),
            add_instructions(wb, case_id="TC05", capability="重复记录标记", prompt="按“客户编号+月份”标记重复记录，写入新表“重复检查”，标记列名为“是否重复”。", configuration=[("来源表", "回访记录"), ("清洗方式", "标记重复记录"), ("判断列", "客户编号、月份"), ("结果 Sheet / 标记列", "重复检查 / 是否重复")], expected=["输出仍为 5 行。", "第 1、2、4、5 行标记为“重复”，第 3 行为空。", "原始工作表不被修改。"]),
        ),
    )

    register(
        "TC06",
        "TC06_移除重复.xlsx",
        "移除重复记录",
        "按客户编号去重，保留最后一条记录，结果写入新表“客户去重”。",
        lambda wb: (
            add_table(wb, "客户名单", ["客户编号", "客户名称", "更新时间"], [["C001", "远航商贸", "2026-08-01"], ["C002", "北辰科技", "2026-08-02"], ["C001", "远航商贸有限公司", "2026-09-01"], ["C003", "星海制造", "2026-09-03"], ["C002", "北辰科技集团", "2026-09-05"]]),
            add_table(wb, "核验答案", ["客户编号", "客户名称", "更新时间"], [["C001", "远航商贸有限公司", "2026-09-01"], ["C003", "星海制造", "2026-09-03"], ["C002", "北辰科技集团", "2026-09-05"]]),
            add_instructions(wb, case_id="TC06", capability="移除重复记录", prompt="按客户编号去重，保留最后一条记录，结果写入新表“客户去重”。", configuration=[("来源表", "客户名单"), ("清洗方式", "移除重复记录"), ("判断列 / 保留", "客户编号 / 最后一条"), ("结果 Sheet", "客户去重")], expected=["输出 3 行，移除 2 条重复记录。", "C001 与 C002 均保留更新时间较晚的那条。", "原始工作表仍保留 5 行。"]),
        ),
    )

    register(
        "TC07",
        "TC07_条件筛选.xlsx",
        "多条件行筛选",
        "筛选状态为“待处理”且金额大于等于 1000 的记录，结果写入“重点待处理”。",
        lambda wb: (
            add_table(wb, "应收明细", ["单号", "客户", "金额", "状态"], [["R01", "甲公司", 1500, "待处理"], ["R02", "乙公司", 800, "待处理"], ["R03", "丙公司", 2000, "已完成"], ["R04", "丁公司", 1000, "待处理"], ["R05", "戊公司", 3000, "待处理"]]),
            add_table(wb, "核验答案", ["单号", "客户", "金额", "状态"], [["R01", "甲公司", 1500, "待处理"], ["R04", "丁公司", 1000, "待处理"], ["R05", "戊公司", 3000, "待处理"]]),
            add_instructions(wb, case_id="TC07", capability="多条件行筛选", prompt="筛选状态为“待处理”且金额大于等于 1000 的记录，结果写入“重点待处理”。", configuration=[("来源表", "应收明细"), ("清洗方式", "按条件筛选"), ("条件 1", "状态 = 待处理"), ("条件 2", "金额 >= 1000"), ("结果 Sheet", "重点待处理")], expected=["输出 R01、R04、R05，共 3 行。", "多个筛选条件按 AND 组合。", "原始工作表不被修改。"]),
        ),
    )

    register(
        "TC08",
        "TC08_重复来源键安全阻止.xlsx",
        "跨表匹配安全边界",
        "按产品编码补充标准单价；来源键重复时选择“阻止并报告”，不要执行有歧义的匹配。",
        lambda wb: (
            add_table(wb, "采购明细", ["采购单", "产品编码", "数量"], [["P01", "SKU-01", 2], ["P02", "SKU-02", 3]]),
            add_table(wb, "产品主数据", ["产品编码", "标准单价", "版本"], [["SKU-01", 100, "旧"], ["SKU-01", 120, "新"], ["SKU-02", 80, "当前"]]),
            add_table(wb, "核验答案", ["核验项", "预期"], [["执行是否应被阻止", "是"], ["重复键", "SKU-01"], ["原表是否改变", "否"]]),
            add_instructions(wb, case_id="TC08", capability="跨表匹配安全边界", prompt="按产品编码补充标准单价；来源键重复时选择“阻止并报告”，不要执行有歧义的匹配。", configuration=[("主表 / 来源表", "采购明细 / 产品主数据"), ("关联列", "产品编码 / 产品编码"), ("返回列", "标准单价"), ("重复策略", "阻止并报告")], expected=["计划校验或执行前明确报告 SKU-01 为重复来源键。", "任务被阻止，不生成含歧义结果的版本。", "将策略切换为“使用最后一条”后，SKU-01 应得到 120，可作为追加测试。"]),
        ),
    )
    return manifest


def create_manifest(cases: list[dict[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "用例总览"
    ws.append(["用例", "文件", "覆盖能力", "可复制的任务描述", "预期性质"])
    for item in cases:
        nature = "安全阻止（预期不执行）" if item["case"] == "TC08" else "成功执行并通过验证"
        ws.append([item["case"], item["file"], item["capability"], item["prompt"], nature])
    style_table(ws)
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 78
    ws.column_dimensions["E"].width = 26
    for cell in ws["D"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    guide = wb.create_sheet("执行清单")
    guide.append(["步骤", "检查内容", "通过标准"])
    checks = [
        (1, "未勾选数据发送同意时尝试建任务", "应用拒绝继续，不上传/分析文件"),
        (2, "上传任一用例并建立只读任务", "显示工作表、表头、行列数、风险提示"),
        (3, "填写或校正任务参数", "计划预览列出修改/新增 Sheet、风险等级与预计范围"),
        (4, "审批并执行", "生成新版本，原始上传文件不被覆盖"),
        (5, "查看验证报告并下载", "状态、计数器、警告与用例预期一致；下载文件可打开"),
        (6, "恢复到版本 0", "当前版本回到原始文件，之后仍可继续创建新版本"),
        (7, "在五种任务类型间切换", "页面不空白，所选表单立即显示"),
    ]
    for row in checks:
        guide.append(row)
    style_table(guide)
    guide.column_dimensions["B"].width = 46
    guide.column_dimensions["C"].width = 68
    wb.save(OUTPUT_DIR / "00_v0.2测试用例总览.xlsx")
    wb.close()


if __name__ == "__main__":
    cases = build_cases()
    create_manifest(cases)
    print(f"created {len(cases) + 1} workbooks in {OUTPUT_DIR}")

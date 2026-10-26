from __future__ import annotations

import sys
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from sheetpilot.config import Settings
from sheetpilot.llm import FakeLLMClient
from sheetpilot.schemas import DuplicatePolicy, LookupTaskSpec, NormalizationMode
from sheetpilot.service import SheetPilotService


def create_workbook(path: Path, *, duplicate: bool = False) -> None:
    workbook = Workbook()
    orders = workbook.active
    orders.title = "订单"
    orders.append(["客户编号", "订单额"])
    orders.append(["001", 100])
    orders.append(["999", 200])
    orders.append([None, 300])
    customers = workbook.create_sheet("客户")
    customers.append(["客户编号", "客户等级", "地区"])
    customers.append(["001", "A", "华东"])
    if duplicate:
        customers.append(["001", "B", "华北"])
    workbook.save(path)
    workbook.close()


def make_spec(**updates) -> LookupTaskSpec:
    values = {
        "main_sheet": "订单",
        "source_sheet": "客户",
        "main_key_column": "客户编号",
        "source_key_column": "客户编号",
        "return_columns": ["客户等级", "地区"],
        "output_column_names": ["客户等级", "地区"],
        "normalization_mode": NormalizationMode.EXACT_TEXT,
        "confidence": 1.0,
    }
    values.update(updates)
    return LookupTaskSpec(**values)


def make_settings(root: Path) -> Settings:
    return Settings(
        project_root=root,
        data_dir=root / "data",
        database_path=root / "data" / "sheetpilot.sqlite3",
        llm_provider="fake",
        api_key="",
        base_url="",
        model="fake",
        max_file_bytes=20 * 1024 * 1024,
        max_sheet_rows=100_000,
    )


def main() -> None:
    with TemporaryDirectory() as directory:
        root = Path(directory)
        input_path = root / "input.xlsx"
        create_workbook(input_path)
        spec = make_spec()
        service = SheetPilotService(make_settings(root), FakeLLMClient(spec))
        profile = service.create_task("input.xlsx", input_path.read_bytes())
        task_id = service.task_id_from_profile(profile)
        parsed = service.parse_request(task_id, "按客户编号补充客户等级和地区")
        _, validation = service.make_plan(task_id, parsed)
        assert validation.valid
        service.approve(task_id, "approve")
        output, report = service.execute(task_id, parsed)
        assert report.status == "passed_with_warnings"
        assert all(check.passed for check in report.checks)
        workbook = load_workbook(output, data_only=True)
        assert workbook["订单"]["C2"].value == "A"
        assert workbook["订单"]["D2"].value == "华东"
        assert workbook["SheetPilot_Unmatched"].max_row == 3
        workbook.close()

    with TemporaryDirectory() as directory:
        root = Path(directory)
        input_path = root / "input.xlsx"
        create_workbook(input_path, duplicate=True)
        spec = make_spec(duplicate_policy=DuplicatePolicy.BLOCK)
        service = SheetPilotService(make_settings(root), FakeLLMClient(spec))
        profile = service.create_task("input.xlsx", input_path.read_bytes())
        task_id = service.task_id_from_profile(profile)
        parsed = service.parse_request(task_id, "按客户编号补充客户等级和地区")
        _, validation = service.make_plan(task_id, parsed)
        assert not validation.valid
        assert validation.duplicate_key_count == 1

    print("SheetPilot core smoke tests: PASS")


if __name__ == "__main__":
    main()

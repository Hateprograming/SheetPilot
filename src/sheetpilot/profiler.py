from __future__ import annotations

import datetime as dt
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .errors import FileRejectedError, WorkbookStructureError
from .files import sha256_file
from .schemas import ColumnProfile, SheetProfile, WorkbookProfile


def _json_value(value: Any) -> Any:
    if isinstance(value, (dt.date, dt.datetime, dt.time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _infer_type(values: list[Any]) -> str:
    actual = [value for value in values if value not in (None, "")]
    if not actual:
        return "empty"
    kinds: set[str] = set()
    for value in actual:
        if isinstance(value, bool):
            kinds.add("boolean")
        elif isinstance(value, (int, float)):
            kinds.add("number")
        elif isinstance(value, (dt.date, dt.datetime)):
            kinds.add("date")
        elif isinstance(value, str) and value.startswith("="):
            kinds.add("formula")
        else:
            kinds.add("text")
    return next(iter(kinds)) if len(kinds) == 1 else "mixed"


def _header_candidates(sheet, limit: int = 10) -> list[tuple[int, float]]:
    candidates: list[tuple[int, float]] = []
    max_column = min(sheet.max_column, 200)
    for row_index in range(1, min(sheet.max_row, limit) + 1):
        values = [sheet.cell(row_index, col).value for col in range(1, max_column + 1)]
        nonempty = [str(value).strip() for value in values if value not in (None, "")]
        if not nonempty:
            continue
        uniqueness = len(set(nonempty)) / len(nonempty)
        text_ratio = sum(not value.replace(".", "", 1).isdigit() for value in nonempty) / len(nonempty)
        next_nonempty = 0
        if row_index < sheet.max_row:
            next_nonempty = sum(
                sheet.cell(row_index + 1, col).value not in (None, "")
                for col in range(1, max_column + 1)
            )
        density = len(nonempty) / max(1, max_column)
        continuation = min(1.0, next_nonempty / max(1, len(nonempty)))
        score = 0.4 * uniqueness + 0.25 * text_ratio + 0.2 * density + 0.15 * continuation
        candidates.append((row_index, round(score, 3)))
    return sorted(candidates, key=lambda item: item[1], reverse=True)


def _column_name(value: Any, index: int) -> str:
    text = str(value).strip() if value not in (None, "") else ""
    return text or f"Column_{get_column_letter(index)}"


def profile_workbook(
    path: Path, original_name: str, max_sheet_rows: int, header_overrides: dict[str, int] | None = None
) -> WorkbookProfile:
    try:
        workbook = load_workbook(path, read_only=False, data_only=False, keep_links=True)
    except Exception as exc:
        raise FileRejectedError("无法打开工作簿，文件可能损坏、加密或不是标准 XLSX。") from exc

    overrides = header_overrides or {}
    external_links = [str(getattr(link, "id", index)) for index, link in enumerate(workbook._external_links)]
    risks: list[str] = []
    if external_links:
        risks.append("检测到外部链接；不处理此类工作簿。")
    sheet_profiles: list[SheetProfile] = []

    for sheet in workbook.worksheets:
        if sheet.max_row > max_sheet_rows + 10:
            workbook.close()
            raise FileRejectedError(f"Sheet“{sheet.title}”超过最大行数限制。")
        candidates = _header_candidates(sheet)
        header_row = overrides.get(sheet.title)
        if header_row is None and candidates:
            header_row = candidates[0][0]
        confidence = candidates[0][1] if candidates and header_row == candidates[0][0] else 1.0 if header_row else 0.0
        headers: list[str] = []
        columns: list[ColumnProfile] = []
        preview_rows: list[dict[str, Any]] = []
        if header_row:
            headers = [_column_name(sheet.cell(header_row, col).value, col) for col in range(1, sheet.max_column + 1)]
            duplicate_headers = [name for name, count in Counter(headers).items() if count > 1]
            if duplicate_headers:
                risks.append(f"Sheet“{sheet.title}”存在重复列名：{', '.join(duplicate_headers[:5])}")
            for col, name in enumerate(headers, 1):
                values = [
                    sheet.cell(row, col).value
                    for row in range(header_row + 1, min(sheet.max_row, header_row + 5000) + 1)
                ]
                normalized = [str(value) for value in values if value not in (None, "")]
                display_header = str(sheet.cell(header_row, col).value).strip() if sheet.cell(header_row, col).value not in (None, "") else ""
                occurrence = sum(
                    1
                    for prior in range(1, col + 1)
                    if str(sheet.cell(header_row, prior).value).strip() == display_header
                    and display_header
                ) or 1
                columns.append(
                    ColumnProfile(
                        name=name,
                        index=col,
                        inferred_type=_infer_type(values),
                        empty_count=sum(value in (None, "") for value in values),
                        unique_count=len(set(normalized)),
                        formula_count=sum(isinstance(value, str) and value.startswith("=") for value in values),
                        letter=get_column_letter(col),
                        display_header=display_header,
                        header_occurrence=occurrence,
                        nonempty_count=len(normalized),
                        active=bool(normalized),
                    )
                )
            for row in range(header_row + 1, min(sheet.max_row, header_row + 3) + 1):
                preview_rows.append(
                    {headers[col - 1]: _json_value(sheet.cell(row, col).value) for col in range(1, len(headers) + 1)}
                )
        merged_ranges = [str(value) for value in sheet.merged_cells.ranges]
        if header_row and any(rng.min_row <= header_row <= rng.max_row for rng in sheet.merged_cells.ranges):
            risks.append(f"Sheet“{sheet.title}”的表头包含合并单元格。")
        sheet_profiles.append(
            SheetProfile(
                name=sheet.title,
                state=sheet.sheet_state,
                max_row=sheet.max_row,
                max_column=sheet.max_column,
                header_row=header_row,
                header_confidence=confidence,
                header_candidates=[row for row, _ in candidates[:3]],
                columns=columns,
                preview_rows=preview_rows,
                merged_ranges=merged_ranges,
                hidden_rows=sum(d.hidden for d in sheet.row_dimensions.values()),
                hidden_columns=sum(d.hidden for d in sheet.column_dimensions.values()),
            )
        )
    workbook.close()
    return WorkbookProfile(
        file_name=original_name,
        size_bytes=path.stat().st_size,
        sha256=sha256_file(path),
        sheets=sheet_profiles,
        external_links=external_links,
        risks=list(dict.fromkeys(risks)),
    )


def ensure_supported_profile(profile: WorkbookProfile) -> None:
    if profile.external_links:
        raise WorkbookStructureError("v0.3 不处理包含外部链接的工作簿。")
    if not profile.sheets:
        raise WorkbookStructureError("工作簿中没有可处理的工作表。")
    for sheet in profile.sheets:
        if sheet.header_row is None:
            raise WorkbookStructureError(f"无法识别Sheet“{sheet.name}”的表头。")
        # Duplicate and blank headers are profile risks, not workbook-wide blockers.
        # Tools resolve a stable column reference and only ask for clarification when
        # the requested header maps to more than one active data column.
        if any(risk == f"Sheet“{sheet.name}”的表头包含合并单元格。" for risk in profile.risks):
            raise WorkbookStructureError(
                f"Sheet“{sheet.name}”表头包含合并单元格；v0.3暂不支持。"
            )

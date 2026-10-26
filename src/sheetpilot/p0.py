"""Operation-scoped delivery, measurements, disclosure and retention.

No workbook rows are copied into operation/measurement records. Authorization
stores only a fingerprint; the preview is regenerated from existing local data.
"""
from __future__ import annotations

import csv
import hashlib
import json
import logging
import shutil
import threading
import time
from contextvars import ContextVar
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any, Literal
from uuid import uuid4

from pydantic import BaseModel, Field

from .errors import IntentValidationError, VerificationError
from .files import sha256_file
from .storage import utc_now


authorized_model_rows: ContextVar[list[dict]] = ContextVar("authorized_model_rows", default=[])


class FeedbackInput(BaseModel):
    outcome: Literal["completed", "partial", "not_completed"]
    reasons: list[str] = Field(default_factory=list, max_length=10)
    note: str = Field(default="", max_length=2000)
    review_minutes: float | None = Field(default=None, ge=0, le=100000)
    baseline_minutes: float | None = Field(default=None, ge=0, le=100000)
    actual_total_minutes: float | None = Field(default=None, ge=0, le=100000)


COUNTERS = {
    "matched": "匹配", "unmatched": "未匹配", "empty_keys": "空关联键",
    "empty_key": "空关联键", "duplicate_keys": "重复关联键", "changed": "修改",
    "failed": "失败", "input_rows": "输入记录", "output_rows": "输出记录",
    "kept": "保留", "removed": "移除", "duplicates": "重复记录", "groups": "分组",
    "imputed": "填补", "mapped": "映射", "mapped_default": "默认映射",
    "valid": "有效", "invalid": "无效", "exceptions": "异常",
    "converted_rows": "转换记录", "non_numeric": "非数值记录",
    "duplicates_removed": "移除重复记录", "duplicate_rows": "重复记录",
    "duplicate_key": "重复关联键",
}


class P0Store:
    def __init__(self, service):
        self.service = service
        self.repo = service.repository
        self.locks: dict[str, threading.RLock] = {}
        self._locks_guard = threading.Lock()
        self._migrate()

    def lock(self, task_id):
        with self._locks_guard:
            return self.locks.setdefault(task_id, threading.RLock())

    def _migrate(self):
        with self.repo.connect() as db:
            db.executescript("""
                CREATE TABLE IF NOT EXISTS p0_operations (
                    operation_id TEXT PRIMARY KEY, task_id TEXT NOT NULL,
                    first_message_id INTEGER, status TEXT NOT NULL,
                    result_version INTEGER, summary_json TEXT,
                    created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS p0_links (
                    kind TEXT NOT NULL, link_id TEXT NOT NULL, task_id TEXT NOT NULL,
                    operation_id TEXT NOT NULL, PRIMARY KEY(kind,link_id));
                CREATE TABLE IF NOT EXISTS p0_measurements (
                    task_id TEXT NOT NULL, operation_id TEXT NOT NULL,
                    measurement_id TEXT NOT NULL, stage TEXT NOT NULL,
                    duration_ms REAL NOT NULL, PRIMARY KEY(operation_id,measurement_id));
                CREATE TABLE IF NOT EXISTS p0_feedback (
                    operation_id TEXT PRIMARY KEY, task_id TEXT NOT NULL,
                    result_version INTEGER, payload_json TEXT NOT NULL, updated_at TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS p0_grants (
                    grant_id TEXT PRIMARY KEY, task_id TEXT NOT NULL,
                    operation_id TEXT NOT NULL, fingerprint TEXT NOT NULL,
                    status TEXT NOT NULL, created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
                CREATE INDEX IF NOT EXISTS idx_p0_operations_task ON p0_operations(task_id,created_at);
            """)
            cols = {row['name'] for row in db.execute('PRAGMA table_info(tasks)')}
            for name in ('last_activity_at', 'expires_at', 'cleanup_status'):
                if name not in cols:
                    db.execute(f'ALTER TABLE tasks ADD COLUMN {name} TEXT')
            now = utc_now()
            expiry = (datetime.now(UTC) + timedelta(days=self.service.settings.retention_days)).isoformat()
            db.execute("UPDATE tasks SET last_activity_at=?,expires_at=?,cleanup_status='active' WHERE expires_at IS NULL", (now, expiry))
            db.execute('INSERT OR IGNORE INTO schema_migrations(version,applied_at) VALUES(5,?)', (now,))

    def touch(self, task_id):
        now = datetime.now(UTC)
        with self.repo.connect() as db:
            db.execute("UPDATE tasks SET last_activity_at=?,expires_at=? WHERE task_id=? AND COALESCE(cleanup_status,'active')='active'",
                       (now.isoformat(), (now + timedelta(days=self.service.settings.retention_days)).isoformat(), task_id))

    def recover_interrupted(self):
        """Called once at startup, after the existing commit recovery."""
        with self.repo.connect() as db:
            ops = db.execute("SELECT * FROM p0_operations WHERE status IN ('planning','executing')").fetchall()
            for op in ops:
                if op['status'] == 'planning':
                    db.execute("UPDATE planning_runs SET status='failed',cancel_requested=1 WHERE job_id IN (SELECT link_id FROM p0_links WHERE operation_id=? AND kind='planning') AND status IN ('queued','running','waiting_for_execution')", (op['operation_id'],))
                    state = 'failed'
                else:
                    db.execute("UPDATE plan_runs SET status='paused' WHERE run_id IN (SELECT link_id FROM p0_links WHERE operation_id=? AND kind='execution') AND status='running'", (op['operation_id'],))
                    state = 'paused'
                db.execute('UPDATE p0_operations SET status=?,updated_at=? WHERE operation_id=?', (state,utc_now(),op['operation_id']))
                newest = db.execute('SELECT operation_id FROM p0_operations WHERE task_id=? ORDER BY created_at DESC,rowid DESC LIMIT 1', (op['task_id'],)).fetchone()
                if newest and newest['operation_id'] == op['operation_id']:
                    error = json.dumps({'safe_message':'应用已重启，本次处理未完成；请检查安全检查点后恢复或重新描述要求。','category':'transient_io'})
                    db.execute('UPDATE tasks SET status=?,error_json=? WHERE task_id=?', (state,error,op['task_id']))

    def current(self, task_id):
        with self.repo.connect() as db:
            row = db.execute('SELECT * FROM p0_operations WHERE task_id=? ORDER BY created_at DESC,rowid DESC LIMIT 1', (task_id,)).fetchone()
        return dict(row) if row else None

    def operation(self, task_id, operation_id):
        with self.repo.connect() as db:
            row = db.execute('SELECT * FROM p0_operations WHERE task_id=? AND operation_id=?', (task_id, operation_id)).fetchone()
        if not row:
            raise IntentValidationError('本次操作不存在。')
        return dict(row)

    def begin(self, task_id, continuation=False):
        current = self.current(task_id)
        if continuation and current:
            self.state(current['operation_id'], 'planning')
            return current['operation_id']
        op_id = uuid4().hex
        now = utc_now()
        with self.repo.connect() as db:
            db.execute("UPDATE p0_operations SET status='superseded',updated_at=? WHERE task_id=? AND status IN ('planning','awaiting_approval','needs_clarification','awaiting_data_authorization')", (now, task_id))
            db.execute('INSERT INTO p0_operations(operation_id,task_id,status,created_at,updated_at) VALUES(?,?,?,?,?)', (op_id, task_id, 'planning', now, now))
        self.touch(task_id)
        return op_id

    def link(self, task_id, operation_id, kind, link_id, message_id=None):
        with self.repo.connect() as db:
            db.execute('INSERT OR REPLACE INTO p0_links VALUES(?,?,?,?)', (kind, link_id, task_id, operation_id))
            if message_id is not None:
                db.execute('UPDATE p0_operations SET first_message_id=COALESCE(first_message_id,?) WHERE operation_id=?', (message_id, operation_id))

    def linked(self, kind, link_id):
        with self.repo.connect() as db:
            row = db.execute('SELECT operation_id FROM p0_links WHERE kind=? AND link_id=?', (kind, link_id)).fetchone()
        return row['operation_id'] if row else None

    @staticmethod
    def plan_key(plan):
        return f'{plan.plan_id}:r{plan.revision}'

    def bind_plan(self, task_id, plan):
        op_id = self.linked('plan', self.plan_key(plan))
        if not op_id:
            current = self.current(task_id)
            op_id = current['operation_id'] if current and current['status'] in ('planning','awaiting_approval','needs_clarification') else self.begin(task_id)
            self.link(task_id, op_id, 'plan', self.plan_key(plan))
        return op_id

    def state(self, op_id, status):
        with self.repo.connect() as db:
            prior = db.execute('SELECT task_id,status,updated_at FROM p0_operations WHERE operation_id=?', (op_id,)).fetchone()
            if prior and prior['status'] in ('needs_clarification','awaiting_approval','awaiting_data_authorization','paused') and prior['status'] != status:
                elapsed = (datetime.now(UTC)-datetime.fromisoformat(prior['updated_at'])).total_seconds()*1000
                db.execute('INSERT OR IGNORE INTO p0_measurements VALUES(?,?,?,?,?)', (prior['task_id'],op_id,'waiting:'+prior['updated_at'],'waiting',max(0,elapsed)))
            db.execute('UPDATE p0_operations SET status=?,updated_at=? WHERE operation_id=?', (status, utc_now(), op_id))

    def measure(self, task_id, op_id, key, stage, milliseconds):
        if not op_id:
            return
        with self.repo.connect() as db:
            db.execute('INSERT OR REPLACE INTO p0_measurements VALUES(?,?,?,?,?)', (task_id, op_id, key, stage, max(0, milliseconds)))

    def complete(self, task_id, op_id, path, report, result):
        task = self.repo.get_task(task_id)
        version = int(task['current_version'])
        record = self.repo.get_version(task_id, version)
        if not record or Path(record['path']) != path or not path.is_file() or sha256_file(path) != record['sha256'] or report.status not in ('passed','passed_with_warnings'):
            self.state(op_id, 'failed')
            raise VerificationError('结果未成为可核验的正式版本，不能交付。')
        summary = self.summary(result.model_dump(mode='json'), report.model_dump(mode='json'))
        summary['exception_artifacts'] = []
        steps = result.details.get('steps') or [result.model_dump(mode='json')]
        for index, step in enumerate(steps):
            name = step.get('details', {}).get('artifact_name')
            if name:
                source = self.service.files.artifact_path(task_id, name)
                if source.is_file() and source.suffix == '.csv':
                    target_name = f'p0_{op_id}_{index}_exceptions.csv'
                    try:
                        shutil.copyfile(source, self.service.files.artifact_path(task_id, target_name))
                        summary['exception_artifacts'].append(target_name)
                    except OSError:
                        summary['warnings'].append('异常分页快照未能保存，请在已交付文件及原异常产物中核对。')
        with self.repo.connect() as db:
            db.execute("UPDATE p0_operations SET status='completed',result_version=?,summary_json=?,updated_at=? WHERE operation_id=?", (version, json.dumps(summary, ensure_ascii=False), utc_now(), op_id))
        self.touch(task_id)

    @staticmethod
    def summary(result, report):
        steps = result.get('details', {}).get('steps') or [result]
        return {
            'scope': list(dict.fromkeys(result.get('modified_sheets', []) + result.get('created_sheets', []))),
            'ranges': result.get('modified_ranges', []),
            'steps': [{'tool': s.get('tool_name'), 'rows_affected': s.get('rows_affected'),
                       'metrics': [{'key': k, 'label': COUNTERS.get(k, k), 'value': v} for k,v in s.get('counters', {}).items()]} for s in steps],
            'row_count_note': '各步骤影响行数分别统计，同一行可能参与多个步骤，不代表去重记录数。',
            'checks': report.get('checks', []),
            'warnings': list(dict.fromkeys(result.get('warnings', []) + report.get('warnings', []))),
            'manual_review_required': report.get('manual_review_required', False),
            'verification_status': report.get('status'),
            'not_verified': ['业务规则是否符合真实业务含义仍需确认。', '未验证外部数据源刷新状态或数据期间是否最新。'] + (['公式结构与计算检查不能证明业务公式正确。'] if any('formula' in str(s.get('tool_name')) for s in steps) else []),
        }

    def workspace(self, task_id):
        op = self.current(task_id)
        task = self.repo.get_task(task_id)
        if not op:
            return {'current_operation': None, 'retention': {'expires_at': task.get('expires_at'), 'days': self.service.settings.retention_days}}
        op_id = op['operation_id']
        with self.repo.connect() as db:
            measures = db.execute('SELECT stage,SUM(duration_ms) AS ms FROM p0_measurements WHERE operation_id=? GROUP BY stage', (op_id,)).fetchall()
            feedback = db.execute('SELECT payload_json FROM p0_feedback WHERE operation_id=?', (op_id,)).fetchone()
            counts = db.execute("SELECT COUNT(*) AS n FROM p0_links WHERE operation_id=? AND kind='planning'", (op_id,)).fetchone()['n']
            transitions = {r['stage']:r['n'] for r in db.execute("SELECT stage,COUNT(*) AS n FROM p0_measurements WHERE operation_id=? AND stage IN ('pause_count','resume_count') GROUP BY stage", (op_id,))}
            execution_ids = [r['link_id'] for r in db.execute("SELECT link_id FROM p0_links WHERE operation_id=? AND kind='execution'", (op_id,))]
        metrics = {r['stage']: round(r['ms']) for r in measures}
        # All timing buckets are disjoint; UI polling never updates them.
        summary = json.loads(op['summary_json']) if op.get('summary_json') else None
        delivery = 'not_ready'
        if op['status'] == 'completed' and op['result_version'] is not None:
            v = self.repo.get_version(task_id, op['result_version'])
            try:
                valid = bool(v and Path(v['path']).is_file() and sha256_file(Path(v['path'])) == v['sha256'])
                if valid and v.get('delivery_json'):
                    export = json.loads(v['delivery_json'])
                    export_path = Path(export['path'])
                    valid = (export_path.resolve() == Path(v['path']).with_suffix('.csv').resolve()
                             and export_path.is_file() and sha256_file(export_path) == export['sha256'])
            except OSError:
                valid = False
            delivery = 'ready' if valid else 'unavailable'
        tool_runs = [r for r in self.repo.list_tool_runs(task_id) if r.get('run_id') in execution_ids]
        active_run = self.repo.latest_plan_run(task_id)
        can_resume = bool(active_run and active_run['run_id'] in execution_ids and active_run['status'] in ('paused','failed'))
        payload = json.loads(feedback['payload_json']) if feedback else None
        net = None if not payload or payload.get('baseline_minutes') is None or payload.get('actual_total_minutes') is None else payload['baseline_minutes'] - payload['actual_total_minutes']
        return {
            'current_operation': {'operation_id': op_id, 'status': op['status'], 'result_version': op['result_version'],
                                  'delivery_status': delivery, 'summary': summary, 'feedback': payload,
                                  'can_resume': can_resume,
                                  'metrics': {'stage_ms': metrics, 'server_ms': sum(v for k,v in metrics.items() if k != 'waiting'),
                                              'waiting_ms': metrics.get('waiting', 0), 'clarification_turns': max(0, counts-1),
                                              'pauses': transitions.get('pause_count', 0), 'resumes': transitions.get('resume_count', 0),
                                              'retries': sum(r.get('status') == 'retrying' for r in tool_runs),
                                              'net_saved_minutes': net}},
            'data_authorization': self.authorization_preview(task_id),
            'retention': {'expires_at': task.get('expires_at'), 'days': self.service.settings.retention_days},
        }

    def save_feedback(self, task_id, op_id, payload: FeedbackInput):
        op = self.operation(task_id, op_id)
        if op['status'] not in ('completed','failed','paused','cancelled','superseded'):
            raise IntentValidationError('请在处理完成或停止后提交反馈。')
        with self.repo.connect() as db:
            db.execute('INSERT OR REPLACE INTO p0_feedback VALUES(?,?,?,?,?)', (op_id, task_id, op['result_version'], payload.model_dump_json(), utc_now()))
        self.touch(task_id)

    def review_rows(self, task_id, op_id):
        op = self.operation(task_id, op_id)
        for message in reversed(self.service.get_conversation(task_id)):
            if op['first_message_id'] and message.id < op['first_message_id']:
                break
            rows = (message.metadata.get('review') or {}).get('rows') or []
            if rows:
                return [{'sheet': r.get('sheet'), 'excel_row': r.get('row'), 'target_values': r.get('raw_values', {}), 'issue_code': r.get('issue_code')} for r in rows[:50]]
        return []

    def fingerprint(self, task_id, op_id, rows):
        task = self.repo.get_task(task_id)
        version = self.repo.get_version(task_id, int(task['current_version']))
        content = json.dumps([op_id, task['current_version'], version['sha256'] if version else None,
                              self.service.settings.llm_provider, self.service.settings.model,
                              self.service.settings.base_url, rows], ensure_ascii=False, sort_keys=True, default=str)
        return hashlib.sha256(content.encode()).hexdigest()

    def authorization(self, task_id, op_id, rows):
        fp = self.fingerprint(task_id, op_id, rows)
        with self.repo.connect() as db:
            row = db.execute('SELECT * FROM p0_grants WHERE task_id=? AND operation_id=? AND fingerprint=? ORDER BY created_at DESC LIMIT 1', (task_id, op_id, fp)).fetchone()
            if row:
                return dict(row)
            grant_id = uuid4().hex
            now = utc_now()
            db.execute("INSERT INTO p0_grants VALUES(?,?,?,?,?,?,?)", (grant_id, task_id, op_id, fp, 'pending', now, now))
        return {'grant_id': grant_id, 'status': 'pending'}

    def authorization_preview(self, task_id):
        op = self.current(task_id)
        if not op or op['status'] != 'awaiting_data_authorization':
            return None
        rows = self.review_rows(task_id, op['operation_id'])
        grant = self.authorization(task_id, op['operation_id'], rows)
        return {'grant_id': grant['grant_id'], 'status': grant['status'], 'count': len(rows), 'rows': rows,
                'purpose': '理解目标列异常并结合你的回答完善处理规则；不发送整行或完整文件。',
                'version': int(self.repo.get_task(task_id)['current_version'])}

    def decide_authorization(self, task_id, grant_id, approved):
        op = self.current(task_id)
        if not op or op['status'] != 'awaiting_data_authorization':
            raise IntentValidationError('授权请求已失效。')
        rows = self.review_rows(task_id, op['operation_id'])
        current = self.authorization(task_id, op['operation_id'], rows)
        if current['grant_id'] != grant_id or current['status'] != 'pending':
            raise IntentValidationError('数据范围已变化或请求已处理，请刷新后重新确认。')
        with self.repo.connect() as db:
            db.execute('UPDATE p0_grants SET status=?,updated_at=? WHERE grant_id=?', ('approved' if approved else 'denied', utc_now(), grant_id))
        self.state(op['operation_id'], 'needs_clarification')
        self.repo.update_task(task_id, status='needs_clarification')
        self.touch(task_id)
        return op

    def busy(self, task_id):
        task = self.repo.get_task(task_id)
        planning = self.repo.latest_planning_run(task_id)
        run = self.repo.latest_plan_run(task_id)
        return bool(task and task['status'] in ('intent_parsed','analyzing_request','approved','executing') or
                    planning and planning['status'] in ('queued','running','waiting_for_execution') or run and run['status'] == 'running')

    def delete(self, task_id):
        if self.busy(task_id):
            raise IntentValidationError('工作簿正在处理或等待安全停止，请停止后再删除。')
        with self.lock(task_id):
            task = self.repo.get_task(task_id)
            if not task:
                return
            if self.busy(task_id):
                raise IntentValidationError('工作簿正在处理或等待安全停止，请停止后再删除。')
            with self.repo.connect() as db:
                db.execute('BEGIN IMMEDIATE')
                active_planning = db.execute("SELECT 1 FROM planning_runs WHERE task_id=? AND status IN ('running','queued','waiting_for_execution')", (task_id,)).fetchone()
                active_execution = db.execute("SELECT 1 FROM plan_runs WHERE task_id=? AND status='running'", (task_id,)).fetchone()
                if active_planning or active_execution:
                    raise IntentValidationError('工作簿正在处理，请先停止。')
                db.execute("UPDATE tasks SET cleanup_status='pending' WHERE task_id=?", (task_id,))
            root = self.service.files.tasks_dir.resolve()
            target = (root / task_id).resolve()
            if target.parent != root or target == root or target != root / task_id or target.is_symlink():
                raise IntentValidationError('任务删除路径不安全。')
            try:
                if target.exists():
                    shutil.rmtree(target)
                # LangGraph uses a separate sqlite checkpoint store. Remove its
                # task threads as well instead of leaving conversation snapshots.
                self._delete_graph_threads(task_id)
                self._scrub_application_backups(task_id)
                self.repo.delete_task(task_id)
            except Exception as exc:
                raise IntentValidationError('清理尚未完成，任务已停止访问；后台将重试。') from exc

    def _delete_graph_threads(self, task_id):
        import sqlite3
        for graph_path in self.service.settings.data_dir.glob('langgraph*.sqlite3'):
            with sqlite3.connect(graph_path) as db:
                tables = {r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
                for table in ('checkpoints','writes'):
                    if table in tables:
                        db.execute(f'DELETE FROM {table} WHERE thread_id=? OR thread_id LIKE ?', (task_id, task_id + ':%'))

    def _scrub_application_backups(self, task_id):
        import sqlite3
        # Only repository-generated migration backups, never arbitrary user paths.
        root = self.service.settings.data_dir / 'backups'
        for path in root.glob(f'{self.repo.path.stem}-pre-v*{self.repo.path.suffix}'):
            with sqlite3.connect(path) as db:
                tables = [r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='table'")]
                for table in tables:
                    if not table.replace('_','').isalnum():
                        continue
                    cols = {r[1] for r in db.execute(f'PRAGMA table_info("{table}")')}
                    if 'task_id' in cols:
                        db.execute(f'DELETE FROM "{table}" WHERE task_id=?', (task_id,))

    def exceptions(self, task_id, op_id, issue=None, offset=0, limit=50):
        from openpyxl import load_workbook
        op = self.operation(task_id, op_id)
        summary = json.loads(op['summary_json']) if op.get('summary_json') else {}
        counts = {}
        page = []
        total = 0
        def accept(row):
            nonlocal total
            code = str(row['issue_code'])
            counts[code] = counts.get(code, 0)+1
            if issue and issue != code:
                return
            if offset <= total < offset+limit:
                page.append(row)
            total += 1
        for name in summary.get('exception_artifacts', []):
            path = self.service.files.artifact_path(task_id, name)
            if path.is_file():
                with path.open(encoding='utf-8-sig', newline='') as stream:
                    for row in csv.DictReader(stream):
                        accept({'sheet':row.get('工作表'), 'excel_row':int(row['行号']), 'column':row.get('源列'),
                                'issue_code':row.get('问题代码'), 'description':row.get('说明'),
                                'value':'（仅在本地异常文件中查看）'})
        if op['result_version'] is not None:
            v = self.repo.get_version(task_id, op['result_version'])
            if v and Path(v['path']).is_file() and sha256_file(Path(v['path'])) == v['sha256']:
                wb = load_workbook(v['path'], read_only=True, data_only=False)
                try:
                    for name in summary.get('scope', []):
                        if name.startswith('SheetPilot_Unmatched') and name in wb.sheetnames:
                            for row in wb[name].iter_rows(min_row=2, values_only=True):
                                if len(row)>=3 and isinstance(row[0], (int,float)):
                                    # The unmatched sheet is an exception output. Its
                                    # source row number is preserved in column A.
                                    source = next((s for s in summary.get('scope', []) if s != name), name)
                                    accept({'sheet':source, 'excel_row':int(row[0]), 'column':'关联键', 'issue_code':row[2],
                                            'description':'关联键为空或未找到来源记录。', 'value':'（关联键已隐藏）'})
                finally:
                    wb.close()
        return {'rows':page, 'total':total, 'counts':counts, 'offset':offset, 'limit':limit, 'version':op['result_version']}

    def cleanup(self):
        with self.repo.connect() as db:
            candidates = db.execute("SELECT task_id FROM tasks WHERE expires_at<=? OR cleanup_status='pending'", (utc_now(),)).fetchall()
        for row in candidates:
            task_id = row['task_id']
            lock = self.lock(task_id)
            if not lock.acquire(blocking=False):
                continue
            try:
                task = self.repo.get_task(task_id)
                if task and (task.get('cleanup_status') == 'pending' or task.get('expires_at','') <= utc_now()) and not self.busy(task_id):
                    self.delete(task_id)
            except (OSError, IntentValidationError):
                logging.getLogger(__name__).warning('Task cleanup pending: %s', task_id)
            finally:
                lock.release()

    def export_metrics(self, path: Path):
        fields = ['operation_id','status','result_version','server_ms','waiting_ms','outcome','baseline_minutes','actual_total_minutes','net_saved_minutes']
        with self.repo.connect() as db:
            ops = db.execute('SELECT operation_id,task_id,status,result_version FROM p0_operations ORDER BY created_at').fetchall()
        with path.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            for op in ops:
                with self.repo.connect() as db:
                    times = db.execute('SELECT stage,SUM(duration_ms) AS ms FROM p0_measurements WHERE operation_id=? GROUP BY stage', (op['operation_id'],)).fetchall()
                    fb = db.execute('SELECT payload_json FROM p0_feedback WHERE operation_id=?', (op['operation_id'],)).fetchone()
                payload = json.loads(fb['payload_json']) if fb else {}
                base, actual = payload.get('baseline_minutes'), payload.get('actual_total_minutes')
                writer.writerow({**{k:op[k] for k in ('operation_id','status','result_version')},
                                 'server_ms': sum(r['ms'] for r in times if r['stage']!='waiting'),
                                 'waiting_ms': sum(r['ms'] for r in times if r['stage']=='waiting'),
                                 'outcome': payload.get('outcome'), 'baseline_minutes':base, 'actual_total_minutes':actual,
                                 'net_saved_minutes': None if base is None or actual is None else base-actual})

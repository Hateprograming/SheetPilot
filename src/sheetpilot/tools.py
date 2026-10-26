from __future__ import annotations

import datetime as dt
import decimal
import csv
import hashlib
import json
import math
import re
import shutil
import statistics
import time
import unicodedata
from itertools import zip_longest
from abc import ABC, abstractmethod
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Generic, TypeVar

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string, get_column_letter
from pydantic import BaseModel

from .calculation import recalculate_workbook, _audit_cached_values
from .conversion import csv_to_xlsx, xlsx_sheet_to_csv
from .files import sha256_file
from .matching import execute_lookup, validate_lookup_spec
from .schemas import (
    AggregateFunction,
    AggregateTaskSpec,
    CleanOperation,
    CleanTaskSpec,
    ConversionDirection,
    ConversionTaskSpec,
    ColumnRef,
    DimensionResolutionAction,
    DerivedFormulaTemplate,
    DerivedResolutionAction,
    FilterOperator,
    FormulaMode,
    FormulaTaskSpec,
    LookupTaskSpec,
    NormalizeOperation,
    NormalizeTaskSpec,
    ParseDimensionsRule,
    RiskLevel,
    SortDirection,
    SortRankTaskSpec,
    TransformOperation,
    TransformTaskSpec,
    ImputationMethod,
    InvalidNumericPolicy,
    TaskValidation,
    ToolResult,
    WorkbookProfile,
)
from .workbook_utils import worksheet_dimensions


SpecT = TypeVar("SpecT", bound=BaseModel)


@dataclass(frozen=True)
class ToolContext:
    task_id: str
    input_path: Path
    output_path: Path
    profile: WorkbookProfile
    version_before: int
    version_after: int
    artifact_dir: Path | None = None
    original_path: Path | None = None
    execution_deadline: float | None = None
    cancelled: Callable[[], bool] | None = None
    report_stage: Callable[[str], None] | None = None


class SpreadsheetTool(ABC, Generic[SpecT]):
    name: str
    description: str
    spec_model: type[SpecT]
    risk_level: RiskLevel

    def parse_spec(self, arguments: dict[str, Any] | SpecT) -> SpecT:
        if isinstance(arguments, self.spec_model):
            return arguments
        return self.spec_model.model_validate(arguments)

    @abstractmethod
    def validate(self, context: ToolContext, spec: SpecT) -> TaskValidation: ...

    @abstractmethod
    def preview(self, context: ToolContext, spec: SpecT) -> dict[str, Any]: ...

    @abstractmethod
    def execute(self, context: ToolContext, spec: SpecT) -> ToolResult: ...

    @abstractmethod
    def verify(
        self, context: ToolContext, spec: SpecT, result: ToolResult
    ) -> TaskValidation: ...


def _sheet_profile(profile: WorkbookProfile, name: str):
    return next((sheet for sheet in profile.sheets if sheet.name == name), None)


def _column_map(profile, sheet_name: str) -> dict[str, int]:
    sheet = _sheet_profile(profile, sheet_name)
    return {column.name: column.index for column in sheet.columns} if sheet else {}


def _normalized_header(value: Any) -> str:
    """Match display headers independently of line wrapping and extra whitespace."""
    return "".join(unicodedata.normalize("NFKC", str(value)).split())


def _resolve_named_column(profile: WorkbookProfile, sheet_name: str, value: str):
    """Resolve legacy string columns by header first, then by Excel letter."""
    sheet = _sheet_profile(profile, sheet_name)
    if not sheet:
        return None, f"工作表“{sheet_name}”不存在。"
    normalized = _normalized_header(value)
    matches = [
        column for column in sheet.columns
        if _normalized_header(column.display_header or column.name) == normalized
    ]
    active = [column for column in matches if column.active]
    if len(active) == 1:
        return active[0], None
    if len(active) > 1:
        return None, f"列“{value}”对应多个活动列，请指定列字母。"
    if len(matches) == 1:
        return matches[0], None
    try:
        index = column_index_from_string(value.strip().upper())
    except ValueError:
        index = 0
    if index:
        column = next((item for item in sheet.columns if item.index == index), None)
        if column:
            return column, None
    return None, f"工作表“{sheet_name}”中不存在列“{value}”。"


def _resolve_formula_column(profile: WorkbookProfile, sheet_name: str, value: str):
    sheet = _sheet_profile(profile, sheet_name)
    if sheet:
        matches = [col for col in sheet.columns if _normalized_header(col.display_header or col.name) == _normalized_header(value)]
        if len(matches) > 1:
            return None, f"目标列“{value}”有多个同名表头，请指定目标列的列字母。"
        named = [col for col in sheet.columns if col.name == value]
        if len(named) == 1:
            return named[0], None
    return _resolve_named_column(profile, sheet_name, value)


def _resolve_column_ref(profile: WorkbookProfile, ref: ColumnRef) -> tuple[Any | None, str | None]:
    sheet = _sheet_profile(profile, ref.sheet)
    if not sheet:
        return None, f"工作表“{ref.sheet}”不存在。"
    by_index = {column.index: column for column in sheet.columns}
    requested_index = ref.column_index
    if not requested_index and ref.column_letter:
        try:
            requested_index = column_index_from_string(ref.column_letter.strip().upper())
        except ValueError:
            return None, f"列字母“{ref.column_letter}”无效。"
    if requested_index:
        column = by_index.get(requested_index)
        if not column:
            return None, f"工作表“{ref.sheet}”中不存在第 {requested_index} 列。"
        if ref.display_header:
            actual = column.display_header or column.name
            if _normalized_header(actual) != _normalized_header(ref.display_header):
                return None, f"第 {requested_index} 列表头已变化：当前为“{actual}”。"
        return column, None
    matches = [
        column for column in sheet.columns
        if _normalized_header(column.display_header or column.name)
        == _normalized_header(ref.display_header)
    ]
    active = [column for column in matches if column.active]
    if len(active) == 1:
        return active[0], None
    if len(active) > 1:
        return None, f"列“{ref.display_header}”对应多个活动列，请指定列字母或列序号。"
    if len(matches) == 1:
        return matches[0], None
    if not matches:
        return None, f"工作表“{ref.sheet}”中不存在列“{ref.display_header}”。"
    return None, f"列“{ref.display_header}”不唯一，请指定列字母或列序号。"


def _source_label(column: Any) -> str:
    return column.display_header or column.name or f"Column_{column.letter or get_column_letter(column.index)}"


def _normalized_text(value: Any, *, casefold: bool = False) -> str:
    text = unicodedata.normalize("NFKC", str(value)).strip()
    return text.casefold() if casefold else text


def _is_missing(value: Any, tokens: list[str]) -> bool:
    if value in (None, ""):
        return True
    normalized = _normalized_text(value)
    return normalized == "" or normalized in {_normalized_text(item) for item in tokens}


def _strict_number(value: Any) -> float | int:
    if isinstance(value, bool):
        raise ValueError("boolean is not numeric input")
    number = decimal.Decimal(str(value).strip())
    if not number.is_finite():
        raise ValueError("number must be finite")
    return int(number) if number == number.to_integral() else float(number)


def _transform_artifact_name(spec: TransformTaskSpec) -> str:
    digest = hashlib.sha256(spec.model_dump_json().encode("utf-8")).hexdigest()[:10]
    return f"transform_{spec.operation.value}_{digest}_exceptions.csv"


def _default_transform_names(spec: TransformTaskSpec, profile: WorkbookProfile) -> list[tuple[str, int | None]]:
    """Return output headers and existing indexes; None means append in this order."""
    names: list[tuple[str, int | None]] = []

    def existing_index(sheet_name: str, name: str) -> int | None:
        sheet = _sheet_profile(profile, sheet_name)
        matches = [column.index for column in (sheet.columns if sheet else []) if column.name == name]
        return matches[0] if len(matches) == 1 else None

    if spec.operation == TransformOperation.IMPUTE_NUMERIC:
        for rule in spec.impute_rules:
            column, _ = _resolve_column_ref(profile, rule.source)
            label = _source_label(column) if column else (rule.source.display_header or "数值")
            value_name = rule.output_column or f"{label}_填补"
            flag_name = rule.indicator_column or f"{value_name}_是否填补"
            target_name = label if rule.overwrite else value_name
            target_index = (
                column.index
                if rule.overwrite and column
                else existing_index(spec.sheet, value_name)
                if rule.overwrite_existing_outputs
                else None
            )
            flag_index = (
                existing_index(spec.sheet, flag_name)
                if rule.overwrite or rule.overwrite_existing_outputs
                else None
            )
            names.append((target_name, target_index))
            names.append((flag_name, flag_index))
    elif spec.operation == TransformOperation.MAP_VALUES:
        for rule in spec.map_rules:
            column, _ = _resolve_column_ref(profile, rule.source)
            label = _source_label(column) if column else (rule.source.display_header or "分类")
            value_name = rule.output_column or f"{label}_映射"
            names.append((label if rule.overwrite else value_name, column.index if rule.overwrite and column else None))
            # Mapping outcomes are audit metadata by default.  Only create a
            # workbook status column when the user explicitly named one.
            if rule.status_column:
                names.append((
                    rule.status_column,
                    existing_index(spec.sheet, rule.status_column) if rule.overwrite else None,
                ))
    elif spec.operation == TransformOperation.PARSE_DIMENSIONS and spec.dimension_rule:
        rule = spec.dimension_rule
        column, _ = _resolve_column_ref(profile, rule.source)
        label = _source_label(column) if column else (rule.source.display_header or "尺寸")
        output_names = [
            rule.length_column or f"{label}_长_mm",
            rule.width_column or f"{label}_宽_mm",
            rule.height_column or f"{label}_高_mm",
            rule.max_dimension_column or f"{label}_最大径_mm",
            rule.status_column or f"{label}_解析状态",
        ]
        names.extend((name, existing_index(spec.sheet, name) if rule.overwrite_existing_outputs else None) for name in output_names)
    elif spec.operation == TransformOperation.DERIVE_METRIC and spec.derive_rule:
        rule = spec.derive_rule
        value_name = rule.output_column or "派生指标_cm3"
        status_name = rule.status_column or f"{value_name}_状态"
        names.extend([
            (value_name, existing_index(spec.sheet, value_name) if rule.overwrite_existing_outputs else None),
            (status_name, existing_index(spec.sheet, status_name) if rule.overwrite_existing_outputs else None),
        ])
    return names


def transform_write_columns(profile: WorkbookProfile, spec: TransformTaskSpec) -> list[tuple[str, int, bool]]:
    sheet = _sheet_profile(profile, spec.sheet)
    next_index = (sheet.max_column if sheet else 0) + 1
    resolved: list[tuple[str, int, bool]] = []
    for name, existing in _default_transform_names(spec, profile):
        if existing:
            resolved.append((name, existing, True))
        else:
            resolved.append((name, next_index, False))
            next_index += 1
    return resolved


def _base_validation(
    profile: WorkbookProfile,
    sheet_name: str,
    required_columns: list[str],
    output_sheet: str | None = None,
) -> TaskValidation:
    errors: list[str] = []
    warnings: list[str] = []
    sheet = _sheet_profile(profile, sheet_name)
    if not sheet:
        errors.append(f"工作表“{sheet_name}”不存在。")
        return TaskValidation(valid=False, errors=errors)
    available = Counter(column.name for column in sheet.columns)
    for column in required_columns:
        if not column:
            errors.append("存在未指定的列。")
        elif available[column] == 0:
            errors.append(f"工作表“{sheet_name}”中不存在列“{column}”。")
        elif available[column] > 1:
            errors.append(f"工作表“{sheet_name}”中的列“{column}”不唯一。")
    if output_sheet and output_sheet in {item.name for item in profile.sheets}:
        errors.append(f"输出工作表“{output_sheet}”已存在。")
    return TaskValidation(valid=not errors, errors=errors, warnings=warnings)


def _copy_input(context: ToolContext) -> None:
    shutil.copy2(context.input_path, context.output_path)


def _filter_matches(value: Any, operator: FilterOperator, expected: Any) -> bool:
    if operator == FilterOperator.IS_BLANK:
        return value in (None, "")
    if operator == FilterOperator.NOT_BLANK:
        return value not in (None, "")
    if operator == FilterOperator.CONTAINS:
        return str(expected).casefold() in str(value or "").casefold()
    if operator == FilterOperator.EQ:
        return value == expected or str(value) == str(expected)
    if operator == FilterOperator.NE:
        return not _filter_matches(value, FilterOperator.EQ, expected)
    try:
        left = decimal.Decimal(str(value))
        right = decimal.Decimal(str(expected))
    except (decimal.InvalidOperation, TypeError, ValueError):
        left, right = str(value), str(expected)
    if operator == FilterOperator.GT:
        return left > right
    if operator == FilterOperator.GTE:
        return left >= right
    if operator == FilterOperator.LT:
        return left < right
    if operator == FilterOperator.LTE:
        return left <= right
    return False


def _row_matches(values: dict[str, Any], filters) -> bool:
    return all(_filter_matches(values.get(item.column), item.operator, item.value) for item in filters)


class LookupTool(SpreadsheetTool[LookupTaskSpec]):
    name = "lookup"
    description = "在同一工作簿的两个工作表之间按单列键执行精确左连接。"
    spec_model = LookupTaskSpec
    risk_level = RiskLevel.MEDIUM

    def validate(self, context: ToolContext, spec: LookupTaskSpec) -> TaskValidation:
        result = validate_lookup_spec(context.input_path, context.profile, spec)
        errors = [*result.missing_fields, *result.ambiguities, *result.output_conflicts]
        if result.duplicate_key_count and spec.duplicate_policy.value == "block":
            errors.append("来源表关联键存在重复，必须选择处理策略。")
        warnings = [result.type_warning] if result.type_warning else []
        return TaskValidation(valid=not errors, errors=errors, warnings=warnings)

    def preview(self, context: ToolContext, spec: LookupTaskSpec) -> dict[str, Any]:
        sheet = _sheet_profile(context.profile, spec.main_sheet)
        return {
            "modified_sheets": [spec.main_sheet],
            "created_sheets": ["SheetPilot_Unmatched（仅有异常时）"],
            "new_columns": spec.output_column_names or spec.return_columns,
            "estimated_rows": max(0, sheet.max_row - (sheet.header_row or 1)) if sheet else 0,
            "destructive": False,
        }

    def execute(self, context: ToolContext, spec: LookupTaskSpec) -> ToolResult:
        change = execute_lookup(
            context.task_id,
            context.input_path,
            context.output_path,
            context.profile,
            spec,
            context.version_before,
            context.version_after,
        )
        return ToolResult(
            tool_name=self.name,
            modified_sheets=change.modified_sheets,
            modified_ranges=change.modified_ranges,
            created_sheets=change.created_sheets,
            rows_affected=change.rows_affected,
            counters={
                "matched": change.matched_count,
                "unmatched": change.unmatched_count,
                "empty_key": change.empty_key_count,
                "duplicate_key": change.duplicate_key_count,
            },
            warnings=change.warnings,
            details={"legacy_changeset": change.model_dump(mode="json")},
        )

    def verify(
        self, context: ToolContext, spec: LookupTaskSpec, result: ToolResult
    ) -> TaskValidation:
        workbook = load_workbook(context.output_path, read_only=True, data_only=False)
        errors: list[str] = []
        try:
            sheet = workbook[spec.main_sheet]
            profile = _sheet_profile(context.profile, spec.main_sheet)
            output_names = spec.output_column_names or spec.return_columns
            actual = [
                sheet.cell(profile.header_row, profile.max_column + 1 + index).value
                for index in range(len(output_names))
            ]
            if actual != output_names:
                errors.append(f"新增列与计划不一致：{actual}")
            output_rows, _ = worksheet_dimensions(
                sheet,
                fallback_max_row=profile.max_row,
                fallback_max_column=profile.max_column + len(output_names),
            )
            if output_rows != profile.max_row:
                errors.append("跨表匹配改变了主表行数。")
        finally:
            workbook.close()
        return TaskValidation(valid=not errors, errors=errors, warnings=result.warnings)


class NormalizeTool(SpreadsheetTool[NormalizeTaskSpec]):
    name = "normalize_column"
    description = "对单列执行确定性的文本、数字、日期或空值标准化。"
    spec_model = NormalizeTaskSpec
    risk_level = RiskLevel.MEDIUM

    def validate(self, context: ToolContext, spec: NormalizeTaskSpec) -> TaskValidation:
        _, column_error = _resolve_named_column(context.profile, spec.sheet, spec.column)
        result = TaskValidation(
            valid=column_error is None,
            errors=[column_error] if column_error else [],
        )
        if not spec.operations:
            result.errors.append("至少选择一种标准化操作。")
        if NormalizeOperation.DATE_FORMAT in spec.operations and not spec.date_format:
            result.errors.append("日期标准化必须指定输出格式。")
        if NormalizeOperation.FILL_BLANK in spec.operations and spec.fill_value is None:
            result.errors.append("填充空值必须指定填充值。")
        if NormalizeOperation.STRIP_SUFFIX in spec.operations and not spec.strip_suffixes:
            result.errors.append("去除后缀必须明确提供后缀文本。")
        sheet = _sheet_profile(context.profile, spec.sheet)
        column, _ = _resolve_named_column(context.profile, spec.sheet, spec.column)
        label = _source_label(column) if column else spec.column
        output_name = label if spec.overwrite else (spec.output_column or f"{label}_标准化")
        if sheet and not spec.overwrite and output_name in {column.name for column in sheet.columns}:
            result.errors.append(f"输出列“{output_name}”已存在。")
        result.valid = not result.errors
        return result

    def preview(self, context: ToolContext, spec: NormalizeTaskSpec) -> dict[str, Any]:
        sheet = _sheet_profile(context.profile, spec.sheet)
        column, _ = _resolve_named_column(context.profile, spec.sheet, spec.column)
        label = _source_label(column) if column else spec.column
        output_name = label if spec.overwrite else (spec.output_column or f"{label}_标准化")
        return {
            "modified_sheets": [spec.sheet],
            "created_sheets": [],
            "new_columns": [] if spec.overwrite else [output_name],
            "estimated_rows": max(0, sheet.max_row - (sheet.header_row or 1)) if sheet else 0,
            "destructive": spec.overwrite,
        }

    @staticmethod
    def _normalize(value: Any, spec: NormalizeTaskSpec) -> tuple[Any, bool]:
        current = value
        try:
            for operation in spec.operations:
                if operation == NormalizeOperation.FILL_BLANK:
                    if current in (None, ""):
                        current = spec.fill_value
                elif current in (None, ""):
                    continue
                elif operation == NormalizeOperation.NFKC:
                    if isinstance(current, str):
                        current = unicodedata.normalize("NFKC", current)
                elif operation == NormalizeOperation.TRIM:
                    if isinstance(current, str):
                        current = current.strip()
                elif operation == NormalizeOperation.LOWER:
                    current = str(current).lower()
                elif operation == NormalizeOperation.UPPER:
                    current = str(current).upper()
                elif operation == NormalizeOperation.TO_TEXT:
                    current = str(current)
                elif operation == NormalizeOperation.TO_NUMBER:
                    number = decimal.Decimal(str(current).strip())
                    current = int(number) if number == number.to_integral() else float(number)
                elif operation == NormalizeOperation.STRIP_SUFFIX:
                    text = str(current)
                    for suffix in sorted(spec.strip_suffixes, key=len, reverse=True):
                        if text.endswith(suffix):
                            text = text[:-len(suffix)]
                            break
                    current = text
                elif operation == NormalizeOperation.DATE_FORMAT:
                    if isinstance(current, dt.datetime):
                        parsed = current
                    elif isinstance(current, dt.date):
                        parsed = dt.datetime.combine(current, dt.time())
                    else:
                        parsed = dt.datetime.fromisoformat(str(current).strip())
                    current = parsed.strftime(spec.date_format or "%Y-%m-%d")
            return current, True
        except (ValueError, TypeError, decimal.InvalidOperation):
            return value, False

    def execute(self, context: ToolContext, spec: NormalizeTaskSpec) -> ToolResult:
        _copy_input(context)
        workbook = load_workbook(context.output_path, data_only=False, keep_links=True)
        sheet = workbook[spec.sheet]
        profile = _sheet_profile(context.profile, spec.sheet)
        column, error = _resolve_named_column(context.profile, spec.sheet, spec.column)
        if column is None:
            workbook.close()
            raise ValueError(error or f"无法解析列“{spec.column}”。")
        source_index = column.index
        target_index = source_index if spec.overwrite else sheet.max_column + 1
        label = _source_label(column)
        output_name = label if spec.overwrite else (spec.output_column or f"{label}_标准化")
        sheet.cell(profile.header_row, target_index, output_name)
        changed = failed = 0
        for row in range(profile.header_row + 1, profile.max_row + 1):
            original = sheet.cell(row, source_index).value
            normalized, ok = self._normalize(original, spec)
            if not ok:
                failed += 1
            if normalized != original:
                changed += 1
            sheet.cell(row, target_index, normalized)
        workbook.save(context.output_path)
        workbook.close()
        warnings = [f"{failed} 个值无法转换，已保留原值。"] if failed else []
        return ToolResult(
            tool_name=self.name,
            modified_sheets=[spec.sheet],
            modified_ranges=[
                f"{spec.sheet}!{get_column_letter(target_index)}{profile.header_row}:"
                f"{get_column_letter(target_index)}{profile.max_row}"
            ],
            rows_affected=max(0, profile.max_row - profile.header_row),
            counters={"changed": changed, "failed": failed},
            warnings=warnings,
            details={"output_column": output_name, "overwrite": spec.overwrite},
        )

    def verify(
        self, context: ToolContext, spec: NormalizeTaskSpec, result: ToolResult
    ) -> TaskValidation:
        profile = _sheet_profile(context.profile, spec.sheet)
        column, _ = _resolve_named_column(context.profile, spec.sheet, spec.column)
        source_index = column.index
        target_index = source_index if spec.overwrite else profile.max_column + 1
        label = _source_label(column)
        output_name = label if spec.overwrite else (spec.output_column or f"{label}_标准化")
        source = load_workbook(context.input_path, read_only=True, data_only=False)
        output = load_workbook(context.output_path, read_only=True, data_only=False)
        errors = []
        try:
            ws = output[spec.sheet]
            headers = next(ws.iter_rows(min_row=profile.header_row, max_row=profile.header_row, values_only=True))
            if len(headers) < target_index or headers[target_index - 1] != output_name:
                errors.append(f"未找到计划中的输出列“{output_name}”。")
            source_rows = source[spec.sheet].iter_rows(min_row=profile.header_row + 1, max_row=profile.max_row, min_col=source_index, max_col=source_index, values_only=True)
            target_rows = ws.iter_rows(min_row=profile.header_row + 1, max_row=profile.max_row, min_col=target_index, max_col=target_index, values_only=True)
            for row, (before, after) in enumerate(zip(source_rows, target_rows), profile.header_row + 1):
                expected, _ = self._normalize(before[0], spec)
                actual = after[0]
                if expected == "" and actual is None:
                    continue  # XLSX encodes empty text as a blank cell.
                if expected != actual or (isinstance(expected, str) and not isinstance(actual, str)):
                    errors.append(f"第 {row} 行标准化结果与计划不符。")
                    break
        finally:
            source.close()
            output.close()
        return TaskValidation(valid=not errors, errors=errors, warnings=result.warnings)


class AggregateTool(SpreadsheetTool[AggregateTaskSpec]):
    name = "aggregate"
    description = "按一至三列分组，并将确定性聚合结果写入新工作表。"
    spec_model = AggregateTaskSpec
    risk_level = RiskLevel.LOW

    def validate(self, context: ToolContext, spec: AggregateTaskSpec) -> TaskValidation:
        columns = [*spec.group_by, *(metric.column for metric in spec.metrics), *(item.column for item in spec.filters)]
        result = _base_validation(context.profile, spec.sheet, columns, spec.output_sheet)
        if not spec.group_by:
            result.errors.append("聚合任务至少需要一个分组列。")
        if not spec.metrics:
            result.errors.append("聚合任务至少需要一个指标。")
        sheet = _sheet_profile(context.profile, spec.sheet)
        types = {column.name: column.inferred_type for column in sheet.columns} if sheet else {}
        for metric in spec.metrics:
            if metric.function != AggregateFunction.COUNT and types.get(metric.column) not in {"number", "empty"}:
                result.errors.append(
                    f"指标列“{metric.column}”不是纯数值列；请先单独执行数据标准化。"
                )
        result.valid = not result.errors
        return result

    def preview(self, context: ToolContext, spec: AggregateTaskSpec) -> dict[str, Any]:
        sheet = _sheet_profile(context.profile, spec.sheet)
        return {
            "modified_sheets": [],
            "created_sheets": [spec.output_sheet],
            "new_columns": [*spec.group_by, *(metric.output_column for metric in spec.metrics)],
            "estimated_rows": max(0, sheet.max_row - (sheet.header_row or 1)) if sheet else 0,
            "destructive": False,
        }

    def execute(self, context: ToolContext, spec: AggregateTaskSpec) -> ToolResult:
        _copy_input(context)
        workbook = load_workbook(context.output_path, data_only=False, keep_links=True)
        values_workbook = load_workbook(context.output_path, data_only=True, keep_links=True)
        profile = _sheet_profile(context.profile, spec.sheet)
        indexes = _column_map(context.profile, spec.sheet)
        sheet = values_workbook[spec.sheet]
        groups: dict[tuple[Any, ...], list[list[Any]]] = defaultdict(
            lambda: [[] for _ in spec.metrics]
        )
        included = 0
        for row in range(profile.header_row + 1, profile.max_row + 1):
            values = {name: sheet.cell(row, index).value for name, index in indexes.items()}
            if spec.filters and not _row_matches(values, spec.filters):
                continue
            included += 1
            key = tuple(values[name] for name in spec.group_by)
            for metric_index, metric in enumerate(spec.metrics):
                groups[key][metric_index].append(values[metric.column])

        output = workbook.create_sheet(spec.output_sheet)
        output.append([*spec.group_by, *(metric.output_column for metric in spec.metrics)])
        non_numeric = 0
        for key in sorted(groups, key=lambda item: tuple(str(value) for value in item)):
            result_values: list[Any] = []
            for metric, raw_values in zip(spec.metrics, groups[key]):
                nonblank = [value for value in raw_values if value not in (None, "")]
                if metric.function == AggregateFunction.COUNT:
                    result_values.append(len(nonblank))
                    continue
                numeric: list[decimal.Decimal] = []
                for value in nonblank:
                    try:
                        numeric.append(decimal.Decimal(str(value)))
                    except decimal.InvalidOperation:
                        non_numeric += 1
                if not numeric:
                    result_values.append(None)
                elif metric.function == AggregateFunction.SUM:
                    result_values.append(float(sum(numeric)))
                elif metric.function == AggregateFunction.MEAN:
                    result_values.append(float(sum(numeric) / len(numeric)))
                elif metric.function == AggregateFunction.MIN:
                    result_values.append(float(min(numeric)))
                elif metric.function == AggregateFunction.MAX:
                    result_values.append(float(max(numeric)))
            output.append([*key, *result_values])
        workbook.save(context.output_path)
        workbook.close()
        values_workbook.close()
        warnings = [f"聚合时忽略了 {non_numeric} 个非数值指标。"] if non_numeric else []
        return ToolResult(
            tool_name=self.name,
            created_sheets=[spec.output_sheet],
            modified_ranges=[f"{spec.output_sheet}!A1:{get_column_letter(len(spec.group_by) + len(spec.metrics))}{len(groups) + 1}"],
            rows_affected=included,
            counters={"input_rows": included, "groups": len(groups), "non_numeric": non_numeric},
            warnings=warnings,
        )

    def verify(
        self, context: ToolContext, spec: AggregateTaskSpec, result: ToolResult
    ) -> TaskValidation:
        workbook = load_workbook(context.output_path, read_only=True, data_only=True)
        errors: list[str] = []
        if spec.output_sheet not in workbook.sheetnames:
            errors.append(f"未生成汇总工作表“{spec.output_sheet}”。")
        else:
            output_rows, _ = worksheet_dimensions(workbook[spec.output_sheet])
            rows = max(0, output_rows - 1)
            if rows != int(result.counters.get("groups", -1)):
                errors.append("汇总结果行数与分组计数不一致。")
        workbook.close()
        return TaskValidation(valid=not errors, errors=errors, warnings=result.warnings)


def _sort_value(value: Any) -> tuple[int, Any]:
    if value is None:
        return (4, "")
    if isinstance(value, bool):
        return (0, int(value))
    if isinstance(value, (int, float, decimal.Decimal)):
        return (1, float(value))
    if isinstance(value, (dt.date, dt.datetime)):
        return (2, value.isoformat())
    return (3, str(value).casefold())


class SortRankTool(SpreadsheetTool[SortRankTaskSpec]):
    name = "sort_rank"
    description = "对完整数据区域执行稳定排序、排名或 Top N，并写入新工作表。"
    spec_model = SortRankTaskSpec
    risk_level = RiskLevel.LOW

    def validate(self, context: ToolContext, spec: SortRankTaskSpec) -> TaskValidation:
        columns = [*(item.column for item in spec.sort_keys), *spec.group_by]
        if spec.rank_column:
            columns.append(spec.rank_column)
        result = _base_validation(context.profile, spec.sheet, columns, spec.output_sheet)
        if not spec.sort_keys and not spec.rank_column:
            result.errors.append("至少指定一个排序列或排名列。")
        result.valid = not result.errors
        return result

    def preview(self, context: ToolContext, spec: SortRankTaskSpec) -> dict[str, Any]:
        sheet = _sheet_profile(context.profile, spec.sheet)
        return {
            "modified_sheets": [],
            "created_sheets": [spec.output_sheet],
            "new_columns": [spec.rank_output_column] if spec.rank_column else [],
            "estimated_rows": min(
                max(0, sheet.max_row - (sheet.header_row or 1)) if sheet else 0,
                spec.top_n or 10**9,
            ),
            "destructive": False,
        }

    def execute(self, context: ToolContext, spec: SortRankTaskSpec) -> ToolResult:
        _copy_input(context)
        workbook = load_workbook(context.output_path, data_only=False, keep_links=True)
        values_workbook = load_workbook(context.output_path, data_only=True, keep_links=True)
        profile = _sheet_profile(context.profile, spec.sheet)
        indexes = _column_map(context.profile, spec.sheet)
        values_sheet = values_workbook[spec.sheet]
        source_sheet = workbook[spec.sheet]
        headers = [source_sheet.cell(profile.header_row, column).value for column in range(1, profile.max_column + 1)]
        rows = [
            [values_sheet.cell(row, column).value for column in range(1, profile.max_column + 1)]
            for row in range(profile.header_row + 1, profile.max_row + 1)
        ]
        rank_values: list[int | None] = [None] * len(rows)
        if spec.rank_column:
            rank_index = indexes[spec.rank_column] - 1
            groups: dict[tuple[Any, ...], list[int]] = defaultdict(list)
            for index, row in enumerate(rows):
                groups[tuple(row[indexes[column] - 1] for column in spec.group_by)].append(index)
            for members in groups.values():
                ordered = sorted(
                    members,
                    key=lambda index: _sort_value(rows[index][rank_index]),
                    reverse=spec.rank_direction == SortDirection.DESC,
                )
                dense_rank = 0
                start = 0
                while start < len(ordered):
                    end = start + 1
                    current = rows[ordered[start]][rank_index]
                    while end < len(ordered) and rows[ordered[end]][rank_index] == current:
                        end += 1
                    dense_rank += 1
                    if spec.rank_method.value == "dense":
                        rank = dense_rank
                    elif spec.rank_method.value == "min":
                        rank = start + 1
                    else:
                        rank = end
                    for position in range(start, end):
                        rank_values[ordered[position]] = rank
                    start = end

        combined = [(row, rank_values[index]) for index, row in enumerate(rows)]
        effective_sort = list(spec.sort_keys)
        if not effective_sort and spec.rank_column:
            from .schemas import SortKey

            effective_sort = [SortKey(column=spec.rank_output_column, direction=SortDirection.ASC)]
        for sort_key in reversed(effective_sort):
            if sort_key.column == spec.rank_output_column and spec.rank_column:
                combined.sort(key=lambda item: _sort_value(item[1]), reverse=sort_key.direction == SortDirection.DESC)
            else:
                index = indexes[sort_key.column] - 1
                combined.sort(key=lambda item: _sort_value(item[0][index]), reverse=sort_key.direction == SortDirection.DESC)

        if spec.top_n:
            if spec.rank_column:
                combined = [item for item in combined if item[1] is not None and item[1] <= spec.top_n]
            else:
                combined = combined[: spec.top_n]
        output = workbook.create_sheet(spec.output_sheet)
        output.append([*headers, *([spec.rank_output_column] if spec.rank_column else [])])
        for row, rank in combined:
            output.append([*row, *([rank] if spec.rank_column else [])])
        workbook.save(context.output_path)
        workbook.close()
        values_workbook.close()
        return ToolResult(
            tool_name=self.name,
            created_sheets=[spec.output_sheet],
            modified_ranges=[f"{spec.output_sheet}!A1:{get_column_letter(len(headers) + (1 if spec.rank_column else 0))}{len(combined) + 1}"],
            rows_affected=len(combined),
            counters={"input_rows": len(rows), "output_rows": len(combined)},
        )

    def verify(
        self, context: ToolContext, spec: SortRankTaskSpec, result: ToolResult
    ) -> TaskValidation:
        workbook = load_workbook(context.output_path, read_only=True, data_only=True)
        errors: list[str] = []
        if spec.output_sheet not in workbook.sheetnames:
            errors.append(f"未生成排序结果工作表“{spec.output_sheet}”。")
        else:
            output_rows, _ = worksheet_dimensions(workbook[spec.output_sheet])
            actual = max(0, output_rows - 1)
            if actual != int(result.counters.get("output_rows", -1)):
                errors.append("排序结果行数与执行记录不一致。")
        workbook.close()
        return TaskValidation(valid=not errors, errors=errors, warnings=result.warnings)


class CleanTool(SpreadsheetTool[CleanTaskSpec]):
    name = "clean"
    description = "识别、标记或移除精确重复记录，按结构化条件筛选记录，或只保留明确指定的列。"
    spec_model = CleanTaskSpec
    risk_level = RiskLevel.MEDIUM

    def validate(self, context: ToolContext, spec: CleanTaskSpec) -> TaskValidation:
        columns = [*spec.key_columns, *(item.column for item in spec.filters)]
        result = _base_validation(
            context.profile,
            spec.sheet,
            columns,
            None if spec.operation == CleanOperation.KEEP_COLUMNS and spec.overwrite else spec.output_sheet,
        )
        if spec.operation in {CleanOperation.MARK_DUPLICATES, CleanOperation.REMOVE_DUPLICATES} and not spec.key_columns:
            result.errors.append("重复记录操作至少需要一个判断列。")
        if spec.operation == CleanOperation.FILTER_ROWS and not spec.filters:
            result.errors.append("筛选任务至少需要一个条件。")
        if spec.operation == CleanOperation.KEEP_COLUMNS:
            if not spec.keep_columns:
                result.errors.append("列保留任务至少需要一个保留列。")
            resolved_indexes: list[int] = []
            for ref in spec.keep_columns:
                if ref.sheet != spec.sheet:
                    result.errors.append(
                        f"保留列“{ref.display_header or ref.column_letter or ref.column_index}”不属于工作表“{spec.sheet}”。"
                    )
                    continue
                column, error = _resolve_column_ref(context.profile, ref)
                if error:
                    result.errors.append(error)
                elif column:
                    resolved_indexes.append(column.index)
            if len(resolved_indexes) != len(set(resolved_indexes)):
                result.errors.append("保留列列表中存在重复列。")
            sheet = _sheet_profile(context.profile, spec.sheet)
            if sheet and len(set(resolved_indexes)) < sheet.max_column:
                if sheet.merged_ranges:
                    result.errors.append("该工作表包含合并单元格，暂不能安全删除其中的列。")
                if any(column.formula_count for column in sheet.columns):
                    result.errors.append("该工作表包含公式，暂不能在不破坏引用的前提下删除列。")
        result.valid = not result.errors
        return result

    def preview(self, context: ToolContext, spec: CleanTaskSpec) -> dict[str, Any]:
        sheet = _sheet_profile(context.profile, spec.sheet)
        if spec.operation == CleanOperation.KEEP_COLUMNS:
            kept = []
            for ref in spec.keep_columns:
                column, _ = _resolve_column_ref(context.profile, ref)
                if column and column.index not in kept:
                    kept.append(column.index)
            deleted = max(0, (sheet.max_column if sheet else 0) - len(kept))
            return {
                "modified_sheets": [spec.sheet] if spec.overwrite else [],
                "created_sheets": [] if spec.overwrite else [spec.output_sheet],
                "new_columns": [],
                "estimated_rows": max(0, sheet.max_row - (sheet.header_row or 1)) if sheet else 0,
                "destructive": bool(spec.overwrite and deleted),
                "summary_counts": {
                    "kept_columns": len(kept),
                    "deleted_columns": deleted,
                },
            }
        return {
            "modified_sheets": [],
            "created_sheets": [spec.output_sheet],
            "new_columns": [spec.duplicate_flag_column]
            if spec.operation == CleanOperation.MARK_DUPLICATES
            else [],
            "estimated_rows": max(0, sheet.max_row - (sheet.header_row or 1)) if sheet else 0,
            "destructive": spec.operation == CleanOperation.REMOVE_DUPLICATES,
        }

    def execute(self, context: ToolContext, spec: CleanTaskSpec) -> ToolResult:
        _copy_input(context)
        workbook = load_workbook(context.output_path, data_only=False, keep_links=True)
        profile = _sheet_profile(context.profile, spec.sheet)
        if spec.operation == CleanOperation.KEEP_COLUMNS:
            resolved = []
            for ref in spec.keep_columns:
                column, error = _resolve_column_ref(context.profile, ref)
                if error or not column:
                    workbook.close()
                    raise ValueError(error or "无法解析保留列。")
                if column.index not in {item.index for item in resolved}:
                    resolved.append(column)
            resolved.sort(key=lambda item: item.index)
            kept_indexes = [item.index for item in resolved]
            source = workbook[spec.sheet]
            row_count = profile.max_row if profile else source.max_row
            original_columns = profile.max_column if profile else source.max_column
            deleted_columns = original_columns - len(kept_indexes)

            if spec.overwrite:
                remove_indexes = [
                    index for index in range(1, original_columns + 1)
                    if index not in set(kept_indexes)
                ]
                groups: list[tuple[int, int]] = []
                for index in remove_indexes:
                    if groups and index == groups[-1][1] + 1:
                        groups[-1] = (groups[-1][0], index)
                    else:
                        groups.append((index, index))
                for start, end in reversed(groups):
                    source.delete_cols(start, end - start + 1)
                target_sheet = spec.sheet
                created_sheets: list[str] = []
                modified_sheets = [spec.sheet]
            else:
                output = workbook.create_sheet(spec.output_sheet)
                for target_column, source_column in enumerate(kept_indexes, 1):
                    source_dimension = source.column_dimensions.get(get_column_letter(source_column))
                    if source_dimension is not None:
                        target_letter = get_column_letter(target_column)
                        output_dimension = copy(source_dimension)
                        output_dimension.index = target_letter
                        output.column_dimensions[target_letter] = output_dimension
                    for row in range(1, row_count + 1):
                        source_cell = source.cell(row, source_column)
                        target_cell = output.cell(row, target_column, source_cell.value)
                        if source_cell.has_style:
                            target_cell._style = copy(source_cell._style)
                        if source_cell.number_format:
                            target_cell.number_format = source_cell.number_format
                        if source_cell.hyperlink:
                            target_cell._hyperlink = copy(source_cell.hyperlink)
                        if source_cell.comment:
                            target_cell.comment = copy(source_cell.comment)
                target_sheet = spec.output_sheet
                created_sheets = [spec.output_sheet]
                modified_sheets = []

            workbook.save(context.output_path)
            workbook.close()
            last_column = get_column_letter(max(1, len(kept_indexes)))
            return ToolResult(
                tool_name=self.name,
                modified_sheets=modified_sheets,
                created_sheets=created_sheets,
                modified_ranges=[f"{target_sheet}!A1:{last_column}{row_count}"],
                rows_affected=max(0, row_count - ((profile.header_row or 1) if profile else 1)),
                counters={
                    "input_columns": original_columns,
                    "output_columns": len(kept_indexes),
                    "deleted_columns": deleted_columns,
                    "output_rows": max(0, row_count - ((profile.header_row or 1) if profile else 1)),
                },
                details={
                    "operation": spec.operation.value,
                    "overwrite": spec.overwrite,
                    "kept_source_columns": kept_indexes,
                    "kept_columns": [
                        item.display_header or item.name or item.letter for item in resolved
                    ],
                },
            )

        values_workbook = load_workbook(context.output_path, data_only=True, keep_links=True)
        indexes = _column_map(context.profile, spec.sheet)
        source = values_workbook[spec.sheet]
        headers = [source.cell(profile.header_row, column).value for column in range(1, profile.max_column + 1)]
        rows = [
            [source.cell(row, column).value for column in range(1, profile.max_column + 1)]
            for row in range(profile.header_row + 1, profile.max_row + 1)
        ]
        selected: list[list[Any]] = []
        duplicate_count = 0
        output_headers = list(headers)
        if spec.operation == CleanOperation.FILTER_ROWS:
            for row in rows:
                values = {str(headers[index]): value for index, value in enumerate(row)}
                if _row_matches(values, spec.filters):
                    selected.append(row)
        else:
            key_indexes = [indexes[column] - 1 for column in spec.key_columns]
            keys = [tuple(row[index] for index in key_indexes) for row in rows]
            counts = Counter(keys)
            duplicate_count = sum(count - 1 for count in counts.values() if count > 1)
            if spec.operation == CleanOperation.MARK_DUPLICATES:
                output_headers.append(spec.duplicate_flag_column)
                selected = [row + ["重复" if counts[key] > 1 else ""] for row, key in zip(rows, keys)]
            else:
                keep_indexes: set[int] = set()
                ordered = range(len(rows)) if spec.keep == "first" else range(len(rows) - 1, -1, -1)
                seen: set[tuple[Any, ...]] = set()
                for index in ordered:
                    if keys[index] not in seen:
                        keep_indexes.add(index)
                        seen.add(keys[index])
                selected = [row for index, row in enumerate(rows) if index in keep_indexes]

        output = workbook.create_sheet(spec.output_sheet)
        output.append(output_headers)
        for row in selected:
            output.append(row)
        workbook.save(context.output_path)
        workbook.close()
        values_workbook.close()
        return ToolResult(
            tool_name=self.name,
            created_sheets=[spec.output_sheet],
            modified_ranges=[f"{spec.output_sheet}!A1:{get_column_letter(len(output_headers))}{len(selected) + 1}"],
            rows_affected=len(selected),
            counters={
                "input_rows": len(rows),
                "output_rows": len(selected),
                "duplicates_removed": duplicate_count
                if spec.operation == CleanOperation.REMOVE_DUPLICATES
                else 0,
                "duplicate_rows": sum(1 for key in (keys if spec.operation != CleanOperation.FILTER_ROWS else []) if counts[key] > 1)
                if spec.operation != CleanOperation.FILTER_ROWS
                else 0,
            },
        )

    def verify(
        self, context: ToolContext, spec: CleanTaskSpec, result: ToolResult
    ) -> TaskValidation:
        if spec.operation == CleanOperation.KEEP_COLUMNS:
            before = load_workbook(context.input_path, read_only=True, data_only=False, keep_links=True)
            after = load_workbook(context.output_path, read_only=True, data_only=False, keep_links=True)
            errors: list[str] = []
            try:
                target_name = spec.sheet if spec.overwrite else spec.output_sheet
                if target_name not in after.sheetnames:
                    errors.append(f"未生成列保留结果工作表“{target_name}”。")
                    return TaskValidation(valid=False, errors=errors, warnings=result.warnings)
                source = before[spec.sheet]
                target = after[target_name]
                indexes: list[int] = []
                for ref in spec.keep_columns:
                    column, error = _resolve_column_ref(context.profile, ref)
                    if error or not column:
                        errors.append(error or "无法解析保留列。")
                    elif column.index not in indexes:
                        indexes.append(column.index)
                indexes.sort()
                if target.max_column != len(indexes):
                    errors.append("列保留结果的列数与计划不一致。")
                if target.max_row != source.max_row:
                    errors.append("列保留操作改变了行数。")
                source_rows = source.iter_rows(
                    min_row=1, max_row=source.max_row, max_col=source.max_column
                )
                target_rows = target.iter_rows(
                    min_row=1, max_row=target.max_row, max_col=max(1, target.max_column)
                )
                for source_row, target_row in zip(source_rows, target_rows):
                    for target_column, source_column in enumerate(indexes, 1):
                        left = source_row[source_column - 1]
                        right = target_row[target_column - 1]
                        if (left.value, left.data_type) != (right.value, right.data_type):
                            errors.append(
                                f"列保留结果与来源值不一致：{target_name}!{right.coordinate}。"
                            )
                            break
                        if (left.number_format or "General") != (right.number_format or "General"):
                            errors.append(
                                f"列保留结果的数字格式与来源不一致：{target_name}!{right.coordinate}。"
                            )
                            break
                        if getattr(left, "_style_id", 0) != getattr(right, "_style_id", 0):
                            errors.append(
                                f"列保留结果的单元格样式与来源不一致：{target_name}!{right.coordinate}。"
                            )
                            break
                    if errors:
                        break
            finally:
                before.close()
                after.close()
            return TaskValidation(valid=not errors, errors=errors, warnings=result.warnings)

        workbook = load_workbook(context.output_path, read_only=True, data_only=True)
        errors: list[str] = []
        if spec.output_sheet not in workbook.sheetnames:
            errors.append(f"未生成清洗结果工作表“{spec.output_sheet}”。")
        else:
            output_rows, _ = worksheet_dimensions(workbook[spec.output_sheet])
            actual = max(0, output_rows - 1)
            if actual != int(result.counters.get("output_rows", -1)):
                errors.append("清洗结果行数与执行记录不一致。")
        workbook.close()
        return TaskValidation(valid=not errors, errors=errors, warnings=result.warnings)


class FormulaTool(SpreadsheetTool[FormulaTaskSpec]):
    name = "formula"
    description = "诊断公式断层、异常引用与模式偏差，并在候选唯一时执行受限修复。"
    spec_model = FormulaTaskSpec
    risk_level = RiskLevel.MEDIUM

    def validate(self, context: ToolContext, spec: FormulaTaskSpec) -> TaskValidation:
        if spec.mode == FormulaMode.RECALCULATE:
            errors = []
            if spec.sheet and not _sheet_profile(context.profile, spec.sheet):
                errors.append(f"工作表“{spec.sheet}”不存在。")
            return TaskValidation(valid=not errors, errors=errors)
        column, column_error = _resolve_formula_column(context.profile, spec.sheet, spec.column)
        result = TaskValidation(valid=column_error is None, errors=[column_error] if column_error else [])
        profile = _sheet_profile(context.profile, spec.sheet)
        if profile:
            start = spec.start_row or profile.header_row + 1
            end = spec.end_row or profile.max_row
            if start <= profile.header_row:
                result.errors.append("公式区域必须从表头下一行开始。")
            if end > profile.max_row:
                result.errors.append("公式区域超过工作表数据范围。")
            if start > end:
                result.errors.append("公式区域起始行不能晚于结束行。")
            if spec.mode in {FormulaMode.SET_FORMULA, FormulaMode.COPY_FORMULA}:
                if spec.start_row is None or spec.end_row is None:
                    result.errors.append("设置或复制公式时必须明确起止行。")
            if spec.mode == FormulaMode.SET_FORMULA:
                if not spec.formula or not spec.formula.startswith("="):
                    result.errors.append("设置公式时必须提供以 = 开头的公式文本。")
            if spec.mode == FormulaMode.COPY_FORMULA:
                if spec.source_row is None:
                    result.errors.append("复制公式时必须指定样板公式所在行。")
                elif spec.source_row <= profile.header_row or spec.source_row > profile.max_row:
                    result.errors.append("样板公式行超过当前数据范围。")
                else:
                    workbook = load_workbook(context.input_path, read_only=True, data_only=False)
                    try:
                        column_index = column.index if column else None
                        source = workbook[spec.sheet].cell(spec.source_row, column_index).value if column_index else None
                        if not self._is_formula(source):
                            result.errors.append("指定样板行中没有可复制的公式。")
                    finally:
                        workbook.close()
        result.valid = not result.errors
        return result

    def preview(self, context: ToolContext, spec: FormulaTaskSpec) -> dict[str, Any]:
        if spec.mode == FormulaMode.RECALCULATE:
            formula_sheets = [
                sheet.name for sheet in context.profile.sheets
                if (not spec.sheet or sheet.name == spec.sheet) and any(column.formula_count for column in sheet.columns)
            ]
            formula_count = sum(
                column.formula_count
                for sheet in context.profile.sheets
                if not spec.sheet or sheet.name == spec.sheet
                for column in sheet.columns
            )
            return {
                "modified_sheets": formula_sheets,
                "created_sheets": [],
                "new_columns": [],
                "estimated_rows": formula_count,
                "destructive": False,
            }
        profile = _sheet_profile(context.profile, spec.sheet)
        start = spec.start_row or ((profile.header_row + 1) if profile else 1)
        end = spec.end_row or (profile.max_row if profile else start)
        return {
            "modified_sheets": [] if spec.mode == FormulaMode.DIAGNOSE else [spec.sheet],
            "created_sheets": [],
            "new_columns": [],
            "estimated_rows": max(0, end - start + 1),
            "destructive": spec.mode == FormulaMode.REPLACE_OUTLIERS or (
                spec.mode in {FormulaMode.SET_FORMULA, FormulaMode.COPY_FORMULA}
                and spec.overwrite_existing
            ),
        }

    @staticmethod
    def _is_formula(value: Any) -> bool:
        return isinstance(value, str) and value.startswith("=")

    @staticmethod
    def _translated(formula: str, source_coordinate: str, target_coordinate: str) -> str | None:
        try:
            return Translator(formula, origin=source_coordinate).translate_formula(target_coordinate)
        except Exception:
            return None

    def _inspect(self, sheet, column_index: int, start: int, end: int) -> dict[str, Any]:
        rows = list(range(start, end + 1))
        values = [sheet.cell(row, column_index).value for row in rows]
        formulas = {
            row: value for row, value in zip(rows, values) if self._is_formula(value)
        }
        broken = [row for row, formula in formulas.items() if "#REF!" in formula.upper()]
        above_rows: list[int | None] = []
        nearest: int | None = None
        for row in rows:
            above_rows.append(nearest)
            if row in formulas:
                nearest = row
        below_rows: list[int | None] = [None] * len(rows)
        nearest = None
        for index in range(len(rows) - 1, -1, -1):
            row = rows[index]
            below_rows[index] = nearest
            if row in formulas:
                nearest = row
        candidates: dict[int, str] = {}
        outliers: list[int] = []
        ambiguous: list[int] = []
        for index, row in enumerate(rows):
            current = values[index]
            above = above_rows[index]
            below = below_rows[index]
            if above is None or below is None:
                continue
            target = sheet.cell(row, column_index).coordinate
            from_above = self._translated(
                formulas[above], sheet.cell(above, column_index).coordinate, target
            )
            from_below = self._translated(
                formulas[below], sheet.cell(below, column_index).coordinate, target
            )
            if not from_above or from_above != from_below:
                if current in (None, ""):
                    ambiguous.append(row)
                continue
            if current in (None, ""):
                candidates[row] = from_above
            elif self._is_formula(current) and current != from_above:
                candidates[row] = from_above
                outliers.append(row)
        return {
            "formula_count": len(formulas),
            "broken_rows": broken,
            "gap_rows": [row for row in candidates if sheet.cell(row, column_index).value in (None, "")],
            "outlier_rows": outliers,
            "ambiguous_rows": sorted(set(ambiguous)),
            "candidates": candidates,
        }

    def execute(self, context: ToolContext, spec: FormulaTaskSpec) -> ToolResult:
        if context.report_stage:
            context.report_stage("正在批量写入与检查公式")
        timings: dict[str, float] = {}
        stage_started = time.perf_counter()
        _copy_input(context)
        workbook = load_workbook(context.output_path, data_only=False, keep_links=True)
        timings["load"] = time.perf_counter() - stage_started
        inspection: dict[str, Any] = {
            "formula_count": 0, "broken_rows": [], "gap_rows": [], "outlier_rows": [],
            "ambiguous_rows": [], "candidates": {},
        }
        repaired: list[int] = []
        generated: list[int] = []
        copied: list[int] = []
        column_index: int | None = None
        inspect_started = time.perf_counter()
        if spec.mode == FormulaMode.RECALCULATE:
            target_formula_sheets = [
                sheet.name for sheet in context.profile.sheets
                if (not spec.sheet or sheet.name == spec.sheet) and any(column.formula_count for column in sheet.columns)
            ]
            inspection["formula_count"] = sum(
                column.formula_count
                for sheet in context.profile.sheets
                if not spec.sheet or sheet.name == spec.sheet
                for column in sheet.columns
            )
        else:
            profile = _sheet_profile(context.profile, spec.sheet)
            column_index = _resolve_formula_column(context.profile, spec.sheet, spec.column)[0].index
            start = spec.start_row or profile.header_row + 1
            end = spec.end_row or profile.max_row
            sheet = workbook[spec.sheet]
            if spec.mode in {
                FormulaMode.DIAGNOSE,
                FormulaMode.FILL_GAPS,
                FormulaMode.REPLACE_OUTLIERS,
            }:
                inspection = self._inspect(sheet, column_index, start, end)
            timings["inspect"] = time.perf_counter() - inspect_started
            target_formula_sheets = [spec.sheet]
            write_started = time.perf_counter()
            if spec.mode in {FormulaMode.FILL_GAPS, FormulaMode.REPLACE_OUTLIERS}:
                selected = inspection["gap_rows"] if spec.mode == FormulaMode.FILL_GAPS else inspection["outlier_rows"]
                for row in selected:
                    sheet.cell(row, column_index, inspection["candidates"][row])
                    repaired.append(row)
            elif spec.mode == FormulaMode.SET_FORMULA:
                assert spec.formula and spec.start_row and spec.end_row
                origin = sheet.cell(spec.start_row, column_index).coordinate
                for row in range(spec.start_row, spec.end_row + 1):
                    cell = sheet.cell(row, column_index)
                    if cell.value not in (None, "") and not spec.overwrite_existing:
                        continue
                    translated = self._translated(spec.formula, origin, cell.coordinate)
                    if translated:
                        cell.value = translated
                        generated.append(row)
            elif spec.mode == FormulaMode.COPY_FORMULA:
                assert spec.source_row and spec.start_row and spec.end_row
                source_cell = sheet.cell(spec.source_row, column_index)
                for row in range(spec.start_row, spec.end_row + 1):
                    cell = sheet.cell(row, column_index)
                    if row == spec.source_row:
                        continue
                    if cell.value not in (None, "") and not spec.overwrite_existing:
                        continue
                    translated = self._translated(source_cell.value, source_cell.coordinate, cell.coordinate)
                    if translated:
                        cell.value = translated
                        copied.append(row)
            timings["write"] = time.perf_counter() - write_started
        if spec.mode == FormulaMode.RECALCULATE:
            timings["inspect"] = time.perf_counter() - inspect_started
            timings["write"] = 0.0
        save_started = time.perf_counter()
        workbook.save(context.output_path)
        workbook.close()
        timings["save"] = time.perf_counter() - save_started
        recalculate_started = time.perf_counter()
        if context.report_stage and spec.recalculate:
            context.report_stage("正在重算公式并检查计算值")
        remaining = max(0, context.execution_deadline - time.monotonic()) if context.execution_deadline else 90
        calculation = recalculate_workbook(context.output_path, timeout=min(90, remaining), cancelled=context.cancelled) if spec.recalculate else None
        timings["recalculate_and_audit"] = time.perf_counter() - recalculate_started
        warnings: list[str] = []
        if inspection["ambiguous_rows"]:
            warnings.append(f"{len(inspection['ambiguous_rows'])} 行无法得到唯一公式候选，未修改。")
        if inspection["broken_rows"]:
            warnings.append(f"检测到 {len(inspection['broken_rows'])} 个包含 #REF! 的公式。")
        if spec.recalculate:
            assert calculation is not None
            if calculation.backend == "none":
                warnings.append("公式重算未完成：" + calculation.message)
            if calculation.error_cells:
                warnings.append(f"重算后仍有 {len(calculation.error_cells)} 个公式错误，需要人工处理。")
            if calculation.blank_cached_cells:
                warnings.append(f"{len(calculation.blank_cached_cells)} 个公式缺少计算缓存，未提交正式版本。")
        changed_rows = [*repaired, *generated, *copied]
        modified_ranges = (
            [f"{spec.sheet}!{get_column_letter(column_index)}{row}" for row in changed_rows]
            if column_index else []
        )
        modified_sheets = [spec.sheet] if changed_rows else (
            target_formula_sheets if calculation and calculation.backend != "none" else []
        )
        return ToolResult(
            tool_name=self.name,
            modified_sheets=modified_sheets,
            modified_ranges=modified_ranges,
            rows_affected=len(changed_rows),
            counters={
                "formula_count": inspection["formula_count"],
                "gaps": len(inspection["gap_rows"]),
                "outliers": len(inspection["outlier_rows"]),
                "broken_refs": len(inspection["broken_rows"]),
                "repaired": len(repaired),
                "generated": len(generated),
                "copied": len(copied),
                "calculated": calculation.calculated_cells if calculation else 0,
                "calculation_errors": len(calculation.error_cells) if calculation else 0,
            },
            warnings=warnings,
            details={
                key: value for key, value in inspection.items() if key != "candidates"
            } | {
                "mode": spec.mode.value,
                "recalculated": bool(calculation and calculation.success),
                "calculation": calculation.model_dump() if calculation else None,
                "timings_seconds": {
                    key: round(value, 6) for key, value in timings.items()
                },
            },
        )

    def verify(
        self, context: ToolContext, spec: FormulaTaskSpec, result: ToolResult
    ) -> TaskValidation:
        started = time.perf_counter()
        workbook = load_workbook(context.output_path, read_only=True, data_only=False)
        errors: list[str] = []
        targets: dict[str, set[str]] = defaultdict(set)
        origin = None
        if spec.mode == FormulaMode.SET_FORMULA:
            column = _resolve_formula_column(context.profile, spec.sheet, spec.column)[0].index
            origin = f"{get_column_letter(column)}{spec.start_row}"
        for address in result.modified_ranges:
            sheet_name, coordinate = address.rsplit("!", 1)
            targets[sheet_name].add(coordinate)
        try:
            for name, pending in targets.items():
                for row in workbook[name].iter_rows():
                    for cell in row:
                        if getattr(cell, "coordinate", None) in pending:
                            pending.remove(cell.coordinate)
                            if not self._is_formula(cell.value):
                                errors.append(f"{name}!{cell.coordinate} 未写入有效公式。")
                            elif spec.mode == FormulaMode.SET_FORMULA:
                                expected = self._translated(spec.formula, origin, cell.coordinate)
                                if cell.value != expected:
                                    errors.append(f"{name}!{cell.coordinate} 公式与批准的计算规则不符。")
                    if not pending:
                        break
                if pending:
                    errors.append(f"{name} 有 {len(pending)} 个目标公式未写入。")
        finally:
            workbook.close()
        calculation = result.details.get("calculation") or {}
        if spec.recalculate and not calculation.get("success"):
            errors.append("公式计算未完成，正式版本未提交：" + str(calculation.get("message") or "计算缓存缺失或存在公式错误"))
        result.details.setdefault("timings_seconds", {})["verify"] = round(time.perf_counter() - started, 6)
        warnings = list(dict.fromkeys(result.warnings))
        return TaskValidation(valid=not errors, errors=errors, warnings=warnings)


_NUMBER = r"[0-9]+(?:\.[0-9]+)?"
_DIMENSION_PATTERN = re.compile(
    rf"(?<![0-9.])(?P<a>{_NUMBER})\s*(?P<ua>mm|cm)?\s*[*xX×/／]\s*"
    rf"(?P<b>{_NUMBER})\s*(?P<ub>mm|cm)?"
    rf"(?:\s*[*xX×/／]\s*(?P<c>{_NUMBER})\s*(?P<uc>mm|cm)?)?",
    re.IGNORECASE,
)
_DIAMETER_PATTERN = re.compile(
    rf"(?:总?直径)\s*(?:约|均|为)?\s*[:：]?\s*(?P<d>{_NUMBER})\s*(?P<u>mm|cm)?",
    re.IGNORECASE,
)
_ADJACENT_UNIT_PATTERN = re.compile(
    rf"(?P<a>{_NUMBER})\s*(?P<ua>mm|cm)\s*"
    rf"(?P<b>{_NUMBER})\s*(?P<ub>mm|cm)"
    rf"(?:\s*(?P<c>{_NUMBER})\s*(?P<uc>mm|cm))?",
    re.IGNORECASE,
)


def _convert_dimension(values: list[float], units: list[str | None], default_unit: str) -> tuple[list[float] | None, str | None]:
    explicit = {unit.casefold() for unit in units if unit}
    if len(explicit) > 1:
        return None, "unit_conflict"
    if explicit and any(unit is None for unit in units) and next(iter(explicit)) != default_unit:
        return None, "unit_conflict"
    unit = next(iter(explicit), default_unit)
    factor = 10.0 if unit == "cm" else 1.0
    converted = [value * factor for value in values]
    if any(not math.isfinite(value) or value <= 0 for value in converted):
        return None, "invalid_format"
    return converted, None


def _dimension_action(
    policies: dict[str, DimensionResolutionAction] | None, status: str
) -> DimensionResolutionAction | None:
    action = (policies or {}).get(status)
    return DimensionResolutionAction(action) if action is not None else None


def _zero_if_authorized(
    policies: dict[str, DimensionResolutionAction] | None, status: str
) -> tuple[list[float | None], str] | None:
    if _dimension_action(policies, status) == DimensionResolutionAction.ZERO:
        return [0.0, 0.0, 0.0, 0.0], "resolved_zero"
    return None


def _dimension_partial_has_context(prefix: str, suffix: str) -> bool:
    """Return whether a 1D/2D measurement has explanatory text around it."""
    surrounding = unicodedata.normalize("NFKC", f"{prefix} {suffix}")
    surrounding = re.sub(r"(?:肿瘤)?(?:大小|尺寸)\s*[:：]?", "", surrounding)
    surrounding = re.sub(r"\b(?:mm|cm)\b", "", surrounding, flags=re.IGNORECASE)
    return bool(re.search(r"[A-Za-z一-鿿]", surrounding))


def _parse_dimensions(
    value: Any,
    default_unit: str,
    resolution_policies: dict[str, DimensionResolutionAction] | None = None,
    value_corrections: dict[str, str] | None = None,
    require_context_for_partial: bool = False,
) -> tuple[list[float | None], str]:
    if value in (None, "") or _normalized_text(value) == "":
        return _zero_if_authorized(resolution_policies, "blank") or (
            [None, None, None, None], "blank"
        )
    text = unicodedata.normalize("NFKC", str(value)).strip()
    corrections = {
        unicodedata.normalize("NFKC", str(source)).strip(): replacement
        for source, replacement in (value_corrections or {}).items()
    }
    if text in corrections:
        corrected, corrected_status = _parse_dimensions(
            corrections[text], default_unit, resolution_policies, None,
            require_context_for_partial,
        )
        if corrected_status in {
            "parsed_3d", "parsed_2d", "diameter_only",
            "resolved_multiple_first", "resolved_multiple_last",
            "resolved_unit_conflict", "resolved_zero",
        }:
            return corrected, "corrected_value"
        return corrected, corrected_status
    adjacent = _ADJACENT_UNIT_PATTERN.fullmatch(text)
    if adjacent:
        raw = [float(adjacent.group("a")), float(adjacent.group("b"))]
        units = [adjacent.group("ua"), adjacent.group("ub")]
        if adjacent.group("c") is not None:
            raw.append(float(adjacent.group("c")))
            units.append(adjacent.group("uc"))
        converted, issue = _convert_dimension(raw, units, default_unit)
        if issue or converted is None:
            return [None, None, None, None], issue or "invalid_format"
        if len(converted) == 3:
            return [*converted, max(converted)], "parsed_3d"
        if require_context_for_partial:
            return _zero_if_authorized(resolution_policies, "invalid_format") or (
                [None, None, None, None], "invalid_format"
            )
        return [converted[0], converted[1], None, max(converted)], "parsed_2d"
    # These shapes can otherwise expose a valid-looking suffix and silently hide
    # a malformed or additional leading dimension.
    four_or_more = re.search(
        rf"{_NUMBER}\s*(?:mm|cm)?\s*[*xX×/／]\s*"
        rf"{_NUMBER}\s*(?:mm|cm)?\s*[*xX×/／]\s*"
        rf"{_NUMBER}\s*(?:mm|cm)?\s*[*xX×/／]\s*{_NUMBER}",
        text,
        re.IGNORECASE,
    )
    if "**" in text or four_or_more or re.search(r"\d+\.\s*[*xX×/／]", text):
        return _zero_if_authorized(resolution_policies, "invalid_format") or (
            [None, None, None, None], "invalid_format"
        )
    candidates = list(_DIMENSION_PATTERN.finditer(text))
    diameters = list(_DIAMETER_PATTERN.finditer(text))
    if len(candidates) + len(diameters) > 1:
        action = _dimension_action(resolution_policies, "multiple_candidates")
        matches = sorted([*candidates, *diameters], key=lambda item: item.start())
        if action in {
            DimensionResolutionAction.FIRST_CANDIDATE,
            DimensionResolutionAction.LAST_CANDIDATE,
        }:
            selected = matches[0] if action == DimensionResolutionAction.FIRST_CANDIDATE else matches[-1]
            parsed, status = _parse_dimensions(selected.group(0), default_unit)
            if status in {"parsed_3d", "parsed_2d", "diameter_only"}:
                suffix = "first" if action == DimensionResolutionAction.FIRST_CANDIDATE else "last"
                return parsed, f"resolved_multiple_{suffix}"
        return _zero_if_authorized(resolution_policies, "multiple_candidates") or (
            [None, None, None, None], "multiple_candidates"
        )
    if len(candidates) == 1:
        match = candidates[0]
        prefix = text[:match.start()].strip()
        suffix = text[match.end():].strip()
        # A unique dimension group remains reliable when it is embedded in a
        # narrative note. Reject numeric material outside the candidate, since
        # that would indicate a malformed or additional measurement.
        prefix_ok = not re.search(r"\d", prefix)
        suffix_without_notes = re.sub(r"(?:\([^()]*\)|（[^（）]*）)", "", suffix)
        suffix_ok = not suffix_without_notes.strip(" \t\r\n,，;；。")
        if not prefix_ok or not suffix_ok:
            return _zero_if_authorized(resolution_policies, "invalid_format") or (
                [None, None, None, None], "invalid_format"
            )
        raw = [float(match.group("a")), float(match.group("b"))]
        units: list[str | None] = [match.group("ua"), match.group("ub")]
        if match.group("c") is not None:
            raw.append(float(match.group("c")))
            units.append(match.group("uc"))
        converted, issue = _convert_dimension(raw, units, default_unit)
        if (
            issue == "unit_conflict"
            and _dimension_action(resolution_policies, "unit_conflict")
            == DimensionResolutionAction.CONVERT_EACH_TO_MM
        ):
            converted = [
                number * (10.0 if (unit or default_unit).casefold() == "cm" else 1.0)
                for number, unit in zip(raw, units)
            ]
            issue = None
            if any(not math.isfinite(number) or number <= 0 for number in converted):
                return [None, None, None, None], "invalid_format"
            if len(converted) == 3:
                return [converted[0], converted[1], converted[2], max(converted)], "resolved_unit_conflict"
            if require_context_for_partial and not _dimension_partial_has_context(prefix, suffix):
                return _zero_if_authorized(resolution_policies, "invalid_format") or (
                    [None, None, None, None], "invalid_format"
                )
            return [converted[0], converted[1], None, max(converted)], "resolved_unit_conflict"
        if issue or converted is None:
            status = issue or "invalid_format"
            return _zero_if_authorized(resolution_policies, status) or (
                [None, None, None, None], status
            )
        if len(converted) == 3:
            return [converted[0], converted[1], converted[2], max(converted)], "parsed_3d"
        if require_context_for_partial and not _dimension_partial_has_context(prefix, suffix):
            return _zero_if_authorized(resolution_policies, "invalid_format") or (
                [None, None, None, None], "invalid_format"
            )
        return [converted[0], converted[1], None, max(converted)], "parsed_2d"
    if len(diameters) == 1:
        match = diameters[0]
        converted, issue = _convert_dimension(
            [float(match.group("d"))], [match.group("u")], default_unit
        )
        if issue or converted is None:
            return [None, None, None, None], issue or "invalid_format"
        return [None, None, None, converted[0]], "diameter_only"
    if not re.search(r"\d", text):
        return _zero_if_authorized(resolution_policies, "no_numeric_measurement") or (
            [None, None, None, None], "no_numeric_measurement"
        )
    return _zero_if_authorized(resolution_policies, "invalid_format") or (
        [None, None, None, None], "invalid_format"
    )


def _derive_metric_value(
    values: list[Any], template: DerivedFormulaTemplate | None
) -> tuple[float | None, str]:
    if all(value in (None, "") for value in values):
        return None, "input_missing"
    if template is None:
        return None, "invalid_input"
    try:
        if template == DerivedFormulaTemplate.RECTANGULAR_OR_DIAMETER_CUBED:
            length, width, height, diameter = values
            dimensions = (length, width, height)
            if all(value not in (None, "") for value in dimensions):
                numbers = [float(_strict_number(value)) for value in dimensions]
                if any(value <= 0 for value in numbers):
                    raise ValueError
                return math.prod(numbers) / 1000, "derived"
            if all(value in (None, "") for value in dimensions) and diameter not in (None, ""):
                number = float(_strict_number(diameter))
                if number <= 0:
                    raise ValueError
                return number ** 3 / 1000, "derived"
            return None, "invalid_input"
        numbers = [float(_strict_number(value)) for value in values]
        if any(value <= 0 for value in numbers):
            raise ValueError
        product = math.prod(numbers)
        if template == DerivedFormulaTemplate.ELLIPSOID_PI_OVER_6:
            return math.pi / 6 * product / 1000, "derived"
        if template == DerivedFormulaTemplate.HALF_PRODUCT:
            return product / 2 / 1000, "derived"
        return product / 1000, "derived"
    except (decimal.InvalidOperation, TypeError, ValueError):
        return None, "invalid_input"


class TransformTool(SpreadsheetTool[TransformTaskSpec]):
    name = "transform_column"
    description = "执行可审计的数值填补、显式映射、尺寸解析或受控派生指标。"
    spec_model = TransformTaskSpec
    risk_level = RiskLevel.MEDIUM

    @staticmethod
    def _parse_dimension_row(raw: Any, row: int, rule: ParseDimensionsRule):
        if row in rule.row_corrections:
            parsed, status = _parse_dimensions(
                rule.row_corrections[row], rule.default_unit,
                rule.resolution_policies, None,
            )
            if status in {
                "parsed_3d", "parsed_2d", "diameter_only",
                "resolved_multiple_first", "resolved_multiple_last",
                "resolved_unit_conflict", "resolved_zero",
            }:
                return parsed, "corrected_value"
            return parsed, status
        return _parse_dimensions(
            raw,
            rule.default_unit,
            rule.resolution_policies,
            rule.value_corrections,
            rule.require_context_for_partial,
        )

    def _source_refs(self, spec: TransformTaskSpec) -> list[ColumnRef]:
        if spec.operation == TransformOperation.IMPUTE_NUMERIC:
            return [rule.source for rule in spec.impute_rules]
        if spec.operation == TransformOperation.MAP_VALUES:
            return [rule.source for rule in spec.map_rules]
        if spec.operation == TransformOperation.PARSE_DIMENSIONS and spec.dimension_rule:
            return [spec.dimension_rule.source]
        if spec.operation == TransformOperation.DERIVE_METRIC and spec.derive_rule:
            return list(spec.derive_rule.sources)
        return []

    def validate(self, context: ToolContext, spec: TransformTaskSpec) -> TaskValidation:
        errors: list[str] = []
        warnings: list[str] = []
        sheet = _sheet_profile(context.profile, spec.sheet)
        if not sheet:
            return TaskValidation(valid=False, errors=[f"工作表“{spec.sheet}”不存在。"])
        for ref in self._source_refs(spec):
            if ref.sheet != spec.sheet:
                errors.append("Transform 的所有来源列必须位于任务指定的同一工作表。")
                continue
            _, error = _resolve_column_ref(context.profile, ref)
            if error:
                errors.append(error)
        if spec.operation == TransformOperation.MAP_VALUES:
            for rule in spec.map_rules:
                if not rule.mapping:
                    errors.append("值映射规则不能为空。")
                normalized: dict[str, Any] = {}
                for key, output in rule.mapping.items():
                    canonical = _normalized_text(key, casefold=rule.casefold)
                    if canonical in normalized and normalized[canonical] != output:
                        errors.append(f"映射键“{key}”标准化后与其他键冲突。")
                    normalized[canonical] = output
        if spec.operation == TransformOperation.DERIVE_METRIC and spec.derive_rule:
            if spec.derive_rule.formula_template is None:
                errors.append("派生指标必须明确选择注册公式模板。")
        layout = transform_write_columns(context.profile, spec)
        existing_names = Counter(column.name for column in sheet.columns)
        seen_outputs: set[str] = set()
        for name, _, overwrite in layout:
            if name in seen_outputs:
                errors.append(f"输出列“{name}”重复。")
            seen_outputs.add(name)
            if not overwrite and existing_names[name]:
                errors.append(f"输出列“{name}”已存在；请选择覆盖或更换列名。")
        if spec.operation == TransformOperation.IMPUTE_NUMERIC:
            layout_cursor = 0
            for rule in spec.impute_rules:
                rule_layout = layout[layout_cursor:layout_cursor + 2]
                layout_cursor += 2
                if rule.overwrite_existing_outputs and any(
                    not overwrite for _, _, overwrite in rule_layout
                ):
                    errors.append("覆盖既有填补输出时，结果列和指示列都必须已存在。")
        if spec.operation == TransformOperation.PARSE_DIMENSIONS and spec.dimension_rule:
            if spec.dimension_rule.overwrite_existing_outputs and any(not overwrite for _, _, overwrite in layout):
                errors.append("覆盖尺寸输出时，所有计划输出列都必须已存在。")
        if spec.operation == TransformOperation.DERIVE_METRIC and spec.derive_rule:
            if spec.derive_rule.overwrite_existing_outputs and any(not overwrite for _, _, overwrite in layout):
                errors.append("覆盖派生输出时，结果列和状态列都必须已存在。")
        return TaskValidation(valid=not errors, errors=list(dict.fromkeys(errors)), warnings=warnings)

    def preview(self, context: ToolContext, spec: TransformTaskSpec) -> dict[str, Any]:
        sheet = _sheet_profile(context.profile, spec.sheet)
        layout = transform_write_columns(context.profile, spec)
        counts: Counter[str] = Counter()
        review_rows: list[dict[str, Any]] = []
        review_totals: Counter[str] = Counter()

        def add_review(
            row: int,
            source_columns: list[str],
            raw_values: dict[str, Any],
            issue_code: str,
            description: str,
        ) -> None:
            review_totals[issue_code] += 1
            if sum(1 for item in review_rows if item["issue_code"] == issue_code) >= 50:
                return
            review_rows.append({
                "sheet": spec.sheet,
                "row": row,
                "source_columns": source_columns,
                "raw_values": raw_values,
                "issue_code": issue_code,
                "description": description,
            })

        if sheet is None:
            return {
                "modified_sheets": [],
                "created_sheets": [],
                "new_columns": [name for name, _, overwrite in layout if not overwrite],
                "estimated_rows": 0,
                "destructive": any(overwrite for _, _, overwrite in layout),
                "summary_counts": {"preview_unavailable": 0},
                "review_rows": [],
                "review_totals": {},
                "review_truncated": {},
            }

        workbook = load_workbook(context.input_path, read_only=False, data_only=False)
        try:
            worksheet = workbook[spec.sheet]
            start = (sheet.header_row or 1) + 1
            end = sheet.max_row
            if spec.operation == TransformOperation.IMPUTE_NUMERIC:
                for rule in spec.impute_rules:
                    source, _ = _resolve_column_ref(context.profile, rule.source)
                    if source is None:
                        counts["preview_unavailable"] += max(0, end - start + 1)
                        continue
                    fill_key = (
                        "planned_fill_value"
                        if len(spec.impute_rules) == 1
                        else f"planned_fill_value_{source.letter or get_column_letter(source.index)}"
                    )
                    observed: list[float | int] = []
                    for row in range(start, end + 1):
                        value = worksheet.cell(row, source.index).value
                        if _is_missing(value, rule.missing_tokens):
                            counts["missing"] += 1
                        else:
                            try:
                                observed.append(_strict_number(value))
                                counts["valid"] += 1
                            except (decimal.InvalidOperation, TypeError, ValueError):
                                label = _source_label(source)
                                add_review(
                                    row,
                                    [label],
                                    {label: value},
                                    "invalid_numeric",
                                    "非空内容无法解析为数值，未自动作为缺失值处理。",
                                )
                                if rule.invalid_policy == InvalidNumericPolicy.IMPUTE:
                                    counts["invalid_imputed"] += 1
                                else:
                                    counts["invalid"] += 1
                    counts["observed_basis"] += len(observed)
                    if rule.method == ImputationMethod.CONSTANT and rule.constant_value is not None:
                        counts[fill_key] = rule.constant_value
                    elif observed and rule.method == ImputationMethod.MEAN:
                        counts[fill_key] = statistics.fmean(observed)
                    elif observed and rule.method == ImputationMethod.MEDIAN:
                        counts[fill_key] = statistics.median(observed)
            elif spec.operation == TransformOperation.MAP_VALUES:
                for rule in spec.map_rules:
                    source, _ = _resolve_column_ref(context.profile, rule.source)
                    if source is None:
                        counts["preview_unavailable"] += max(0, end - start + 1)
                        continue
                    mapping = {_normalized_text(key, casefold=rule.casefold) for key in rule.mapping}
                    missing = {_normalized_text(item, casefold=rule.casefold) for item in rule.missing_tokens}
                    for row in range(start, end + 1):
                        raw = worksheet.cell(row, source.index).value
                        value = "" if raw is None else _normalized_text(raw, casefold=rule.casefold)
                        if value == "" or value in missing:
                            counts["mapped_blank" if rule.blank_value is not None else "blank"] += 1
                        elif value in mapping:
                            counts["mapped"] += 1
                        else:
                            counts["mapped_default" if rule.unmapped_value is not None else "unknown"] += 1
                eligible = counts["mapped"] + counts["mapped_default"] + counts["unknown"]
                counts["mapping_coverage_pct"] = round(
                    (counts["mapped"] + counts["mapped_default"]) * 100 / eligible, 2
                ) if eligible else 100.0
            elif spec.operation == TransformOperation.PARSE_DIMENSIONS and spec.dimension_rule:
                source, _ = _resolve_column_ref(context.profile, spec.dimension_rule.source)
                if source is None:
                    counts["preview_unavailable"] += max(0, end - start + 1)
                    return {
                        "modified_sheets": [spec.sheet],
                        "created_sheets": [],
                        "new_columns": [name for name, _, overwrite in layout if not overwrite],
                        "estimated_rows": max(0, sheet.max_row - (sheet.header_row or 1)),
                        "destructive": any(overwrite for _, _, overwrite in layout),
                        "summary_counts": dict(counts),
                        "review_rows": [],
                        "review_totals": {},
                        "review_truncated": {},
                    }
                for row in range(start, end + 1):
                    raw = worksheet.cell(row, source.index).value
                    _, status = self._parse_dimension_row(raw, row, spec.dimension_rule)
                    counts[status] += 1
                    if status in {
                        "multiple_candidates", "unit_conflict", "invalid_format",
                        "no_numeric_measurement",
                    }:
                        label = _source_label(source)
                        add_review(
                            row,
                            [label],
                            {label: raw},
                            status,
                            "尺寸无法在不猜测的情况下可靠解析。",
                        )
            elif spec.operation == TransformOperation.DERIVE_METRIC and spec.derive_rule:
                sources = [_resolve_column_ref(context.profile, ref)[0] for ref in spec.derive_rule.sources]
                if any(column is None for column in sources):
                    counts["preview_unavailable"] = max(0, end - start + 1)
                elif any(column.index > worksheet.max_column for column in sources):
                    # These sources are outputs of an earlier composite step. Their
                    # values do not exist in the base workbook yet and are verified
                    # after that dependency executes.
                    counts["deferred_to_dependency"] = max(0, end - start + 1)
                else:
                    for row in range(start, end + 1):
                        values = [worksheet.cell(row, column.index).value for column in sources]
                        _, status = _derive_metric_value(
                            values, spec.derive_rule.formula_template
                        )
                        if spec.derive_rule.resolution_policies.get(status) == DerivedResolutionAction.ZERO:
                            status = "resolved_missing_zero" if status == "input_missing" else "resolved_invalid_zero"
                        counts[status] += 1
                        if status in {"input_missing", "invalid_input"}:
                            labels = [_source_label(column) for column in sources]
                            add_review(
                                row, labels, dict(zip(labels, values)), status,
                                "派生指标输入不足，无法按已选公式可靠计算。",
                            )
        finally:
            workbook.close()
        return {
            "modified_sheets": [spec.sheet],
            "created_sheets": [],
            "new_columns": [name for name, _, overwrite in layout if not overwrite],
            "estimated_rows": max(0, sheet.max_row - (sheet.header_row or 1)) if sheet else 0,
            "destructive": any(overwrite for _, _, overwrite in layout),
            "summary_counts": dict(counts),
            "review_rows": sorted(review_rows, key=lambda item: (item["issue_code"], item["row"])),
            "review_totals": dict(review_totals),
            "review_truncated": {
                code: total > 50 for code, total in review_totals.items()
            },
        }

    @staticmethod
    def _write_exception_report(path: Path, exceptions: list[dict[str, Any]]) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=["工作表", "行号", "源列", "原始值", "问题代码", "说明"],
            )
            writer.writeheader()
            for item in exceptions:
                writer.writerow({
                    "工作表": item["sheet"],
                    "行号": item["row"],
                    "源列": item["source_column"],
                    "原始值": item["raw_value"],
                    "问题代码": item["issue_code"],
                    "说明": item["message"],
                })

    def execute(self, context: ToolContext, spec: TransformTaskSpec) -> ToolResult:
        _copy_input(context)
        workbook = load_workbook(context.output_path, data_only=False, keep_links=True)
        sheet = workbook[spec.sheet]
        profile = _sheet_profile(context.profile, spec.sheet)
        layout = transform_write_columns(context.profile, spec)
        for name, index, _ in layout:
            sheet.cell(profile.header_row, index, name)
        output_indexes = [index for _, index, _ in layout]
        exceptions: list[dict[str, Any]] = []
        counters: Counter[str] = Counter()
        lineage: list[dict[str, Any]] = []

        def lineage_source(column: Any) -> dict[str, Any]:
            return {
                "column_index": column.index,
                "column_letter": column.letter or get_column_letter(column.index),
                "display_header": column.display_header or column.name,
                "header_occurrence": column.header_occurrence,
            }

        def add_exception(row: int, source: str, raw: Any, code: str, message: str) -> None:
            exceptions.append({
                "sheet": spec.sheet,
                "row": row,
                "source_column": source,
                "raw_value": "" if raw is None else str(raw),
                "issue_code": code,
                "message": message,
            })

        cursor = 0
        if spec.operation == TransformOperation.IMPUTE_NUMERIC:
            for rule in spec.impute_rules:
                source, _ = _resolve_column_ref(context.profile, rule.source)
                target_index, flag_index = output_indexes[cursor:cursor + 2]
                cursor += 2
                observed: list[float | int] = []
                values: dict[int, tuple[str, Any]] = {}
                for row in range(profile.header_row + 1, profile.max_row + 1):
                    raw = sheet.cell(row, source.index).value
                    if _is_missing(raw, rule.missing_tokens):
                        values[row] = ("missing", raw)
                    else:
                        try:
                            number = _strict_number(raw)
                            observed.append(number)
                            values[row] = ("valid", raw)
                        except (decimal.InvalidOperation, TypeError, ValueError):
                            values[row] = ("invalid", raw)
                if not observed and rule.method != ImputationMethod.CONSTANT:
                    workbook.close()
                    raise ValueError(f"列“{_source_label(source)}”没有可用于填补的有效数值。")
                if rule.method == ImputationMethod.MEAN:
                    fill_value: float | int = statistics.fmean(observed)
                elif rule.method == ImputationMethod.MEDIAN:
                    fill_value = statistics.median(observed)
                else:
                    assert rule.constant_value is not None
                    fill_value = rule.constant_value
                for row, (status, value) in values.items():
                    if status == "valid":
                        sheet.cell(row, target_index, value)
                        sheet.cell(row, flag_index, 0)
                        counters["valid"] += 1
                    elif status == "missing":
                        sheet.cell(row, target_index, fill_value)
                        sheet.cell(row, flag_index, 1)
                        counters["imputed"] += 1
                    else:
                        if rule.invalid_policy == InvalidNumericPolicy.IMPUTE:
                            sheet.cell(row, target_index, fill_value)
                            sheet.cell(row, flag_index, 1)
                            counters["invalid_imputed"] += 1
                        else:
                            if not rule.overwrite:
                                sheet.cell(row, target_index, None)
                            sheet.cell(row, flag_index, None)
                            counters["invalid"] += 1
                            add_exception(row, _source_label(source), value, "invalid_numeric", "非空值无法按严格数值规则解析，未填补。")
                counters["observed_basis"] += len(observed)
                lineage.append({
                    "output_column": sheet.cell(profile.header_row, target_index).value,
                    "output_column_index": target_index,
                    "source_columns": [lineage_source(source)],
                    "operation": spec.operation.value,
                    "parameters": {
                        "method": rule.method.value,
                        "fill_value": fill_value,
                        "invalid_policy": rule.invalid_policy.value,
                    },
                })

        elif spec.operation == TransformOperation.MAP_VALUES:
            for rule in spec.map_rules:
                source, _ = _resolve_column_ref(context.profile, rule.source)
                target_index = output_indexes[cursor]
                cursor += 1
                status_index = None
                if rule.status_column:
                    status_index = output_indexes[cursor]
                    cursor += 1
                mapping = {
                    _normalized_text(key, casefold=rule.casefold): output
                    for key, output in rule.mapping.items()
                }
                missing = {_normalized_text(item, casefold=rule.casefold) for item in rule.missing_tokens}
                for row in range(profile.header_row + 1, profile.max_row + 1):
                    raw = sheet.cell(row, source.index).value
                    canonical = "" if raw is None else _normalized_text(raw, casefold=rule.casefold)
                    if canonical == "" or canonical in missing:
                        if rule.blank_value is not None:
                            sheet.cell(row, target_index, rule.blank_value)
                            if status_index is not None:
                                sheet.cell(row, status_index, "mapped_blank")
                            counters["mapped_blank"] += 1
                        else:
                            sheet.cell(row, target_index, None)
                            if status_index is not None:
                                sheet.cell(row, status_index, "blank")
                            counters["blank"] += 1
                    elif canonical in mapping:
                        sheet.cell(row, target_index, mapping[canonical])
                        if status_index is not None:
                            sheet.cell(row, status_index, "mapped")
                        counters["mapped"] += 1
                    elif rule.unmapped_value is not None:
                        sheet.cell(row, target_index, rule.unmapped_value)
                        if status_index is not None:
                            sheet.cell(row, status_index, "mapped_default")
                        counters["mapped_default"] += 1
                    else:
                        sheet.cell(row, target_index, None)
                        if status_index is not None:
                            sheet.cell(row, status_index, "unknown")
                        counters["unknown"] += 1
                        add_exception(row, _source_label(source), raw, "unknown_mapping", "值不在用户确认的映射表中，未映射。")
                lineage.append({
                    "output_column": sheet.cell(profile.header_row, target_index).value,
                    "output_column_index": target_index,
                    "source_columns": [lineage_source(source)],
                    "operation": spec.operation.value,
                    "parameters": {
                        "mapping": rule.mapping,
                        "casefold": rule.casefold,
                        "blank_value": rule.blank_value,
                        "unmapped_value": rule.unmapped_value,
                        "status_column_emitted": status_index is not None,
                    },
                })

        elif spec.operation == TransformOperation.PARSE_DIMENSIONS and spec.dimension_rule:
            rule = spec.dimension_rule
            source, _ = _resolve_column_ref(context.profile, rule.source)
            length_index, width_index, height_index, max_index, status_index = output_indexes
            for row in range(profile.header_row + 1, profile.max_row + 1):
                raw = sheet.cell(row, source.index).value
                parsed, status = self._parse_dimension_row(raw, row, rule)
                for index, value in zip((length_index, width_index, height_index, max_index), parsed):
                    sheet.cell(row, index, value)
                sheet.cell(row, status_index, status)
                counters[status] += 1
                if status not in {
                    "parsed_3d", "parsed_2d", "diameter_only", "blank",
                    "resolved_multiple_first", "resolved_multiple_last",
                    "resolved_unit_conflict", "resolved_zero", "corrected_value",
                }:
                    add_exception(row, _source_label(source), raw, status, "尺寸无法在不猜测的情况下可靠解析。")
            for index in (length_index, width_index, height_index, max_index):
                lineage.append({
                    "output_column": sheet.cell(profile.header_row, index).value,
                    "output_column_index": index,
                    "source_columns": [lineage_source(source)],
                    "operation": spec.operation.value,
                    "parameters": {
                        "default_unit": rule.default_unit,
                        "require_context_for_partial": rule.require_context_for_partial,
                        "resolution_policies": {
                            key: value.value for key, value in rule.resolution_policies.items()
                        },
                        "value_corrections": rule.value_corrections,
                        "row_corrections": rule.row_corrections,
                    },
                })

        elif spec.operation == TransformOperation.DERIVE_METRIC and spec.derive_rule:
            rule = spec.derive_rule
            sources = [_resolve_column_ref(context.profile, ref)[0] for ref in rule.sources]
            target_index, status_index = output_indexes
            assert rule.formula_template is not None
            for row in range(profile.header_row + 1, profile.max_row + 1):
                raw_values = [sheet.cell(row, column.index).value for column in sources]
                result_value, status = _derive_metric_value(raw_values, rule.formula_template)
                if rule.resolution_policies.get(status) == DerivedResolutionAction.ZERO:
                    status = "resolved_missing_zero" if status == "input_missing" else "resolved_invalid_zero"
                    result_value = 0.0
                sheet.cell(row, target_index, result_value)
                sheet.cell(row, status_index, status)
                counters[status] += 1
                if status not in {"derived", "resolved_invalid_zero", "resolved_missing_zero"}:
                    add_exception(row, ",".join(_source_label(column) for column in sources), " | ".join("" if value is None else str(value) for value in raw_values), status, "派生指标缺少三个有效正数输入。")
            lineage.append({
                "output_column": sheet.cell(profile.header_row, target_index).value,
                "output_column_index": target_index,
                "source_columns": [lineage_source(column) for column in sources],
                "operation": spec.operation.value,
                "parameters": {
                    "formula_template": rule.formula_template.value,
                    "output_unit": rule.output_unit,
                    "resolution_policies": {
                        key: value.value for key, value in rule.resolution_policies.items()
                    },
                },
            })

        artifact_name = _transform_artifact_name(spec)
        artifact_path = (context.artifact_dir or context.output_path.parent) / artifact_name
        self._write_exception_report(artifact_path, exceptions)
        workbook.save(context.output_path)
        workbook.close()
        ranges = [
            f"{spec.sheet}!{get_column_letter(index)}{profile.header_row}:{get_column_letter(index)}{profile.max_row}"
            for index in sorted(set(output_indexes))
        ]
        warnings = [f"{len(exceptions)} 条记录需要后续处理。"] if exceptions else []
        counters["exceptions"] = len(exceptions)
        return ToolResult(
            tool_name=self.name,
            modified_sheets=[spec.sheet],
            modified_ranges=ranges,
            rows_affected=max(0, profile.max_row - profile.header_row),
            counters=dict(counters),
            warnings=warnings,
            details={
                "operation": spec.operation.value,
                "artifact_path": str(artifact_path),
                "artifact_name": artifact_name,
                "output_columns": [name for name, _, _ in layout],
                "lineage": lineage,
                "follow_up_available": bool(exceptions),
            },
        )

    def verify(self, context: ToolContext, spec: TransformTaskSpec, result: ToolResult) -> TaskValidation:
        errors: list[str] = []
        workbook = load_workbook(context.output_path, read_only=False, data_only=False)
        source_workbook = load_workbook(context.input_path, read_only=False, data_only=False)
        try:
            sheet = workbook[spec.sheet]
            source_sheet = source_workbook[spec.sheet]
            profile = _sheet_profile(context.profile, spec.sheet)
            layout = transform_write_columns(context.profile, spec)
            for name, index, _ in layout:
                if sheet.cell(profile.header_row, index).value != name:
                    errors.append(f"未找到计划中的输出列“{name}”。")
            records = max(0, profile.max_row - profile.header_row)
            if spec.operation == TransformOperation.IMPUTE_NUMERIC:
                total = int(
                    result.counters.get("valid", 0)
                    + result.counters.get("imputed", 0)
                    + result.counters.get("invalid", 0)
                    + result.counters.get("invalid_imputed", 0)
                )
                if total != records * len(spec.impute_rules):
                    errors.append("填补结果计数无法与输入行数核对。")
                cursor = 0
                for rule in spec.impute_rules:
                    source, _ = _resolve_column_ref(context.profile, rule.source)
                    target_index, flag_index = layout[cursor][1], layout[cursor + 1][1]
                    cursor += 2
                    observed: list[float | int] = []
                    for row in range(profile.header_row + 1, profile.max_row + 1):
                        raw = source_sheet.cell(row, source.index).value
                        if not _is_missing(raw, rule.missing_tokens):
                            try:
                                observed.append(_strict_number(raw))
                            except (decimal.InvalidOperation, TypeError, ValueError):
                                pass
                    expected_fill = (
                        statistics.fmean(observed)
                        if rule.method == ImputationMethod.MEAN
                        else statistics.median(observed)
                        if rule.method == ImputationMethod.MEDIAN
                        else rule.constant_value
                    )
                    for row in range(profile.header_row + 1, profile.max_row + 1):
                        raw = source_sheet.cell(row, source.index).value
                        target = sheet.cell(row, target_index).value
                        flag = sheet.cell(row, flag_index).value
                        if _is_missing(raw, rule.missing_tokens):
                            fill_matches = (
                                isinstance(target, (int, float))
                                and expected_fill is not None
                                and math.isclose(
                                    float(target), float(expected_fill),
                                    rel_tol=1e-12, abs_tol=1e-12,
                                )
                            )
                            if not fill_matches or flag != 1:
                                errors.append(f"第 {row} 行缺失值未按重算统计量填补。")
                                break
                        else:
                            invalid_source = False
                            try:
                                _strict_number(raw)
                                expected = raw
                                expected_flag = 0
                            except (decimal.InvalidOperation, TypeError, ValueError):
                                invalid_source = True
                                if rule.invalid_policy == InvalidNumericPolicy.IMPUTE:
                                    expected = expected_fill
                                    expected_flag = 1
                                else:
                                    expected = raw if rule.overwrite else None
                                    expected_flag = None
                            if (
                                invalid_source
                                and rule.invalid_policy == InvalidNumericPolicy.IMPUTE
                            ):
                                value_matches = (
                                    isinstance(target, (int, float))
                                    and expected is not None
                                    and math.isclose(
                                        float(target), float(expected),
                                        rel_tol=1e-12, abs_tol=1e-12,
                                    )
                                )
                            else:
                                value_matches = target == expected
                            if not value_matches or flag != expected_flag:
                                errors.append(f"第 {row} 行非缺失原值被意外改变。")
                                break
                        if not rule.overwrite and sheet.cell(row, source.index).value != raw:
                            errors.append(f"第 {row} 行原列被意外改变。")
                            break
            elif spec.operation == TransformOperation.MAP_VALUES:
                total = int(
                    result.counters.get("mapped", 0)
                    + result.counters.get("mapped_default", 0)
                    + result.counters.get("mapped_blank", 0)
                    + result.counters.get("unknown", 0)
                    + result.counters.get("blank", 0)
                )
                if total != records * len(spec.map_rules):
                    errors.append("映射结果计数无法与输入行数核对。")
            elif spec.operation == TransformOperation.PARSE_DIMENSIONS:
                statuses = {
                    "parsed_3d", "parsed_2d", "diameter_only",
                    "multiple_candidates", "unit_conflict", "invalid_format",
                    "no_numeric_measurement", "blank",
                    "resolved_multiple_first", "resolved_multiple_last",
                    "resolved_unit_conflict", "resolved_zero", "corrected_value",
                }
                if sum(int(result.counters.get(status, 0)) for status in statuses) != records:
                    errors.append("尺寸解析状态未覆盖全部输入行。")
                _, indexes = None, [index for _, index, _ in layout]
                for row in range(profile.header_row + 1, profile.max_row + 1):
                    length, width, height, maximum = [sheet.cell(row, index).value for index in indexes[:4]]
                    present = [value for value in (length, width, height) if value not in (None, "")]
                    status = sheet.cell(row, indexes[4]).value
                    if any(
                        not isinstance(value, (int, float))
                        or (value < 0 if status == "resolved_zero" else value <= 0)
                        for value in present
                    ):
                        errors.append(f"第 {row} 行解析尺寸不是正数。")
                        break
                    if present and maximum != max(present):
                        errors.append(f"第 {row} 行最大径与解析维度不一致。")
                        break
            elif spec.operation == TransformOperation.DERIVE_METRIC and spec.derive_rule:
                target_index = layout[0][1]
                status_index = layout[1][1]
                sources = [_resolve_column_ref(context.profile, ref)[0] for ref in spec.derive_rule.sources]
                for row in range(profile.header_row + 1, profile.max_row + 1):
                    actual = sheet.cell(row, target_index).value
                    status = sheet.cell(row, status_index).value
                    if status in {"resolved_invalid_zero", "resolved_missing_zero"}:
                        if actual != 0:
                            errors.append(f"第 {row} 行经授权置零的派生结果不是 0。")
                            break
                        continue
                    values = [sheet.cell(row, column.index).value for column in sources]
                    template = spec.derive_rule.formula_template
                    expected, expected_status = _derive_metric_value(values, template)
                    if actual in (None, ""):
                        if expected_status == "derived":
                            errors.append(f"第 {row} 行输入完整但派生结果为空。")
                            break
                        continue
                    if expected is None:
                        errors.append(f"第 {row} 行无效输入产生了派生结果。")
                        break
                    if not math.isclose(float(actual), expected, rel_tol=1e-12, abs_tol=1e-12):
                        errors.append(f"第 {row} 行派生结果无法独立复算。")
                        break
        finally:
            workbook.close()
            source_workbook.close()
        artifact = Path(result.details.get("artifact_path", ""))
        if not artifact.is_file():
            errors.append("异常记录附件不存在。")
        return TaskValidation(valid=not errors, errors=errors, warnings=result.warnings)


class ConversionTool(SpreadsheetTool[ConversionTaskSpec]):
    name = "conversion"
    description = "Convert staged CSV to XLSX or export one XLSX sheet to CSV."
    spec_model = ConversionTaskSpec
    risk_level = RiskLevel.LOW

    def validate(self, context: ToolContext, spec: ConversionTaskSpec) -> TaskValidation:
        errors: list[str] = []
        warnings: list[str] = []
        if (
            spec.direction == ConversionDirection.CSV_TO_XLSX
            and "CSV_DELIMITER_UNCERTAIN" in context.profile.risks
            and spec.delimiter == "auto"
        ):
            errors.append("CSV 分隔符无法可靠识别，必须明确指定分隔符。")
        if spec.direction == ConversionDirection.XLSX_TO_CSV:
            if not spec.source_sheet:
                errors.append("XLSX 导出 CSV 必须指定来源工作表。")
            elif not _sheet_profile(context.profile, spec.source_sheet):
                errors.append(f"工作表“{spec.source_sheet}”不存在。")
        if spec.encoding == "auto" and spec.direction == ConversionDirection.XLSX_TO_CSV:
            warnings.append("CSV 输出将使用 UTF-8 BOM。")
        return TaskValidation(valid=not errors, errors=errors, warnings=warnings)

    def preview(self, context: ToolContext, spec: ConversionTaskSpec) -> dict[str, Any]:
        rows = 0
        if spec.source_sheet:
            sheet = _sheet_profile(context.profile, spec.source_sheet)
            rows = max(0, sheet.max_row - (sheet.header_row or 1)) if sheet else 0
        return {
            "modified_sheets": [],
            "created_sheets": ["CSV_Data"] if spec.direction == ConversionDirection.CSV_TO_XLSX else [],
            "new_columns": [],
            "estimated_rows": rows,
            "destructive": False,
        }

    def execute(self, context: ToolContext, spec: ConversionTaskSpec) -> ToolResult:
        details: dict[str, Any] = {
            "direction": spec.direction.value,
            "preserve_text": spec.preserve_text,
        }
        created_sheets: list[str] = []
        rows = 0
        if spec.direction == ConversionDirection.CSV_TO_XLSX:
            if context.original_path is None or context.original_path.suffix.casefold() != ".csv":
                raise ValueError("CSV 原始文件不可用。")
            detection = csv_to_xlsx(
                context.original_path,
                context.output_path,
                encoding=spec.encoding,
                delimiter=spec.delimiter,
                infer_types=spec.infer_types and not spec.preserve_text,
            )
            created_sheets = [context.profile.sheets[0].name] if context.profile.sheets else []
            rows = max(0, context.profile.sheets[0].max_row - 1) if context.profile.sheets else 0
            details.update({
                "output_format": "xlsx",
                "encoding": detection.encoding,
                "delimiter": detection.delimiter,
            })
        else:
            _copy_input(context)
            audit = _audit_cached_values(context.output_path)
            if not audit.success:
                remaining = max(0, context.execution_deadline - time.monotonic()) if context.execution_deadline else 90
                audit = recalculate_workbook(context.output_path, timeout=min(90, remaining), cancelled=context.cancelled)
                details["calculation"] = audit.model_dump()
                if not audit.success:
                    return ToolResult(tool_name=self.name, details=details, warnings=["公式计算未完成，未导出 CSV。"])
            artifact_dir = context.artifact_dir or context.output_path.parent
            artifact_dir.mkdir(parents=True, exist_ok=True)
            output_name = Path(spec.output_name or f"{spec.source_sheet}.csv").name
            if not output_name.casefold().endswith(".csv"):
                output_name += ".csv"
            artifact = artifact_dir / output_name
            encoding = "utf-8-sig" if spec.encoding == "auto" else spec.encoding
            delimiter = "," if spec.delimiter == "auto" else spec.delimiter
            rows = xlsx_sheet_to_csv(
                context.output_path,
                artifact,
                sheet_name=spec.source_sheet,
                encoding=encoding,
                delimiter=delimiter,
            )
            details.update({
                "output_format": "csv",
                "artifact_path": str(artifact),
                "artifact_sha256": sha256_file(artifact),
                "encoding": encoding,
                "delimiter": delimiter,
            })
        return ToolResult(
            tool_name=self.name,
            created_sheets=created_sheets,
            rows_affected=rows,
            counters={"converted_rows": rows},
            details=details,
        )

    def verify(
        self, context: ToolContext, spec: ConversionTaskSpec, result: ToolResult
    ) -> TaskValidation:
        errors: list[str] = []
        if spec.direction == ConversionDirection.CSV_TO_XLSX:
            try:
                workbook = load_workbook(context.output_path, read_only=True)
                workbook.close()
            except Exception as exc:
                errors.append(f"转换后的 XLSX 无法打开：{exc}")
        else:
            calculation = result.details.get("calculation")
            if calculation and not calculation.get("success"):
                return TaskValidation(valid=False, errors=["公式计算未完成，未导出 CSV：" + calculation.get("message", "")])
            artifact = Path(result.details.get("artifact_path", ""))
            if not artifact.is_file():
                errors.append("CSV 导出文件不存在。")
            elif sha256_file(artifact) != result.details.get("artifact_sha256"):
                errors.append("CSV 导出文件完整性校验失败。")
            else:
                workbook = load_workbook(context.output_path, read_only=True, data_only=True)
                try:
                    with artifact.open("r", encoding=result.details["encoding"], newline="") as handle:
                        actual_rows = csv.reader(handle, delimiter=result.details["delimiter"])
                        rows = workbook[spec.source_sheet].iter_rows(values_only=True)
                        for expected, actual in zip_longest(rows, actual_rows):
                            if expected is None or actual is None or ["" if value is None else str(value) for value in expected] != actual:
                                errors.append("CSV 内容与最终工作表不一致。")
                                break
                finally:
                    workbook.close()
        return TaskValidation(valid=not errors, errors=errors)


class ToolRegistry:
    def __init__(self, tools: list[SpreadsheetTool[Any]] | None = None):
        items = tools or [
            LookupTool(), NormalizeTool(), AggregateTool(), SortRankTool(), CleanTool(), FormulaTool(),
            ConversionTool(), TransformTool(),
        ]
        self._tools = {tool.name: tool for tool in items}

    def get(self, name: str) -> SpreadsheetTool[Any]:
        try:
            return self._tools[name]
        except KeyError as exc:
            raise KeyError(f"unknown spreadsheet tool: {name}") from exc

    def names(self) -> list[str]:
        return sorted(self._tools)

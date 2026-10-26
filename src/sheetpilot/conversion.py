from __future__ import annotations

import csv
import datetime as dt
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .workbook_utils import worksheet_dimensions


SUPPORTED_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030")
SUPPORTED_DELIMITERS = (",", "\t", ";", "|")


@dataclass(frozen=True)
class CsvDetection:
    encoding: str
    delimiter: str
    confident: bool


def detect_csv(path: Path) -> CsvDetection:
    raw = path.read_bytes()
    decoded = None
    encoding = ""
    for candidate in SUPPORTED_ENCODINGS:
        try:
            decoded = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise ValueError("CSV 编码无法识别；仅支持 UTF-8、UTF-8 BOM 和 GB18030。")
    sample = decoded[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters="".join(SUPPORTED_DELIMITERS))
        delimiter = dialect.delimiter
        confident = delimiter in SUPPORTED_DELIMITERS
    except csv.Error:
        delimiter = ","
        confident = False
    return CsvDetection(encoding=encoding, delimiter=delimiter, confident=confident)


def csv_to_xlsx(
    source: Path,
    target: Path,
    *,
    encoding: str = "auto",
    delimiter: str = "auto",
    sheet_name: str = "CSV_Data",
    infer_types: bool = False,
) -> CsvDetection:
    detection = detect_csv(source)
    selected_encoding = detection.encoding if encoding == "auto" else encoding
    selected_delimiter = detection.delimiter if delimiter == "auto" else delimiter
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31] or "CSV_Data"
    with source.open("r", encoding=selected_encoding, newline="") as handle:
        for row_index, row in enumerate(csv.reader(handle, delimiter=selected_delimiter)):
            if infer_types and row_index > 0:
                sheet.append([_infer_csv_value(value) for value in row])
            else:
                sheet.append(row)
    workbook.save(target)
    workbook.close()
    return CsvDetection(
        encoding=selected_encoding,
        delimiter=selected_delimiter,
        confident=detection.confident or delimiter != "auto",
    )


def _infer_csv_value(value: str):
    text = value.strip()
    if not text:
        return None
    try:
        if not (len(text) > 1 and text.startswith("0") and text[1].isdigit()):
            return int(text)
    except ValueError:
        pass
    try:
        return float(text)
    except ValueError:
        pass
    try:
        return dt.date.fromisoformat(text)
    except ValueError:
        return value


def xlsx_sheet_to_csv(
    source: Path,
    target: Path,
    *,
    sheet_name: str,
    encoding: str = "utf-8-sig",
    delimiter: str = ",",
) -> int:
    workbook = load_workbook(source, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise KeyError(f"工作表“{sheet_name}”不存在。")
    sheet = workbook[sheet_name]
    max_row, max_column = worksheet_dimensions(sheet)
    count = 0
    with target.open("w", encoding=encoding, newline="") as handle:
        writer = csv.writer(handle, delimiter=delimiter)
        for row in sheet.iter_rows(
            min_row=1, max_row=max_row, min_col=1, max_col=max_column, values_only=True
        ):
            writer.writerow(["" if value is None else value for value in row])
            count += 1
    workbook.close()
    return count

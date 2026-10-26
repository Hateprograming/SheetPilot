from __future__ import annotations

import hashlib
import random
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from .files import sha256_file
from .profiler import profile_workbook
from .schemas import (
    TaskPlan,
    VerificationCheck,
    VerificationReport,
    WorkbookChangeSet,
    WorkbookProfile,
)


def _split_range(address: str) -> tuple[str, tuple[int, int, int, int]]:
    sheet, cells = address.rsplit("!", 1)
    return sheet.strip("'"), range_boundaries(cells)


def _range_map(ranges: Iterable[str]) -> dict[str, list[tuple[int, int, int, int]]]:
    result: dict[str, list[tuple[int, int, int, int]]] = {}
    for address in ranges:
        try:
            sheet, bounds = _split_range(address)
        except (ValueError, TypeError):
            continue
        result.setdefault(sheet, []).append(bounds)
    return result


def _allowed(
    allowed: dict[str, list[tuple[int, int, int, int]]],
    sheet: str,
    row: int,
    column: int,
) -> bool:
    return any(
        min_col <= column <= max_col and min_row <= row <= max_row
        for min_col, min_row, max_col, max_row in allowed.get(sheet, [])
    )


def _hidden_rows(sheet) -> tuple[int, ...]:
    return tuple(sorted(index for index, item in sheet.row_dimensions.items() if item.hidden))


def _hidden_columns(sheet) -> tuple[str, ...]:
    return tuple(sorted(index for index, item in sheet.column_dimensions.items() if item.hidden))


def _column_layout(sheet) -> tuple[object, tuple[tuple[str, float | None, bool], ...]]:
    return (
        sheet.sheet_format.defaultColWidth,
        tuple(sorted(
            (name, dimension.width, bool(dimension.hidden))
            for name, dimension in sheet.column_dimensions.items()
        )),
    )


def _table_names(sheet) -> tuple[str, ...]:
    return tuple(sorted(sheet.tables.keys()))


def _defined_names(workbook) -> tuple[str, ...]:
    names = []
    for item in workbook.defined_names.values():
        names.append(repr((item.name, item.attr_text, item.localSheetId, item.hidden)))
    return tuple(sorted(names))


def _style_signature(cell) -> tuple[str, ...]:
    return (
        str(cell.number_format),
        repr(cell.font),
        repr(cell.fill),
        repr(cell.border),
        repr(cell.alignment),
        repr(cell.protection),
    )


def _sample_coordinates(sheet_name: str, rows: int, columns: int, limit: int = 200):
    total = max(0, rows) * max(0, columns)
    if not total:
        return []
    fixed = {0, max(0, total - 1)}
    for column in range(min(columns, limit)):
        fixed.add(column)
    remaining = max(0, limit - len(fixed))
    seed = int.from_bytes(hashlib.sha256(sheet_name.encode("utf-8")).digest()[:8], "big")
    rng = random.Random(seed)
    population = range(total)
    sampled = set(rng.sample(population, min(remaining, total))) if remaining else set()
    return [divmod(index, columns) for index in sorted(fixed | sampled)]


class WorkbookDiffVerifier:
    """Independently compare workbook contents with the approved write contract."""

    def compare(self, before_path: Path, after_path: Path, plan: TaskPlan) -> WorkbookChangeSet:
        approved_column_projection_sheets = {
            str(step.arguments.get("sheet", ""))
            for step in plan.steps
            if step.tool_name == "clean"
            and step.arguments.get("operation") == "keep_columns"
            and bool(step.arguments.get("overwrite"))
        }
        allowed_ranges = _range_map(
            address
            for step in plan.steps
            for address in step.expected_changes.allowed_ranges
        )
        modified_sheets = {
            name for step in plan.steps for name in step.expected_changes.modified_sheets
        }
        allowed_created = {
            name
            for step in plan.steps
            for name in (
                step.expected_changes.allowed_created_sheets
                or step.expected_changes.created_sheets
            )
        }
        changed: list[str] = []
        unexpected: list[str] = []
        structures: list[str] = []
        style_warnings: list[str] = []

        before_struct = load_workbook(before_path, data_only=False, keep_links=True)
        after_struct = load_workbook(after_path, data_only=False, keep_links=True)
        try:
            before_names = list(before_struct.sheetnames)
            after_names = list(after_struct.sheetnames)
            created = [name for name in after_names if name not in before_names]
            deleted = [name for name in before_names if name not in after_names]
            if deleted:
                unexpected.append("删除了工作表：" + "、".join(deleted))
            extra = [name for name in created if name not in allowed_created]
            if extra:
                unexpected.append("创建了未批准的工作表：" + "、".join(extra))
            existing_after_order = [name for name in after_names if name in before_names]
            if existing_after_order != before_names:
                unexpected.append("原有工作表的顺序发生变化。")
            if _defined_names(before_struct) != _defined_names(after_struct):
                unexpected.append("工作簿名称范围发生了计划外变化。")
            if len(before_struct._external_links) != len(after_struct._external_links):
                unexpected.append("工作簿外部链接发生了计划外变化。")

            for name in before_names:
                if name not in after_struct.sheetnames:
                    continue
                left = before_struct[name]
                right = after_struct[name]
                if name not in modified_sheets and _column_layout(left) != _column_layout(right):
                    unexpected.append(f"{name} 的列宽设置发生了计划外变化。")
                structural_pairs = {
                    "可见状态": (left.sheet_state, right.sheet_state),
                    "合并区域": (
                        tuple(sorted(str(item) for item in left.merged_cells.ranges)),
                        tuple(sorted(str(item) for item in right.merged_cells.ranges)),
                    ),
                    "隐藏行": (_hidden_rows(left), _hidden_rows(right)),
                    "隐藏列": (_hidden_columns(left), _hidden_columns(right)),
                    "冻结窗格": (str(left.freeze_panes), str(right.freeze_panes)),
                    "表": (_table_names(left), _table_names(right)),
                    "自动筛选": (str(left.auto_filter.ref), str(right.auto_filter.ref)),
                }
                for label, (before_value, after_value) in structural_pairs.items():
                    if before_value != after_value:
                        unexpected.append(f"{name} 的{label}发生了计划外变化。")
                        structures.append(f"{name}:{label}")

                max_row = max(left.max_row, right.max_row)
                max_column = max(left.max_column, right.max_column)
                permitted = allowed_ranges.get(name, [])
                permitted_max_row = max((item[3] for item in permitted), default=left.max_row)
                permitted_max_col = max((item[2] for item in permitted), default=left.max_column)
                if right.max_row > max(left.max_row, permitted_max_row):
                    unexpected.append(f"{name} 出现了未批准的新增行。")
                if right.max_column > max(left.max_column, permitted_max_col):
                    unexpected.append(f"{name} 出现了未批准的新增列。")
                if right.max_row < left.max_row:
                    unexpected.append(f"{name} 的数据区域被计划外缩小。")
                if (
                    right.max_column < left.max_column
                    and name not in approved_column_projection_sheets
                ):
                    unexpected.append(f"{name} 的数据区域被计划外缩小。")

                for row_offset, column_offset in _sample_coordinates(name, max_row, max_column):
                    row = row_offset + 1
                    column = column_offset + 1
                    if _allowed(allowed_ranges, name, row, column):
                        continue
                    before_cell = left.cell(row, column)
                    after_cell = right.cell(row, column)
                    if _style_signature(before_cell) != _style_signature(after_cell):
                        style_warnings.append(f"{name}!{after_cell.coordinate} 的样式发生变化。")
                    before_row = left.row_dimensions.get(row)
                    before_column = left.column_dimensions.get(before_cell.column_letter)
                    after_row = right.row_dimensions.get(row)
                    after_column = right.column_dimensions.get(after_cell.column_letter)
                    before_dimension = (
                        before_row.height if before_row else None,
                        before_column.width if before_column else None,
                    )
                    after_dimension = (
                        after_row.height if after_row else None,
                        after_column.width if after_column else None,
                    )
                    if before_dimension != after_dimension:
                        style_warnings.append(f"{name}!{after_cell.coordinate} 的行高或列宽发生变化。")
                if len(left._charts) != len(right._charts) or len(left._images) != len(right._images):
                    style_warnings.append(f"{name} 的图表或图片数量发生变化。")
        finally:
            before_struct.close()
            after_struct.close()

        before_values = load_workbook(before_path, read_only=True, data_only=False, keep_links=True)
        after_values = load_workbook(after_path, read_only=True, data_only=False, keep_links=True)
        try:
            for name in before_values.sheetnames:
                if name not in after_values.sheetnames:
                    continue
                left = before_values[name]
                right = after_values[name]
                max_row = max(left.max_row or 1, right.max_row or 1)
                max_column = max(left.max_column or 1, right.max_column or 1)
                left_rows = left.iter_rows(min_row=1, max_row=max_row, max_col=max_column)
                right_rows = right.iter_rows(min_row=1, max_row=max_row, max_col=max_column)
                for row_number, (left_row, right_row) in enumerate(zip(left_rows, right_rows), 1):
                    for column_number, (before_cell, after_cell) in enumerate(zip(left_row, right_row), 1):
                        coordinate = f"{get_column_letter(column_number)}{row_number}"
                        before_value = before_cell.value
                        after_value = after_cell.value
                        before_type = before_cell.data_type
                        after_type = after_cell.data_type
                        permitted_change = _allowed(
                            allowed_ranges, name, row_number, column_number
                        )
                        # Read-only EmptyCell reports no number format while a newly
                        # materialised cell reports Excel's equivalent default,
                        # ``General``.  Treat those as the same format so that an
                        # approved value write does not create a false warning.
                        before_format = before_cell.number_format or "General"
                        after_format = after_cell.number_format or "General"
                        if before_format != after_format and not permitted_change:
                            style_warnings.append(
                                f"{name}!{coordinate} 的数字格式发生变化。"
                            )
                        if (before_value, before_type) == (after_value, after_type):
                            continue
                        address = f"{name}!{coordinate}"
                        if len(changed) < 200:
                            changed.append(address)
                        if not permitted_change:
                            if len(unexpected) < 200:
                                unexpected.append(f"未批准的单元格变化：{address}")
                        if (
                            isinstance(after_value, str)
                            and after_value.startswith("=")
                            and "#REF!" in after_value.upper()
                            and not (isinstance(before_value, str) and "#REF!" in before_value.upper())
                        ):
                            unexpected.append(f"新增了损坏公式引用：{address}")
        finally:
            before_values.close()
            after_values.close()

        return WorkbookChangeSet(
            changed_ranges=list(dict.fromkeys(changed)),
            created_sheets=[name for name in after_names if name not in before_names],
            deleted_sheets=[name for name in before_names if name not in after_names],
            structural_changes=list(dict.fromkeys(structures)),
            unexpected_changes=list(dict.fromkeys(unexpected)),
            style_warnings=list(dict.fromkeys(style_warnings))[:200],
        )


class FinalVerifier:
    def __init__(self, max_sheet_rows: int):
        self.max_sheet_rows = max_sheet_rows
        self.diff = WorkbookDiffVerifier()

    def verify(
        self,
        *,
        task_id: str,
        original_path: Path,
        original_sha256: str,
        base_path: Path,
        base_version: int,
        base_sha256: str,
        candidate_path: Path,
        original_name: str,
        plan: TaskPlan,
        step_reports: list[VerificationReport],
        artifact_dir: Path,
    ) -> tuple[VerificationReport, WorkbookProfile, WorkbookChangeSet]:
        checks: list[VerificationCheck] = []
        current_original_hash = sha256_file(original_path) if original_path.is_file() else ""
        checks.append(VerificationCheck(
            name="original_upload_hash",
            passed=current_original_hash == original_sha256,
            detail="原始上传文件与入库哈希一致",
        ))
        checks.append(VerificationCheck(
            name="base_version_hash",
            passed=base_path.is_file() and sha256_file(base_path) == base_sha256,
            detail=f"正式基线 V{base_version} 未变化",
        ))
        checks.append(VerificationCheck(
            name="all_steps_verified",
            passed=bool(step_reports) and all(item.status != "failed_rolled_back" for item in step_reports),
            detail="所有步骤均已完成独立步骤核验",
        ))
        try:
            profile = profile_workbook(candidate_path, original_name, self.max_sheet_rows)
            checks.append(VerificationCheck(name="final_file_reopen", passed=True, detail="最终候选文件可重新打开"))
        except Exception as exc:
            empty = WorkbookChangeSet(unexpected_changes=[f"最终文件无法打开：{exc}"])
            report = VerificationReport(
                task_id=task_id,
                status="failed_rolled_back",
                checks=[*checks, VerificationCheck(name="final_file_reopen", passed=False, detail=str(exc))],
                warnings=[],
                manual_review_required=True,
                unexpected_changes=empty.unexpected_changes,
                verification_scope="final",
            )
            raise ValueError(report.model_dump_json()) from exc

        changes = self.diff.compare(base_path, candidate_path, plan)
        checks.append(VerificationCheck(
            name="planned_changes_only",
            passed=not changes.unexpected_changes,
            detail=("未发现计划外变更" if not changes.unexpected_changes else "；".join(changes.unexpected_changes[:10])),
        ))
        profile_columns = {
            sheet.name: {column.name for column in sheet.columns}
            for sheet in profile.sheets
        }
        # ``allowed_created_sheets`` is an authorization ceiling, not a promise
        # that a conditional report sheet will always be produced.  Require only
        # unconditional declarations; task-specific verifiers decide whether a
        # conditional output was necessary for the actual data.
        expected_sheets = {
            name
            for step in plan.steps
            for name in step.expected_changes.created_sheets
            if "（" not in name and "(" not in name
        }
        missing_sheets = sorted(expected_sheets - set(profile_columns))
        checks.append(VerificationCheck(
            name="planned_sheets_created",
            passed=not missing_sheets,
            detail=("计划工作表均存在" if not missing_sheets else "缺失：" + "、".join(missing_sheets)),
        ))
        missing_columns: list[str] = []
        for step in plan.steps:
            target_sheets = list(dict.fromkeys([
                *step.expected_changes.modified_sheets,
                *step.expected_changes.allowed_created_sheets,
                *step.expected_changes.created_sheets,
            ]))
            for column in step.expected_changes.new_columns:
                if not any(column in profile_columns.get(sheet, set()) for sheet in target_sheets):
                    missing_columns.append(column)
        checks.append(VerificationCheck(
            name="planned_columns_created",
            passed=not missing_columns,
            detail=(
                "计划新增列均存在"
                if not missing_columns
                else "缺失：" + "、".join(dict.fromkeys(missing_columns))
            ),
        ))
        expected_artifacts = {
            Path(name).name
            for step in plan.steps
            for name in step.expected_changes.allowed_artifacts
        }
        missing_artifacts = [name for name in expected_artifacts if not (artifact_dir / name).is_file()]
        checks.append(VerificationCheck(
            name="planned_artifacts_created",
            passed=not missing_artifacts,
            detail=("计划附件均存在" if not missing_artifacts else "缺失：" + "、".join(missing_artifacts)),
        ))
        final_hash = sha256_file(candidate_path)
        warnings = list(dict.fromkeys([
            *(warning for item in step_reports for warning in item.warnings),
            *changes.style_warnings,
        ]))
        failed = any(not item.passed for item in checks)
        report = VerificationReport(
            task_id=task_id,
            status="failed_rolled_back" if failed else "passed_with_warnings" if warnings else "passed",
            checks=checks,
            warnings=warnings,
            manual_review_required=bool(failed or warnings),
            output_path=str(candidate_path) if not failed else None,
            unexpected_changes=changes.unexpected_changes,
            style_warnings=changes.style_warnings,
            final_sha256=final_hash,
            verification_scope="final",
        )
        return report, profile, changes

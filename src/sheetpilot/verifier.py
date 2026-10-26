from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .files import sha256_file
from .schemas import (
    ChangeSet,
    LookupTaskSpec,
    VerificationCheck,
    VerificationReport,
    WorkbookProfile,
    TaskValidation,
    ToolResult,
)


def _cell_signature(value: Any, data_type: str, number_format: str) -> str:
    return repr((value, data_type, number_format))


def _original_region_hash(sheet, rows: int, columns: int) -> str:
    digest = hashlib.sha256()
    for row in range(1, rows + 1):
        for column in range(1, columns + 1):
            cell = sheet.cell(row, column)
            digest.update(_cell_signature(cell.value, cell.data_type, cell.number_format).encode("utf-8"))
    return digest.hexdigest()


def verify_output(
    task_id: str,
    original_path: Path,
    output_path: Path,
    original_hash: str,
    profile: WorkbookProfile,
    spec: LookupTaskSpec,
    changeset: ChangeSet,
) -> VerificationReport:
    checks: list[VerificationCheck] = []
    checks.append(
        VerificationCheck(
            name="original_file_unchanged",
            passed=sha256_file(original_path) == original_hash,
            detail="原文件哈希保持不变",
        )
    )
    try:
        original = load_workbook(original_path, data_only=False, keep_links=True)
        output = load_workbook(output_path, data_only=False, keep_links=True)
        checks.append(VerificationCheck(name="file_reopen", passed=True, detail="结果文件可以重新打开"))
    except Exception as exc:
        return VerificationReport(
            task_id=task_id,
            status="failed_rolled_back",
            checks=checks + [VerificationCheck(name="file_reopen", passed=False, detail=str(exc))],
            warnings=["结果文件无法打开。"],
            manual_review_required=True,
        )

    original_names = original.sheetnames
    preserved = all(name in output.sheetnames for name in original_names)
    checks.append(VerificationCheck(name="original_sheets_preserved", passed=preserved))
    main_profile = next(sheet for sheet in profile.sheets if sheet.name == spec.main_sheet)
    original_main = original[spec.main_sheet]
    output_main = output[spec.main_sheet]
    checks.append(
        VerificationCheck(
            name="row_count_preserved",
            passed=original_main.max_row == output_main.max_row,
            detail=f"{original_main.max_row} → {output_main.max_row}",
        )
    )
    original_region_same = _original_region_hash(
        original_main, original_main.max_row, main_profile.max_column
    ) == _original_region_hash(output_main, original_main.max_row, main_profile.max_column)
    checks.append(
        VerificationCheck(
            name="original_region_preserved",
            passed=original_region_same,
            detail="目标新增列以外的主表区域未变化",
        )
    )
    output_names = spec.output_column_names or spec.return_columns
    actual_headers = [
        output_main.cell(main_profile.header_row, main_profile.max_column + 1 + offset).value
        for offset in range(len(output_names))
    ]
    checks.append(
        VerificationCheck(
            name="new_columns_match_plan",
            passed=actual_headers == output_names,
            detail=f"实际：{actual_headers}",
        )
    )
    expected_rows = max(0, main_profile.max_row - main_profile.header_row)
    checks.append(
        VerificationCheck(
            name="write_count_complete",
            passed=changeset.rows_affected == expected_rows,
            detail=f"写入 {changeset.rows_affected} / {expected_rows}",
        )
    )
    accounted = changeset.matched_count + changeset.unmatched_count + changeset.empty_key_count
    checks.append(
        VerificationCheck(
            name="row_accounting",
            passed=accounted == expected_rows,
            detail=f"统计覆盖 {accounted} / {expected_rows}",
        )
    )
    if changeset.created_sheets:
        unmatched_sheet = output[changeset.created_sheets[0]]
        actual_unmatched = max(0, unmatched_sheet.max_row - 1)
        expected_unmatched = changeset.unmatched_count + changeset.empty_key_count
        checks.append(
            VerificationCheck(
                name="unmatched_report_count",
                passed=actual_unmatched == expected_unmatched,
                detail=f"异常清单 {actual_unmatched} / {expected_unmatched}",
            )
        )
    original.close()
    output.close()
    failed = [check for check in checks if not check.passed]
    if failed:
        status = "failed_rolled_back"
    elif changeset.warnings:
        status = "passed_with_warnings"
    else:
        status = "passed"
    return VerificationReport(
        task_id=task_id,
        status=status,
        checks=checks,
        warnings=changeset.warnings,
        manual_review_required=bool(changeset.warnings or failed),
        output_path=str(output_path) if not failed else None,
    )


def verify_tool_output(
    task_id: str,
    original_path: Path,
    original_hash: str,
    input_path: Path,
    input_hash: str,
    output_path: Path,
    result: ToolResult,
    task_verification: TaskValidation,
) -> VerificationReport:
    """Apply invariants shared by every v0.2 spreadsheet tool."""
    checks: list[VerificationCheck] = [
        VerificationCheck(
            name="original_file_unchanged",
            passed=sha256_file(original_path) == original_hash,
            detail="原始上传文件哈希保持不变",
        ),
        VerificationCheck(
            name="input_version_unchanged",
            passed=sha256_file(input_path) == input_hash,
            detail="执行期间输入版本未发生变化",
        ),
    ]
    try:
        workbook = load_workbook(output_path, read_only=True, data_only=False, keep_links=True)
        checks.append(VerificationCheck(name="file_reopen", passed=True, detail="结果文件可以重新打开"))
        missing_sheets = [name for name in result.created_sheets if name not in workbook.sheetnames]
        checks.append(
            VerificationCheck(
                name="planned_sheets_created",
                passed=not missing_sheets,
                detail="缺失：" + ", ".join(missing_sheets) if missing_sheets else "计划中新建的工作表均存在",
            )
        )
        workbook.close()
    except Exception as exc:
        checks.append(VerificationCheck(name="file_reopen", passed=False, detail=str(exc)))

    checks.append(
        VerificationCheck(
            name="task_specific_verification",
            passed=task_verification.valid,
            detail="；".join(task_verification.errors) if task_verification.errors else "任务专用检查通过",
        )
    )
    failed = [check for check in checks if not check.passed]
    warnings = list(dict.fromkeys([*result.warnings, *task_verification.warnings]))
    status = "failed_rolled_back" if failed else "passed_with_warnings" if warnings else "passed"
    return VerificationReport(
        task_id=task_id,
        status=status,
        checks=checks,
        warnings=warnings,
        manual_review_required=bool(failed or warnings),
        output_path=str(output_path) if not failed else None,
    )

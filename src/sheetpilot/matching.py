from __future__ import annotations

import decimal
import json
import unicodedata
from collections import Counter
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .errors import DuplicateKeyError, ExecutionError, IntentValidationError
from .schemas import (
    ChangeSet,
    DuplicatePolicy,
    ExecutionPlan,
    LookupTaskSpec,
    NormalizationMode,
    PlanStep,
    ValidationResult,
    WorkbookProfile,
)
from .workbook_utils import worksheet_dimensions


def normalize_key(value: Any, mode: NormalizationMode) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFKC", str(value)).strip()
    if not text:
        return None
    if mode == NormalizationMode.EXACT_TEXT:
        return text
    try:
        number = decimal.Decimal(text)
    except decimal.InvalidOperation:
        return text
    if not number.is_finite():
        return text
    normalized = format(number.normalize(), "f")
    if "." in normalized:
        normalized = normalized.rstrip("0").rstrip(".")
    return "0" if normalized in {"-0", ""} else normalized


def _sheet(profile: WorkbookProfile, name: str):
    return next((sheet for sheet in profile.sheets if sheet.name == name), None)


def _column_index(sheet_profile, name: str) -> int | None:
    matches = [column.index for column in sheet_profile.columns if column.name == name]
    return matches[0] if len(matches) == 1 else None


def validate_lookup_spec(path: Path, profile: WorkbookProfile, spec: LookupTaskSpec) -> ValidationResult:
    missing: list[str] = []
    ambiguities = list(spec.ambiguities)
    for field in ("main_sheet", "source_sheet", "main_key_column", "source_key_column"):
        if not getattr(spec, field):
            missing.append(field)
    if not spec.return_columns:
        missing.append("return_columns")
    output_column_names = spec.output_column_names or list(spec.return_columns)
    main = _sheet(profile, spec.main_sheet)
    source = _sheet(profile, spec.source_sheet)
    if spec.main_sheet and not main:
        ambiguities.append(f"主表“{spec.main_sheet}”不存在")
    if spec.source_sheet and not source:
        ambiguities.append(f"来源表“{spec.source_sheet}”不存在")
    if main and _column_index(main, spec.main_key_column) is None:
        ambiguities.append(f"主表关联列“{spec.main_key_column}”不存在或不唯一")
    if source and _column_index(source, spec.source_key_column) is None:
        ambiguities.append(f"来源表关联列“{spec.source_key_column}”不存在或不唯一")
    if source:
        for column in spec.return_columns:
            if _column_index(source, column) is None:
                ambiguities.append(f"返回列“{column}”不存在或不唯一")
    conflicts = []
    if main:
        existing = {column.name for column in main.columns}
        conflicts = [name for name in output_column_names if name in existing]
    duplicate_examples: list[dict[str, Any]] = []
    duplicate_key_count = 0
    type_warning = None
    if not missing and not ambiguities and main and source:
        workbook = load_workbook(path, read_only=True, data_only=True)
        source_ws = workbook[spec.source_sheet]
        source_key_index = _column_index(source, spec.source_key_column)
        assert source_key_index and source.header_row
        source_max_row, _ = worksheet_dimensions(
            source_ws,
            fallback_max_row=source.max_row,
            fallback_max_column=source.max_column,
        )
        keys = [
            normalize_key(source_ws.cell(row, source_key_index).value, spec.normalization_mode)
            for row in range(source.header_row + 1, source_max_row + 1)
        ]
        counts = Counter(key for key in keys if key is not None)
        all_duplicates = [
            {"key": key, "count": count} for key, count in counts.items() if count > 1
        ]
        duplicate_key_count = len(all_duplicates)
        duplicate_examples = all_duplicates[:20]
        workbook.close()
        main_type = next(c.inferred_type for c in main.columns if c.name == spec.main_key_column)
        source_type = next(c.inferred_type for c in source.columns if c.name == spec.source_key_column)
        if main_type != source_type:
            type_warning = f"两侧关联列类型不同：{main_type} / {source_type}"
    return ValidationResult(
        valid=not missing and not ambiguities and not conflicts and not (
        duplicate_key_count and spec.duplicate_policy == DuplicatePolicy.BLOCK
    ),
        missing_fields=missing,
        ambiguities=ambiguities,
        duplicate_key_count=duplicate_key_count,
        duplicate_examples=duplicate_examples,
        output_conflicts=conflicts,
        type_warning=type_warning,
    )


def build_plan(task_id: str, profile: WorkbookProfile, spec: LookupTaskSpec, validation: ValidationResult) -> ExecutionPlan:
    main = _sheet(profile, spec.main_sheet)
    if not main:
        raise IntentValidationError("无法为不存在的主表生成计划。")
    output_names = spec.output_column_names or spec.return_columns
    mapping = [f"{source} → {output}" for source, output in zip(spec.return_columns, output_names)]
    risk = "high" if validation.output_conflicts else "medium"
    return ExecutionPlan(
        task_id=task_id,
        summary=f"从“{spec.source_sheet}”匹配数据并追加到“{spec.main_sheet}”",
        steps=[
            PlanStep(order=1, action="validate", description="检查工作表、列、空键和重复键"),
            PlanStep(order=2, action="normalize", description=f"使用 {spec.normalization_mode.value} 标准化关联键"),
            PlanStep(order=3, action="left_join", description="按主表顺序执行左连接"),
            PlanStep(order=4, action="append", description="在主表末尾追加静态值列"),
            PlanStep(order=5, action="report", description="为未匹配记录生成异常清单"),
            PlanStep(order=6, action="verify", description="重新打开文件并检查结构与计数"),
        ],
        main_sheet=spec.main_sheet,
        source_sheet=spec.source_sheet,
        key_mapping=f"{spec.main_key_column} → {spec.source_key_column}",
        return_mapping=mapping,
        expected_new_columns=output_names,
        estimated_rows=max(0, main.max_row - (main.header_row or 1)),
        duplicate_key_count=validation.duplicate_key_count,
        output_conflicts=validation.output_conflicts,
        risk_level=risk,
    )


def _copy_header_style(source_cell, target_cell) -> None:
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.protection = copy(source_cell.protection)
        target_cell.number_format = source_cell.number_format


def execute_lookup(
    task_id: str,
    input_path: Path,
    output_path: Path,
    profile: WorkbookProfile,
    spec: LookupTaskSpec,
    version_before: int,
    version_after: int,
) -> ChangeSet:
    validation = validate_lookup_spec(input_path, profile, spec)
    if validation.missing_fields or validation.ambiguities or validation.output_conflicts:
        raise IntentValidationError(
            "任务信息不完整或与工作簿不一致。",
            detail=validation.model_dump(mode="json"),
        )
    if validation.duplicate_key_count and spec.duplicate_policy == DuplicatePolicy.BLOCK:
        raise DuplicateKeyError(
            "来源表关联键存在重复；请选择取消、第一条或最后一条。",
            detail=validation.model_dump(mode="json"),
        )

    workbook = load_workbook(input_path, data_only=False, keep_links=True)
    value_workbook = load_workbook(input_path, data_only=True, keep_links=True)
    main_profile = _sheet(profile, spec.main_sheet)
    source_profile = _sheet(profile, spec.source_sheet)
    assert main_profile and source_profile and main_profile.header_row and source_profile.header_row
    main_ws = workbook[spec.main_sheet]
    main_ws_values = value_workbook[spec.main_sheet]
    source_ws_values = value_workbook[spec.source_sheet]
    source_ws_formulas = workbook[spec.source_sheet]
    main_key_index = _column_index(main_profile, spec.main_key_column)
    source_key_index = _column_index(source_profile, spec.source_key_column)
    return_indexes = [_column_index(source_profile, name) for name in spec.return_columns]
    assert main_key_index and source_key_index and all(return_indexes)

    lookup: dict[str, tuple[Any, ...]] = {}
    duplicates: set[str] = set()
    for row in range(source_profile.header_row + 1, source_ws_values.max_row + 1):
        key = normalize_key(source_ws_values.cell(row, source_key_index).value, spec.normalization_mode)
        if key is None:
            continue
        values: list[Any] = []
        for index in return_indexes:
            formula_cell = source_ws_formulas.cell(row, index)
            value = source_ws_values.cell(row, index).value
            if formula_cell.data_type == "f" and value is None:
                workbook.close()
                value_workbook.close()
                raise ExecutionError("来源返回列包含没有缓存结果的公式，v0.3 无法安全写入静态值。")
            values.append(value)
        record = tuple(values)
        if key in lookup:
            duplicates.add(key)
            if spec.duplicate_policy == DuplicatePolicy.LAST:
                lookup[key] = record
        else:
            lookup[key] = record

    output_names = spec.output_column_names or spec.return_columns
    start_column = main_ws.max_column + 1
    template_header = main_ws.cell(main_profile.header_row, max(1, main_ws.max_column))
    for offset, name in enumerate(output_names):
        target = main_ws.cell(main_profile.header_row, start_column + offset, name)
        _copy_header_style(template_header, target)

    unmatched_rows: list[tuple[int, Any, str]] = []
    matched = unmatched = empty = 0
    for row in range(main_profile.header_row + 1, main_ws.max_row + 1):
        raw_key = main_ws_values.cell(row, main_key_index).value
        key = normalize_key(raw_key, spec.normalization_mode)
        if key is None:
            empty += 1
            unmatched_rows.append((row, raw_key, "empty_key"))
            values = (None,) * len(output_names)
        elif key not in lookup:
            unmatched += 1
            unmatched_rows.append((row, raw_key, "not_found"))
            values = (None,) * len(output_names)
        else:
            matched += 1
            values = lookup[key]
        for offset, value in enumerate(values):
            main_ws.cell(row, start_column + offset, value)

    created_sheets: list[str] = []
    if unmatched_rows:
        base = "SheetPilot_Unmatched"
        name = base
        suffix = 2
        while name in workbook.sheetnames:
            name = f"{base}_{suffix}"
            suffix += 1
        unmatched_ws = workbook.create_sheet(name)
        unmatched_ws.append(["主表行号", "关联键", "失败原因"])
        for record in unmatched_rows:
            unmatched_ws.append(list(record))
        created_sheets.append(name)

    try:
        workbook.save(output_path)
    except Exception as exc:
        raise ExecutionError("保存结果文件失败。") from exc
    finally:
        workbook.close()
        value_workbook.close()

    end_column = start_column + len(output_names) - 1
    warnings = []
    if unmatched_rows:
        warnings.append(f"{len(unmatched_rows)} 条记录需要人工检查。")
    if duplicates:
        warnings.append(f"按 {spec.duplicate_policy.value} 处理了 {len(duplicates)} 个重复来源键。")
    return ChangeSet(
        task_id=task_id,
        version_before=version_before,
        version_after=version_after,
        modified_sheets=[spec.main_sheet],
        modified_ranges=[
            f"{spec.main_sheet}!{get_column_letter(start_column)}{main_profile.header_row}:"
            f"{get_column_letter(end_column)}{main_ws.max_row}"
        ],
        created_sheets=created_sheets,
        rows_affected=max(0, main_profile.max_row - main_profile.header_row),
        matched_count=matched,
        unmatched_count=unmatched,
        empty_key_count=empty,
        duplicate_key_count=len(duplicates),
        warnings=warnings,
    )


def write_conflict_report(path: Path, validation: ValidationResult) -> None:
    path.write_text(json.dumps(validation.model_dump(mode="json"), ensure_ascii=False, indent=2), encoding="utf-8")

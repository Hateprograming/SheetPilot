from __future__ import annotations

import json
import os
import shutil
import subprocess
import tempfile
import time
import signal
import zipfile
import xml.etree.ElementTree as ET
from contextvars import ContextVar
from typing import Callable
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties


EXCEL_ERROR_PREFIXES = (
    "#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#SPILL!", "#CALC!",
)


@dataclass(slots=True)
class RecalculationResult:
    success: bool
    backend: str = "none"
    formula_cells: int = 0
    calculated_cells: int = 0
    empty_string_cells: int = 0
    blank_cached_cells: list[str] = field(default_factory=list)
    error_cells: dict[str, str] = field(default_factory=dict)
    message: str = ""
    timings_seconds: dict[str, float] = field(default_factory=dict)

    def model_dump(self) -> dict[str, Any]:
        return asdict(self)


def _set_automatic_calculation(path: Path) -> None:
    """Tell spreadsheet applications to perform a full calculation when the file is opened."""
    workbook = load_workbook(path, data_only=False, keep_links=True)
    try:
        if workbook.calculation is None:
            workbook.calculation = CalcProperties()
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(path)
    finally:
        workbook.close()


def _audit_cached_values(path: Path) -> RecalculationResult:
    """One streaming pass per sheet; distinguish absent caches from empty strings."""
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=True)
    result = RecalculationResult(success=True)
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    try:
        with zipfile.ZipFile(path) as archive:
            for sheet in workbook.worksheets:
                with archive.open(sheet._worksheet_path.lstrip("/")) as handle:
                    for _, node in ET.iterparse(handle, events=("end",)):
                        if node.tag == ns + "c":
                            if node.find(ns + "f") is not None:
                                result.formula_cells += 1
                                address = f"{sheet.title}!{node.get('r')}"
                                cached = node.find(ns + "v")
                                value = cached.text if cached is not None else None
                                if node.get("t") == "e":
                                    result.error_cells[address] = value or "未知公式错误"
                                elif value is not None:
                                    result.calculated_cells += 1
                                elif node.get("t") == "str" and cached is not None:
                                    result.calculated_cells += 1
                                    result.empty_string_cells += 1
                                else:
                                    result.blank_cached_cells.append(address)
                            node.clear()
                        elif node.tag == ns + "row":
                            node.clear()
    finally:
        workbook.close()
    result.success = not result.error_cells and not result.blank_cached_cells
    return result


_control: ContextVar[Callable[[], bool] | None] = ContextVar("recalc_control", default=None)


def _run_backend(command, *, timeout, owner_pid_file: Path | None = None, **kwargs):
    """Poll only this invocation's process; cancellation never kills unrelated Excel."""
    if os.name != "nt":
        kwargs["start_new_session"] = True
    process = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, **kwargs)
    deadline = time.monotonic() + timeout
    try:
        while True:
            cancelled = _control.get()
            if cancelled and cancelled():
                raise InterruptedError("公式重算已中断，正式版本未改变。")
            remaining = deadline - time.monotonic()
            if remaining <= 0:
                raise subprocess.TimeoutExpired(command, timeout)
            try:
                stdout, stderr = process.communicate(timeout=min(0.2, remaining))
                return subprocess.CompletedProcess(command, process.returncode, stdout, stderr)
            except subprocess.TimeoutExpired:
                continue
    finally:
        if process.poll() is None:
            if os.name == "nt":
                subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"], capture_output=True,
                               creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            else:
                try:
                    os.killpg(process.pid, signal.SIGKILL)
                except ProcessLookupError:
                    pass
            if process.poll() is None:
                process.kill()
            process.communicate()
            if owner_pid_file and owner_pid_file.is_file() and os.name == "nt":
                # The script writes only a newly created Excel instance's PID.
                pid = owner_pid_file.read_text(encoding="utf-8-sig").strip()
                if pid.isdigit():
                    subprocess.run(["taskkill", "/PID", pid, "/T", "/F"], capture_output=True,
                                   creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))


def _excel_script_path() -> Path:
    return Path(__file__).resolve().parent / "scripts" / "recalculate_excel.ps1"


def _recalculate_with_excel(path: Path, timeout: int) -> tuple[bool, str]:
    if os.name != "nt":
        return False, "Microsoft Excel COM 仅适用于 Windows。"
    powershell = shutil.which("powershell.exe") or shutil.which("powershell")
    script = _excel_script_path()
    if not powershell or not script.exists():
        return False, "未找到 PowerShell 或 Excel 重算脚本。"
    owner_handle = tempfile.NamedTemporaryFile(prefix="sheetpilot-excel-", suffix=".pid", delete=False)
    owner_file = Path(owner_handle.name)
    owner_handle.close()
    try:
        completed = _run_backend(
            [
                powershell,
                "-NoProfile",
                "-NonInteractive",
                "-ExecutionPolicy", "Bypass",
                "-File", str(script),
                "-WorkbookPath", str(path.resolve()),
                "-OwnerPidFile", str(owner_file),
            ],
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout, owner_pid_file=owner_file,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    except subprocess.TimeoutExpired:
        return False, f"Microsoft Excel 重算超过 {timeout} 秒。"
    finally:
        owner_file.unlink(missing_ok=True)
    if completed.returncode != 0:
        detail = (completed.stderr or completed.stdout).strip()
        return False, detail or f"Microsoft Excel 重算进程退出码为 {completed.returncode}。"
    lines = [line.strip() for line in completed.stdout.splitlines() if line.strip()]
    if not lines:
        return False, "Microsoft Excel 重算没有返回结果。"
    try:
        payload = json.loads(lines[-1])
    except json.JSONDecodeError:
        return False, lines[-1]
    return bool(payload.get("success")), str(payload.get("message", ""))


def _recalculate_with_libreoffice(path: Path, timeout: int) -> tuple[bool, str]:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return False, "未找到 LibreOffice。"
    with tempfile.TemporaryDirectory(prefix="sheetpilot-recalc-") as temporary:
        root = Path(temporary)
        source = root / path.name
        output_dir = root / "output"
        output_dir.mkdir()
        shutil.copy2(path, source)
        try:
            completed = _run_backend(
                [
                    soffice, f"-env:UserInstallation={(root / 'profile').resolve().as_uri()}",
                    "--headless", "--norestore", "--convert-to", "xlsx",
                    "--outdir", str(output_dir), str(source),
                ],
                text=True,
                timeout=timeout,
            )
        except subprocess.TimeoutExpired:
            return False, f"LibreOffice 重算超过 {timeout} 秒。"
        converted = output_dir / path.name
        if completed.returncode != 0 or not converted.exists():
            detail = (completed.stderr or completed.stdout).strip()
            return False, detail or "LibreOffice 未生成重算后的工作簿。"
        shutil.copy2(converted, path)
    return True, "LibreOffice 已完成工作簿重算。"


def recalculate_workbook(path: Path, timeout: float = 90, *, cancelled: Callable[[], bool] | None = None) -> RecalculationResult:
    """Recalculate formulas with Excel/LibreOffice and verify cached values.

    The workbook is always marked for automatic full calculation. Excel is preferred on
    Windows because it provides the most faithful formula compatibility; LibreOffice is the
    portable fallback. If neither backend is available, formulas remain intact and the caller
    receives an explicit non-success result.
    """
    path = Path(path)
    deadline = time.monotonic() + max(0, timeout)
    if cancelled and cancelled():
        return RecalculationResult(success=False, message="公式重算已中断，正式版本未改变。")
    if timeout <= 0:
        return RecalculationResult(success=False, message="公式重算达到执行时间预算。")
    timings: dict[str, float] = {}
    stage_started = time.perf_counter()
    _set_automatic_calculation(path)
    timings["set_automatic_calculation"] = time.perf_counter() - stage_started
    messages: list[str] = []
    for backend, runner in (("excel", _recalculate_with_excel), ("libreoffice", _recalculate_with_libreoffice)):
        stage_started = time.perf_counter()
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            messages.append("公式重算达到时间预算。")
            break
        token = _control.set(cancelled)
        try:
            ok, message = runner(path, remaining)
        except InterruptedError as exc:
            messages.append(str(exc))
            break
        finally:
            _control.reset(token)
        timings[f"backend_{backend}"] = time.perf_counter() - stage_started
        messages.append(f"{backend}: {message}")
        if not ok:
            continue
        stage_started = time.perf_counter()
        audit = _audit_cached_values(path)
        timings["cache_audit"] = time.perf_counter() - stage_started
        audit.backend = backend
        audit.message = message
        audit.timings_seconds = {
            key: round(value, 6) for key, value in timings.items()
        }
        audit.success = not audit.error_cells and not audit.blank_cached_cells
        if not audit.success:
            audit.message += f"；{len(audit.blank_cached_cells)} 个公式缺少缓存，{len(audit.error_cells)} 个计算错误。"
        if time.monotonic() >= deadline:
            audit.success = False
            audit.message = "公式重算及核验超过执行时间预算。"
        return audit
    stage_started = time.perf_counter()
    audit = _audit_cached_values(path)
    timings["cache_audit"] = time.perf_counter() - stage_started
    audit.success = False
    audit.backend = "none"
    audit.message = "；".join(messages)
    audit.timings_seconds = {
        key: round(value, 6) for key, value in timings.items()
    }
    return audit

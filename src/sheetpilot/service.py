from __future__ import annotations

import json
import logging
import re
import shutil
import time
import threading
import unicodedata
from collections import Counter
from concurrent.futures import Future, ThreadPoolExecutor
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from .conversation import clarification_for, plan_text, result_text, workbook_greeting
from .commit import VersionCommitter
from .conversion import csv_to_xlsx
from .context import build_context, guess_task_type
from .config import Settings
from .errors import IntentValidationError, LLMError, SheetPilotError, VerificationError
from .files import TaskFileManager, sha256_file
from .llm import LLMClient, OpenAICompatibleClient
from .integrity import FinalVerifier, WorkbookDiffVerifier
from .matching import build_plan, validate_lookup_spec, write_conflict_report
from .planner import build_task_plan, validate_plan_structure
from .policy import PolicyEngine
from .profiler import ensure_supported_profile, profile_workbook
from .schemas import (
    AgentTurnResult,
    AggregateTaskSpec,
    CleanOperation,
    CleanTaskSpec,
    ColumnRef,
    ConversationMessage,
    ConversationRole,
    ConversationState,
    DimensionResolutionAction,
    CompositeTaskSpec,
    ConversionTaskSpec,
    ConversionDirection,
    ExecutionResult,
    ExecutionPlan,
    FailureCategory,
    FailureRecord,
    FormulaMode,
    FormulaTaskSpec,
    InvalidNumericPolicy,
    LookupTaskSpec,
    MessageKind,
    NormalizeTaskSpec,
    NormalizeOperation,
    SortDirection,
    SortKey,
    SortRankTaskSpec,
    TaskEnvelope,
    TaskPlan,
    TaskValidation,
    TransformOperation,
    TransformTaskSpec,
    ToolResult,
    ValidationResult,
    VerificationReport,
    VerificationCheck,
    WorkbookProfile,
)
from .p0 import P0Store, authorized_model_rows
from .storage import Repository
from .skills import SkillRegistry, capability_refusal
from .tools import ToolContext, ToolRegistry, transform_write_columns
from .transaction import TransactionExecutor
from .verifier import verify_tool_output
from .workbook_utils import worksheet_dimensions


LOGGER = logging.getLogger(__name__)


class SheetPilotService:
    def __init__(self, settings: Settings, llm: LLMClient | None = None):
        settings.ensure_dirs()
        self.settings = settings
        self.repository = Repository(settings.database_path)
        self.files = TaskFileManager(settings.data_dir, settings.max_file_bytes)
        self.llm = llm
        self._policy_sources: dict[str, str] = {}
        self.tools = ToolRegistry()
        self.skills = SkillRegistry()
        self.policy = PolicyEngine(self.skills)
        self.final_verifier = FinalVerifier(settings.max_sheet_rows)
        self.diff_verifier = WorkbookDiffVerifier()
        self.committer = VersionCommitter(
            self.repository, self.files, self.final_verifier
        )
        self.transaction_executor = TransactionExecutor(self)
        self._background_executor = ThreadPoolExecutor(
            max_workers=4, thread_name_prefix="sheetpilot"
        )
        self._background_futures: dict[str, Future[Any]] = {}
        self.p0 = P0Store(self)
        self.committer.recover_incomplete_commits()
        self.p0.recover_interrupted()

    def configured_llm(self) -> LLMClient:
        if self.llm is None:
            self.llm = OpenAICompatibleClient(self.settings)
        return self.llm

    def create_task(
        self, original_name: str, content: bytes, *, owner_id: str | None = None
    ) -> WorkbookProfile:
        task_id, original_path = self.files.ingest(original_name, content)
        self.repository.create_task(
            task_id, original_name, original_path, sha256_file(original_path), owner_id
        )
        self.p0.touch(task_id)
        try:
            current_path = original_path
            csv_detection = None
            if original_path.suffix.casefold() == ".csv":
                current_path = self.files.staging_path(task_id)
                csv_detection = csv_to_xlsx(original_path, current_path)
            profile = profile_workbook(
                current_path, original_name, self.settings.max_sheet_rows
            )
            ensure_supported_profile(profile)
            if csv_detection and not csv_detection.confident:
                profile.risks.append("CSV_DELIMITER_UNCERTAIN")
        except Exception as exc:
            self.repository.update_task(
                task_id,
                status="rejected",
                error_json=json.dumps({"type": type(exc).__name__, "message": str(exc)}, ensure_ascii=False),
            )
            raise
        self.repository.update_task(
            task_id,
            status="needs_conversion" if original_path.suffix.casefold() == ".csv" else "profiled",
            current_path=str(current_path),
            profile_json=profile.model_dump_json(),
        )
        self.repository.add_version(
            task_id, 0, current_path, profile.sha256, None, profile.model_dump_json()
        )
        self.repository.log_event(task_id, "profiled", {"risks": profile.risks})
        self._add_message(
            task_id,
            ConversationRole.ASSISTANT,
            MessageKind.WORKBOOK,
            (
                "CSV 已在本地完成安全预读。请告诉我将它转换为 XLSX；转换计划获批前，其他 Skill 不会执行。"
                if original_path.suffix.casefold() == ".csv"
                else workbook_greeting(profile)
            ),
            {"profile": profile.model_dump(mode="json"), "source_format": original_path.suffix.casefold()},
        )
        self.p0.touch(task_id)
        return profile.model_copy(update={"file_name": f"{task_id}:{profile.file_name}"})

    def task_id_from_profile(self, profile: WorkbookProfile) -> str:
        return profile.file_name.split(":", 1)[0]

    def load_profile(self, task_id: str) -> WorkbookProfile:
        row = self._task(task_id)
        return WorkbookProfile.model_validate_json(row["profile_json"])

    def update_headers(self, task_id: str, overrides: dict[str, int]) -> WorkbookProfile:
        row = self._task(task_id)
        profile = profile_workbook(
            Path(row["current_path"]), row["original_name"], self.settings.max_sheet_rows, overrides
        )
        ensure_supported_profile(profile)
        self.repository.update_task(task_id, profile_json=profile.model_dump_json())
        self.repository.update_version_profile(
            task_id, int(row["current_version"]), profile.model_dump_json()
        )
        self.repository.log_event(task_id, "headers_updated", {"overrides": overrides})
        return profile

    def parse_request(self, task_id: str, request: str) -> LookupTaskSpec:
        profile = self.load_profile(task_id)
        spec = self.configured_llm().parse_lookup(request, profile)
        self.repository.update_task(task_id, status="intent_parsed", spec_json=spec.model_dump_json())
        self.repository.log_event(
            task_id, "intent_parsed",
            {"task_type": spec.task_type, "confidence": spec.confidence},
        )
        return spec

    def save_spec(self, task_id: str, spec: LookupTaskSpec) -> None:
        self.repository.update_task(task_id, status="intent_parsed", spec_json=spec.model_dump_json())

    def validate(self, task_id: str, spec: LookupTaskSpec) -> ValidationResult:
        row = self._task(task_id)
        profile = self.load_profile(task_id)
        result = validate_lookup_spec(Path(row["current_path"]), profile, spec)
        self.repository.log_event(
            task_id,
            "intent_validated",
            {
                "valid": result.valid,
                "missing_field_count": len(result.missing_fields),
                "ambiguity_count": len(result.ambiguities),
                "duplicate_key_count": result.duplicate_key_count,
                "output_conflict_count": len(result.output_conflicts),
                "has_type_warning": bool(result.type_warning),
            },
        )
        if result.duplicate_key_count:
            write_conflict_report(
                self.files.artifact_path(task_id, "duplicate_keys.json"), result
            )
        return result

    def make_plan(self, task_id: str, spec: LookupTaskSpec) -> tuple[ExecutionPlan, ValidationResult]:
        profile = self.load_profile(task_id)
        validation = self.validate(task_id, spec)
        plan = build_plan(task_id, profile, spec, validation)
        self.repository.update_task(
            task_id,
            status="awaiting_approval",
            spec_json=spec.model_dump_json(),
            plan_json=plan.model_dump_json(),
        )
        self.repository.log_event(task_id, "plan_built", plan.model_dump(mode="json"))
        return plan, validation

    def approve(self, task_id: str, decision: str, payload: dict[str, Any] | None = None) -> None:
        if decision not in {"approve", "edit", "reject"}:
            raise ValueError("invalid approval decision")
        approval_payload = payload or {}
        if decision == "approve":
            row = self._task(task_id)
            if row.get("plan_json"):
                stored_plan = json.loads(row["plan_json"])
                acknowledged = set(approval_payload.get("acknowledgements", []))
                missing = [
                    item for item in stored_plan.get("confirmation_requirements", [])
                    if item not in acknowledged
                ]
                if missing:
                    raise IntentValidationError(
                        "该计划包含覆盖已有列操作，必须先明确确认覆盖风险。",
                        detail={"missing_acknowledgements": missing},
                    )
        self.repository.add_approval(task_id, decision, approval_payload)
        self.repository.update_task(
            task_id, status="approved" if decision == "approve" else "cancelled" if decision == "reject" else "editing"
        )

    def execute(self, task_id: str, spec: LookupTaskSpec) -> tuple[Path, VerificationReport]:
        row = self._task(task_id)
        if row["status"] != "approved":
            raise IntentValidationError("任务尚未获得用户批准。")
        # Compatibility entry point for the pre-Skill lookup API.  Translate the
        # already-approved lookup spec into the current deterministic plan so it
        # cannot bypass leases, independent diff verification, or atomic commit.
        profile = self.load_profile(task_id)
        current_version = int(row["current_version"])
        context = ToolContext(
            task_id=task_id,
            input_path=Path(row["current_path"]),
            output_path=self.files.task_dir(task_id) / "runs" / "legacy-preview.xlsx",
            profile=profile,
            version_before=current_version,
            version_after=current_version + 1,
            artifact_dir=self.files.task_dir(task_id) / "artifacts",
            original_path=Path(row["original_path"]),
        )
        envelope = TaskEnvelope(goal="执行已批准的精确关联计划", spec=spec)
        safe_plan, validation = build_task_plan(context, envelope, self.tools, self.skills)
        policy_validation = self.policy.validate_envelope(profile, envelope)
        errors = [*validation.errors, *policy_validation.errors]
        if errors:
            raise IntentValidationError("旧版关联计划无法通过当前安全策略。", detail={"errors": errors})
        self.repository.update_task(
            task_id,
            status="approved",
            spec_json=envelope.model_dump_json(),
            plan_json=safe_plan.model_dump_json(),
        )
        output, report, result = self.execute_task_plan(task_id, envelope, safe_plan)
        # Preserve the old result_payload shape for callers that still read
        # lookup counters directly from changeset_json.
        step_details = (result.details.get("steps") or [{}])[0].get("details", {})
        legacy_changeset = step_details.get("legacy_changeset") or dict(result.counters)
        if "legacy_changeset" not in step_details:
            legacy_changeset.update({"warnings": result.warnings, "details": result.details})
        self.repository.update_task(
            task_id, changeset_json=json.dumps(legacy_changeset, ensure_ascii=False)
        )
        return output, report

    def save_task_envelope(self, task_id: str, envelope: TaskEnvelope) -> None:
        self.repository.update_task(
            task_id,
            status="intent_parsed",
            spec_json=envelope.model_dump_json(),
        )

    @staticmethod
    def prepare_model_envelope(
        envelope: TaskEnvelope, user_source_text: str,
        profile: WorkbookProfile | None = None,
    ) -> TaskEnvelope:
        """Conservatively normalize model output before deterministic planning.

        Model-authored assumptions are never treated as user authorization. Unsafe
        exception policies are removed so preflight can ask the user with real rows.
        """
        prepared = envelope.model_copy(deep=True)
        prepared.assumptions = []
        normalized = unicodedata.normalize("NFKC", user_source_text).casefold()
        compact = re.sub(r"\s+", "", normalized)
        compact_segments = [
            re.sub(r"\s+", "", item)
            for item in re.split(r"[。；;\n]+", normalized)
            if item.strip()
        ]
        specs = (
            list(prepared.spec.operations)
            if isinstance(prepared.spec, CompositeTaskSpec)
            else [prepared.spec]
        )
        if isinstance(prepared.spec, CompositeTaskSpec):
            exports = [spec for spec in specs if isinstance(spec, ConversionTaskSpec) and spec.direction == ConversionDirection.XLSX_TO_CSV]
            if len(exports) == 1:
                specs = [spec for spec in specs if spec is not exports[0]] + exports
                prepared.spec.operations = specs

        # A single explicit column expression has an unambiguous local route,
        # even if the model initially chose a transform or diagnostic skill.
        from .formula_intent import column_formula
        latest_request = user_source_text.splitlines()[-1] if user_source_text.splitlines() else user_source_text
        if profile is not None and not isinstance(prepared.spec, CompositeTaskSpec) and not any(
            word in latest_request.casefold() for word in ("然后", "同时", "并", "csv", "转换为")
        ):
            named_sheets = [s for s in profile.sheets if s.name in latest_request]
            candidates = []
            for sheet in named_sheets or profile.sheets:
                for col in sheet.columns:
                    target = col.display_header or col.name
                    formula = column_formula(profile, sheet.name, target, latest_request)
                    if formula:
                        candidate = prepared.spec.model_copy(deep=True) if isinstance(prepared.spec, FormulaTaskSpec) else FormulaTaskSpec()
                        candidate.sheet, candidate.column = sheet.name, col.name
                        candidate.formula = formula
                        candidate.mode, candidate.confidence = FormulaMode.SET_FORMULA, 1
                        candidates.append(candidate)
            if len(candidates) == 1:
                prepared.spec = candidates[0]
                specs = [prepared.spec]

        if profile is not None:
            sheet_by_name = {sheet.name: sheet for sheet in profile.sheets}

            def unique_sheet_for_columns(names: list[str]) -> str | None:
                wanted = [name for name in names if name]
                if not wanted:
                    return None
                matches = []
                for sheet in profile.sheets:
                    headers = {
                        column.display_header or column.name
                        for column in sheet.columns
                        if column.display_header or column.name
                    }
                    if all(name in headers for name in wanted):
                        matches.append(sheet.name)
                return matches[0] if len(matches) == 1 else None

            def clear_resolved_sheet_issues(spec: Any) -> None:
                spec.missing_fields = [
                    item for item in spec.missing_fields
                    if str(item).casefold() not in {"sheet", "工作表", "未指定工作表"}
                ]
                spec.ambiguities = [
                    item for item in spec.ambiguities
                    if not (
                        any(token in str(item).casefold() for token in ("sheet", "工作表"))
                        and any(token in str(item) for token in ("缺少", "未指定", "不确定", "无法确定"))
                    )
                ]

            for spec in specs:
                names: list[str] = []
                refs: list[ColumnRef] = []
                if isinstance(spec, NormalizeTaskSpec):
                    names = [spec.column]
                elif isinstance(spec, FormulaTaskSpec):
                    names = [spec.column]
                elif isinstance(spec, TransformTaskSpec):
                    if spec.operation == TransformOperation.IMPUTE_NUMERIC:
                        refs = [rule.source for rule in spec.impute_rules]
                    elif spec.operation == TransformOperation.MAP_VALUES:
                        refs = [rule.source for rule in spec.map_rules]
                    elif spec.operation == TransformOperation.PARSE_DIMENSIONS and spec.dimension_rule:
                        refs = [spec.dimension_rule.source]
                    elif spec.operation == TransformOperation.DERIVE_METRIC and spec.derive_rule:
                        refs = list(spec.derive_rule.sources)
                    names = [ref.display_header for ref in refs if ref.display_header]
                if hasattr(spec, "sheet") and (
                    not spec.sheet or spec.sheet not in sheet_by_name
                ):
                    inferred_sheet = unique_sheet_for_columns(names)
                    if inferred_sheet:
                        spec.sheet = inferred_sheet
                        clear_resolved_sheet_issues(spec)
                        for ref in refs:
                            ref.sheet = inferred_sheet
                if isinstance(spec, LookupTaskSpec) and not spec.main_sheet and spec.source_sheet:
                    candidates = [
                        sheet.name for sheet in profile.sheets
                        if sheet.name != spec.source_sheet
                        and any(
                            (column.display_header or column.name) == spec.main_key_column
                            for column in sheet.columns
                        )
                    ]
                    if len(candidates) == 1:
                        spec.main_sheet = candidates[0]
                        spec.missing_fields = [
                            item for item in spec.missing_fields if item != "main_sheet"
                        ]

            known_columns = {
                column.display_header or column.name
                for sheet in profile.sheets for column in sheet.columns
                if column.display_header or column.name
            }
            mentioned = [name for name in known_columns if name in user_source_text]
            if len(mentioned) == 1:
                explicit_column = mentioned[0]
                for spec in specs:
                    if isinstance(spec, NormalizeTaskSpec):
                        spec.column = explicit_column
                    elif isinstance(spec, TransformTaskSpec):
                        refs = []
                        if spec.operation == TransformOperation.IMPUTE_NUMERIC:
                            refs = [rule.source for rule in spec.impute_rules]
                        elif spec.operation == TransformOperation.MAP_VALUES:
                            refs = [rule.source for rule in spec.map_rules]
                        elif spec.operation == TransformOperation.PARSE_DIMENSIONS and spec.dimension_rule:
                            refs = [spec.dimension_rule.source]
                        if len(refs) == 1:
                            match = next(
                                (
                                    (sheet, column)
                                    for sheet in profile.sheets for column in sheet.columns
                                    if (column.display_header or column.name) == explicit_column
                                ),
                                None,
                            )
                            if match:
                                sheet, column = match
                                refs[0].sheet = sheet.name
                                refs[0].column_index = column.index
                                refs[0].column_letter = column.letter
                                refs[0].display_header = column.display_header or column.name
                                refs[0].header_occurrence = column.header_occurrence
            for spec in specs:
                if isinstance(spec, LookupTaskSpec):
                    authorized_returns = [
                        column for column in spec.return_columns
                        if column and column in user_source_text
                    ]
                    if spec.return_columns and not authorized_returns:
                        spec.return_columns = []
                        spec.output_column_names = []
                        if "return_columns" not in spec.missing_fields:
                            spec.missing_fields.append("return_columns")
                elif (
                    isinstance(spec, CleanTaskSpec)
                    and spec.operation in {
                        CleanOperation.MARK_DUPLICATES,
                        CleanOperation.REMOVE_DUPLICATES,
                    }
                    and spec.key_columns
                    and not any(key in user_source_text for key in spec.key_columns)
                ):
                    spec.key_columns = []
                    if "key_columns" not in spec.missing_fields:
                        spec.missing_fields.append("key_columns")

        # Ground text normalization in the user's operation, never model defaults.
        for spec in specs:
            if not isinstance(spec, NormalizeTaskSpec):
                continue
            segments = [part for part in compact_segments if spec.column and spec.column.casefold() in part]
            relevant = "；".join(segments or compact_segments)
            explicit = []
            prohibited = set()
            cues = (
                (NormalizeOperation.NFKC, ("nfkc", "全角", "半角")),
                (NormalizeOperation.TRIM, ("去空格", "去除空格", "去掉空格", "清除空格", "首尾空格", "首尾空白", "两端空格", "两端空白", "两边空格", "前后空格", "trim")),
                (NormalizeOperation.LOWER, ("小写",)),
                (NormalizeOperation.UPPER, ("大写",)),
                (NormalizeOperation.TO_TEXT, ("转成文本", "转为文本", "文本格式")),
                (NormalizeOperation.TO_NUMBER, ("转成数值", "转为数值", "转成数字", "转为数字")),
            )
            for operation, words in cues:
                positive = False
                for word in words:
                    for match in re.finditer(re.escape(word), relevant):
                        if re.search(r"(?:不|不要|无需|不用|不需要|不是)(?:做|执行|进行|使用|统一|去|去除|去掉|清除)?$", relevant[:match.start()]):
                            prohibited.add(operation)
                        else:
                            positive = True
                if positive and operation not in prohibited:
                    explicit.append(operation)
            # A literal before/after example also authorizes trimming.
            for before, after in re.findall(r'[“\"‘]([^”\"’]*)[”\"’]\s*(?:→|->|变为|变成|改为|转换为)\s*[“\"‘]([^”\"’]*)[”\"’]', user_source_text):
                if before != after and before.strip() == after and NormalizeOperation.TRIM not in explicit:
                    explicit.append(NormalizeOperation.TRIM)
            specialized = any(word in relevant for word in ("日期", "填充", "后缀"))
            if explicit:
                if specialized:
                    allowed = {*explicit, NormalizeOperation.DATE_FORMAT, NormalizeOperation.FILL_BLANK, NormalizeOperation.STRIP_SUFFIX}
                    spec.operations = list(dict.fromkeys([op for op in spec.operations if op in allowed] + explicit))
                    if NormalizeOperation.STRIP_SUFFIX in spec.operations and NormalizeOperation.TO_NUMBER in spec.operations:
                        spec.operations.remove(NormalizeOperation.TO_NUMBER)
                        spec.operations.append(NormalizeOperation.TO_NUMBER)
                else:
                    spec.operations = explicit
            elif not specialized:
                spec.operations = [op for op in (NormalizeOperation.NFKC, NormalizeOperation.TRIM) if op not in prohibited]
            if spec.operations:
                spec.missing_fields = [item for item in spec.missing_fields if "operation" not in str(item).casefold() and "标准化操作" not in str(item)]
                spec.ambiguities = [item for item in spec.ambiguities if not any(cue in str(item) for cue in ("标准化规则", "标准化操作", "标准化方式"))]
            if any(word in relevant for word in ("新列", "新增列", "保留原列", "不要覆盖", "不覆盖", "不改原列", "不修改原列", "保持原列")):
                spec.overwrite = False
            elif any(word in relevant for word in ("原列", "这一列", "这列", "全部改为")):
                spec.overwrite = True

        for spec in specs:
            if isinstance(spec, ConversionTaskSpec) and spec.direction == ConversionDirection.XLSX_TO_CSV and profile is not None:
                mentioned_sheets = [s.name for s in profile.sheets if s.name in user_source_text]
                if len(mentioned_sheets) == 1:
                    spec.source_sheet = mentioned_sheets[0]
                elif len(profile.sheets) == 1:
                    spec.source_sheet = profile.sheets[0].name
                elif not mentioned_sheets:
                    # An LLM-selected first sheet is not authorization to discard the others.
                    spec.source_sheet = ""
                if spec.source_sheet:
                    spec.missing_fields = [v for v in spec.missing_fields if v != "source_sheet"]

        # Presentation-only names are deterministic unless the user explicitly
        # supplied one. This prevents model wording variance from changing files.
        explicit_result_column = bool(re.search(r"结果列(?:叫|名为)", normalized))
        explicit_new_column = bool(re.search(r"(?:新列|新增列)(?:叫|名为)", normalized))
        explicit_impute_names = bool(re.search(r"新列(?:输出)?[^，。；;]+和[^，。；;]+", normalized))
        for spec in specs:
            if isinstance(spec, AggregateTaskSpec) and not explicit_result_column:
                for metric in spec.metrics:
                    metric.output_column = f"{metric.column}_{metric.function.value}"
            if not isinstance(spec, TransformTaskSpec):
                continue
            if spec.operation == TransformOperation.MAP_VALUES and not explicit_new_column:
                for rule in spec.map_rules:
                    rule.output_column = ""
                    rule.status_column = ""
            if spec.operation == TransformOperation.IMPUTE_NUMERIC and not explicit_impute_names:
                for rule in spec.impute_rules:
                    rule.output_column = ""
                    rule.indicator_column = ""

        # Column arithmetic is an explicit business rule, not an invented formula.
        from .formula_intent import column_expression, column_formula
        inference_formula = any(cue in compact for cue in ("填补公式空白", "公式空白", "公式断层", "修复公式", "异常公式", "公式异常"))
        formula_match = re.search(r"(?P<formula>=[^\s，；;。]+)", user_source_text)
        for spec in specs:
            if not isinstance(spec, FormulaTaskSpec) or inference_formula:
                continue
            derived = column_formula(profile, spec.sheet, spec.column, user_source_text, spec.start_row) if profile is not None else None
            explicit = spec.formula if spec.formula and spec.formula in user_source_text else (formula_match.group("formula") if formula_match else None)
            expression = column_expression(spec.column, user_source_text)
            # Unresolved column arithmetic must not masquerade as an Excel named range.
            if expression and not derived and re.search(r"[\u4e00-\u9fff]", expression) and not any(c in expression for c in ('!', '"')):
                explicit = None
            if derived or explicit or expression or any(cue in compact for cue in ("填公式", "设置公式", "写入公式", "填入公式")):
                spec.mode = FormulaMode.SET_FORMULA
                spec.formula = derived or explicit
                if spec.formula:
                    spec.missing_fields = [v for v in spec.missing_fields if "formula" not in str(v).casefold() and "公式" not in str(v)]
                    spec.ambiguities = [v for v in spec.ambiguities if "公式" not in str(v)]
                elif "formula" not in spec.missing_fields:
                    spec.missing_fields.append("formula")
                if profile is not None:
                    sheet = next((s for s in profile.sheets if s.name == spec.sheet), None)
                    if sheet:
                        spec.start_row = spec.start_row or (sheet.header_row or 1) + 1
                        spec.end_row = spec.end_row or sheet.max_row
                spec.overwrite_existing = bool(re.search(r"(?<!不)(?<!不要)(?<!禁止)(?<!不可)(?<!不允许)(?:覆盖|重算已有结果|重算现有结果)", compact))

        dimension_cues = {
            DimensionResolutionAction.FIRST_CANDIDATE: ("取前", "第一组", "前者"),
            DimensionResolutionAction.LAST_CANDIDATE: ("取后", "最后一组", "后者"),
            DimensionResolutionAction.CONVERT_EACH_TO_MM: ("换算为mm", "统一为mm", "转为mm"),
            DimensionResolutionAction.ZERO: ("填0", "为0", "统一为0", "置零"),
        }
        map_specs: list[TransformTaskSpec] = []
        require_partial_context = (
            any(cue in compact for cue in ("1个或两个数", "一个或两个数", "1个或2个数", "一个或二个数"))
            and "文字" in compact
            and any(cue in compact for cue in ("错误", "无效", "不合法"))
        )
        for spec in specs:
            if isinstance(spec, FormulaTaskSpec):
                spec.recalculate = not any(cue in compact for cue in ("不重算", "仅检查公式文本", "只检查公式文本"))
            if isinstance(spec, FormulaTaskSpec) and spec.mode in {
                FormulaMode.DIAGNOSE,
                FormulaMode.FILL_GAPS,
                FormulaMode.REPLACE_OUTLIERS,
            }:
                # These modes deterministically inspect neighbouring formulas.
                # "No formula text supplied" is not an unresolved user choice.
                spec.ambiguities = [
                    item for item in spec.ambiguities
                    if not (
                        any(token in str(item) for token in ("未提供", "没有提供"))
                        and "公式" in str(item)
                    )
                ]
            if isinstance(spec, FormulaTaskSpec) and any(
                cue in compact for cue in ("不重算", "仅检查公式文本", "只检查公式文本")
            ):
                spec.recalculate = False
            if not isinstance(spec, TransformTaskSpec):
                continue
            if spec.operation == TransformOperation.IMPUTE_NUMERIC:
                for rule in spec.impute_rules:
                    if not any(
                        cue in compact
                        for cue in ("覆盖现有输出", "覆盖已有输出", "重算已有结果", "重算现有结果")
                    ):
                        rule.overwrite_existing_outputs = False
                    source_name = unicodedata.normalize(
                        "NFKC", rule.source.display_header or ""
                    ).strip().casefold()
                    relevant = [
                        segment for segment in compact_segments
                        if source_name and source_name in segment
                    ]
                    if len(spec.impute_rules) == 1 and not relevant:
                        relevant = compact_segments
                    rule.invalid_policy_confirmed = any(
                        any(token in segment for token in ("错误值", "异常值", "无法解析", "非数值"))
                        and any(token in segment for token in (
                            "错误值也填", "异常值也填", "保持留空", "错误值留空", "异常值留空",
                        ))
                        for segment in relevant
                    )
            elif spec.operation == TransformOperation.PARSE_DIMENSIONS and spec.dimension_rule:
                if require_partial_context:
                    spec.dimension_rule.require_context_for_partial = True
                policies = dict(spec.dimension_rule.resolution_policies)
                for status, action in list(policies.items()):
                    if action == DimensionResolutionAction.LEAVE_BLANK:
                        continue
                    if not any(cue in compact for cue in dimension_cues[action]):
                        policies.pop(status, None)
                spec.dimension_rule.resolution_policies = policies
            elif spec.operation == TransformOperation.MAP_VALUES:
                map_specs.append(spec)
                for rule in spec.map_rules:
                    normalized_mapping = {
                        unicodedata.normalize("NFKC", str(key)).strip().casefold(): value
                        for key, value in rule.mapping.items()
                    }
                    keys = {
                        unicodedata.normalize("NFKC", str(key)).strip().casefold()
                        for key in rule.mapping
                    }
                    if keys and keys <= {"有", "无", "是", "否"}:
                        # "未记录算无" authorizes blanks, not arbitrary unknown text.
                        rule.unmapped_value = None
                    elif keys:
                        # A blank category is not the same as an explicitly "other" category.
                        rule.blank_value = None

        # Bind names only when the wording and destination fields have a unique
        # correspondence. This repairs model-added prefixes without guessing.
        impute_specs = [
            spec for spec in specs
            if isinstance(spec, TransformTaskSpec)
            and spec.operation == TransformOperation.IMPUTE_NUMERIC
        ]
        if len(impute_specs) == 1 and len(impute_specs[0].impute_rules) == 1:
            name_matches = list(re.finditer(
                r"新列(?:输出)?\s*[‘'\"“]?"
                r"(?P<output>[^\s，,；;。和‘’'\"“”]+)[’'\"”]?\s*和\s*"
                r"[‘'\"“]?(?P<indicator>[^\s，,；;。‘’'\"“”]+)[’'\"”]?",
                user_source_text,
            ))
            if len(name_matches) == 1:
                match = name_matches[0]
                rule = impute_specs[0].impute_rules[0]
                rule.output_column = match.group("output")
                rule.indicator_column = match.group("indicator")

        output_specs = [spec for spec in specs if hasattr(spec, "output_sheet")]
        named_sheets: list[str] = []
        for pattern in (
            r"(?:写入|输出到|筛到)?新(?:工作)?表(?:叫|名为)?\s*[‘'\"“]?"
            r"(?P<name>[^\s，,；;。的‘’'\"“”]+)",
            r"新建\s*[‘'\"“]?(?P<name>[^\s，,；;。‘’'\"“”]+?)表(?=[，,；;。]|$)",
            r"输出\s*[‘'\"“]?(?P<name>[^\s，,；;。‘’'\"“”]+?)表(?=[，,；;。]|$)",
        ):
            named_sheets.extend(
                match.group("name").strip(" \t'\"“”")
                for match in re.finditer(pattern, user_source_text)
            )
        named_sheets = list(dict.fromkeys(named_sheets))
        if len(output_specs) == 1 and len(named_sheets) == 1:
            output_specs[0].output_sheet = named_sheets[0]

        # Repair common structured-output defects in explicit one-hot requests.
        # A model may collapse several requested indicator columns into one combined
        # flag, or emit one combined rule alongside some correctly separated rules.
        if "indicator" in compact and any(token in compact for token in ("分别", "各", "每个")):
            for spec in map_specs:
                groups: dict[tuple[str, str, int], list[Any]] = {}
                for rule in spec.map_rules:
                    ref = rule.source
                    key = (ref.sheet, ref.column_letter or ref.display_header, ref.column_index or 0)
                    groups.setdefault(key, []).append(rule)
                for rules in groups.values():
                    single_positives = {
                        key
                        for rule in rules
                        for key, value in rule.mapping.items()
                        if value in (1, True)
                        and sum(item in (1, True) for item in rule.mapping.values()) == 1
                    }
                    for rule in rules:
                        positives = {
                            key for key, value in rule.mapping.items() if value in (1, True)
                        }
                        if len(positives) > 1:
                            remaining = positives - single_positives
                            if len(rules) > 1 and len(remaining) == 1:
                                rule.mapping = {next(iter(remaining)): 1}
                expanded_rules = []
                for rule in spec.map_rules:
                    positives = [
                        key for key, value in rule.mapping.items() if value in (1, True)
                    ]
                    if len(positives) <= 1 or rule.unmapped_value not in (0, False):
                        expanded_rules.append(rule)
                        continue
                    for category in positives:
                        separated = rule.model_copy(deep=True)
                        separated.mapping = {category: 1}
                        separated.output_column = f"{category}_Indicator"
                        separated.status_column = ""
                        separated.blank_value = None
                        expanded_rules.append(separated)
                spec.map_rules = expanded_rules

        if isinstance(prepared.spec, CompositeTaskSpec):
            if prepared.spec.confidence < 0.6 and specs and all(
                item.confidence >= 0.6 for item in specs
            ):
                prepared.spec.confidence = min(item.confidence for item in specs)
        return prepared

    def parse_task(self, task_id: str, request: str) -> TaskEnvelope:
        profile = self.load_profile(task_id)
        envelope = self.configured_llm().parse_task(
            request, profile, capability_manifest=self.skills.manifest()
        )
        envelope = self.prepare_model_envelope(envelope, request, profile)
        self._policy_sources[task_id] = request
        self.save_task_envelope(task_id, envelope)
        self.repository.log_event(
            task_id,
            "task_intent_parsed",
            {
                "task_type": envelope.spec.task_type,
                "confidence": envelope.spec.confidence,
            },
        )
        return envelope

    def enqueue_message(self, task_id: str, text: str) -> dict[str, Any]:
        """Persist a user turn immediately and plan it in the background."""
        request = text.strip()
        if not request:
            raise IntentValidationError("请输入你希望怎样处理这个工作簿。")
        row = self._task(task_id)
        compact_request = re.sub(r"\s+", "", request.casefold())
        if row.get("status") == "needs_clarification" and any(
            token in compact_request
            for token in ("缺少什么", "需要什么信息", "具体缺什么", "要确认什么")
        ):
            self._add_message(
                task_id, ConversationRole.USER, MessageKind.TEXT, request
            )
            clarification = next(
                (
                    item for item in reversed(self.get_conversation(task_id))
                    if item.kind == MessageKind.CLARIFICATION
                ),
                None,
            )
            if clarification:
                content = clarification.content
                if content.startswith("我还缺少一个"):
                    errors = clarification.metadata.get("validation_errors", [])
                    envelope_data = clarification.metadata.get("envelope")
                    if envelope_data:
                        prompt = clarification_for(
                            TaskValidation(valid=False, errors=list(errors)),
                            TaskEnvelope.model_validate(envelope_data),
                            self.load_profile(task_id),
                        )
                        content = prompt.question
                        choices = prompt.choices
                    else:
                        choices = clarification.metadata.get("choices", [])
                else:
                    choices = clarification.metadata.get("choices", [])
                self._add_message(
                    task_id,
                    ConversationRole.ASSISTANT,
                    MessageKind.CLARIFICATION,
                    f"当前需要你确认的是：{content}",
                    {
                        **clarification.metadata,
                        "choices": choices,
                        "repeated_clarification": True,
                    },
                )
            return {"job_id": f"clarification-help:{uuid4().hex}", "status": "waiting_for_user"}
        op_id = self.p0.begin(task_id, continuation=row["status"] in {"needs_clarification", "awaiting_approval", "awaiting_data_authorization"})
        active_run = self.repository.latest_plan_run(task_id)
        waiting_for_execution = bool(
            active_run and active_run.get("status") == "running"
        ) or row.get("status") in {"approved", "executing"}
        if active_run and active_run.get("status") == "running":
            self.repository.request_run_control(active_run["run_id"], "cancel")
            self.repository.add_activity_event(
                task_id, "execution", active_run["run_id"], "cancel_requested",
                "已收到新要求，正在完成当前安全步骤后停止", "running",
            )
        run = self.repository.create_planning_run(
            task_id,
            ConversationRole.USER.value,
            MessageKind.TEXT.value,
            request,
            waiting_for_execution=waiting_for_execution,
        )
        self.p0.link(task_id, op_id, "planning", run["job_id"], run["request_message_id"])
        future = self._background_executor.submit(
            self._run_planning_job, run["job_id"], run["fencing_token"]
        )
        self._background_futures[run["job_id"]] = future
        return run

    def cancel_planning(self, task_id: str) -> None:
        run = self.repository.latest_planning_run(task_id)
        if not run or run["status"] not in {"queued", "waiting_for_execution", "running"}:
            raise IntentValidationError("当前没有正在生成的回复。")
        self.repository.update_planning_run(
            run["job_id"], status="cancelled", stage="cancelled", cancel_requested=1
        )
        self.repository.add_activity_event(
            task_id, "planning", run["job_id"], "cancelled",
            "已停止当前规划，旧结果不会被采用", "cancelled",
        )
        task = self._task(task_id)
        self.repository.update_task(
            task_id,
            status=(
                "needs_conversion"
                if Path(task["original_path"]).suffix.casefold() == ".csv"
                and int(task["current_version"]) == 0
                else "profiled"
            ),
        )

    def select_custom_clarification_input(self, task_id: str, label: str) -> None:
        """Record an input-mode choice without launching a redundant planning pass."""
        task = self._task(task_id)
        if task["status"] != "needs_clarification":
            raise IntentValidationError("当前没有等待补充的计划问题。")
        self._add_message(task_id, ConversationRole.USER, MessageKind.TEXT, label)
        self.repository.add_activity_event(
            task_id, "planning", f"clarification:{task_id}", "awaiting_rule",
            "等待你输入具体修正规则", "waiting",
        )

    @staticmethod
    def _model_conversation(history: list[ConversationMessage]) -> list[dict[str, Any]]:
        return [{"role": "user", "content": item.content}
                for item in history if item.role == ConversationRole.USER][-8:]

    @staticmethod
    def _is_column_projection_request(text: str) -> bool:
        compact = re.sub(r"\s+", "", unicodedata.normalize("NFKC", text).casefold())
        return (
            any(token in compact for token in ("只保留", "仅保留"))
            or (
                (
                    any(token in compact for token in ("其他列", "其余列", "未提及的列", "剩余列"))
                    or ("其他" in compact and "列" in compact)
                    or ("其余" in compact and "列" in compact)
                )
                and any(token in compact for token in ("删除", "删掉", "删去", "移除"))
            )
        )

    def _build_followup_column_projection(
        self,
        previous: TaskEnvelope | None,
        request: str,
        history: list[ConversationMessage],
        profile: WorkbookProfile,
    ) -> TaskEnvelope | None:
        """Resolve “keep what I mentioned/generated” from the last completed run.

        This reference cannot be reconstructed from the latest user sentence alone.
        The prior result contains the exact source lineage and written ranges, which
        are safer than asking the model to guess identifiers or regenerated columns.
        """
        current_request_is_projection = self._is_column_projection_request(request)
        continuing_unresolved_projection = bool(
            previous
            and isinstance(previous.spec, CleanTaskSpec)
            and previous.spec.operation != CleanOperation.KEEP_COLUMNS
            and previous.spec.missing_fields
            and self._is_column_projection_request(
                previous.goal + "\n" + "\n".join(previous.constraints)
            )
        )
        if not current_request_is_projection and not continuing_unresolved_projection:
            return None

        sheets = {sheet.name: sheet for sheet in profile.sheets}
        selected_sheet = next(
            (
                name for name in sheets
                if unicodedata.normalize("NFKC", name).casefold()
                in unicodedata.normalize("NFKC", request).casefold()
            ),
            "",
        )
        if (
            not selected_sheet
            and previous
            and isinstance(previous.spec, CleanTaskSpec)
            and previous.spec.sheet in sheets
        ):
            selected_sheet = previous.spec.sheet
        if not selected_sheet and len(sheets) == 1:
            selected_sheet = next(iter(sheets))
        if not selected_sheet:
            return None

        result_message = next(
            (item for item in reversed(history) if item.kind == MessageKind.RESULT),
            None,
        )
        if not result_message:
            return None
        result = result_message.metadata.get("result") or {}
        steps = (result.get("details") or {}).get("steps") or []
        keep_indexes: set[int] = set()

        for step in steps:
            details = step.get("details") or {}
            for lineage in details.get("lineage") or []:
                for source in lineage.get("source_columns") or []:
                    try:
                        index = int(source.get("column_index"))
                    except (TypeError, ValueError):
                        continue
                    if index > 0:
                        keep_indexes.add(index)
            for address in step.get("modified_ranges") or []:
                try:
                    sheet_name, cells = str(address).rsplit("!", 1)
                    if sheet_name.strip("'") != selected_sheet:
                        continue
                    start, end = cells.split(":", 1)
                    start_match = re.match(r"\$?([A-Za-z]+)", start)
                    end_match = re.match(r"\$?([A-Za-z]+)", end)
                    if not start_match or not end_match:
                        continue
                    first = column_index_from_string(start_match.group(1).upper())
                    last = column_index_from_string(end_match.group(1).upper())
                    keep_indexes.update(range(min(first, last), max(first, last) + 1))
                except (ValueError, TypeError):
                    continue

        sheet = sheets[selected_sheet]
        columns = [
            column for column in sheet.columns if column.index in keep_indexes
        ]
        columns.sort(key=lambda item: item.index)
        if not columns:
            return None
        refs = [
            ColumnRef(
                sheet=selected_sheet,
                column_index=column.index,
                column_letter=column.letter or get_column_letter(column.index),
                display_header=column.display_header,
                header_occurrence=column.header_occurrence,
            )
            for column in columns
        ]
        return TaskEnvelope(
            goal=(
                f"在工作表“{selected_sheet}”中，仅保留上一轮用户明确涉及的来源列"
                "和上一轮 Agent 生成的结果/状态列，删除其他列"
            ),
            spec=CleanTaskSpec(
                sheet=selected_sheet,
                operation=CleanOperation.KEEP_COLUMNS,
                keep_columns=refs,
                overwrite=True,
                confidence=1.0,
            ),
            constraints=[
                "preserve_row_order",
                "preserve_retained_values_and_formats",
                "do_not_invent_identifier_columns",
            ],
        )

    @staticmethod
    def _is_indicator_status_cleanup_request(text: str) -> bool:
        compact = re.sub(r"\s+", "", unicodedata.normalize("NFKC", text).casefold())
        delete_requested = any(
            token in compact for token in ("删除", "删掉", "删去", "移除", "不要", "去掉")
        )
        mapping_status_named = (
            "indicator" in compact
            and any(token in compact for token in ("状态", "mapped", "mapped_default"))
        )
        return delete_requested and mapping_status_named

    def _build_indicator_status_cleanup(
        self,
        previous: TaskEnvelope | None,
        request: str,
        profile: WorkbookProfile,
    ) -> TaskEnvelope | None:
        """Build an exact projection that removes only mapping audit-status columns.

        Indicator value columns remain intact. Dimension, derived-metric and imputation
        status columns are deliberately excluded because the user referred specifically
        to the columns whose values are ``mapped``/``mapped_default``.
        """
        if not self._is_indicator_status_cleanup_request(request):
            return None
        sheets = {sheet.name: sheet for sheet in profile.sheets}
        folded_request = unicodedata.normalize("NFKC", request).casefold()
        selected_sheet = next(
            (
                name for name in sheets
                if unicodedata.normalize("NFKC", name).casefold() in folded_request
            ),
            "",
        )
        if not selected_sheet and previous:
            candidate = getattr(previous.spec, "sheet", "")
            if candidate in sheets:
                selected_sheet = candidate
        if not selected_sheet and len(sheets) == 1:
            selected_sheet = next(iter(sheets))
        if not selected_sheet:
            return None

        sheet = sheets[selected_sheet]
        status_values = {"mapped", "mapped_default", "blank", "unknown"}
        preview_values: dict[str, set[str]] = {}
        for row in sheet.preview_rows:
            for header, value in row.items():
                if value is None or str(value).strip() == "":
                    continue
                preview_values.setdefault(str(header), set()).add(
                    unicodedata.normalize("NFKC", str(value)).strip().casefold()
                )

        def is_mapping_status(column) -> bool:
            header = column.display_header or column.name
            folded = unicodedata.normalize("NFKC", header).strip().casefold()
            if "indicator" not in folded:
                return False
            values = preview_values.get(header, set()) | preview_values.get(column.name, set())
            status_header = "状态" in folded or folded.endswith("_status")
            status_content = bool(values) and values <= status_values and bool(
                values & {"mapped", "mapped_default"}
            )
            return status_header or status_content

        removed = [column for column in sheet.columns if is_mapping_status(column)]
        if not removed:
            return None
        kept = [column for column in sheet.columns if column not in removed]
        refs = [
            ColumnRef(
                sheet=selected_sheet,
                column_index=column.index,
                column_letter=column.letter or get_column_letter(column.index),
                display_header=column.display_header,
                header_occurrence=column.header_occurrence,
            )
            for column in kept
        ]
        removed_names = [column.display_header or column.name for column in removed]
        return TaskEnvelope(
            goal=(
                f"在工作表“{selected_sheet}”中删除 {len(removed_names)} 个 Indicator 映射状态列"
                f"（{'、'.join(removed_names)}），保留 Indicator 数值列和其他列"
            ),
            spec=CleanTaskSpec(
                sheet=selected_sheet,
                operation=CleanOperation.KEEP_COLUMNS,
                keep_columns=refs,
                overwrite=True,
                confidence=1.0,
            ),
            constraints=[
                "preserve_row_order",
                "preserve_retained_values_and_formats",
                "remove_only_indicator_mapping_status_columns",
            ],
        )

    def _resume_clarification_envelope(
        self,
        previous: TaskEnvelope | None,
        request: str,
        history: list[ConversationMessage],
        profile: WorkbookProfile | None = None,
    ) -> TaskEnvelope | None:
        """Deterministically merge common clarification answers into the draft."""
        if not previous:
            return None
        draft = previous.model_copy(deep=True)
        specs = (
            list(draft.spec.operations)
            if isinstance(draft.spec, CompositeTaskSpec)
            else [draft.spec]
        )
        clarification = next(
            (item for item in reversed(history[:-1]) if item.kind == MessageKind.CLARIFICATION),
            None,
        )
        if not clarification:
            return None
        review = clarification.metadata.get("review") or {}
        folded = request.casefold().replace(" ", "")

        def clear_answered_issue(spec: Any, code: str, cues: tuple[str, ...]) -> None:
            """Remove only the question answered by this turn from a draft spec."""
            markers = (code.casefold(), *(cue.casefold() for cue in cues))
            spec.missing_fields = [
                item for item in spec.missing_fields
                if not any(marker in str(item).casefold() for marker in markers)
            ]
            spec.ambiguities = [
                item for item in spec.ambiguities
                if not any(marker in str(item).casefold() for marker in markers)
            ]

        field = clarification.metadata.get("field")
        issue_code = clarification.metadata.get("issue_code") or review.get("issue_code")

        if field == "source_sheet":
            conversions = [item for item in specs if isinstance(item, ConversionTaskSpec)]
            answer = request.strip(" ‘’'\"“”")
            if len(conversions) == 1 and profile is not None and answer in {sheet.name for sheet in profile.sheets}:
                conversions[0].source_sheet = answer
                clear_answered_issue(conversions[0], "source_sheet", ("来源工作表",))
                return draft

        if field == "column" and issue_code == "ambiguous_formula_column" and profile is not None:
            formulas = [item for item in specs if isinstance(item, FormulaTaskSpec)]
            match = re.match(r"\s*([A-Za-z]{1,3})(?:\s*列.*)?\s*$", request)
            if len(formulas) == 1 and match:
                from .formula_intent import column_formula
                target = formulas[0]
                target.column = match.group(1).upper()
                source = "\n".join(item.content for item in history if item.role == ConversationRole.USER)
                target.formula = column_formula(profile, target.sheet, target.column, source, target.start_row)
                if target.formula:
                    target.mode = FormulaMode.SET_FORMULA
                    clear_answered_issue(target, "formula", ("公式",))
                    clear_answered_issue(target, "column", ("目标列", "列歧义", "同名表头"))
                    return draft

        if field == "formula" or issue_code == "missing_formula":
            match = re.search(r"=[^，；;。\n]+", request)
            formula_specs = [item for item in specs if isinstance(item, FormulaTaskSpec)]
            if len(formula_specs) != 1:
                return None
            formula_spec = formula_specs[0]
            from .formula_intent import column_formula
            derived = column_formula(
                profile, formula_spec.sheet, formula_spec.column,
                f"在“{formula_spec.column}”列填入{request}", formula_spec.start_row,
            ) if profile is not None else None
            if not derived and not match:
                return None
            formula_spec.mode = FormulaMode.SET_FORMULA
            formula_spec.formula = derived or match.group(0).strip()
            clear_answered_issue(formula_spec, "formula", ("公式",))
            if isinstance(draft.spec, CompositeTaskSpec):
                clear_answered_issue(draft.spec, "formula", ("公式",))
            return draft

        if field == "operations" and any(isinstance(item, NormalizeTaskSpec) for item in specs):
            operations = []
            folded_answer = unicodedata.normalize("NFKC", request).casefold().replace(" ", "")
            for operation, cues in (
                ("trim", ("去空格", "首尾空格", "trim")),
                ("nfkc", ("nfkc", "全角", "半角")),
                ("lower", ("小写",)),
                ("upper", ("大写",)),
                ("to_text", ("转成文本", "转为文本")),
                ("to_number", ("转成数值", "转为数值", "转成数字", "转为数字")),
            ):
                if any(cue in folded_answer for cue in cues):
                    operations.append(NormalizeOperation(operation))
            if not operations:
                return None
            target = next(item for item in specs if isinstance(item, NormalizeTaskSpec))
            target.operations = operations
            clear_answered_issue(target, "operations", ("标准化操作",))
            return draft

        if field == "column" and any(isinstance(item, NormalizeTaskSpec) for item in specs):
            match = re.search(r"(?:处理|选择|使用)?[‘'\"“]?([^，,；;。‘’'\"“”]+?)[’'\"”]?列", request)
            target = next(item for item in specs if isinstance(item, NormalizeTaskSpec))
            answer = match.group(1).strip() if match else request.strip(" ‘’'\"“”")
            if not answer or len(answer) > 100 or any(c in answer for c in "，；。\n"):
                return None
            target.column = answer
            clear_answered_issue(target, "column", ("目标列", "具体目标列", "列歧义", "列无法确定"))
            return draft

        if field == "rank_column" and any(isinstance(item, SortRankTaskSpec) for item in specs):
            match = re.search(r"按\s*([^，,；;。]+?)\s*(升序|降序)", request)
            if not match:
                return None
            target = next(item for item in specs if isinstance(item, SortRankTaskSpec))
            column = match.group(1).strip(" ‘'\"“”")
            direction = SortDirection.DESC if match.group(2) == "降序" else SortDirection.ASC
            target.sort_keys = [SortKey(column=column, direction=direction)]
            if any(cue in draft.goal for cue in ("排名", "排个名", "名次")):
                target.rank_column = column
                target.rank_direction = direction
            clear_answered_issue(target, "sort_keys", ("排序列", "排序方向"))
            clear_answered_issue(target, "rank_column", ("排名列", "按哪一列排名"))
            return draft

        if issue_code == "invalid_numeric":
            impute_specs = [
                spec for spec in specs
                if isinstance(spec, TransformTaskSpec)
                and spec.operation == TransformOperation.IMPUTE_NUMERIC
            ]
            if not impute_specs:
                return None
            applied = False
            if "留空" in folded:
                for item in impute_specs:
                    for impute_rule in item.impute_rules:
                        impute_rule.invalid_policy = InvalidNumericPolicy.LEAVE_BLANK
                        impute_rule.invalid_policy_confirmed = True
                applied = True
            elif any(token in folded for token in ("均值填", "平均值填", "用同一均值")):
                for item in impute_specs:
                    for impute_rule in item.impute_rules:
                        impute_rule.invalid_policy = InvalidNumericPolicy.IMPUTE
                        impute_rule.invalid_policy_confirmed = True
                applied = True
            if not applied:
                return None
            for item in impute_specs:
                clear_answered_issue(item, "invalid_numeric", ("错误值", "异常值", "非数值"))
            if isinstance(draft.spec, CompositeTaskSpec):
                clear_answered_issue(
                    draft.spec, "invalid_numeric", ("错误值", "异常值", "非数值")
                )
            return draft

        candidates = [
            spec for spec in specs
            if isinstance(spec, TransformTaskSpec)
            and spec.operation == TransformOperation.PARSE_DIMENSIONS
            and spec.dimension_rule is not None
        ]
        if not candidates:
            return None
        allowed_rows = {
            int(item["row"]): item
            for item in review.get("rows", [])
            if item.get("row") is not None
        }
        spec = candidates[0]
        rule = spec.dimension_rule
        assert rule is not None
        applied = False
        action_map = {
            "multiple_candidates": (
                ("第一组", DimensionResolutionAction.FIRST_CANDIDATE),
                ("最后一组", DimensionResolutionAction.LAST_CANDIDATE),
                ("填0", DimensionResolutionAction.ZERO),
                ("留空", DimensionResolutionAction.LEAVE_BLANK),
            ),
            "unit_conflict": (
                ("换算为mm", DimensionResolutionAction.CONVERT_EACH_TO_MM),
                ("填0", DimensionResolutionAction.ZERO),
                ("留空", DimensionResolutionAction.LEAVE_BLANK),
            ),
            "invalid_format": (
                ("填0", DimensionResolutionAction.ZERO),
                ("留空", DimensionResolutionAction.LEAVE_BLANK),
            ),
            "no_numeric_measurement": (
                ("填0", DimensionResolutionAction.ZERO),
                ("留空", DimensionResolutionAction.LEAVE_BLANK),
            ),
        }
        for token, action in action_map.get(str(issue_code), ()):
            if token in folded:
                rule.resolution_policies[str(issue_code)] = action
                applied = True
                break
        correction_pattern = re.compile(
            r"第\s*(\d+)\s*行(?:数据)?\s*(?:改为|修改为|设为)\s*[:：]?\s*"
            r"(.+?)(?=\s*第\s*\d+\s*行|$)",
            re.IGNORECASE,
        )
        for match in correction_pattern.finditer(request):
            row = int(match.group(1))
            replacement = match.group(2).strip(" \t\r\n,，;；。")
            if row in allowed_rows and replacement:
                rule.row_corrections[row] = replacement
                applied = True
        if not applied:
            return None
        issue_cues = {
            "multiple_candidates": ("多组候选", "多个候选"),
            "unit_conflict": ("单位冲突", "混合单位"),
            "invalid_format": ("格式错误", "无效格式"),
            "no_numeric_measurement": ("没有数值", "无数值"),
        }.get(str(issue_code), ())
        clear_answered_issue(spec, str(issue_code), issue_cues)
        if isinstance(draft.spec, CompositeTaskSpec):
            clear_answered_issue(draft.spec, str(issue_code), issue_cues)
        return draft

    def _planning_event(
        self, task_id: str, job_id: str, stage: str, label: str, status: str = "running",
        metadata: dict[str, Any] | None = None,
    ) -> None:
        self.repository.update_planning_run(job_id, stage=stage)
        self.repository.add_activity_event(
            task_id, "planning", job_id, stage, label, status, metadata
        )

    @staticmethod
    def _clarification_review(plan: TaskPlan, validation: TaskValidation) -> dict[str, Any]:
        errors = "；".join(validation.errors)
        if any(token in errors for token in (
            "体积公式未在用户原始要求中明确选择",
            "派生指标必须明确选择注册公式模板",
        )):
            return {}
        issue_code = None
        for prefix in (
            "DERIVE_RESOLUTION_REQUIRED:",
            "DIMENSION_RESOLUTION_REQUIRED:",
            "IMPUTE_RESOLUTION_REQUIRED:",
        ):
            marker = next((item for item in validation.errors if item.startswith(prefix)), None)
            if marker:
                issue_code = marker.split(":", 2)[1]
                break
        if not issue_code:
            return {}
        rows: list[dict[str, Any]] = []
        total = 0
        truncated = False
        for step in plan.steps:
            rows.extend(
                item for item in step.expected_changes.review_rows
                if item.get("issue_code") == issue_code
            )
            total += int(step.expected_changes.review_totals.get(issue_code, 0))
            truncated = truncated or bool(
                step.expected_changes.review_truncated.get(issue_code, False)
            )
        rows = sorted(rows, key=lambda item: (item.get("sheet", ""), int(item.get("row", 0))))[:50]
        return {
            "issue_code": issue_code,
            "rows": rows,
            "total": total,
            "truncated": truncated or total > len(rows),
        }

    def _run_planning_job(self, job_id: str, fencing_token: str) -> None:
        run = self.repository.get_planning_run(job_id)
        if not run:
            return
        task_id = run["task_id"]
        op_id = self.p0.linked("planning", job_id)
        started = time.perf_counter()
        def expire():
            if not self.repository.planning_run_is_current(job_id, fencing_token):
                return
            self.repository.update_planning_run(job_id, status="running")
            published = self.repository.publish_planning_result(
                job_id=job_id, fencing_token=fencing_token,
                task_status="failed", spec_json=None, plan_json=None,
                message_role="assistant", message_kind="error",
                message_content="计划生成超过时间预算，已停止采用本次回复；文件未修改。可发送要求重新尝试。",
                message_metadata={"timeout": True}, run_status="failed", run_stage="timeout")
            if published and op_id:
                self.p0.state(op_id, "failed")
        timer = threading.Timer(self.settings.planning_budget_seconds, expire)
        timer.daemon = True
        timer.start()
        try:
            self._run_planning_job_body(job_id, fencing_token)
        finally:
            timer.cancel()
            authorized_model_rows.set([])
            if op_id:
                with self.repository.connect() as db:
                    precheck = db.execute("SELECT duration_ms FROM p0_measurements WHERE operation_id=? AND measurement_id=?", (op_id, f"precheck:{job_id}")).fetchone()
                self.p0.measure(task_id, op_id, f"planning:{job_id}", "planning", max(0, (time.perf_counter()-started)*1000-(precheck[0] if precheck else 0)))
                self.p0.touch(task_id)

    def _run_planning_job_body(self, job_id: str, fencing_token: str) -> None:
        run = self.repository.get_planning_run(job_id)
        if not run:
            return
        task_id = run["task_id"]
        try:
            if run["status"] == "waiting_for_execution":
                self._planning_event(
                    task_id, job_id, "waiting_for_execution",
                    "正在安全停止上一计划，随后按新要求重新规划",
                )
                while self.repository.planning_run_is_current(job_id, fencing_token):
                    active = self.repository.latest_plan_run(task_id)
                    if not active or active.get("status") != "running":
                        break
                    time.sleep(0.1)
            if not self.repository.planning_run_is_current(job_id, fencing_token):
                return
            self.repository.update_planning_run(job_id, status="running", stage="context")
            message = self.repository.get_message(int(run["request_message_id"]))
            if not message:
                raise IntentValidationError("找不到本次用户要求。")
            request = message["content"]
            refusal = capability_refusal(request)
            if refusal:
                content = refusal.reason
                if refusal.alternatives:
                    content += "\n\n可以改为：" + "；".join(refusal.alternatives)
                self.repository.publish_planning_result(
                    job_id=job_id, fencing_token=fencing_token,
                    task_status="profiled", spec_json=None, plan_json=None,
                    message_role=ConversationRole.ASSISTANT.value,
                    message_kind=MessageKind.REFUSAL.value,
                    message_content=content,
                    message_metadata={"refusal": refusal.model_dump(mode="json"), "activity_run_id": job_id},
                )
                op_id = self.p0.linked("planning", job_id)
                if op_id:
                    self.p0.state(op_id, "cancelled")
                return
            row = self._task(task_id)
            profile = self.load_profile(task_id)
            previous = None
            if row.get("spec_json"):
                try:
                    previous = TaskEnvelope.model_validate_json(row["spec_json"])
                except Exception:
                    previous = None
            history = self.get_conversation(task_id)
            previous_errors: list[str] = []
            for item in reversed(history[:-1]):
                if item.kind == MessageKind.CLARIFICATION:
                    previous_errors = list(item.metadata.get("validation_errors", []))
                    break
            conversation = self._model_conversation(history)
            self._planning_event(task_id, job_id, "context", "正在读取工作簿结构", "completed")
            hint = guess_task_type(request)
            skill_context = build_context(profile, request)
            if hint and hint.value != "composite":
                try:
                    skill_context = self.skills.get(hint).build_context(profile, request)
                except (KeyError, ValueError):
                    pass
            self._planning_event(task_id, job_id, "intent", "正在理解你的处理要求")
            intent_started = time.perf_counter()
            envelope = self._build_indicator_status_cleanup(previous, request, profile)
            if envelope is None:
                envelope = self._build_followup_column_projection(
                    previous, request, history, profile
                )
            if envelope is None:
                envelope = self._resume_clarification_envelope(previous, request, history, profile)
            if envelope is None:
                op_id = self.p0.linked("planning", job_id)
                rows = self.p0.review_rows(task_id, op_id) if op_id else []
                grant = self.p0.authorization(task_id, op_id, rows) if rows else None
                if grant and grant["status"] == "pending":
                    published = self.repository.publish_planning_result(
                        job_id=job_id, fencing_token=fencing_token,
                        task_status="awaiting_data_authorization",
                        spec_json=row.get("spec_json"), plan_json=row.get("plan_json"),
                        message_role="assistant", message_kind="clarification",
                        message_content="需要你确认是否向模型发送相关异常目标值。可在数据授权卡查看范围，或拒绝并手动指定规则。",
                        message_metadata={"authorization_request": grant["grant_id"]},
                        run_status="waiting_for_user", run_stage="data_authorization")
                    if published:
                        self.p0.state(op_id, "awaiting_data_authorization")
                    return
                authorized_model_rows.set(rows if grant and grant["status"] == "approved" else [])
                envelope = self.configured_llm().parse_task(
                    request,
                    profile,
                    previous_envelope=previous,
                    validation_errors=previous_errors,
                    conversation=conversation,
                    capability_manifest=self.skills.manifest(),
                    workbook_context=skill_context,
                )
                source_text = "\n".join(
                    item.content for item in history if item.role == ConversationRole.USER
                )
                envelope = self.prepare_model_envelope(
                    envelope, source_text or request, profile
                )
            if not self.repository.planning_run_is_current(job_id, fencing_token):
                return
            self._planning_event(
                task_id, job_id, "intent", "已识别处理目标", "completed",
                {
                    "task_type": envelope.spec.task_type.value,
                    "duration_ms": round((time.perf_counter() - intent_started) * 1000),
                },
            )
            self._planning_event(task_id, job_id, "preview", "正在预检受影响的列和异常记录")
            preview_started = time.perf_counter()
            source_text = "\n".join(
                item.content for item in history if item.role == ConversationRole.USER
            )
            previous_plan_payload = json.loads(row["plan_json"]) if row.get("plan_json") else None
            plan, validation = self.make_task_plan(
                task_id, envelope,
                policy_source_text=source_text or request,
                persist=False,
            )
            self._planning_event(
                task_id, job_id, "preview", "已完成数据预检", "completed",
                {
                    "step_count": len(plan.steps),
                    "valid": validation.valid,
                    "duration_ms": round((time.perf_counter() - preview_started) * 1000),
                },
            )
            op_id = self.p0.linked("planning", job_id)
            self.p0.measure(task_id, op_id, f"precheck:{job_id}", "precheck", (time.perf_counter()-preview_started)*1000)
            self._planning_event(task_id, job_id, "safety", "正在检查风险和写入边界")
            if not self.repository.planning_run_is_current(job_id, fencing_token):
                return
            if validation.valid:
                content = plan_text(envelope, plan, validation)
                metadata = {
                    "plan": plan.model_dump(mode="json"),
                    "envelope": envelope.model_dump(mode="json"),
                    "warnings": validation.warnings,
                    "previous_plan": previous_plan_payload if plan.revision > 1 else None,
                    "activity_run_id": job_id,
                }
                task_status = "awaiting_approval"
                kind = MessageKind.PLAN
            else:
                clarification = clarification_for(validation, envelope, profile)
                content = clarification.question
                metadata = {
                    "choices": clarification.choices,
                    "field": clarification.field,
                    "issue_code": clarification.issue_code,
                    "validation_errors": validation.errors,
                    "envelope": envelope.model_dump(mode="json"),
                    "review": self._clarification_review(plan, validation),
                    "activity_run_id": job_id,
                }
                task_status = "needs_clarification"
                kind = MessageKind.CLARIFICATION
            published = self.repository.publish_planning_result(
                job_id=job_id, fencing_token=fencing_token,
                task_status=task_status,
                spec_json=envelope.model_dump_json(),
                plan_json=plan.model_dump_json(),
                message_role=ConversationRole.ASSISTANT.value,
                message_kind=kind.value,
                message_content=content,
                message_metadata=metadata,
                run_status="completed" if validation.valid else "waiting_for_user",
                run_stage="completed" if validation.valid else "waiting_for_user",
            )
            if published:
                op_id = self.p0.linked("planning", job_id)
                if op_id:
                    self.p0.link(task_id, op_id, "plan", self.p0.plan_key(plan))
                    self.p0.state(op_id, task_status)
                self.repository.add_activity_event(
                    task_id, "planning", job_id,
                    "completed" if validation.valid else "waiting_for_user",
                    "计划已生成" if validation.valid else "需要你确认异常处理规则",
                    "completed" if validation.valid else "waiting",
                    {"valid": validation.valid},
                )
        except Exception as exc:
            LOGGER.exception(
                "Planning run failed (task_id=%s, job_id=%s, error_type=%s)",
                task_id, job_id, type(exc).__name__,
            )
            safe_message = (
                "模型没有返回有效的任务结构，请稍后重试；你的工作簿没有被修改。"
                if isinstance(exc, LLMError)
                else "构建执行计划时发生内部错误，工作簿没有被修改。"
            )
            published = self.repository.publish_planning_result(
                job_id=job_id, fencing_token=fencing_token,
                task_status="failed", spec_json=None, plan_json=None,
                message_role=ConversationRole.ASSISTANT.value,
                message_kind=MessageKind.ERROR.value,
                message_content=safe_message,
                message_metadata={"error_type": type(exc).__name__, "activity_run_id": job_id},
                error_json=json.dumps({"type": type(exc).__name__, "message": str(exc)}, ensure_ascii=False),
            )
            if published:
                op_id = self.p0.linked("planning", job_id)
                if op_id:
                    self.p0.state(op_id, "failed")
                self.repository.add_activity_event(
                    task_id, "planning", job_id, "failed", "计划生成失败", "failed",
                    {"error_type": type(exc).__name__},
                )

    def submit_message(self, task_id: str, text: str) -> AgentTurnResult:
        request = text.strip()
        if not request:
            raise IntentValidationError("请输入你希望怎样处理这个工作簿。")
        row = self._task(task_id)
        refusal = capability_refusal(request)
        if refusal:
            self._add_message(task_id, ConversationRole.USER, MessageKind.TEXT, request)
            content = refusal.reason
            if refusal.alternatives:
                content += "\n\n可以改为：" + "；".join(refusal.alternatives)
            reply = self._add_message(
                task_id,
                ConversationRole.ASSISTANT,
                MessageKind.REFUSAL,
                content,
                {"refusal": refusal.model_dump(mode="json")},
            )
            return AgentTurnResult(
                state=ConversationState.WORKBOOK_READY,
                reply=reply,
                refusal=refusal,
            )
        previous_plan_payload = None
        if row.get("plan_json"):
            try:
                previous_plan_payload = json.loads(row["plan_json"])
            except Exception:
                previous_plan_payload = None
        if request in {"停止处理", "取消", "取消任务"} and row["plan_json"]:
            self.cancel_current_plan(task_id)
            reply = self.get_conversation(task_id)[-1]
            return AgentTurnResult(state=ConversationState.WORKBOOK_READY, reply=reply)

        if row["status"] == "awaiting_approval":
            self.approve(task_id, "edit", {"source": "conversation", "message": request})
        elif row["status"] == "paused":
            active_run = self.repository.latest_plan_run(task_id)
            if active_run and active_run["status"] in {"paused", "failed"}:
                self.repository.update_plan_run(active_run["run_id"], status="superseded")
                self.files.remove_run_checkpoints(task_id, active_run["run_id"])
                self.repository.log_event(
                    task_id, "run_superseded",
                    {"run_id": active_run["run_id"], "reason": "plan_revision"},
                )
        self._add_message(task_id, ConversationRole.USER, MessageKind.TEXT, request)
        self.repository.update_task(task_id, status="analyzing_request", error_json=None)
        profile = self.load_profile(task_id)
        previous: TaskEnvelope | None = None
        if row.get("spec_json"):
            try:
                previous = TaskEnvelope.model_validate_json(row["spec_json"])
            except Exception:
                previous = None
        history = self.get_conversation(task_id)
        previous_errors: list[str] = []
        for message in reversed(history[:-1]):
            if message.kind == MessageKind.CLARIFICATION:
                previous_errors = list(message.metadata.get("validation_errors", []))
                break
        conversation = self._model_conversation(history)
        try:
            envelope = self._build_indicator_status_cleanup(previous, request, profile)
            if envelope is None:
                envelope = self._build_followup_column_projection(
                    previous, request, history, profile
                )
            if envelope is None:
                envelope = self._resume_clarification_envelope(previous, request, history, profile)
            if envelope is not None:
                source_text = "\n".join(
                    item.content for item in history if item.role == ConversationRole.USER
                )
                plan, validation = self.make_task_plan(
                    task_id, envelope, policy_source_text=source_text or request
                )
            else:
                from .graph import build_planning_workflow

                workflow = build_planning_workflow(self)
                graph_result = workflow.invoke(
                    {
                        "task_id": task_id,
                        "user_request": request,
                        "previous_spec": previous.model_dump(mode="json") if previous else None,
                        "validation_errors": previous_errors,
                        "conversation": conversation,
                    }
                )
                envelope = TaskEnvelope.model_validate(graph_result["spec"])
                plan = TaskPlan.model_validate(graph_result["plan"])
                validation = TaskValidation.model_validate(graph_result["validation"])
        except Exception as exc:
            LOGGER.exception(
                "Synchronous planning failed (task_id=%s, error_type=%s)",
                task_id, type(exc).__name__,
            )
            failure_message = (
                "模型没有返回有效的任务结构，请稍后重试；你的工作簿没有被修改。"
                if isinstance(exc, LLMError)
                else "我已经理解了你的要求，但构建执行计划时发生内部错误。无需改写要求，请重试或查看后台日志。"
            )
            self.repository.update_task(
                task_id,
                status="failed",
                error_json=json.dumps({"type": type(exc).__name__, "message": str(exc)}, ensure_ascii=False),
            )
            self.repository.log_event(
                task_id,
                "planning_failed",
                {"type": type(exc).__name__, "message": str(exc)},
            )
            reply = self._add_message(
                task_id,
                ConversationRole.ASSISTANT,
                MessageKind.ERROR,
                failure_message,
                {"error_type": type(exc).__name__},
            )
            return AgentTurnResult(state=ConversationState.WORKBOOK_READY, reply=reply)

        if not validation.valid:
            clarification = clarification_for(validation, envelope, profile)
            reply = self._add_message(
                task_id,
                ConversationRole.ASSISTANT,
                MessageKind.CLARIFICATION,
                clarification.question,
                {
                    "choices": clarification.choices,
                    "field": clarification.field,
                    "issue_code": clarification.issue_code,
                    "validation_errors": validation.errors,
                    "envelope": envelope.model_dump(mode="json"),
                    "review": self._clarification_review(plan, validation),
                },
            )
            return AgentTurnResult(
                state=ConversationState.NEEDS_CLARIFICATION,
                reply=reply,
                needs_clarification=True,
                envelope=envelope,
                clarification=clarification,
            )

        reply = self._add_message(
            task_id,
            ConversationRole.ASSISTANT,
            MessageKind.PLAN,
            plan_text(envelope, plan, validation),
            {
                "plan": plan.model_dump(mode="json"),
                "envelope": envelope.model_dump(mode="json"),
                "warnings": validation.warnings,
                "previous_plan": previous_plan_payload if plan.revision > 1 else None,
            },
        )
        return AgentTurnResult(
            state=ConversationState.PLAN_READY,
            reply=reply,
            envelope=envelope,
            plan=plan,
            actions=["approve", "edit", "cancel"],
            selection=(
                self.skills.select(envelope)[0]
                if envelope.spec.task_type.value != "composite"
                else None
            ),
        )

    def approve_current_plan(
        self, task_id: str, acknowledgements: list[str] | None = None
    ) -> ExecutionResult:
        envelope, plan = self._approve_plan_for_execution(task_id, acknowledgements)
        return self._execute_approved_plan(task_id, envelope, plan)

    def enqueue_current_plan_execution(
        self, task_id: str, acknowledgements: list[str] | None = None
    ) -> dict[str, Any]:
        """Approve immediately and execute without blocking the chat UI."""
        envelope, plan = self._approve_plan_for_execution(task_id, acknowledgements)
        activity_run_id = f"{plan.plan_id}:r{plan.revision}"
        self.repository.add_activity_event(
            task_id, "execution", activity_run_id, "queued",
            "已确认计划，正在准备安全工作副本", "completed",
            {"step_count": len(plan.steps)},
        )
        future = self._background_executor.submit(
            self._run_execution_job, task_id, envelope, plan, activity_run_id
        )
        self._background_futures[f"execution:{activity_run_id}"] = future
        return {"run_id": activity_run_id, "task_id": task_id, "status": "queued"}

    def _approve_plan_for_execution(
        self, task_id: str, acknowledgements: list[str] | None
    ) -> tuple[TaskEnvelope, TaskPlan]:
        row = self._task(task_id)
        if row["status"] != "awaiting_approval" or not row["spec_json"] or not row["plan_json"]:
            raise IntentValidationError("当前没有可以执行的已确认计划。")
        envelope = TaskEnvelope.model_validate_json(row["spec_json"])
        plan = TaskPlan.model_validate_json(row["plan_json"])
        self.approve(
            task_id,
            "approve",
            {
                "source": "conversation",
                "plan": plan.model_dump(mode="json"),
                "acknowledgements": acknowledgements or [],
            },
        )
        self._add_message(
            task_id,
            ConversationRole.ASSISTANT,
            MessageKind.PROGRESS,
            "计划已确认。正在处理工作副本，并在交付前重新打开文件完成核验。",
            {"state": ConversationState.EXECUTING.value},
        )
        return envelope, plan

    def _run_execution_job(
        self, task_id: str, envelope: TaskEnvelope, plan: TaskPlan, activity_run_id: str
    ) -> None:
        started = time.perf_counter()
        self.repository.add_activity_event(
            task_id, "execution", activity_run_id, "running",
            "正在执行已批准的处理步骤", "running",
        )
        try:
            result = self._execute_approved_plan(task_id, envelope, plan)
            self.repository.add_activity_event(
                task_id, "execution", activity_run_id, "completed",
                "处理完成，已核验并生成新版本", "completed",
                {
                    "version": result.version,
                    "rows_affected": result.result.rows_affected,
                    "duration_ms": round((time.perf_counter() - started) * 1000),
                },
            )
        except Exception as exc:
            active = self.repository.latest_plan_run(task_id)
            cancelled = bool(active and active.get("status") == "cancelled")
            self.repository.add_activity_event(
                task_id, "execution", activity_run_id,
                "cancelled" if cancelled else "failed",
                "已安全停止，上一正式版本保持不变" if cancelled else "执行未完成，上一正式版本保持不变",
                "cancelled" if cancelled else "failed",
                {
                    "error_type": type(exc).__name__,
                    "duration_ms": round((time.perf_counter() - started) * 1000),
                },
            )

    def _execute_approved_plan(
        self, task_id: str, envelope: TaskEnvelope, plan: TaskPlan
    ) -> ExecutionResult:
        try:
            from .graph import sqlite_execution_workflow

            with sqlite_execution_workflow(self) as workflow:
                graph_result = workflow.invoke(
                    {
                        "task_id": task_id,
                        "spec": envelope.model_dump(mode="json"),
                        "plan": plan.model_dump(mode="json"),
                        "status": "approved",
                    },
                    config={"configurable": {"thread_id": f"{task_id}:{plan.plan_id}:r{plan.revision}"}},
                )
            if graph_result.get("status") != "completed":
                raise IntentValidationError(
                    graph_result.get("error", {}).get("message", "执行图未能完成任务。")
                )
            payload = graph_result["result"]
            output = Path(payload["output_path"])
            report = VerificationReport.model_validate(payload["verification"])
            result = ToolResult.model_validate(payload["tool_result"])
        except Exception as exc:
            active = self.repository.latest_plan_run(task_id)
            cancelled = bool(active and active.get("status") == "cancelled")
            self._add_message(
                task_id,
                ConversationRole.ASSISTANT,
                MessageKind.TEXT if cancelled else MessageKind.ERROR,
                (
                    "已在安全步骤边界停止旧处理。上一正式版本保持不变；现在会按你的新要求继续规划。"
                    if cancelled
                    else "执行没有产生可交付的新版本。原文件和上一版本保持不变，请检查要求后重试。"
                ),
                {"error_type": type(exc).__name__, "cancelled": cancelled},
            )
            raise
        task = self._task(task_id)
        version = int(task["current_version"])
        preview_sheets = [*result.modified_sheets, *result.created_sheets]
        reply = self._add_message(
            task_id,
            ConversationRole.ASSISTANT,
            MessageKind.RESULT,
            result_text(version, result, report),
            {
                "version": version,
                "result": result.model_dump(mode="json"),
                "verification": report.model_dump(mode="json"),
                "preview_sheets": list(dict.fromkeys(preview_sheets)),
                "activity_run_id": f"{plan.plan_id}:r{plan.revision}",
            },
        )
        return ExecutionResult(
            state=ConversationState.RESULT_READY,
            version=version,
            output_path=str(output),
            report=report,
            result=result,
            reply=reply,
        )

    def cancel_current_plan(self, task_id: str) -> None:
        row = self._task(task_id)
        op = self.p0.current(task_id)
        if op:
            self.p0.state(op["operation_id"], "cancelled")
        self.p0.touch(task_id)
        if row["plan_json"]:
            self.approve(task_id, "reject", {"source": "conversation"})
        self.repository.update_task(
            task_id,
            status=(
                "needs_conversion"
                if Path(row["original_path"]).suffix.casefold() == ".csv"
                and int(row["current_version"]) == 0
                else "profiled"
            ),
            spec_json=None,
            plan_json=None,
            error_json=None,
        )
        self._add_message(
            task_id,
            ConversationRole.ASSISTANT,
            MessageKind.TEXT,
            "已取消这份计划，没有修改文件。你可以直接告诉我新的处理要求。",
        )

    def get_conversation(self, task_id: str) -> list[ConversationMessage]:
        self._task(task_id)
        messages: list[ConversationMessage] = []
        for row in self.repository.list_messages(task_id):
            messages.append(
                ConversationMessage(
                    id=row["id"],
                    task_id=row["task_id"],
                    role=row["role"],
                    kind=row["kind"],
                    content=row["content"],
                    metadata=row["metadata"],
                    created_at=row["created_at"],
                )
            )
        return messages

    def get_workspace(self, task_id: str) -> dict[str, Any]:
        task = self._task(task_id)
        profile = self.load_profile(task_id)
        active_run = self.repository.latest_plan_run(task_id)
        versions = []
        for version in self.repository.list_versions(task_id):
            version_profile_json = version.get("profile_json") or task["profile_json"]
            try:
                version_profile = WorkbookProfile.model_validate_json(version_profile_json)
                version["sheets"] = [sheet.name for sheet in version_profile.sheets]
            except (TypeError, ValueError):
                version["sheets"] = []
            versions.append(version)
        return {
            "task": task,
            "profile": profile.model_dump(mode="json"),
            "messages": [item.model_dump(mode="json") for item in self.get_conversation(task_id)],
            "versions": versions,
            "result": json.loads(task["changeset_json"]) if task["changeset_json"] else None,
            "verification": json.loads(task["verification_json"]) if task["verification_json"] else None,
            "active_run": active_run,
            "tool_runs": self.repository.list_tool_runs(task_id),
            "planning_run": self.repository.latest_planning_run(task_id),
            "activity_events": self.repository.list_activity_events(task_id),
            **self.p0.workspace(task_id),
        }

    def recent_workspaces(
        self, limit: int | None = 12, *, owner_id: str | None = None
    ) -> list[dict[str, Any]]:
        return [
            {
                "task_id": row["task_id"],
                "name": row["original_name"],
                "version": row["current_version"],
                "status": row["status"],
                "updated_at": row["updated_at"],
            }
            for row in self.repository.recent_tasks(limit, owner_id=owner_id)
            if row.get("profile_json") and row.get("cleanup_status") != "pending"
        ]

    def rename_workbook(self, task_id: str, name: str) -> str:
        task = self._task(task_id)
        raw = name.strip()
        if not raw or len(raw) > 180 or raw != Path(raw).name or any(ord(char) < 32 for char in raw):
            raise IntentValidationError("请输入有效的工作簿名称。")
        suffix = Path(task["original_name"]).suffix
        stem = raw[:-len(suffix)] if suffix and raw.casefold().endswith(suffix.casefold()) else raw
        stem = stem.strip().rstrip(".")
        if not stem:
            raise IntentValidationError("工作簿名称不能为空。")
        renamed = f"{stem}{suffix}"
        self.repository.update_task(task_id, original_name=renamed)
        self.p0.touch(task_id)
        return renamed

    def delete_workbook(self, task_id: str) -> None:
        self.p0.delete(task_id)

    def get_workbook_preview(
        self, task_id: str, version: int | None, sheet_name: str, limit: int = 50, *, start_row: int | None = None
    ) -> dict[str, Any]:
        task = self._task(task_id)
        selected_version = int(task["current_version"] if version is None else version)
        version_row = self.repository.get_version(task_id, selected_version)
        if not version_row:
            raise KeyError(f"unknown version {selected_version}")
        profile_json = version_row.get("profile_json") or task["profile_json"]
        profile = WorkbookProfile.model_validate_json(profile_json)
        sheet_profile = next((item for item in profile.sheets if item.name == sheet_name), None)
        if not sheet_profile:
            raise KeyError(f"unknown sheet: {sheet_name}")
        try:
            requested_limit = int(limit)
        except (TypeError, ValueError):
            requested_limit = 50
        safe_limit = max(1, min(requested_limit, 500))

        workbook = load_workbook(Path(version_row["path"]), read_only=True, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                raise KeyError(f"unknown sheet in workbook: {sheet_name}")
            sheet = workbook[sheet_name]
            actual_max_row, actual_max_column = worksheet_dimensions(
                sheet,
                fallback_max_row=sheet_profile.max_row,
                fallback_max_column=sheet_profile.max_column,
            )
            try:
                candidate_header = int(sheet_profile.header_row)
            except (TypeError, ValueError):
                candidate_header = 1
            header_row = candidate_header if 1 <= candidate_header <= actual_max_row else 1
            max_column = max(1, min(int(sheet_profile.max_column or actual_max_column), actual_max_column))
            first_data_row = max(header_row + 1, min(int(start_row or header_row+1), max(header_row+1, actual_max_row)))
            last_row = min(actual_max_row, first_data_row + safe_limit - 1)
            # Read-only worksheets are streaming XML readers. Calling ``cell``
            # repeatedly makes openpyxl rescan the worksheet for every cell and
            # turns a small 50 x N preview into a minutes-long operation. Read
            # the rectangular preview once, in row order, instead.
            preview_values = sheet.iter_rows(
                min_row=header_row,
                max_row=last_row,
                min_col=1,
                max_col=max_column,
                values_only=True,
            )
            header_values = next(preview_values, tuple())
            base_headers = [
                str(value).strip() if value not in (None, "") else f"列{index}"
                for index, value in enumerate(header_values, 1)
            ]
            header_counts = Counter(base_headers)
            headers = [
                f"{name}（{get_column_letter(index)}列）"
                if header_counts[name] > 1
                else name
                for index, name in enumerate(base_headers, 1)
            ]
            rows = [
                {
                    headers[index]: value
                    for index, value in enumerate(row_values)
                }
                for excel_row, row_values in enumerate(preview_values, header_row+1)
                if excel_row >= first_data_row
            ]
            total_rows = max(0, actual_max_row - header_row)
        finally:
            workbook.close()
        return {
            "sheet": sheet_name,
            "version": selected_version,
            "headers": headers,
            "rows": rows,
            "total_rows": total_rows,
            "start_row": first_data_row,
        }

    def get_download(self, task_id: str, *, original: bool = False) -> tuple[str, bytes]:
        task = self._task(task_id)
        if not original:
            return self.get_version_download(task_id, int(task["current_version"]))
        path = Path(task["original_path"])
        self.p0.touch(task_id)
        return task["original_name"], path.read_bytes()

    def get_version_download(self, task_id: str, version: int) -> tuple[str, bytes]:
        task = self._task(task_id)
        record = self.repository.get_version(task_id, int(version))
        if not record:
            raise KeyError(f"unknown version {version}")
        path = Path(record["path"])
        if not path.is_file() or sha256_file(path) != record["sha256"]:
            raise VerificationError("正式版本完整性校验失败，这个版本暂时不可下载。")
        delivery = json.loads(record["delivery_json"]) if record.get("delivery_json") else None
        if delivery:
            path = Path(delivery["path"])
            expected = self.files.version_path(task_id, int(version)).with_suffix(".csv").resolve()
            if path.resolve() != expected or not path.is_file() or sha256_file(path) != delivery["sha256"]:
                raise VerificationError("交付文件完整性校验失败，这个版本暂时不可下载。")
        name = delivery["name"] if delivery else f"SheetPilot_v{version}_{Path(task['original_name']).stem}{path.suffix}"
        self.p0.touch(task_id)
        return name, path.read_bytes()

    def get_artifact_download(self, task_id: str, artifact_path: str) -> tuple[str, bytes]:
        root = (self.files.task_dir(task_id) / "artifacts").resolve()
        path = Path(artifact_path).resolve()
        if path.parent != root or not path.is_file():
            raise KeyError("unknown task artifact")
        return path.name, path.read_bytes()

    def _add_message(
        self,
        task_id: str,
        role: ConversationRole,
        kind: MessageKind,
        content: str,
        metadata: dict[str, Any] | None = None,
    ) -> ConversationMessage:
        row = self.repository.add_message(
            task_id, role.value, kind.value, content, metadata
        )
        return ConversationMessage.model_validate(row)

    def make_task_plan(
        self, task_id: str, envelope: TaskEnvelope,
        *,
        policy_source_text: str | None = None,
        persist: bool = True,
    ) -> tuple[TaskPlan, TaskValidation]:
        row = self._task(task_id)
        profile = self.load_profile(task_id)
        current_version = int(row["current_version"])
        context = ToolContext(
            task_id=task_id,
            input_path=Path(row["current_path"]),
            output_path=self.files.version_path(task_id, current_version + 1),
            profile=profile,
            version_before=current_version,
            version_after=current_version + 1,
        )
        plan, validation = build_task_plan(context, envelope, self.tools, self.skills)
        policy_validation = self.policy.validate_envelope(
            profile, envelope,
            user_source_text=policy_source_text or self._policy_sources.get(task_id),
        )
        validation = TaskValidation(
            valid=validation.valid and policy_validation.valid,
            errors=list(dict.fromkeys([*validation.errors, *policy_validation.errors])),
            warnings=list(dict.fromkeys([*validation.warnings, *policy_validation.warnings])),
        )
        dependency_errors = self._lineage_dependency_errors(
            task_id, current_version, profile, envelope
        )
        if dependency_errors:
            validation = TaskValidation(
                valid=False,
                errors=list(dict.fromkeys([*validation.errors, *dependency_errors])),
                warnings=validation.warnings,
            )
        if row.get("plan_json"):
            try:
                previous = TaskPlan.model_validate_json(row["plan_json"])
                if previous.base_version == current_version:
                    plan.plan_id = previous.plan_id
                    plan.revision = previous.revision + 1
            except Exception:
                pass
        if persist:
            self.repository.update_task(
                task_id,
                status="awaiting_approval" if validation.valid else "needs_clarification",
                spec_json=envelope.model_dump_json(),
                plan_json=plan.model_dump_json(),
                error_json=None,
            )
            self.repository.log_event(
                task_id,
                "task_plan_built",
                {
                    "task_type": envelope.spec.task_type,
                    "skills": plan.skill_ids,
                    "skill_versions": plan.skill_versions,
                    "valid": validation.valid,
                    "errors": validation.errors,
                    "warnings": validation.warnings,
                },
            )
        return plan, validation

    def _lineage_dependency_errors(
        self,
        task_id: str,
        version: int,
        profile: WorkbookProfile,
        envelope: TaskEnvelope,
    ) -> list[str]:
        specs = (
            list(envelope.spec.operations)
            if isinstance(envelope.spec, CompositeTaskSpec)
            else [envelope.spec]
        )
        overwritten: set[tuple[str, int]] = set()
        recomputed_outputs: set[tuple[str, str]] = set()

        def ref_index(ref) -> int | None:
            if ref.column_index:
                return ref.column_index
            if ref.column_letter:
                from openpyxl.utils.cell import column_index_from_string
                try:
                    return column_index_from_string(ref.column_letter)
                except ValueError:
                    return None
            sheet = next((item for item in profile.sheets if item.name == ref.sheet), None)
            matches = [
                item for item in (sheet.columns if sheet else [])
                if (item.display_header or item.name) == ref.display_header and item.active
            ]
            return matches[0].index if len(matches) == 1 else None

        for spec in specs:
            if not isinstance(spec, TransformTaskSpec):
                continue
            for output_name, _, _ in transform_write_columns(profile, spec):
                recomputed_outputs.add((spec.sheet, output_name))
            if spec.operation == TransformOperation.IMPUTE_NUMERIC:
                for rule in spec.impute_rules:
                    index = ref_index(rule.source)
                    if rule.overwrite and index:
                        overwritten.add((spec.sheet, index))
            elif spec.operation == TransformOperation.MAP_VALUES:
                for rule in spec.map_rules:
                    index = ref_index(rule.source)
                    if rule.overwrite and index:
                        overwritten.add((spec.sheet, index))
        if not overwritten:
            return []
        errors: list[str] = []
        for lineage in self.repository.list_lineage(task_id, version):
            sources = {
                int(item["column_index"])
                for item in lineage["source_columns"]
                if item.get("column_index") is not None
            }
            if (lineage["sheet_name"], lineage["output_column"]) in recomputed_outputs:
                continue
            if any((lineage["sheet_name"], index) in overwritten for index in sources):
                errors.append(
                    f"覆盖的来源列被派生列“{lineage['output_column']}”依赖；"
                    "新计划必须同时重算该依赖列。"
                )
        return errors

    def _record_transform_lineage(
        self,
        task_id: str,
        source_version: int,
        target_version: int,
        results: list[ToolResult],
    ) -> None:
        self.repository.copy_lineage(task_id, source_version, target_version)
        for result in results:
            if result.tool_name != "transform_column":
                continue
            sheet_name = result.modified_sheets[0] if result.modified_sheets else ""
            for item in result.details.get("lineage", []):
                self.repository.add_lineage(
                    task_id,
                    target_version,
                    sheet_name=sheet_name,
                    output_column=str(item["output_column"]),
                    output_column_index=item.get("output_column_index"),
                    source_columns=list(item.get("source_columns", [])),
                    operation=str(item["operation"]),
                    parameters=dict(item.get("parameters", {})),
                )

    def execute_task_plan(
        self,
        task_id: str,
        envelope: TaskEnvelope,
        plan: TaskPlan,
        *,
        run_id: str | None = None,
    ) -> tuple[Path, VerificationReport, ToolResult]:
        return self.transaction_executor.execute(
            task_id, envelope, plan, run_id=run_id
        )

    def _execute_task_plan_impl(self, task_id, envelope, plan, *, run_id=None):
        with self.p0.lock(task_id):
            self._task(task_id)
            op_id = self.p0.bind_plan(task_id, plan)
            if self.p0.operation(task_id, op_id)['status'] == 'completed':
                raise IntentValidationError('本计划已经执行完成，请为新要求生成新的计划。')
            segment_id = uuid4().hex
            segment_started = time.perf_counter()
            with self.repository.connect() as db:
                measured_before = db.execute("SELECT COALESCE(SUM(duration_ms),0) FROM p0_measurements WHERE operation_id=? AND stage!='waiting'", (op_id,)).fetchone()[0]
            if run_id is not None:
                self.p0.measure(task_id, op_id, f"resume:{uuid4().hex}", "resume_count", 0)
            self.p0.state(op_id, "executing")
            try:
                output, report, result = self._execute_task_plan_body(task_id, envelope, plan, run_id=run_id)
                self.p0.complete(task_id, op_id, output, report, result)
                return output, report, result
            except Exception:
                run = self.repository.latest_plan_run(task_id)
                state = run.get("status") if run else "failed"
                if state == "paused":
                    self.p0.measure(task_id, op_id, f"pause:{uuid4().hex}", "pause_count", 0)
                self.p0.state(op_id, state if state in {"paused", "cancelled"} else "failed")
                raise
            finally:
                with self.repository.connect() as db:
                    measured_after = db.execute("SELECT COALESCE(SUM(duration_ms),0) FROM p0_measurements WHERE operation_id=? AND stage!='waiting'", (op_id,)).fetchone()[0]
                overhead = max(0, (time.perf_counter()-segment_started)*1000-(measured_after-measured_before))
                self.p0.measure(task_id, op_id, f'overhead:{segment_id}', 'preparation', overhead)
                self.p0.touch(task_id)

    def _execute_task_plan_body(
        self,
        task_id: str,
        envelope: TaskEnvelope,
        plan: TaskPlan,
        *,
        run_id: str | None = None,
    ) -> tuple[Path, VerificationReport, ToolResult]:
        deadline = time.monotonic() + self.settings.execution_budget_seconds
        row = self._task(task_id)
        if row["status"] != "approved" and run_id is None:
            raise IntentValidationError("任务尚未获得用户批准。")
        current_path = Path(row["current_path"])
        current_version = int(row["current_version"])
        profile = self.load_profile(task_id)
        original_path = Path(row["original_path"])
        original_sha256 = row.get("original_sha256") or ""
        if not original_path.is_file() or not original_sha256:
            failure = FailureRecord(
                category=FailureCategory.SAFETY_VIOLATION,
                step_id="integrity",
                phase="prepare",
                safe_message="原始文件完整性基线不可用，请重新上传文件。",
                technical_detail="missing original file or baseline hash",
            )
            self.repository.update_task(task_id, status="failed", error_json=failure.model_dump_json())
            raise IntentValidationError(failure.safe_message, detail=failure.model_dump(mode="json"))
        if sha256_file(original_path) != original_sha256:
            failure = FailureRecord(
                category=FailureCategory.SAFETY_VIOLATION,
                step_id="integrity",
                phase="prepare",
                safe_message="原始文件完整性校验失败，任务已停止且不会创建新版本。",
                technical_detail="original upload hash mismatch",
            )
            self.repository.update_task(task_id, status="failed", error_json=failure.model_dump_json())
            self.repository.log_event(task_id, "original_integrity_failed", {})
            raise IntentValidationError(failure.safe_message, detail=failure.model_dump(mode="json"))
        if plan.base_version != current_version or plan.base_sha256 != sha256_file(current_path):
            raise IntentValidationError("文件版本已经变化，请重新生成并批准执行计划。")
        if not 1 <= len(plan.steps) <= 5:
            raise IntentValidationError("执行计划必须包含 1 到 5 个业务步骤。")
        structure_errors = validate_plan_structure(plan, self.tools, self.skills)
        if structure_errors:
            raise IntentValidationError("执行计划结构无效。", detail={"errors": structure_errors})
        if row.get("plan_json"):
            stored_plan = TaskPlan.model_validate_json(row["plan_json"])
            if stored_plan.plan_id != plan.plan_id or stored_plan.revision != plan.revision:
                raise IntentValidationError("计划已经修订，旧审批不再有效。")
        active_run_id = run_id or uuid4().hex
        existing_run = self.repository.get_plan_run(active_run_id)
        if existing_run:
            if existing_run["plan_id"] != plan.plan_id or int(existing_run["plan_revision"]) != plan.revision:
                raise IntentValidationError("恢复请求与已批准的计划版本不一致。")
            start_index = int(existing_run["current_step"])
            input_path = Path(existing_run["last_checkpoint_path"] or current_path)
            self.repository.update_plan_run(active_run_id, status="running", error_json=None)
            prior_rows = [
                item for item in self.repository.list_tool_runs(task_id)
                if item.get("run_id") == active_run_id and item["status"] == "completed"
            ]
            step_results = [
                ToolResult.model_validate_json(item["result_json"])
                for item in prior_rows if item.get("result_json")
            ]
            step_reports = []
            fencing_token = existing_run.get("fencing_token") or ""
            if not fencing_token or not self.repository.validate_execution_lease(
                task_id, active_run_id, fencing_token
            ):
                raise IntentValidationError("恢复运行的执行租约已经失效。")
            for completed_step in plan.steps[:start_index]:
                report_path = self.files.run_report_path(
                    task_id, active_run_id, f"step_{completed_step.order:02d}.json"
                )
                if report_path.exists():
                    step_reports.append(VerificationReport.model_validate_json(report_path.read_text(encoding="utf-8")))
        else:
            try:
                lease = self.repository.acquire_plan_run(
                    active_run_id, task_id, plan.plan_id, plan.revision,
                    current_version, plan.base_sha256, sha256_file(current_path),
                )
            except ValueError as exc:
                raise IntentValidationError(str(exc)) from exc
            fencing_token = lease["fencing_token"]
            start_index = 0
            input_path = current_path
            step_results = []
            step_reports = []
        op_id = self.p0.bind_plan(task_id, plan)
        self.p0.link(task_id, op_id, "execution", active_run_id)
        def check_budget():
            if time.monotonic() >= deadline:
                failure = FailureRecord(category=FailureCategory.TRANSIENT_IO, step_id="budget", phase="prepare",
                                        safe_message="达到执行时间预算，已在安全检查点暂停；可主动恢复。")
                self.repository.update_plan_run(active_run_id, status="paused", error_json=failure.model_dump_json())
                self.repository.update_task(task_id, status="paused", error_json=failure.model_dump_json())
                raise IntentValidationError(failure.safe_message)
        output_version = current_version + 1
        self.repository.update_task(task_id, status="executing", error_json=None)
        try:
            completed_ids = {step.id for step in plan.steps[:start_index]}
            for index in range(start_index, len(plan.steps)):
                check_budget()
                if not self.repository.validate_execution_lease(
                    task_id, active_run_id, fencing_token
                ):
                    raise IntentValidationError("执行租约已经失效，任务已停止。")
                if sha256_file(original_path) != original_sha256:
                    failure = FailureRecord(
                        category=FailureCategory.SAFETY_VIOLATION,
                        step_id=plan.steps[index].id,
                        phase="prepare",
                        safe_message="原始文件完整性校验失败，任务已停止。",
                        technical_detail="original upload hash mismatch at step boundary",
                    )
                    self.repository.update_plan_run(
                        active_run_id, status="failed", error_json=failure.model_dump_json()
                    )
                    self.repository.update_task(
                        task_id, status="failed", error_json=failure.model_dump_json()
                    )
                    raise VerificationError(failure.safe_message, detail=failure.model_dump(mode="json"))
                control = self.repository.get_plan_run(active_run_id)
                if control and control["cancel_requested"]:
                    self.repository.update_plan_run(active_run_id, status="cancelled")
                    self.files.remove_run_files(task_id, active_run_id)
                    self.repository.update_task(task_id, status="profiled", error_json=None)
                    raise IntentValidationError("任务已在步骤边界取消，正式版本未改变。")
                if control and control["pause_requested"]:
                    self.repository.update_plan_run(active_run_id, status="paused")
                    self.repository.update_task(task_id, status="paused")
                    raise IntentValidationError("任务已在步骤边界暂停。")
                step = plan.steps[index]
                if any(dependency not in completed_ids for dependency in step.depends_on):
                    raise IntentValidationError(f"步骤 {step.id} 的前置依赖尚未完成。")
                input_profile = profile_workbook(
                    input_path, row["original_name"], self.settings.max_sheet_rows
                )
                checkpoint = self.files.checkpoint_path(task_id, active_run_id, step.order)
                context = ToolContext(
                    task_id=task_id, input_path=input_path, output_path=checkpoint,
                    profile=input_profile, version_before=current_version,
                    version_after=output_version,
                    artifact_dir=self.files.task_dir(task_id) / "artifacts",
                    original_path=Path(row["original_path"]),
                    execution_deadline=deadline,
                    cancelled=lambda: bool((self.repository.get_plan_run(active_run_id) or {}).get("cancel_requested")),
                    report_stage=lambda label: self.repository.add_activity_event(
                        task_id, "execution", active_run_id, "formula", label, "running",
                    ),
                )
                tool = self.tools.get(step.tool_name)
                spec = tool.parse_spec(step.arguments)
                validation = tool.validate(context, spec)
                if not validation.valid:
                    failure = FailureRecord(
                        category=FailureCategory.USER_INPUT_REQUIRED, step_id=step.id,
                        phase="prepare", repairable=True, requires_user_input=True,
                        safe_message="当前数据与计划不再一致，需要修改后续计划。",
                        technical_detail="；".join(validation.errors),
                    )
                    self.repository.update_plan_run(
                        active_run_id, status="paused", current_step=index,
                        last_checkpoint_path=str(input_path) if input_path != current_path else None,
                        error_json=failure.model_dump_json(),
                    )
                    self.repository.update_task(task_id, status="paused", error_json=failure.model_dump_json())
                    raise IntentValidationError(failure.safe_message, detail=failure.model_dump(mode="json"))
                attempts = 2 if step.retry_policy == "transient_once" else 1
                result = None
                for attempt in range(1, attempts + 1):
                    tool_run_id = self.repository.start_tool_run(
                        task_id, step.id, step.tool_name, step.arguments,
                        run_id=active_run_id, step_order=step.order, attempt=attempt,
                        input_sha256=sha256_file(input_path),
                    )
                    try:
                        execution_started = time.perf_counter()
                        try:
                            result = tool.execute(context, spec)
                        finally:
                            self.p0.measure(task_id, op_id, f"execute:{tool_run_id}", "execution", (time.perf_counter()-execution_started)*1000)
                        break
                    except (OSError, TimeoutError) as exc:
                        checkpoint.unlink(missing_ok=True)
                        retrying = attempt < attempts
                        self.repository.finish_tool_run(
                            tool_run_id, "retrying" if retrying else "failed",
                            error={"type": type(exc).__name__, "message": str(exc)},
                            failure_category=FailureCategory.TRANSIENT_IO.value,
                        )
                        if not retrying:
                            raise
                assert result is not None
                verification_started = time.perf_counter()
                self.repository.add_activity_event(task_id, "execution", active_run_id, "verification", "正在核验本步骤的实际结果", "running")
                task_verification = tool.verify(context, spec, result)
                report = verify_tool_output(
                    task_id=task_id, original_path=original_path,
                    original_hash=original_sha256,
                    input_path=input_path, input_hash=input_profile.sha256,
                    output_path=checkpoint, result=result,
                    task_verification=task_verification,
                )
                step_plan = plan.model_copy(update={"steps": [step]})
                change_set = self.diff_verifier.compare(input_path, checkpoint, step_plan)
                report.unexpected_changes = change_set.unexpected_changes
                report.style_warnings = change_set.style_warnings
                report.warnings = list(dict.fromkeys([
                    *report.warnings, *change_set.style_warnings,
                ]))
                report.checks.append(VerificationCheck(
                    name="planned_changes_only",
                    passed=not change_set.unexpected_changes,
                    detail=(
                        "未发现计划外变更"
                        if not change_set.unexpected_changes
                        else "；".join(change_set.unexpected_changes[:10])
                    ),
                ))
                if change_set.unexpected_changes:
                    report.status = "failed_rolled_back"
                    report.manual_review_required = True
                    report.output_path = None
                elif report.warnings and report.status == "passed":
                    report.status = "passed_with_warnings"
                    report.manual_review_required = True
                self.files.run_report_path(task_id, active_run_id, f"step_{step.order:02d}.json").write_text(
                    report.model_dump_json(indent=2), encoding="utf-8"
                )
                if report.status == "failed_rolled_back":
                    self.p0.measure(task_id, op_id, f"verify:{tool_run_id}", "verification", (time.perf_counter()-verification_started)*1000)
                    failure_category = (
                        FailureCategory.SAFETY_VIOLATION
                        if report.unexpected_changes
                        else FailureCategory.VERIFICATION_FAILED
                    )
                    self.repository.finish_tool_run(
                        tool_run_id, "failed", result=result.model_dump(mode="json"),
                        error={"verification": report.model_dump(mode="json")},
                        checkpoint_path=str(checkpoint),
                        failure_category=failure_category.value,
                    )
                    failure = FailureRecord(
                        category=failure_category, step_id=step.id,
                        phase="verify", repairable=True, requires_user_input=True,
                        safe_message=(
                            "检测到计划外变更，已阻止提交正式版本。"
                            if report.unexpected_changes
                            else "；".join(task_verification.errors[:2]) or "步骤结果未通过核验，已暂停且未提交正式版本。"
                        ),
                        technical_detail="；".join(check.detail for check in report.checks if not check.passed),
                    )
                    self.repository.update_plan_run(
                        active_run_id, status="paused", current_step=index,
                        last_checkpoint_path=str(input_path) if input_path != current_path else None,
                        error_json=failure.model_dump_json(),
                    )
                    self.repository.update_task(task_id, status="paused", error_json=failure.model_dump_json())
                    raise VerificationError(failure.safe_message, detail=failure.model_dump(mode="json"))
                self.p0.measure(task_id, op_id, f"verify:{tool_run_id}", "verification", (time.perf_counter()-verification_started)*1000)
                checkpoint_hash = sha256_file(checkpoint)
                self.repository.finish_tool_run(
                    tool_run_id, "completed", result=result.model_dump(mode="json"),
                    output_sha256=checkpoint_hash, checkpoint_path=str(checkpoint),
                )
                step_results.append(result)
                step_reports.append(report)
                completed_ids.add(step.id)
                input_path = checkpoint
                self.repository.update_plan_run(
                    active_run_id, current_step=index + 1,
                    last_checkpoint_path=str(checkpoint), status="running",
                )

            combined_result = ToolResult(
                tool_name=step_results[0].tool_name if len(step_results) == 1 else "transaction",
                modified_sheets=list(dict.fromkeys(name for item in step_results for name in item.modified_sheets)),
                modified_ranges=list(dict.fromkeys(name for item in step_results for name in item.modified_ranges)),
                created_sheets=list(dict.fromkeys(name for item in step_results for name in item.created_sheets)),
                rows_affected=sum(item.rows_affected for item in step_results),
                counters=(
                    dict(step_results[0].counters)
                    if len(step_results) == 1
                    else {f"step_{index + 1}_{key}": value for index, item in enumerate(step_results) for key, value in item.counters.items()}
                ),
                warnings=list(dict.fromkeys([
                    *(message for item in step_results for message in item.warnings),
                    *(message for item in step_reports for message in item.warnings),
                ])),
                details={"run_id": active_run_id, "steps": [item.model_dump(mode="json") for item in step_results]},
            )
            control = self.repository.get_plan_run(active_run_id)
            if control and control["cancel_requested"]:
                self.repository.update_plan_run(active_run_id, status="cancelled")
                self.files.remove_run_files(task_id, active_run_id)
                self.repository.update_task(task_id, status="profiled", error_json=None)
                raise IntentValidationError("任务已在最终提交前取消，正式版本未改变。")
            if not self.repository.validate_execution_lease(
                task_id, active_run_id, fencing_token
            ):
                raise IntentValidationError("执行租约已经失效，无法提交正式版本。")
            check_budget()
            row = self._task(task_id)
            commit_started = time.perf_counter()
            self.repository.add_activity_event(task_id, "execution", active_run_id, "commit", "正在终检并保存正式结果", "running")
            final_path, report, output_profile = self.committer.commit(
                task=row,
                plan=plan,
                run_id=active_run_id,
                fencing_token=fencing_token,
                source_path=input_path,
                step_reports=step_reports,
                combined_result=combined_result,
            )
            self.p0.measure(task_id, op_id, f"commit:{active_run_id}", "commit", (time.perf_counter()-commit_started)*1000)
            self.files.run_report_path(task_id, active_run_id, "final.json").write_text(
                report.model_dump_json(indent=2), encoding="utf-8"
            )
            self.repository.log_event(
                task_id,
                "transaction_completed",
                {
                    "version": output_version,
                    "run_id": active_run_id,
                    "plan_id": plan.plan_id,
                    "plan_revision": plan.revision,
                    "commit_id": report.commit_id,
                    "steps": len(plan.steps),
                    "status": report.status,
                },
            )
            self._record_transform_lineage(
                task_id, current_version, output_version, step_results
            )
            self.files.remove_run_checkpoints(task_id, active_run_id)
            return final_path, report, combined_result
        except Exception as exc:
            run = self.repository.get_plan_run(active_run_id)
            if run and run.get("cancel_requested"):
                self.repository.update_plan_run(active_run_id, status="cancelled")
                self.files.remove_run_files(task_id, active_run_id)
                self.repository.update_task(task_id, status="profiled", error_json=None)
            if run and run["status"] == "running" and not run.get("cancel_requested"):
                current_step = int(run["current_step"])
                at_commit = current_step >= len(plan.steps)
                category = (
                    FailureCategory.VERIFICATION_FAILED
                    if isinstance(exc, VerificationError)
                    else FailureCategory.FATAL
                )
                failure = FailureRecord(
                    category=category,
                    step_id=plan.steps[current_step].id if current_step < len(plan.steps) else "commit",
                    phase="commit" if at_commit else "execute",
                    rollback_required=True,
                    safe_message=(
                        "最终验证或提交失败，正式版本未改变。"
                        if at_commit else "执行失败，正式版本未改变。"
                    ),
                    technical_detail=f"{type(exc).__name__}: {exc}",
                )
                self.repository.update_plan_run(active_run_id, status="failed", error_json=failure.model_dump_json())
                self.repository.update_task(task_id, status="failed", error_json=failure.model_dump_json())
            self.repository.log_event(
                task_id, "transaction_failed", {"run_id": active_run_id, "type": type(exc).__name__, "message": str(exc)}
            )
            raise

    def resume_current_run(self, task_id: str) -> tuple[Path, VerificationReport, ToolResult]:
        run = self.repository.latest_plan_run(task_id)
        row = self._task(task_id)
        if not run or run["status"] not in {"paused", "failed"}:
            raise IntentValidationError("当前没有可恢复的运行。")
        if not row["spec_json"] or not row["plan_json"]:
            raise IntentValidationError("恢复所需的计划数据不存在。")
        self.repository.request_run_control(run["run_id"], "resume")
        self.repository.update_task(task_id, status="approved", error_json=None)
        envelope = TaskEnvelope.model_validate_json(row["spec_json"])
        plan = TaskPlan.model_validate_json(row["plan_json"])
        from .graph import sqlite_execution_workflow

        try:
            with sqlite_execution_workflow(self) as workflow:
                graph_result = workflow.invoke(
                    {
                        "task_id": task_id,
                        "spec": envelope.model_dump(mode="json"),
                        "plan": plan.model_dump(mode="json"),
                        "run_id": run["run_id"],
                        "status": "approved",
                    },
                    config={"configurable": {"thread_id": f"{task_id}:{plan.plan_id}:r{plan.revision}"}},
                )
        except SheetPilotError:
            raise
        except Exception as exc:
            self.repository.log_event(
                task_id,
                "resume_failed",
                {"run_id": run["run_id"], "type": type(exc).__name__, "message": str(exc)},
            )
            raise IntentValidationError(
                "从检查点重试时发生内部错误，正式版本没有改变。请查看后台日志或取消后重新生成计划。"
            ) from exc
        if graph_result.get("status") != "completed":
            raise IntentValidationError(
                graph_result.get("error", {}).get("message", "恢复执行未能完成。")
            )
        payload = graph_result["result"]
        output = Path(payload["output_path"])
        report = VerificationReport.model_validate(payload["verification"])
        result = ToolResult.model_validate(payload["tool_result"])
        version = int(self._task(task_id)["current_version"])
        self._add_message(
            task_id, ConversationRole.ASSISTANT, MessageKind.RESULT,
            result_text(version, result, report),
            {
                "version": version,
                "result": result.model_dump(mode="json"),
                "verification": report.model_dump(mode="json"),
                "preview_sheets": list(dict.fromkeys([*result.modified_sheets, *result.created_sheets])),
                "resumed_run": run["run_id"],
            },
        )
        return output, report, result

    def control_current_run(self, task_id: str, action: str) -> None:
        run = self.repository.latest_plan_run(task_id)
        if not run or run["status"] not in {"running", "paused", "failed"}:
            raise IntentValidationError("当前没有可控制的运行。")
        if action == "cancel" and run["status"] in {"paused", "failed"}:
            task = self._task(task_id)
            self.repository.update_plan_run(run["run_id"], status="cancelled", cancel_requested=1)
            op_id = self.p0.linked("execution", run["run_id"])
            if op_id:
                self.p0.state(op_id, "cancelled")
            self.files.remove_run_files(task_id, run["run_id"])
            self.repository.update_task(
                task_id,
                status=(
                    "needs_conversion"
                    if Path(task["original_path"]).suffix.casefold() == ".csv"
                    and int(task["current_version"]) == 0
                    else "profiled"
                ),
                error_json=None,
            )
            self._add_message(
                task_id, ConversationRole.ASSISTANT, MessageKind.TEXT,
                "已取消未完成的运行，上一正式版本保持不变。",
            )
            return
        self.repository.request_run_control(run["run_id"], action)

    def skip_optional_step(self, task_id: str) -> tuple[Path, VerificationReport, ToolResult]:
        run = self.repository.latest_plan_run(task_id)
        row = self._task(task_id)
        if not run or run["status"] not in {"paused", "failed"}:
            raise IntentValidationError("当前没有可跳过步骤的运行。")
        plan = TaskPlan.model_validate_json(row["plan_json"])
        index = int(run["current_step"])
        if index >= len(plan.steps) or plan.steps[index].required:
            raise IntentValidationError("当前步骤是必要步骤，不能跳过。")
        self.repository.update_plan_run(run["run_id"], current_step=index + 1, status="paused")
        return self.resume_current_run(task_id)

    def restore_version(self, task_id: str, version: int) -> WorkbookProfile:
        version_row = self.repository.get_version(task_id, version)
        if not version_row:
            raise KeyError(f"unknown version {version} for task {task_id}")
        task = self._task(task_id)
        profile_json = version_row.get("profile_json")
        if profile_json:
            profile = WorkbookProfile.model_validate_json(profile_json)
        else:
            profile = profile_workbook(
                Path(version_row["path"]), task["original_name"], self.settings.max_sheet_rows
            )
        self.repository.update_task(
            task_id,
            status=(
                "needs_conversion"
                if Path(task["original_path"]).suffix.casefold() == ".csv" and version == 0
                else "profiled"
            ),
            current_path=version_row["path"],
            current_version=version,
            profile_json=profile.model_dump_json(),
            spec_json=None,
            plan_json=None,
            changeset_json=None,
            verification_json=None,
            error_json=None,
        )
        op_id = self.p0.begin(task_id)
        self.p0.state(op_id, "restored")
        self.repository.log_event(task_id, "version_restored", {"version": version})
        self._add_message(
            task_id,
            ConversationRole.ASSISTANT,
            MessageKind.TEXT,
            f"已恢复到 V{version}。接下来的要求会基于这个版本处理。",
            {"version": version},
        )
        return profile

    def result_payload(self, task_id: str) -> dict[str, Any]:
        row = self._task(task_id)
        return {
            "task": row,
            "changeset": json.loads(row["changeset_json"]) if row["changeset_json"] else None,
            "verification": json.loads(row["verification_json"]) if row["verification_json"] else None,
            "versions": self.repository.list_versions(task_id),
            "tool_runs": self.repository.list_tool_runs(task_id),
        }

    def _task(self, task_id: str) -> dict[str, Any]:
        row = self.repository.get_task(task_id)
        if row is None or row.get("cleanup_status") == "pending":
            raise KeyError(f"unknown task: {task_id}")
        return row

from __future__ import annotations

from pathlib import Path

from openpyxl.utils import get_column_letter

from .schemas import (
    CleanOperation, ColumnProfile, CompositeTaskSpec, ConversionDirection, ExpectedChanges, FormulaMode,
    RiskLevel, SheetProfile, TaskEnvelope, TaskPlan, TaskType, TaskValidation,
    ToolStep, TransformOperation, TransformTaskSpec, WorkbookProfile,
)
from .skills import SkillRegistry
from .tools import (
    ToolContext, ToolRegistry, _resolve_column_ref, _resolve_named_column, _resolve_formula_column, _transform_artifact_name,
    transform_write_columns,
)


def _operations(envelope: TaskEnvelope):
    return list(envelope.spec.operations) if isinstance(envelope.spec, CompositeTaskSpec) else [envelope.spec]


def _summary(envelope: TaskEnvelope) -> str:
    if isinstance(envelope.spec, CompositeTaskSpec):
        return envelope.goal or f"依次完成 {len(envelope.spec.operations)} 个表格处理步骤"
    spec = envelope.spec
    if spec.task_type == TaskType.LOOKUP:
        return f"从“{spec.source_sheet}”匹配字段并追加到“{spec.main_sheet}”"
    if spec.task_type == TaskType.NORMALIZE:
        return f"标准化“{spec.sheet}”中的“{spec.column}”列"
    if spec.task_type == TaskType.AGGREGATE:
        return f"按 {', '.join(spec.group_by)} 汇总“{spec.sheet}”并写入“{spec.output_sheet}”"
    if spec.task_type == TaskType.SORT_RANK:
        return f"对“{spec.sheet}”排序或排名并写入“{spec.output_sheet}”"
    if spec.task_type == TaskType.CLEAN:
        if spec.operation == CleanOperation.KEEP_COLUMNS:
            target = "覆盖当前工作表" if spec.overwrite else f"写入“{spec.output_sheet}”"
            return f"在“{spec.sheet}”中保留指定列并{target}"
        return f"清洗“{spec.sheet}”并写入“{spec.output_sheet}”"
    if spec.task_type == TaskType.FORMULA:
        return f"诊断或修复“{spec.sheet}”中“{spec.column}”列的公式"
    if spec.task_type == TaskType.CONVERSION:
        return "将 CSV 转换为 XLSX" if spec.direction == ConversionDirection.CSV_TO_XLSX else f"将“{spec.source_sheet}”导出为 CSV"
    if spec.task_type == TaskType.TRANSFORM:
        labels = {
            TransformOperation.IMPUTE_NUMERIC: "填补数值缺失",
            TransformOperation.MAP_VALUES: "映射分类值",
            TransformOperation.PARSE_DIMENSIONS: "解析尺寸字段",
            TransformOperation.DERIVE_METRIC: "计算受控派生指标",
        }
        return f"在“{spec.sheet}”中{labels[spec.operation]}"
    return envelope.goal


def _description(spec) -> str:
    if spec.task_type == TaskType.LOOKUP:
        return "检查关联键后执行左连接，追加静态值，并为异常记录生成清单。"
    if spec.task_type == TaskType.NORMALIZE:
        destination = f"覆盖原列“{spec.column}”" if spec.overwrite else f"写入新列“{spec.output_column or spec.column + '_标准化'}”"
        return f"依次执行：{', '.join(item.value for item in spec.operations)}；{destination}。"
    if spec.task_type == TaskType.AGGREGATE:
        metrics = ", ".join(f"{item.column}:{item.function.value}" for item in spec.metrics)
        return f"按 {', '.join(spec.group_by)} 分组，计算 {metrics}。"
    if spec.task_type == TaskType.SORT_RANK:
        return "对完整数据区域执行稳定排序/排名，不改动来源工作表。"
    if spec.task_type == TaskType.CLEAN:
        if spec.operation == CleanOperation.KEEP_COLUMNS:
            labels = "、".join(
                item.column_letter or item.display_header or str(item.column_index or "?")
                for item in spec.keep_columns
            )
            return (
                f"仅保留 {len(spec.keep_columns)} 列（{labels}），按原工作簿顺序排列；"
                f"{'覆盖当前工作表' if spec.overwrite else '写入新工作表'}。"
            )
        return f"执行 {spec.operation.value}，结果写入新工作表。"
    if spec.task_type == TaskType.CONVERSION:
        return "保留原文件并生成新的转换产物。"
    if spec.task_type == TaskType.TRANSFORM:
        base = {
            TransformOperation.IMPUTE_NUMERIC: (
                "填补明确缺失值；格式错误按用户确认的规则处理。"
                if any(rule.invalid_policy.value == "impute" for rule in spec.impute_rules)
                else "仅填补明确缺失值；格式错误保持原样并进入异常清单。"
            ),
            TransformOperation.MAP_VALUES: "按用户确认的映射表转换；未知值留空并进入异常清单。",
            TransformOperation.PARSE_DIMENSIONS: "提取长、宽、高和最大径；多候选或单位冲突不猜测。",
            TransformOperation.DERIVE_METRIC: "使用用户明确选择的注册公式模板计算派生指标。",
        }[spec.operation]
        if spec.operation == TransformOperation.IMPUTE_NUMERIC:
            refs = [rule.source for rule in spec.impute_rules]
            overwrite = any(rule.overwrite for rule in spec.impute_rules)
            detail = ", ".join(rule.method.value for rule in spec.impute_rules)
        elif spec.operation == TransformOperation.MAP_VALUES:
            refs = [rule.source for rule in spec.map_rules]
            overwrite = any(rule.overwrite for rule in spec.map_rules)
            detail = "trim + NFKC" + (" + casefold" if any(rule.casefold for rule in spec.map_rules) else "")
        elif spec.operation == TransformOperation.PARSE_DIMENSIONS:
            refs = [spec.dimension_rule.source] if spec.dimension_rule else []
            overwrite = bool(spec.dimension_rule and spec.dimension_rule.overwrite_existing_outputs)
            detail = f"输入单位 {spec.dimension_rule.default_unit}" if spec.dimension_rule else ""
        else:
            refs = list(spec.derive_rule.sources) if spec.derive_rule else []
            overwrite = bool(spec.derive_rule and spec.derive_rule.overwrite_existing_outputs)
            template = spec.derive_rule.formula_template.value if spec.derive_rule and spec.derive_rule.formula_template else "待确认"
            detail = f"公式 {template}，mm³ 转 cm³"
        sources = "、".join(
            ref.column_letter or ref.display_header or str(ref.column_index or "?") for ref in refs
        )
        return f"{base} 来源列：{sources}；写入：{'覆盖' if overwrite else '新增列'}；{detail}。"
    labels = {
        FormulaMode.DIAGNOSE: "诊断公式断层、异常引用与模式偏差。",
        FormulaMode.FILL_GAPS: "只填充可由上下相邻公式唯一推断的空白断层。",
        FormulaMode.REPLACE_OUTLIERS: "替换可由上下相邻公式唯一推断的异常公式。",
        FormulaMode.SET_FORMULA: "将用户明确提供的公式按相对引用平移到目标范围。",
        FormulaMode.COPY_FORMULA: "从指定样板行复制公式到目标范围。",
        FormulaMode.RECALCULATE: "使用可用的 Excel/LibreOffice 引擎重算公式并刷新缓存值。",
    }
    return labels[spec.mode]


def _risk(spec) -> RiskLevel:
    if spec.task_type == TaskType.TRANSFORM:
        destructive = (
            any(rule.overwrite for rule in spec.impute_rules)
            or any(rule.overwrite for rule in spec.map_rules)
            or bool(spec.dimension_rule and spec.dimension_rule.overwrite_existing_outputs)
            or bool(spec.derive_rule and spec.derive_rule.overwrite_existing_outputs)
        )
        return RiskLevel.HIGH if destructive else RiskLevel.MEDIUM
    if spec.task_type == TaskType.NORMALIZE and spec.overwrite:
        return RiskLevel.HIGH
    if spec.task_type == TaskType.FORMULA:
        if spec.mode == FormulaMode.REPLACE_OUTLIERS:
            return RiskLevel.HIGH
        if spec.mode == FormulaMode.SET_FORMULA and spec.overwrite_existing:
            return RiskLevel.HIGH
        if spec.mode == FormulaMode.COPY_FORMULA and spec.overwrite_existing:
            return RiskLevel.HIGH
        if spec.mode in {FormulaMode.FILL_GAPS, FormulaMode.SET_FORMULA, FormulaMode.COPY_FORMULA}:
            return RiskLevel.MEDIUM
        return RiskLevel.LOW
    if spec.task_type == TaskType.CLEAN:
        if spec.operation == CleanOperation.KEEP_COLUMNS and spec.overwrite:
            return RiskLevel.HIGH
        if spec.operation == CleanOperation.REMOVE_DUPLICATES:
            return RiskLevel.MEDIUM
    if spec.task_type in {TaskType.AGGREGATE, TaskType.SORT_RANK}:
        return RiskLevel.LOW
    if spec.task_type == TaskType.CONVERSION:
        return RiskLevel.LOW
    return RiskLevel.MEDIUM


def _risk_max(levels: list[RiskLevel]) -> RiskLevel:
    order = {RiskLevel.LOW: 0, RiskLevel.MEDIUM: 1, RiskLevel.HIGH: 2}
    return max(levels, key=order.get)


def _simulate_profile(profile: WorkbookProfile, spec, changes: ExpectedChanges) -> WorkbookProfile:
    predicted = profile.model_copy(deep=True)
    by_name = {sheet.name: sheet for sheet in predicted.sheets}
    if spec.task_type == TaskType.LOOKUP:
        sheet = by_name.get(spec.main_sheet)
        if sheet:
            for name in changes.new_columns:
                sheet.max_column += 1
                sheet.columns.append(ColumnProfile(name=name, index=sheet.max_column, inferred_type="mixed", empty_count=0, unique_count=0, formula_count=0))
    elif spec.task_type == TaskType.NORMALIZE and not spec.overwrite:
        sheet = by_name.get(spec.sheet)
        if sheet:
            column, _ = _resolve_named_column(profile, spec.sheet, spec.column)
            label = (column.display_header or column.name) if column else spec.column
            name = spec.output_column or f"{label}_标准化"
            sheet.max_column += 1
            inferred = "number" if any(item.value == "to_number" for item in spec.operations) else "mixed"
            sheet.columns.append(ColumnProfile(name=name, index=sheet.max_column, inferred_type=inferred, empty_count=0, unique_count=0, formula_count=0))
    elif spec.task_type == TaskType.TRANSFORM:
        sheet = by_name.get(spec.sheet)
        if sheet:
            for name in changes.new_columns:
                sheet.max_column += 1
                inferred = "text" if name.endswith("状态") else "number"
                sheet.columns.append(ColumnProfile(
                    name=name,
                    index=sheet.max_column,
                    inferred_type=inferred,
                    empty_count=0,
                    unique_count=0,
                    formula_count=0,
                    letter=get_column_letter(sheet.max_column),
                    display_header=name,
                    nonempty_count=max(0, sheet.max_row - (sheet.header_row or 1)),
                    active=True,
                ))
    elif (
        spec.task_type == TaskType.CLEAN
        and spec.operation == CleanOperation.KEEP_COLUMNS
        and spec.overwrite
    ):
        sheet = by_name.get(spec.sheet)
        if sheet:
            selected: list[ColumnProfile] = []
            for ref in spec.keep_columns:
                column, _ = _resolve_column_ref(predicted, ref)
                if column and column.index not in {item.index for item in selected}:
                    selected.append(column)
            selected.sort(key=lambda item: item.index)
            for index, column in enumerate(selected, 1):
                column.index = index
                column.letter = get_column_letter(index)
            sheet.columns = selected
            sheet.max_column = max(1, len(selected))
    for name in changes.created_sheets:
        if "（" in name or name in by_name:
            continue
        column_specs: list[tuple[str, str]] = []
        if spec.task_type == TaskType.AGGREGATE:
            column_specs = [*( (item, "mixed") for item in spec.group_by ), *( (item.output_column, "number") for item in spec.metrics )]
        elif spec.task_type in {TaskType.SORT_RANK, TaskType.CLEAN}:
            source = by_name.get(spec.sheet)
            if spec.task_type == TaskType.CLEAN and spec.operation == CleanOperation.KEEP_COLUMNS:
                requested = {
                    column.index
                    for ref in spec.keep_columns
                    for column, error in [_resolve_column_ref(predicted, ref)]
                    if column is not None and error is None
                }
                column_specs = [
                    (item.name, item.inferred_type)
                    for item in (source.columns if source else [])
                    if item.index in requested
                ]
            else:
                column_specs = [(item.name, item.inferred_type) for item in (source.columns if source else [])]
            if spec.task_type == TaskType.SORT_RANK and spec.rank_column:
                column_specs.append((spec.rank_output_column, "number"))
            if spec.task_type == TaskType.CLEAN and spec.operation == CleanOperation.MARK_DUPLICATES:
                column_specs.append((spec.duplicate_flag_column, "text"))
        else:
            column_specs = [(item, "mixed") for item in changes.new_columns]
        columns = [
            ColumnProfile(name=column_name, index=index, inferred_type=inferred_type, empty_count=0, unique_count=0, formula_count=0)
            for index, (column_name, inferred_type) in enumerate(column_specs, 1)
        ]
        predicted.sheets.append(SheetProfile(name=name, state="visible", max_row=max(1, changes.estimated_rows + 1), max_column=max(1, len(columns)), header_row=1, header_confidence=1.0, columns=columns))
    return predicted


def _authorize_changes(
    profile: WorkbookProfile, spec, changes: ExpectedChanges
) -> ExpectedChanges:
    """Turn a preview into the exact write contract enforced after execution."""
    authorized = changes.model_copy(deep=True)
    authorized.allowed_created_sheets = list(dict.fromkeys(
        name.split("（", 1)[0] for name in changes.created_sheets
    ))
    by_name = {sheet.name: sheet for sheet in profile.sheets}
    if spec.task_type == TaskType.LOOKUP:
        sheet = by_name.get(spec.main_sheet)
        if sheet:
            names = spec.output_column_names or spec.return_columns
            start = sheet.max_column + 1
            end = start + max(0, len(names) - 1)
            authorized.allowed_ranges = [
                f"{spec.main_sheet}!{get_column_letter(start)}{sheet.header_row}:"
                f"{get_column_letter(end)}{sheet.max_row}"
            ] if names else []
    elif spec.task_type == TaskType.NORMALIZE:
        sheet = by_name.get(spec.sheet)
        if sheet:
            column, _ = _resolve_named_column(profile, spec.sheet, spec.column)
            index = column.index if spec.overwrite and column else sheet.max_column + 1
            if index:
                authorized.allowed_ranges = [
                    f"{spec.sheet}!{get_column_letter(index)}{sheet.header_row}:"
                    f"{get_column_letter(index)}{sheet.max_row}"
                ]
    elif spec.task_type == TaskType.FORMULA:
        sheet = by_name.get(spec.sheet)
        authorized.allow_formula_recalculation = bool(spec.recalculate)
        if sheet and spec.mode not in {FormulaMode.DIAGNOSE, FormulaMode.RECALCULATE}:
            column, _ = _resolve_formula_column(profile, spec.sheet, spec.column)
            index = column.index if column else None
            if index:
                start = spec.start_row or (sheet.header_row or 1) + 1
                end = spec.end_row or sheet.max_row
                authorized.allowed_ranges = [
                    f"{spec.sheet}!{get_column_letter(index)}{start}:"
                    f"{get_column_letter(index)}{end}"
                ]
    elif spec.task_type == TaskType.CONVERSION:
        if spec.direction == ConversionDirection.XLSX_TO_CSV:
            authorized.allow_formula_recalculation = True
            output_name = Path(spec.output_name or f"{spec.source_sheet}.csv").name
            if not output_name.casefold().endswith(".csv"):
                output_name += ".csv"
            authorized.allowed_artifacts = [output_name]
    elif spec.task_type == TaskType.TRANSFORM:
        sheet = by_name.get(spec.sheet)
        if sheet:
            authorized.allowed_ranges = [
                f"{spec.sheet}!{get_column_letter(index)}{sheet.header_row}:"
                f"{get_column_letter(index)}{sheet.max_row}"
                for _, index, _ in transform_write_columns(profile, spec)
            ]
            authorized.allowed_artifacts = [_transform_artifact_name(spec)]
    elif (
        spec.task_type == TaskType.CLEAN
        and spec.operation == CleanOperation.KEEP_COLUMNS
        and spec.overwrite
    ):
        sheet = by_name.get(spec.sheet)
        if sheet:
            authorized.allowed_ranges = [
                f"{spec.sheet}!A{sheet.header_row or 1}:"
                f"{get_column_letter(sheet.max_column)}{sheet.max_row}"
            ]
    return authorized


def build_task_plan(
    context: ToolContext,
    envelope: TaskEnvelope,
    registry: ToolRegistry,
    skills: SkillRegistry | None = None,
) -> tuple[TaskPlan, TaskValidation]:
    skill_registry = skills or SkillRegistry()
    operations = _operations(envelope)
    errors: list[str] = []
    warnings: list[str] = []
    steps: list[ToolStep] = []
    predicted_profile = context.profile.model_copy(deep=True)
    touched_sheets: set[str] = set()
    total_rows = 0
    errors.extend(skill_registry.validate_composition(envelope))
    selected_skills = []
    for index, spec in enumerate(operations, 1):
        try:
            skill = skill_registry.get(spec.task_type)
        except KeyError:
            errors.append(f"步骤 {index} 使用了未知任务类型“{spec.task_type}”。")
            continue
        selected_skills.append(skill)
        if spec.confidence < 0.6:
            errors.append(f"步骤 {index} 的 Skill 选择置信度过低，需要进一步澄清。")
        tool_name = skill.tool_name
        if tool_name not in registry.names():
            errors.append(f"Skill“{skill.id}”允许的工具“{tool_name}”未注册。")
            continue
        tool = registry.get(tool_name)
        step_context = ToolContext(task_id=context.task_id, input_path=context.input_path, output_path=context.output_path, profile=predicted_profile, version_before=context.version_before, version_after=context.version_after)
        validation = tool.validate(step_context, spec)
        errors.extend(f"步骤 {index}：{item}" for item in validation.errors)
        warnings.extend(f"步骤 {index}：{item}" for item in validation.warnings)
        # FormulaTool deterministically validates every field it needs. In inference modes,
        # an absent user-supplied formula is intentional: candidates come from neighboring
        # formulas, and non-unique candidates are reported without writing. LLM-authored
        # missing_fields/ambiguities must therefore not create false clarification loops.
        if spec.task_type != TaskType.FORMULA:
            errors.extend(f"步骤 {index} 缺少信息：{item}" for item in spec.missing_fields)
            errors.extend(f"步骤 {index} 存在歧义：{item}" for item in spec.ambiguities)
        preview = ExpectedChanges.model_validate(tool.preview(step_context, spec))
        preview = _authorize_changes(predicted_profile, spec, preview)
        if (
            isinstance(spec, TransformTaskSpec)
            and spec.operation == TransformOperation.IMPUTE_NUMERIC
        ):
            count = int(preview.review_totals.get("invalid_numeric", 0) or 0)
            if count >= 2 and any(
                not rule.invalid_policy_confirmed for rule in spec.impute_rules
            ):
                errors.append(f"IMPUTE_RESOLUTION_REQUIRED:invalid_numeric:{count}")
        if (
            isinstance(spec, TransformTaskSpec)
            and spec.operation == TransformOperation.PARSE_DIMENSIONS
            and spec.dimension_rule is not None
        ):
            for status in (
                "multiple_candidates",
                "unit_conflict",
                "invalid_format",
                "no_numeric_measurement",
            ):
                count = int(preview.summary_counts.get(status, 0) or 0)
                if count >= 2 and status not in spec.dimension_rule.resolution_policies:
                    errors.append(f"DIMENSION_RESOLUTION_REQUIRED:{status}:{count}")
                    break
        if (
            isinstance(spec, TransformTaskSpec)
            and spec.operation == TransformOperation.DERIVE_METRIC
            and spec.derive_rule is not None
        ):
            count = int(preview.summary_counts.get("invalid_input", 0) or 0)
            if count >= 2 and "invalid_input" not in spec.derive_rule.resolution_policies:
                errors.append(f"DERIVE_RESOLUTION_REQUIRED:invalid_input:{count}")
        step_id = f"step_{index}"
        steps.append(ToolStep(id=step_id, order=index, tool_name=tool_name, arguments=spec.model_dump(mode="json"), description=_description(spec), risk_level=_risk(spec), expected_changes=preview, depends_on=[f"step_{index - 1}"] if index > 1 else [], preconditions=["input_hash_matches", "dependencies_completed"], verification_rules=list(skill.verification_rules)))
        touched_sheets.update(preview.modified_sheets)
        touched_sheets.update(preview.created_sheets)
        total_rows += preview.estimated_rows
        predicted_profile = _simulate_profile(predicted_profile, spec, preview)
    risk = _risk_max([step.risk_level for step in steps]) if steps else RiskLevel.HIGH
    if len({name for name in touched_sheets if "（" not in name}) > 3 or total_rows > 50_000:
        risk = {RiskLevel.LOW: RiskLevel.MEDIUM, RiskLevel.MEDIUM: RiskLevel.HIGH, RiskLevel.HIGH: RiskLevel.HIGH}[risk]
        warnings.append("计划影响多个工作表或超过 50,000 行，整体风险已提升一级。")
    confirmation_requirements = (
        ["overwrite_source_columns"]
        if any(
            step.expected_changes.destructive
            and (
                step.tool_name in {"transform_column", "normalize_column"}
                or (
                    step.tool_name == "clean"
                    and step.arguments.get("operation") == CleanOperation.KEEP_COLUMNS.value
                    and bool(step.arguments.get("overwrite"))
                )
            )
            for step in steps
        )
        else []
    )
    plan = TaskPlan(task_id=context.task_id, task_type=envelope.spec.task_type, input_version=context.version_before, input_sha256=context.profile.sha256, base_version=context.version_before, base_sha256=context.profile.sha256, summary=_summary(envelope), steps=steps, risk_level=risk, invariants=["original_file_unchanged", "base_version_unchanged_until_commit", "every_step_verified", "single_final_version_commit"], skill_ids=[skill.id for skill in selected_skills], skill_versions={skill.id: skill.version for skill in selected_skills}, plan_assumptions=list(envelope.assumptions), confirmation_requirements=confirmation_requirements)
    errors.extend(validate_plan_structure(plan, registry, skill_registry))
    return plan, TaskValidation(valid=not errors, errors=errors, warnings=list(dict.fromkeys(warnings)))


def validate_plan_structure(
    plan: TaskPlan, registry: ToolRegistry, skills: SkillRegistry | None = None
) -> list[str]:
    errors: list[str] = []
    ids = [step.id for step in plan.steps]
    if len(ids) != len(set(ids)):
        errors.append("计划步骤 ID 必须唯一。")
    if [step.order for step in plan.steps] != list(range(1, len(plan.steps) + 1)):
        errors.append("计划步骤顺序必须从 1 连续递增。")
    known = set(registry.names())
    seen: set[str] = set()
    skill_registry = skills or SkillRegistry()
    if plan.skill_ids and len(plan.skill_ids) != len(plan.steps):
        errors.append("计划中的 Skill 数量必须与步骤数量一致。")
    for step in plan.steps:
        if step.tool_name not in known:
            errors.append(f"步骤 {step.id} 使用了未知工具“{step.tool_name}”。")
        for dependency in step.depends_on:
            if dependency not in ids:
                errors.append(f"步骤 {step.id} 依赖不存在的步骤“{dependency}”。")
            elif dependency not in seen:
                errors.append(f"步骤 {step.id} 的依赖“{dependency}”未在其之前完成，计划可能存在循环。")
        seen.add(step.id)
    for index, skill_id in enumerate(plan.skill_ids):
        try:
            skill = skill_registry.get_by_id(skill_id)
        except KeyError:
            errors.append(f"计划引用了未注册的 Skill“{skill_id}”。")
            continue
        if plan.skill_versions.get(skill_id) != skill.version:
            errors.append(f"Skill“{skill_id}”版本与当前注册表不一致。")
        if index < len(plan.steps) and plan.steps[index].tool_name != skill.tool_name:
            errors.append(f"Skill“{skill_id}”不允许工具“{plan.steps[index].tool_name}”。")
    return errors

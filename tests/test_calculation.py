from __future__ import annotations

import re
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

import sheetpilot.calculation as calculation


def remove_calculation_properties(path: Path) -> None:
    rewritten = path.with_name(f"{path.stem}_without_calc_properties.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        rewritten, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/workbook.xml":
                data = re.sub(br"<calcPr\b[^>]*/>", b"", data)
            target.writestr(item, data)
    rewritten.replace(path)


def test_recalculation_initializes_missing_calculation_properties(tmp_path, monkeypatch):
    path = tmp_path / "without_calc_properties.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "计算"
    sheet.append(["数量", "单价", "金额"])
    sheet.append([2, 15, "=A2*B2"])
    workbook.save(path)
    workbook.close()
    remove_calculation_properties(path)

    before = load_workbook(path, data_only=False)
    try:
        assert before.calculation is None
    finally:
        before.close()

    monkeypatch.setattr(
        calculation, "_recalculate_with_excel", lambda path, timeout: (False, "Excel unavailable")
    )
    monkeypatch.setattr(
        calculation,
        "_recalculate_with_libreoffice",
        lambda path, timeout: (False, "LibreOffice unavailable"),
    )

    result = calculation.recalculate_workbook(path)

    after = load_workbook(path, data_only=False)
    try:
        assert after.calculation is not None
        assert after.calculation.calcMode == "auto"
        assert after.calculation.fullCalcOnLoad is True
        assert after.calculation.forceFullCalc is True
        assert after["计算"]["C2"].value == "=A2*B2"
    finally:
        after.close()
    assert result.backend == "none"
    assert result.formula_cells == 1

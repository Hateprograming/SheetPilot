from __future__ import annotations

import copy
import json
import sys
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from evals.phase1.cases import base_book,prepare
from evals.phase1.common import APP,DEFAULT_OUT,Case,protected_output,load_cases,sha
from evals.phase1.oracle import compare,snapshot,unchanged_metadata
from evals.phase1.runner import Calls,ScriptedClient,answer_clarification,guard_plan,run_suite
from sheetpilot.schemas import TaskEnvelope,AgentTurnResult,ConversationState,ConversationMessage,ConversationRole,MessageKind
from sheetpilot.profiler import profile_workbook
from sheetpilot.skills import capability_refusal
TEST_OUT=DEFAULT_OUT/'unit-tests'


@pytest.fixture(scope='module')
def suite():
    prepare(TEST_OUT)
    return {c.id:c for c in load_cases(TEST_OUT/'cases')}


def turn(case,spec=None,clarify=False):
    env=TaskEnvelope(goal=case.request,spec=spec or case.spec)
    return AgentTurnResult(state=ConversationState.NEEDS_CLARIFICATION if clarify else ConversationState.PLAN_READY,
                          needs_clarification=clarify,envelope=env,
                          reply=ConversationMessage(task_id='test',role=ConversationRole.ASSISTANT,kind=MessageKind.TEXT,content='test'))


def test_preparation_is_reproducible_and_sources_unchanged(suite):
    assert len(suite)==280
    before={k:v.input_sha256 for k,v in suite.items()}
    prepare(TEST_OUT)
    assert before=={c.id:c.input_sha256 for c in load_cases(TEST_OUT/'cases')}
    assert sum(c.source['type']=='spreadsheetbench_adapted' for c in suite.values())==28
    assert len({c.group for c in suite.values() if c.source['type']=='spreadsheetbench_adapted'})==14


def test_coverage_and_ambiguity_are_real(suite):
    from collections import Counter
    normals=[c for c in suite.values() if c.kind=='normal']
    assert len(normals)==184
    assert sorted(Counter(c.capability for c in normals).values())==[23]*8
    assert sum(c.spec['task_type']=='composite' for c in normals)==11
    q=suite['Q1'];book=snapshot(TEST_OUT/'cases'/q.id/q.input)
    assert len(book['客户'][0])==3
    expanded=[c for c in suite.values() if c.id.startswith('P-')]
    assert len(expanded)==240
    assert Counter(c.kind for c in expanded)==Counter(normal=160,clarification=40,refusal=20,fault=20)
    assert Counter(c.source['type'] for c in expanded)==Counter(synthetic=168,research_reconstructed=48,spreadsheetbench_adapted=24)
    assert len({c.source['source_item_id'] for c in expanded if c.source['type']=='research_reconstructed'})==48
    assert all(c.modes==['offline'] for c in expanded if c.kind=='fault')


def test_unsupported_requests_are_refused_without_blocking_local_work(suite):
    for case in suite.values():
        if case.kind=='refusal':
            assert capability_refusal(case.request) is not None, case.id
    for request in (
        '不要联网，只处理本地表格。',
        '不用插件，按客户编号精确匹配当前工作簿的两张表。',
        '请把当前工作簿的订单表和客户表按编号关联。',
    ):
        assert capability_refusal(request) is None


def test_unapproved_source_sheet_column_width_is_a_blocking_change(suite,tmp_path):
    from sheetpilot.config import Settings
    from sheetpilot.integrity import WorkbookDiffVerifier
    from sheetpilot.service import SheetPilotService
    case=suite['P-CLEAN-15']
    source=TEST_OUT/'cases'/case.id/case.input
    settings=Settings(project_root=tmp_path,data_dir=tmp_path/'data',database_path=tmp_path/'data/db',
                      llm_provider='fake',api_key='',base_url='',model='fake',
                      max_file_bytes=20_000_000,max_sheet_rows=100000)
    service=SheetPilotService(settings,ScriptedClient(case))
    try:
        profile=service.create_task(source.name,source.read_bytes())
        turn=service.submit_message(service.task_id_from_profile(profile),case.request)
        assert turn.plan is not None
        changed=tmp_path/'changed.xlsx'
        workbook=load_workbook(source)
        workbook['交易'].column_dimensions['B'].width=13
        workbook.save(changed)
        workbook.close()
        changes=WorkbookDiffVerifier().compare(source,changed,turn.plan)
        assert any('交易' in item and '列宽' in item for item in changes.unexpected_changes)
    finally:
        service._background_executor.shutdown(wait=False)


def test_batch_budget_cannot_be_raised_and_unrun_cases_are_recorded(suite,monkeypatch):
    from evals.phase1.common import dump
    import evals.phase1.runner as runner
    out=DEFAULT_OUT/'tests-budget';prepare(out);used=[]
    class Child:
        returncode=0
        def __init__(self,cmd,**kwargs):
            folder=Path(cmd[cmd.index('--folder')+1]);limit=int(cmd[cmd.index('--call-limit')+1])
            calls=min(limit,4 if not used else 6);used.append(calls)
            case=Case.model_validate_json(Path(cmd[cmd.index('--case')+1]).read_text(encoding='utf-8'))
            dump(folder/'calls.json',[{'status':200,'usage':None}]*calls)
            dump(folder/'result.json',{'case_id':case.id,'kind':case.kind,'capability':case.capability,'status':'passed','mode':'live','calls':calls})
        def wait(self,timeout=None):return 0
    monkeypatch.setattr(runner.subprocess,'Popen',Child)
    # git discovery uses subprocess.run; isolate that too so it cannot consume
    # the worker substitute. No requests leave this unit test.
    from types import SimpleNamespace
    monkeypatch.setattr(runner.subprocess,'run',lambda *a,**k:SimpleNamespace(stdout='unit-test'))
    run=run_suite(out,'live',max_tasks=100,max_calls=100,rerun=True)
    meta=json.loads((run/'run.json').read_text(encoding='utf-8'))
    assert sum(used)==60 and meta['max_calls']==60 and meta['selected_count']==20
    records=[json.loads(p.read_text(encoding='utf-8')) for p in run.glob('*/result.json')]
    assert len(records)==20 and any(r['status']=='blocked' for r in records)


def test_live_batches_resume_without_repeating_completed_cases(suite,monkeypatch):
    from evals.phase1.common import dump
    import evals.phase1.runner as runner
    from types import SimpleNamespace
    selected=[]
    class Child:
        returncode=0
        def __init__(self,cmd,**kwargs):
            folder=Path(cmd[cmd.index('--folder')+1]);case=Case.model_validate_json(Path(cmd[cmd.index('--case')+1]).read_text(encoding='utf-8'))
            selected.append(case.id)
            dump(folder/'calls.json',[{'status':200,'usage':None}])
            dump(folder/'result.json',{'case_id':case.id,'kind':case.kind,'capability':case.capability,'status':'passed','stage':'checking','mode':'live','calls':1})
        def wait(self,timeout=None):return 0
    monkeypatch.setattr(runner.subprocess,'Popen',Child)
    monkeypatch.setattr(runner.subprocess,'run',lambda *a,**k:SimpleNamespace(stdout='unit-test'))
    first=run_suite(TEST_OUT,'live',max_tasks=5,max_calls=20)
    second=run_suite(TEST_OUT,'live',max_tasks=5,max_calls=20)
    assert first and second and len(selected)==10 and len(set(selected))==10
    data=json.loads((TEST_OUT/'suite-summary.json').read_text(encoding='utf-8'))
    assert data['live_remaining']<=256-10


@pytest.mark.parametrize('change',['value','type','row','sheet','outside'])
def test_checker_detects_wrong_output(suite,change):
    e=suite['L1'].checks['workbook'];bad=copy.deepcopy(e)
    if change=='value':bad['订单'][1][-1]='WRONG'
    if change=='type':bad['订单'][1][1]=1
    if change=='row':bad['订单'].pop()
    if change=='sheet':bad['extra']=[['oops']]
    if change=='outside':bad['客户'][1][1]='WRONG'
    assert compare(e,bad)
    assert not compare(e,e)


def test_number_tolerance_does_not_coerce_blank_or_text():
    assert not compare({'a':[[1.0]]},{'a':[[1+1e-11]]})
    for v in [None,'1',True]:assert compare({'a':[[1]]},{'a':[[v]]})
    assert compare({'a':[[None]]},{'a':[[0]]})


def test_style_and_layout_mutation_detected(tmp_path):
    source=tmp_path/'in.xlsx';base_book(source);out=tmp_path/'out.xlsx'
    w=load_workbook(source);w['订单']['B2'].number_format='@';w['订单'].freeze_panes=None;w.save(out);w.close()
    diffs=unchanged_metadata(source,out,None)
    assert {d['scope'] for d in diffs}=={'style','layout'}


def test_unknown_clarification_is_never_answered(suite):
    case=suite['Q1'];t=turn(case,case.initial_spec,True)
    assert answer_clarification(case,t,False)==case.reply['answer']
    assert answer_clarification(case,t,True) is None
    assert answer_clarification(suite['L1'],t,False) is None
    assert answer_clarification(case,turn(case),False) is None


def test_call_cap_includes_failed_requests(tmp_path):
    calls=Calls(tmp_path/'calls.json',2);network=[]
    def fail(request):network.append(1);raise httpx.ConnectError('controlled failure')
    with httpx.Client(transport=httpx.MockTransport(fail),event_hooks={'request':[calls.before]}) as client:
        for _ in range(4):
            with pytest.raises((httpx.ConnectError,RuntimeError)):client.post('https://test.invalid',json={'model':'test'})
    assert len(network)==2 and len(calls.events)==2
    assert all(e['usage'] is None for e in calls.events)


def test_output_paths_cannot_reach_user_workspaces(tmp_path):
    with pytest.raises(ValueError):protected_output(APP/'data')
    with pytest.raises(ValueError):protected_output(DEFAULT_OUT/'..'/'data_web')
    assert protected_output(DEFAULT_OUT/'runs')==DEFAULT_OUT.resolve()/'runs'


def test_case_schema_rejects_unknown_and_absolute_input(suite):
    data=suite['L1'].model_dump();data['input']='../data/input.xlsx'
    with pytest.raises(ValueError):Case(**data)
    data['input']='input.xlsx';data['execute_python']='danger'
    with pytest.raises(ValueError):Case(**data)


def test_isolated_timeouts_do_not_stop_next_case(suite):
    out=DEFAULT_OUT/'tests-timeout';prepare(out)
    run=run_suite(out,'offline',ids=['L1','N1'],case_seconds=.001)
    results=[json.loads(p.read_text(encoding='utf-8')) for p in run.glob('*/result.json')]
    assert len(results)==2
    assert all(r['stage']=='runner_timeout' for r in results)


def test_plan_guard_blocks_changed_semantics(suite,tmp_path):
    from sheetpilot.service import SheetPilotService
    from sheetpilot.config import Settings
    case=suite['L1'];settings=Settings(project_root=tmp_path,data_dir=tmp_path/'data',database_path=tmp_path/'data/db',llm_provider='fake',api_key='',base_url='',model='fake',max_file_bytes=20_000_000,max_sheet_rows=100000)
    service=SheetPilotService(settings,ScriptedClient(case))
    try:
        source=TEST_OUT/'cases/L1'/case.input
        profile=service.create_task(source.name,source.read_bytes());task_id=service.task_id_from_profile(profile)
        t=service.submit_message(task_id,case.request)
        assert not guard_plan(case,t,profile)
        t.envelope.spec.constraints=['保留行序']
        assert not guard_plan(case,t,profile)
        original=copy.deepcopy(t.plan.steps[0].arguments)
        t.plan.steps[0].arguments['return_columns']=['未授权列']
        assert guard_plan(case,t,profile)
        t.plan.steps[0].arguments=original
        from sheetpilot.schemas import DuplicatePolicy
        t.envelope.spec.duplicate_policy=DuplicatePolicy.LAST
        assert guard_plan(case,t,profile)
        assert len(service.repository.list_versions(task_id))==1
    finally:service._background_executor.shutdown(wait=False)

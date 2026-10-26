from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from sheetpilot.errors import DuplicateKeyError
from sheetpilot.matching import execute_lookup, normalize_key, validate_lookup_spec
from sheetpilot.profiler import profile_workbook
from sheetpilot.schemas import DuplicatePolicy, LookupTaskSpec, NormalizationMode


def spec(**updates):
    values = dict(
        main_sheet="订单",
        source_sheet="客户",
        main_key_column="客户编号",
        source_key_column="客户编号",
        return_columns=["客户等级", "地区"],
        output_column_names=["客户等级", "地区"],
        confidence=1.0,
    )
    values.update(updates)
    return LookupTaskSpec(**values)


def test_exact_normalization_preserves_leading_zeros():
    assert normalize_key(" ００１ ", NormalizationMode.EXACT_TEXT) == "001"
    assert normalize_key("001", NormalizationMode.EXACT_TEXT) != normalize_key(1, NormalizationMode.EXACT_TEXT)
    assert normalize_key("001", NormalizationMode.NUMERIC_EQUIVALENCE) == normalize_key(1, NormalizationMode.NUMERIC_EQUIVALENCE)


def test_executes_lookup_and_preserves_original(workbook_factory, tmp_path):
    source = workbook_factory()
    original_bytes = source.read_bytes()
    output = tmp_path / "result.xlsx"
    profile = profile_workbook(source, "input.xlsx", 100_000)
    changes = execute_lookup("task", source, output, profile, spec(), 0, 1)
    assert source.read_bytes() == original_bytes
    assert changes.matched_count == 3
    wb = load_workbook(output, data_only=True)
    ws = wb["订单"]
    assert [ws.cell(1, c).value for c in range(1, 5)] == ["客户编号", "订单额", "客户等级", "地区"]
    assert ws["C2"].value == "A"
    assert "SheetPilot_Unmatched" not in wb.sheetnames
    wb.close()


def test_unmatched_and_empty_keys_create_report(workbook_factory, tmp_path):
    source = workbook_factory(main_rows=[("001", 10), ("999", 20), (None, 30)])
    output = tmp_path / "result.xlsx"
    profile = profile_workbook(source, "input.xlsx", 100_000)
    changes = execute_lookup("task", source, output, profile, spec(), 0, 1)
    assert changes.matched_count == 1
    assert changes.unmatched_count == 1
    assert changes.empty_key_count == 1
    wb = load_workbook(output, data_only=True)
    assert wb["SheetPilot_Unmatched"].max_row == 3
    wb.close()


def test_duplicate_source_keys_block_by_default(workbook_factory, tmp_path):
    source = workbook_factory(source_rows=[("001", "A", "华东"), ("001", "B", "华北")])
    profile = profile_workbook(source, "input.xlsx", 100_000)
    validation = validate_lookup_spec(source, profile, spec())
    assert validation.duplicate_key_count == 1
    with pytest.raises(DuplicateKeyError):
        execute_lookup("task", source, tmp_path / "result.xlsx", profile, spec(), 0, 1)


def test_duplicate_source_keys_can_use_last(workbook_factory, tmp_path):
    source = workbook_factory(source_rows=[("001", "A", "华东"), ("001", "B", "华北")], main_rows=[("001", 10)])
    profile = profile_workbook(source, "input.xlsx", 100_000)
    output = tmp_path / "result.xlsx"
    execute_lookup("task", source, output, profile, spec(duplicate_policy=DuplicatePolicy.LAST), 0, 1)
    wb = load_workbook(output, data_only=True)
    assert wb["订单"]["C2"].value == "B"
    wb.close()


def test_numeric_equivalence_requires_explicit_mode(workbook_factory, tmp_path):
    source = workbook_factory(main_rows=[("001", 10)], source_rows=[(1, "A", "华东")])
    profile = profile_workbook(source, "input.xlsx", 100_000)
    exact_output = tmp_path / "exact.xlsx"
    exact = execute_lookup("exact", source, exact_output, profile, spec(), 0, 1)
    assert exact.unmatched_count == 1
    numeric_output = tmp_path / "numeric.xlsx"
    numeric = execute_lookup(
        "numeric",
        source,
        numeric_output,
        profile,
        spec(normalization_mode=NormalizationMode.NUMERIC_EQUIVALENCE),
        0,
        1,
    )
    assert numeric.matched_count == 1


def test_duplicate_count_is_not_truncated_to_examples(workbook_factory):
    source_rows = []
    for index in range(21):
        source_rows.extend([(f"K{index}", "A", "华东"), (f"K{index}", "B", "华北")])
    source = workbook_factory(source_rows=source_rows, main_rows=[("K0", 10)])
    profile = profile_workbook(source, "input.xlsx", 100_000)
    validation = validate_lookup_spec(source, profile, spec())
    assert validation.duplicate_key_count == 21
    assert len(validation.duplicate_examples) == 20

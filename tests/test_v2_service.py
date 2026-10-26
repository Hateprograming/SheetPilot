from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from sheetpilot.config import Settings
from sheetpilot.context import build_context
from sheetpilot.profiler import profile_workbook
from sheetpilot.schemas import (
    AggregateFunction,
    AggregateMetric,
    AggregateTaskSpec,
    CleanOperation,
    CleanTaskSpec,
    ColumnRef,
    ConversationRole,
    LookupTaskSpec,
    NormalizeOperation,
    NormalizeTaskSpec,
    MessageKind,
    RankMethod,
    SortDirection,
    SortKey,
    SortRankTaskSpec,
    TaskEnvelope,
)
from sheetpilot.errors import IntentValidationError
from sheetpilot.service import SheetPilotService


def make_service(tmp_path: Path) -> SheetPilotService:
    settings = Settings(
        project_root=tmp_path,
        data_dir=tmp_path / "data",
        database_path=tmp_path / "data" / "sheetpilot.sqlite3",
        llm_provider="fake",
        api_key="",
        base_url="",
        model="fake",
        max_file_bytes=20 * 1024 * 1024,
        max_sheet_rows=100_000,
    )
    return SheetPilotService(settings)


def create_task(service: SheetPilotService, path: Path) -> str:
    profile = service.create_task("input.xlsx", path.read_bytes())
    return service.task_id_from_profile(profile)


def remove_worksheet_dimensions(path: Path) -> None:
    rewritten = path.with_name(f"{path.stem}_without_dimensions.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        rewritten, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                data = re.sub(br"<dimension\s+ref=\"[^\"]+\"\s*/>", b"", data)
            target.writestr(item, data)
    rewritten.replace(path)


@pytest.mark.parametrize("legacy_header", [None, 0, 999])
def test_workbook_preview_recovers_from_invalid_legacy_header_row(
    workbook_factory, tmp_path, legacy_header
):
    path = workbook_factory()
    service = make_service(tmp_path)
    task_id = create_task(service, path)
    profile = service.load_profile(task_id)
    legacy_profile = profile.model_copy(deep=True)
    legacy_profile.sheets[0].header_row = legacy_header
    service.repository.update_version_profile(task_id, 0, legacy_profile.model_dump_json())

    preview = service.get_workbook_preview(task_id, 0, "订单", 20)

    assert preview["headers"] == ["客户编号", "订单额"]
    assert preview["total_rows"] == 3
    assert preview["rows"][0] == {"客户编号": "001", "订单额": 10}


def test_workbook_preview_sanitizes_invalid_limit(workbook_factory, tmp_path):
    path = workbook_factory()
    service = make_service(tmp_path)
    task_id = create_task(service, path)

    preview = service.get_workbook_preview(task_id, 0, "订单", None)

    assert len(preview["rows"]) == 3


def test_workbook_preview_streams_wide_sheet_and_preserves_duplicate_headers(tmp_path):
    path = tmp_path / "wide.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "19-20"
    sheet.append(["性别", "性别", None, *[f"列{i}" for i in range(4, 93)]])
    for row_index in range(1, 61):
        sheet.append(["男", "女", row_index, *range(4, 93)])
    workbook.save(path)

    service = make_service(tmp_path)
    task_id = create_task(service, path)
    preview = service.get_workbook_preview(task_id, 0, "19-20", 50)

    assert preview["total_rows"] == 60
    assert len(preview["rows"]) == 50
    assert len(preview["headers"]) == 92
    assert preview["headers"][:3] == ["性别（A列）", "性别（B列）", "列3"]
    assert preview["rows"][0]["性别（A列）"] == "男"
    assert preview["rows"][0]["性别（B列）"] == "女"


def test_dimensionless_xlsx_can_be_previewed_and_planned(workbook_factory, tmp_path):
    path = workbook_factory()
    remove_worksheet_dimensions(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        assert workbook["客户"].max_row is None
    finally:
        workbook.close()

    service = make_service(tmp_path)
    task_id = create_task(service, path)
    preview = service.get_workbook_preview(task_id, 0, "订单", 20)
    envelope = TaskEnvelope(
        goal="按客户编号补充客户等级和地区",
        spec=LookupTaskSpec(
            main_sheet="订单",
            source_sheet="客户",
            main_key_column="客户编号",
            source_key_column="客户编号",
            return_columns=["客户等级", "地区"],
            output_column_names=["客户等级", "地区"],
            confidence=1,
        ),
    )

    plan, validation = service.make_task_plan(task_id, envelope)

    assert preview["total_rows"] == 3
    assert validation.valid, validation.errors
    assert plan.steps[0].tool_name == "lookup"

    service.approve(task_id, "approve")
    output, report, result = service.execute_task_plan(task_id, envelope, plan)
    output_book = load_workbook(output, data_only=True)
    try:
        assert output_book["订单"]["C1"].value == "客户等级"
        assert output_book["订单"]["D1"].value == "地区"
    finally:
        output_book.close()
    assert report.status == "passed"
    assert result.counters["matched"] == 3


def execute(service: SheetPilotService, task_id: str, envelope: TaskEnvelope):
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid, validation.errors
    service.approve(task_id, "approve")
    return service.execute_task_plan(task_id, envelope, plan)


def test_v2_lookup_uses_registry(workbook_factory, tmp_path):
    path = workbook_factory(main_rows=[("001", 10), ("999", 20)])
    service = make_service(tmp_path)
    task_id = create_task(service, path)
    envelope = TaskEnvelope(
        goal="补充客户等级",
        spec=LookupTaskSpec(
            main_sheet="订单",
            source_sheet="客户",
            main_key_column="客户编号",
            source_key_column="客户编号",
            return_columns=["客户等级"],
            output_column_names=["客户等级"],
            confidence=1,
        ),
    )
    output, report, result = execute(service, task_id, envelope)
    assert report.status == "passed_with_warnings"
    assert result.counters["matched"] == 1
    assert load_workbook(output, data_only=True)["订单"]["C2"].value == "A"


def test_v2_normalize_then_aggregate_uses_current_profile(workbook_factory, tmp_path):
    path = workbook_factory(main_rows=[(" ００１ ", 10), ("００１", 20), ("002", 5)])
    service = make_service(tmp_path)
    task_id = create_task(service, path)
    normalize = TaskEnvelope(
        goal="统一客户编号",
        spec=NormalizeTaskSpec(
            sheet="订单",
            column="客户编号",
            operations=[NormalizeOperation.NFKC, NormalizeOperation.TRIM],
            output_column="客户编号_标准化",
            confidence=1,
        ),
    )
    first_output, first_report, _ = execute(service, task_id, normalize)
    assert first_report.status == "passed"
    assert "客户编号_标准化" in [
        item.name for item in service.load_profile(task_id).sheets[0].columns
    ]

    aggregate = TaskEnvelope(
        goal="按标准化编号汇总订单额",
        spec=AggregateTaskSpec(
            sheet="订单",
            group_by=["客户编号_标准化"],
            metrics=[
                AggregateMetric(
                    column="订单额",
                    function=AggregateFunction.SUM,
                    output_column="订单额合计",
                )
            ],
            output_sheet="编号汇总",
            confidence=1,
        ),
    )
    second_output, second_report, second_result = execute(service, task_id, aggregate)
    assert second_report.status == "passed"
    assert second_result.counters["groups"] == 2
    assert first_output != second_output
    assert service.repository.get_task(task_id)["current_version"] == 2
    wb = load_workbook(second_output, data_only=True)
    values = {wb["编号汇总"].cell(row, 1).value: wb["编号汇总"].cell(row, 2).value for row in range(2, 4)}
    assert values["001"] == 30
    wb.close()


def test_v2_sort_rank_and_top_n(workbook_factory, tmp_path):
    path = workbook_factory(main_rows=[("001", 10), ("002", 30), ("003", 20)])
    service = make_service(tmp_path)
    task_id = create_task(service, path)
    envelope = TaskEnvelope(
        goal="订单额排名并保留前两名",
        spec=SortRankTaskSpec(
            sheet="订单",
            sort_keys=[SortKey(column="订单额", direction=SortDirection.DESC)],
            rank_column="订单额",
            rank_method=RankMethod.DENSE,
            top_n=2,
            output_sheet="订单排名",
            confidence=1,
        ),
    )
    output, report, result = execute(service, task_id, envelope)
    assert report.status == "passed"
    assert result.counters["output_rows"] == 2
    wb = load_workbook(output, data_only=True)
    assert wb["订单排名"]["A2"].value == "002"
    assert wb["订单排名"]["C2"].value == 1
    wb.close()


def test_v2_clean_remove_duplicates_and_restore(workbook_factory, tmp_path):
    path = workbook_factory(main_rows=[("001", 10), ("001", 20), ("002", 30)])
    service = make_service(tmp_path)
    task_id = create_task(service, path)
    envelope = TaskEnvelope(
        goal="按客户编号去重",
        spec=CleanTaskSpec(
            sheet="订单",
            operation=CleanOperation.REMOVE_DUPLICATES,
            key_columns=["客户编号"],
            keep="first",
            output_sheet="订单去重",
            confidence=1,
        ),
    )
    output, report, result = execute(service, task_id, envelope)
    assert report.status == "passed"
    assert result.counters["duplicates_removed"] == 1
    wb = load_workbook(output, data_only=True)
    assert wb["订单去重"].max_row == 3
    wb.close()
    profile = service.restore_version(task_id, 0)
    task = service.repository.get_task(task_id)
    assert task["current_version"] == 0
    assert "订单去重" not in [sheet.name for sheet in profile.sheets]


def test_clean_keep_columns_overwrite_requires_confirmation_and_preserves_values(
    workbook_factory, tmp_path
):
    path = workbook_factory(
        main_headers=("患者编号", "年龄", "备注", "年龄_数值"),
        main_rows=[("P001", "60岁", "复查", 60), ("P002", 58, "", 58)],
    )
    service = make_service(tmp_path)
    task_id = create_task(service, path)
    envelope = TaskEnvelope(
        goal="只保留年龄及生成的年龄数值列，删除其他列",
        spec=CleanTaskSpec(
            sheet="订单",
            operation=CleanOperation.KEEP_COLUMNS,
            keep_columns=[
                ColumnRef(sheet="订单", column_index=2, column_letter="B", display_header="年龄"),
                ColumnRef(sheet="订单", column_index=4, column_letter="D", display_header="年龄_数值"),
            ],
            overwrite=True,
            confidence=1,
        ),
    )

    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid, validation.errors
    assert plan.risk_level.value == "high"
    assert plan.confirmation_requirements == ["overwrite_source_columns"]
    with pytest.raises(IntentValidationError):
        service.approve(task_id, "approve")

    service.approve(
        task_id, "approve", {"acknowledgements": ["overwrite_source_columns"]}
    )
    output, report, result = service.execute_task_plan(task_id, envelope, plan)
    assert report.status == "passed"
    assert result.counters["deleted_columns"] == 2
    workbook = load_workbook(output, data_only=False)
    try:
        sheet = workbook["订单"]
        assert sheet.max_column == 2
        assert [sheet.cell(1, column).value for column in range(1, 3)] == ["年龄", "年龄_数值"]
        assert [sheet.cell(2, column).value for column in range(1, 3)] == ["60岁", 60]
        assert workbook["客户"]["A1"].value == "客户编号"
    finally:
        workbook.close()


def test_followup_projection_recovers_from_sheet_name_clarification(workbook_factory, tmp_path):
    path = workbook_factory(
        main_headers=("患者编号", "年龄", "备注", "年龄_数值"),
        main_rows=[("P001", "60岁", "复查", 60)],
    )
    service = make_service(tmp_path)
    task_id = create_task(service, path)
    service.repository.add_message(
        task_id,
        ConversationRole.ASSISTANT.value,
        MessageKind.RESULT.value,
        "上一轮处理完成",
        {
            "result": {
                "details": {
                    "steps": [
                        {
                            "modified_ranges": ["订单!D1:D2"],
                            "details": {
                                "lineage": [
                                    {
                                        "source_columns": [{"column_index": 2}],
                                        "output_column": "年龄_数值",
                                    }
                                ]
                            },
                        }
                    ]
                }
            }
        },
    )
    previous = TaskEnvelope(
        goal="除了我提到的列和系统生成的列，删除其他列",
        spec=CleanTaskSpec(
            sheet="订单",
            operation=CleanOperation.FILTER_ROWS,
            missing_fields=["需要列投影能力"],
            confidence=0.4,
        ),
    )
    service.repository.add_message(
        task_id, ConversationRole.USER.value, MessageKind.TEXT.value, "订单", {}
    )

    envelope = service._build_followup_column_projection(
        previous, "订单", service.get_conversation(task_id), service.load_profile(task_id)
    )

    assert envelope is not None
    assert isinstance(envelope.spec, CleanTaskSpec)
    assert envelope.spec.operation == CleanOperation.KEEP_COLUMNS
    assert envelope.spec.overwrite is True
    assert [item.column_index for item in envelope.spec.keep_columns] == [2, 4]


def test_indicator_status_cleanup_removes_only_mapping_audit_columns(
    workbook_factory, tmp_path
):
    path = workbook_factory(
        main_headers=(
            "患者编号",
            "性别_indicator",
            "性别_indicator_状态",
            "高血压_indicator",
            "高血压_indicator_状态",
            "Tumor_Size_Status",
            "BMI_was_imputed",
        ),
        main_rows=[
            ("P001", 1, "mapped", 0, "mapped_default", "parsed_3d", 0),
            ("P002", 0, "mapped", 1, "mapped", "diameter_only", 1),
        ],
    )
    service = make_service(tmp_path)
    task_id = create_task(service, path)
    previous = TaskEnvelope(
        goal="仅保留上一轮涉及的列，删除其他列",
        spec=CleanTaskSpec(
            sheet="订单",
            operation=CleanOperation.KEEP_COLUMNS,
            keep_columns=[ColumnRef(sheet="订单", column_index=1)],
            overwrite=True,
            confidence=1,
        ),
    )

    envelope = service._build_indicator_status_cleanup(
        previous,
        "生成 indicator 时多出的 mapped 和 mapped_default 状态列也全部删去",
        service.load_profile(task_id),
    )

    assert envelope is not None
    assert envelope.spec.operation == CleanOperation.KEEP_COLUMNS
    kept_headers = [item.display_header for item in envelope.spec.keep_columns]
    assert kept_headers == [
        "患者编号",
        "性别_indicator",
        "高血压_indicator",
        "Tumor_Size_Status",
        "BMI_was_imputed",
    ]
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid, validation.errors
    assert plan.steps[0].expected_changes.summary_counts == {
        "kept_columns": 5,
        "deleted_columns": 2,
    }
    service.approve(
        task_id, "approve", {"acknowledgements": ["overwrite_source_columns"]}
    )
    output, report, result = service.execute_task_plan(task_id, envelope, plan)
    assert report.status == "passed"
    assert result.counters["deleted_columns"] == 2
    workbook = load_workbook(output, data_only=True)
    try:
        headers = [workbook["订单"].cell(1, column).value for column in range(1, 6)]
        assert headers == kept_headers
    finally:
        workbook.close()


def test_completed_projection_does_not_hijack_unrelated_followup(workbook_factory, tmp_path):
    path = workbook_factory()
    service = make_service(tmp_path)
    task_id = create_task(service, path)
    previous = TaskEnvelope(
        goal="只保留我提到的列，删除其他列",
        spec=CleanTaskSpec(
            sheet="订单",
            operation=CleanOperation.KEEP_COLUMNS,
            keep_columns=[ColumnRef(sheet="订单", column_index=1)],
            overwrite=True,
            confidence=1,
        ),
    )
    envelope = service._build_followup_column_projection(
        previous,
        "把 BMI 空值用平均值填充",
        service.get_conversation(task_id),
        service.load_profile(task_id),
    )
    assert envelope is None


def test_rank_max_assigns_same_end_position_to_ties(workbook_factory, tmp_path):
    path = workbook_factory(main_rows=[("001", 30), ("002", 30), ("003", 20)])
    service = make_service(tmp_path)
    task_id = create_task(service, path)
    envelope = TaskEnvelope(
        goal="订单额并列末位排名",
        spec=SortRankTaskSpec(
            sheet="订单",
            rank_column="订单额",
            rank_method=RankMethod.MAX,
            output_sheet="并列排名",
            confidence=1,
        ),
    )
    output, _, _ = execute(service, task_id, envelope)
    wb = load_workbook(output, data_only=True)
    ranks = [wb["并列排名"].cell(row, 3).value for row in range(2, 5)]
    wb.close()
    assert ranks == [2, 2, 3]


def test_context_builder_sends_structure_not_preview_rows(workbook_factory):
    path = workbook_factory()
    profile = profile_workbook(path, "input.xlsx", 100_000)
    context = build_context(profile, "按地区汇总订单额")
    assert context["task_type_hint"] == "aggregate"
    assert "preview_rows" not in str(context)
    assert context["sheets"][0]["columns"][1]["type"] == "number"

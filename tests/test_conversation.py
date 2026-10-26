from __future__ import annotations

from pathlib import Path

from sheetpilot.config import Settings
from sheetpilot.llm import FakeLLMClient
from sheetpilot.schemas import (
    ConversationState,
    DuplicatePolicy,
    FormulaMode,
    FormulaTaskSpec,
    LookupTaskSpec,
    MessageKind,
    TaskEnvelope,
)
from sheetpilot.service import SheetPilotService


def make_service(tmp_path: Path, envelope: TaskEnvelope) -> SheetPilotService:
    settings = Settings(
        project_root=tmp_path,
        data_dir=tmp_path / "data",
        database_path=tmp_path / "data" / "sheetpilot.sqlite3",
        llm_provider="fake",
        api_key="",
        base_url="",
        model="fake",
        max_file_bytes=20 * 1024 * 1024,
        max_sheet_rows=100_000,
    )
    return SheetPilotService(settings, FakeLLMClient(envelope))


def lookup_envelope(*, return_columns=None, duplicate_policy=DuplicatePolicy.BLOCK):
    return TaskEnvelope(
        goal="把客户等级补充到订单表",
        spec=LookupTaskSpec(
            main_sheet="订单",
            source_sheet="客户",
            main_key_column="客户编号",
            source_key_column="客户编号",
            return_columns=return_columns or [],
            output_column_names=return_columns or [],
            duplicate_policy=duplicate_policy,
            confidence=1,
        ),
    )


def test_upload_creates_persistent_workbook_message(workbook_factory, tmp_path):
    service = make_service(tmp_path, lookup_envelope(return_columns=["客户等级"]))
    source = workbook_factory()
    profile = service.create_task(source.name, source.read_bytes())
    task_id = service.task_id_from_profile(profile)

    messages = service.get_conversation(task_id)

    assert len(messages) == 1
    assert messages[0].kind == MessageKind.WORKBOOK
    assert "告诉我你希望怎样处理" in messages[0].content


def test_valid_natural_language_turn_creates_plan_and_can_execute(workbook_factory, tmp_path):
    service = make_service(tmp_path, lookup_envelope(return_columns=["客户等级"]))
    source = workbook_factory()
    profile = service.create_task(source.name, source.read_bytes())
    task_id = service.task_id_from_profile(profile)

    turn = service.submit_message(task_id, "按客户编号补充客户等级")
    assert turn.state == ConversationState.PLAN_READY
    assert turn.actions == ["approve", "edit", "cancel"]
    assert "return_columns" not in turn.reply.content

    executed = service.approve_current_plan(task_id)
    assert executed.state == ConversationState.RESULT_READY
    assert executed.version == 1
    assert service.get_conversation(task_id)[-1].kind == MessageKind.RESULT


def test_missing_return_columns_becomes_chinese_clarification(workbook_factory, tmp_path):
    service = make_service(tmp_path, lookup_envelope())
    source = workbook_factory()
    profile = service.create_task(source.name, source.read_bytes())
    task_id = service.task_id_from_profile(profile)

    turn = service.submit_message(task_id, "按客户编号匹配")

    assert turn.state == ConversationState.NEEDS_CLARIFICATION
    assert turn.clarification is not None
    assert turn.clarification.question == "你希望从来源表补充哪些列？"
    assert "客户等级" in turn.clarification.choices
    assert "return_columns" not in turn.reply.content


def test_cancel_plan_does_not_create_version(workbook_factory, tmp_path):
    service = make_service(tmp_path, lookup_envelope(return_columns=["客户等级"]))
    source = workbook_factory()
    profile = service.create_task(source.name, source.read_bytes())
    task_id = service.task_id_from_profile(profile)
    service.submit_message(task_id, "按客户编号补充客户等级")

    service.cancel_current_plan(task_id)

    workspace = service.get_workspace(task_id)
    assert workspace["task"]["current_version"] == 0
    assert workspace["task"]["status"] == "profiled"
    assert workspace["messages"][-1]["content"].startswith("已取消")


def test_natural_language_revision_receives_previous_complete_envelope(workbook_factory, tmp_path):
    first = lookup_envelope(return_columns=["客户等级"])
    revised = lookup_envelope(
        return_columns=["地区"], duplicate_policy=DuplicatePolicy.LAST
    )

    class RecordingLLM:
        def __init__(self):
            self.calls = []

        def parse_task(self, request, profile, **kwargs):
            self.calls.append((request, kwargs))
            return first if len(self.calls) == 1 else revised

    settings = Settings(
        project_root=tmp_path,
        data_dir=tmp_path / "data",
        database_path=tmp_path / "data" / "sheetpilot.sqlite3",
        llm_provider="fake",
        api_key="",
        base_url="",
        model="fake",
        max_file_bytes=20 * 1024 * 1024,
        max_sheet_rows=100_000,
    )
    llm = RecordingLLM()
    service = SheetPilotService(settings, llm)
    source = workbook_factory()
    profile = service.create_task(source.name, source.read_bytes())
    task_id = service.task_id_from_profile(profile)
    service.submit_message(task_id, "补充客户等级")

    turn = service.submit_message(task_id, "改成补充地区，重复编号使用最后一条")

    previous = llm.calls[1][1]["previous_envelope"]
    assert previous.spec.return_columns == ["客户等级"]
    assert turn.envelope.spec.return_columns == ["地区"]
    assert turn.envelope.spec.duplicate_policy == DuplicatePolicy.LAST
    assert turn.state == ConversationState.PLAN_READY


def test_formula_inference_does_not_ask_for_formula_text(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "计算"
    sheet.append(["数量", "单价", "金额"])
    sheet.append([2, 10, "=A2*B2"])
    sheet.append([3, 10, "=A3*B3"])
    sheet.append([4, 10, None])
    sheet.append([5, 10, "=A5+B5"])
    sheet.append([6, 10, "=A6+B6"])
    workbook.create_sheet("参数").append(["参数", "值"])
    workbook.save(source)
    workbook.close()

    envelope = TaskEnvelope(
        goal="填补计算表金额列第2到第6行的公式空白",
        spec=FormulaTaskSpec(
            sheet="计算",
            column="金额",
            start_row=2,
            end_row=6,
            mode=FormulaMode.FILL_GAPS,
            ambiguities=["用户未提供具体公式文本，系统将根据上下文推断候选公式"],
            confidence=0.95,
        ),
    )
    service = make_service(tmp_path / "service", envelope)
    profile = service.create_task(source.name, source.read_bytes())
    task_id = service.task_id_from_profile(profile)

    turn = service.submit_message(
        task_id, "尝试填补‘计算’表‘金额’列第 2 到第 6 行的公式空白"
    )

    assert turn.state == ConversationState.PLAN_READY
    assert turn.needs_clarification is False
    assert turn.plan is not None
    assert turn.plan.steps[0].tool_name == "formula"


def test_formula_inference_still_blocks_a_missing_real_column(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "formula_missing_column.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "计算"
    sheet.append(["数量", "单价", "金额"])
    sheet.append([2, 10, "=A2*B2"])
    workbook.create_sheet("参数").append(["参数", "值"])
    workbook.save(source)
    workbook.close()
    envelope = TaskEnvelope(
        goal="填补公式",
        spec=FormulaTaskSpec(
            sheet="计算", column="不存在的列", start_row=2, end_row=2,
            mode=FormulaMode.FILL_GAPS, confidence=0.8,
        ),
    )
    service = make_service(tmp_path / "service_missing", envelope)
    profile = service.create_task(source.name, source.read_bytes())
    task_id = service.task_id_from_profile(profile)

    turn = service.submit_message(task_id, "填补公式")

    assert turn.state == ConversationState.NEEDS_CLARIFICATION

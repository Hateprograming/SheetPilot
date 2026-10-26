from __future__ import annotations

import os

import pytest
from openpyxl import Workbook, load_workbook

from sheetpilot.calculation import recalculate_workbook
from sheetpilot.profiler import profile_workbook
from sheetpilot.schemas import FormulaTaskSpec
from sheetpilot.tools import FormulaTool, ToolContext


@pytest.mark.skipif(
    os.environ.get("SHEETPILOT_EXCEL_INTEGRATION") != "1",
    reason="requires a local Microsoft Excel/LibreOffice calculation engine",
)
def test_real_calculation_engine_writes_cached_values(tmp_path):
    path = tmp_path / "calculation.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "计算"
    sheet.append(["数量", "单价", "金额"])
    sheet.append([2, 15, "=A2*B2"])
    sheet.append([4, 12.5, "=A3*B3"])
    workbook.create_sheet("说明").append(["项目", "内容"])
    workbook.save(path)
    workbook.close()

    result = recalculate_workbook(path)
    assert result.success, result.message
    assert result.calculated_cells == 2
    values = load_workbook(path, data_only=True)
    assert values["计算"]["C2"].value == 30
    assert values["计算"]["C3"].value == 50
    values.close()


@pytest.mark.skipif(os.environ.get("SHEETPILOT_EXCEL_INTEGRATION") != "1", reason="requires Excel/LibreOffice")
@pytest.mark.parametrize("rows", [3_000, 5_000, 40_000])
def test_real_engine_full_formula_pipeline(tmp_path, rows):
    path = tmp_path / "source.xlsx"
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("计算")
    sheet.append(["数量", "单价", "金额"])
    for _ in range(rows):
        sheet.append([2, 15, None])
    workbook.save(path)
    profile = profile_workbook(path, path.name, 100_000)
    context = ToolContext(task_id="real-engine", input_path=path, output_path=tmp_path / "result.xlsx", profile=profile, version_before=0, version_after=1)
    spec = FormulaTaskSpec(sheet="计算", column="金额", start_row=2, end_row=rows+1, mode="set_formula", formula="=A2*B2", recalculate=True, confidence=1)
    tool = FormulaTool()
    result = tool.execute(context, spec)
    assert tool.verify(context, spec, result).valid, result.warnings
    assert result.counters["calculated"] == rows
    values = load_workbook(context.output_path, read_only=True, data_only=True)
    try:
        assert all(row[2] == 30 for row in values["计算"].iter_rows(min_row=2, values_only=True))
    finally:
        values.close()

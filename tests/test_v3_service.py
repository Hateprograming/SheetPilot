from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from sheetpilot.calculation import RecalculationResult
from sheetpilot.config import Settings
from sheetpilot.errors import IntentValidationError
from sheetpilot.schemas import (
    AggregateFunction,
    AggregateMetric,
    AggregateTaskSpec,
    CompositeTaskSpec,
    FormulaMode,
    FormulaTaskSpec,
    NormalizeOperation,
    NormalizeTaskSpec,
    SortDirection,
    SortKey,
    SortRankTaskSpec,
    TaskEnvelope,
)
from sheetpilot.service import SheetPilotService


def make_service(tmp_path: Path) -> SheetPilotService:
    return SheetPilotService(Settings(
        project_root=tmp_path,
        data_dir=tmp_path / "data",
        database_path=tmp_path / "data" / "sheetpilot.sqlite3",
        llm_provider="fake", api_key="", base_url="", model="fake",
        max_file_bytes=20 * 1024 * 1024, max_sheet_rows=100_000,
    ))


def create_task(service: SheetPilotService, source: Path) -> str:
    profile = service.create_task(source.name, source.read_bytes())
    return service.task_id_from_profile(profile)


def composite_envelope() -> TaskEnvelope:
    return TaskEnvelope(
        goal="统一编号，然后汇总，最后排序",
        spec=CompositeTaskSpec(
            confidence=1,
            operations=[
                NormalizeTaskSpec(
                    sheet="订单", column="客户编号",
                    operations=[NormalizeOperation.NFKC, NormalizeOperation.TRIM],
                    output_column="客户编号_标准化", confidence=1,
                ),
                AggregateTaskSpec(
                    sheet="订单", group_by=["客户编号_标准化"],
                    metrics=[AggregateMetric(column="订单额", function=AggregateFunction.SUM, output_column="订单额合计")],
                    output_sheet="编号汇总", confidence=1,
                ),
                SortRankTaskSpec(
                    sheet="编号汇总",
                    sort_keys=[SortKey(column="订单额合计", direction=SortDirection.DESC)],
                    output_sheet="汇总排名", confidence=1,
                ),
            ],
        ),
    )


def test_composite_task_commits_one_visible_version(workbook_factory, tmp_path):
    source = workbook_factory(main_rows=[(" ００１ ", 10), ("００１", 20), ("002", 5)])
    service = make_service(tmp_path)
    task_id = create_task(service, source)
    envelope = composite_envelope()
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid, validation.errors
    assert len(plan.steps) == 3
    service.approve(task_id, "approve")
    output, report, result = service.execute_task_plan(task_id, envelope, plan)

    assert report.status == "passed"
    assert result.tool_name == "transaction"
    assert service.repository.get_task(task_id)["current_version"] == 1
    assert len(service.repository.list_versions(task_id)) == 2
    assert service.repository.latest_plan_run(task_id)["status"] == "completed"
    wb = load_workbook(output, data_only=True)
    assert wb["编号汇总"]["B2"].value == 30
    assert wb["汇总排名"]["B2"].value == 30
    wb.close()


def test_transient_second_step_retries_without_repeating_first(workbook_factory, tmp_path, monkeypatch):
    source = workbook_factory(main_rows=[("001", 10), ("001", 20)])
    service = make_service(tmp_path)
    task_id = create_task(service, source)
    envelope = TaskEnvelope(
        goal="标准化后汇总",
        spec=CompositeTaskSpec(confidence=1, operations=composite_envelope().spec.operations[:2]),
    )
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid, validation.errors
    aggregate = service.tools.get("aggregate")
    original_execute = aggregate.execute
    calls = {"count": 0}

    def flaky(context, spec):
        calls["count"] += 1
        if calls["count"] == 1:
            raise OSError("temporary file lock")
        return original_execute(context, spec)

    monkeypatch.setattr(aggregate, "execute", flaky)
    service.approve(task_id, "approve")
    service.execute_task_plan(task_id, envelope, plan)
    runs = service.repository.list_tool_runs(task_id)
    assert [item["step_order"] for item in runs] == [1, 2, 2]
    assert [item["attempt"] for item in runs] == [1, 1, 2]
    assert runs[1]["status"] == "retrying"
    assert runs[2]["status"] == "completed"


def test_failed_run_resumes_after_service_restart(workbook_factory, tmp_path, monkeypatch):
    source = workbook_factory(main_rows=[("001", 10), ("001", 20)])
    service = make_service(tmp_path)
    task_id = create_task(service, source)
    envelope = TaskEnvelope(
        goal="标准化后汇总",
        spec=CompositeTaskSpec(confidence=1, operations=composite_envelope().spec.operations[:2]),
    )
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid
    aggregate = service.tools.get("aggregate")

    def unavailable(context, spec):
        raise OSError("temporary service outage")

    monkeypatch.setattr(aggregate, "execute", unavailable)
    service.approve(task_id, "approve")
    with pytest.raises(OSError):
        service.execute_task_plan(task_id, envelope, plan)
    failed_run = service.repository.latest_plan_run(task_id)
    assert failed_run["status"] == "failed"
    assert failed_run["current_step"] == 1

    restarted = make_service(tmp_path)
    output, report, _ = restarted.resume_current_run(task_id)
    assert report.status == "passed"
    assert output.exists()
    assert restarted.repository.get_task(task_id)["current_version"] == 1
    runs = restarted.repository.list_tool_runs(task_id)
    assert sum(item["step_order"] == 1 for item in runs) == 1


def test_revised_plan_invalidates_old_approval(workbook_factory, tmp_path):
    source = workbook_factory()
    service = make_service(tmp_path)
    task_id = create_task(service, source)
    first = TaskEnvelope(
        goal="标准化编号",
        spec=NormalizeTaskSpec(
            sheet="订单", column="客户编号", operations=[NormalizeOperation.TRIM],
            output_column="编号A", confidence=1,
        ),
    )
    old_plan, validation = service.make_task_plan(task_id, first)
    assert validation.valid
    revised = first.model_copy(deep=True)
    revised.spec.output_column = "编号B"
    new_plan, validation = service.make_task_plan(task_id, revised)
    assert validation.valid
    assert new_plan.plan_id == old_plan.plan_id
    assert new_plan.revision == old_plan.revision + 1
    service.approve(task_id, "approve")
    with pytest.raises(IntentValidationError, match="旧审批"):
        service.execute_task_plan(task_id, first, old_plan)


def test_plan_cycle_is_rejected(workbook_factory, tmp_path):
    source = workbook_factory()
    service = make_service(tmp_path)
    task_id = create_task(service, source)
    envelope = composite_envelope()
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid
    plan.steps[0].depends_on = [plan.steps[-1].id]
    service.approve(task_id, "approve")
    with pytest.raises(IntentValidationError, match="结构无效"):
        service.execute_task_plan(task_id, envelope, plan)


def formula_workbook(path: Path, *, ambiguous: bool = False) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "计算"
    ws.append(["数量", "单价", "金额"])
    for row in range(2, 7):
        ws.cell(row, 1, row)
        ws.cell(row, 2, 10)
    ws["C2"] = "=A2*B2"
    ws["C3"] = "=A3*B3"
    ws["C5"] = "=A5*B5" if not ambiguous else "=A5+B5"
    ws["C6"] = "=A6*B6" if not ambiguous else "=A6+B6"
    wb.create_sheet("辅助").append(["占位"])
    wb.save(path)
    wb.close()
    return path


def test_formula_gap_is_filled_and_recalculation_is_recorded(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "sheetpilot.tools.recalculate_workbook",
        lambda path, **kwargs: RecalculationResult(
            success=True, backend="test", formula_cells=5, calculated_cells=5,
            message="test recalculation",
        ),
    )
    source = formula_workbook(tmp_path / "formula.xlsx")
    service = make_service(tmp_path)
    task_id = create_task(service, source)
    envelope = TaskEnvelope(goal="填补金额公式断层", spec=FormulaTaskSpec(
        sheet="计算", column="金额", start_row=2, end_row=6,
        mode=FormulaMode.FILL_GAPS, confidence=1,
    ))
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid, validation.errors
    service.approve(task_id, "approve")
    output, report, result = service.execute_task_plan(task_id, envelope, plan)
    wb = load_workbook(output, data_only=False)
    assert wb["计算"]["C4"].value == "=A4*B4"
    wb.close()
    assert result.counters["repaired"] == 1
    assert result.counters["calculated"] == 5
    assert result.details["steps"][0]["details"]["recalculated"] is True
    assert report.status == "passed"


def test_ambiguous_formula_candidate_is_not_written(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "sheetpilot.tools.recalculate_workbook",
        lambda path, **kwargs: RecalculationResult(
            success=True, backend="test", formula_cells=4, calculated_cells=4,
            message="test recalculation",
        ),
    )
    source = formula_workbook(tmp_path / "ambiguous.xlsx", ambiguous=True)
    service = make_service(tmp_path)
    task_id = create_task(service, source)
    envelope = TaskEnvelope(goal="尝试填补公式", spec=FormulaTaskSpec(
        sheet="计算", column="金额", start_row=2, end_row=6,
        mode=FormulaMode.FILL_GAPS, confidence=1,
    ))
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid
    service.approve(task_id, "approve")
    output, _, result = service.execute_task_plan(task_id, envelope, plan)
    wb = load_workbook(output, data_only=False)
    assert wb["计算"]["C4"].value is None
    wb.close()
    assert result.counters["repaired"] == 0
    assert result.details["steps"][0]["details"]["ambiguous_rows"] == [4]


def test_formula_set_copy_and_recalculate_modes(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "sheetpilot.tools.recalculate_workbook",
        lambda path, **kwargs: RecalculationResult(
            success=True, backend="test", formula_cells=5, calculated_cells=5,
            message="test recalculation",
        ),
    )
    source = formula_workbook(tmp_path / "formula_modes.xlsx")
    service = make_service(tmp_path / "service")
    task_id = create_task(service, source)
    envelope = TaskEnvelope(goal="重建金额公式", spec=FormulaTaskSpec(
        sheet="计算", column="金额", start_row=2, end_row=6,
        mode=FormulaMode.SET_FORMULA, formula="=A2*B2",
        overwrite_existing=True, confidence=1,
    ))
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid, validation.errors
    assert plan.risk_level.value == "high"
    service.approve(task_id, "approve")
    output, _, result = service.execute_task_plan(task_id, envelope, plan)
    wb = load_workbook(output, data_only=False)
    assert [wb["计算"][f"C{row}"].value for row in range(2, 7)] == [
        "=A2*B2", "=A3*B3", "=A4*B4", "=A5*B5", "=A6*B6",
    ]
    wb.close()
    assert result.counters["generated"] == 5

    recalc = TaskEnvelope(goal="重算工作簿", spec=FormulaTaskSpec(
        mode=FormulaMode.RECALCULATE, confidence=1,
    ))
    recalc_plan, validation = service.make_task_plan(task_id, recalc)
    assert validation.valid, validation.errors
    assert recalc_plan.risk_level.value == "low"
    service.approve(task_id, "approve")
    _, _, recalc_result = service.execute_task_plan(task_id, recalc, recalc_plan)
    assert recalc_result.details["steps"][0]["details"]["recalculated"] is True

    copy_source = formula_workbook(tmp_path / "formula_copy.xlsx")
    copy_service = make_service(tmp_path / "copy_service")
    copy_task_id = create_task(copy_service, copy_source)
    copy_envelope = TaskEnvelope(goal="复制样板公式", spec=FormulaTaskSpec(
        sheet="计算", column="金额", start_row=2, end_row=6,
        mode=FormulaMode.COPY_FORMULA, source_row=3,
        overwrite_existing=False, confidence=1,
    ))
    copy_plan, validation = copy_service.make_task_plan(copy_task_id, copy_envelope)
    assert validation.valid, validation.errors
    assert copy_plan.risk_level.value == "medium"
    copy_service.approve(copy_task_id, "approve")
    copy_output, _, copy_result = copy_service.execute_task_plan(
        copy_task_id, copy_envelope, copy_plan
    )
    copy_workbook = load_workbook(copy_output, data_only=False)
    assert copy_workbook["计算"]["C4"].value == "=A4*B4"
    copy_workbook.close()
    assert copy_result.counters["copied"] == 1

from __future__ import annotations

import csv
import io
import sys
import time
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from sheetpilot import calculation
from sheetpilot.errors import VerificationError
from sheetpilot.formula_intent import column_formula
from sheetpilot.profiler import profile_workbook
from sheetpilot.schemas import (
    CompositeTaskSpec, ConversionTaskSpec, FormulaTaskSpec, NormalizeTaskSpec,
    TaskEnvelope,
)
from sheetpilot.service import SheetPilotService
from sheetpilot.tools import FormulaTool, ToolContext
from sheetpilot.web_api import _download_response, _safe_workspace
from test_v1_skills import make_service


def workbook_bytes(rows=None, multi=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet.append(["编号", "数量", "单价", "总额"])
    for row in rows or [[" 001 ", 2, 15, None], ["　００２ ", 4, 12.5, None]]:
        sheet.append(row)
    if multi:
        book.create_sheet("其他").append(["备注"])
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def upload(service, content=None):
    profile = service.create_task("测试.xlsx", content or workbook_bytes())
    return service.task_id_from_profile(profile)


def execute(service, task_id, envelope):
    plan, valid = service.make_task_plan(task_id, envelope)
    assert valid.valid, valid.errors
    service.approve(task_id, "approve", {"acknowledgements": plan.confirmation_requirements})
    return service.execute_task_plan(task_id, envelope, plan)


def cache_products(path, timeout):
    """Deterministic stand-in for an engine, exercising real XLSX cache auditing."""
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    updated = path.with_suffix(".cached.xlsx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(updated, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.startswith("xl/worksheets/sheet") and info.filename.endswith(".xml"):
                root = ET.fromstring(data)
                for row in root.iter(ns + "row"):
                    cells = {c.get("r"): c for c in row.findall(ns + "c")}
                    number = row.get("r")
                    for cell in cells.values():
                        if cell.find(ns + "f") is not None:
                            left = float(cells[f"B{number}"].find(ns + "v").text)
                            right = float(cells[f"C{number}"].find(ns + "v").text)
                            value = cell.find(ns + "v")
                            if value is None:
                                value = ET.SubElement(cell, ns + "v")
                            value.text = str(left * right)
                data = ET.tostring(root, encoding="utf-8")
            target.writestr(info, data)
    updated.replace(path)
    return True, "测试引擎完成计算"


@pytest.mark.parametrize("query, expected", [
    ("标准化数据表的编号这一列", ["001", "002"]),
    ("编号这一列去掉两端空格", ["001", "００２"]),
    ("编号这一列清除首尾空白", ["001", "００２"]),
    ('编号这一列“ 001 ”→“001”', ["001", "００２"]),
    ("编号原列仅做 NFKC", [" 001 ", " 002 "]),
    ("编号这一列标准化，保留前导零，不要转为数字", ["001", "002"]),
])
def test_normalization_query_to_actual_download(tmp_path, query, expected):
    model = TaskEnvelope(goal=query, spec=NormalizeTaskSpec(
        sheet="数据", column="编号", operations=["nfkc", "to_number"], confidence=1,
    ))
    service = make_service(tmp_path, model)
    task_id = upload(service)
    turn = service.submit_message(task_id, query)
    assert turn.plan is not None, turn
    assert turn.plan.steps[0].arguments["overwrite"]
    service.approve_current_plan(task_id, acknowledgements=turn.plan.confirmation_requirements)
    name, content = service.get_download(task_id)
    assert name.endswith(".xlsx")
    book = load_workbook(io.BytesIO(content))
    assert [book["数据"].cell(row, 1).value for row in (2, 3)] == expected
    assert book["数据"]["A2"].data_type == "s"
    book.close()


def test_csv_primary_delivery_is_persisted_per_version(tmp_path):
    service = make_service(tmp_path)
    original = workbook_bytes([["001", 2, 15, None], ['含逗号,换行\n中文', 4, 12.5, None]])
    task_id = upload(service, original)
    export = TaskEnvelope(goal="导出 CSV", spec=ConversionTaskSpec(direction="xlsx_to_csv", source_sheet="数据", confidence=1))
    execute(service, task_id, export)
    first_name, first = service.get_download(task_id)
    assert first_name.endswith(".csv") and first.startswith(b"\xef\xbb\xbf")
    assert list(csv.reader(io.StringIO(first.decode("utf-8-sig"))))[2][0] == '含逗号,换行\n中文'
    response = _download_response(first_name, first)
    assert response.headers["content-type"].startswith("text/csv")
    assert ".csv" in response.headers["content-disposition"]
    assert service.get_download(task_id, original=True)[1] == original
    changed = TaskEnvelope(goal="标准化", spec=NormalizeTaskSpec(sheet="数据", column="编号", operations=["trim"], overwrite=True, confidence=1))
    execute(service, task_id, changed)
    assert service.get_download(task_id)[0].endswith(".xlsx")
    execute(service, task_id, export)  # Same attachment name, distinct immutable delivery.
    service = SheetPilotService(service.settings)
    assert service.get_version_download(task_id, 1) == (first_name, first)
    service.restore_version(task_id, 1)
    assert service.get_download(task_id) == (first_name, first)
    record = service.repository.get_version(task_id, 1)
    import json
    delivery = json.loads(record["delivery_json"])
    Path(delivery["path"]).write_bytes(b"corrupt")
    with pytest.raises(VerificationError):
        service.get_version_download(task_id, 1)
    assert "delivery_json" not in _safe_workspace({"delivery_json": record["delivery_json"]})


def test_conversion_sheet_selection_is_grounded(tmp_path):
    for multi in (False, True):
        model = TaskEnvelope(goal="转换 CSV", spec=ConversionTaskSpec(direction="xlsx_to_csv", source_sheet="数据", confidence=1))
        service = make_service(tmp_path / str(multi), model)
        task_id = upload(service, workbook_bytes(multi=multi))
        turn = service.submit_message(task_id, "转换成 CSV")
        if multi:
            assert turn.clarification and turn.clarification.field == "source_sheet"
        else:
            assert turn.plan is not None


def test_csv_commit_interruption_can_be_retried(tmp_path, monkeypatch):
    service = make_service(tmp_path)
    task_id = upload(service)
    envelope = TaskEnvelope(goal="CSV", spec=ConversionTaskSpec(direction="xlsx_to_csv", source_sheet="数据", output_name="交付.csv", confidence=1))
    def fail_commit(**kwargs):
        raise RuntimeError("模拟数据库提交中断")
    monkeypatch.setattr(service.repository, "complete_version_commit", fail_commit)
    with pytest.raises(RuntimeError):
        execute(service, task_id, envelope)
    assert service.repository.get_task(task_id)["current_version"] == 0
    service = SheetPilotService(service.settings)
    service.resume_current_run(task_id)
    assert service.get_download(task_id)[0] == "交付.csv"
    assert service.get_workspace(task_id)["current_operation"]["delivery_status"] == "ready"
    export = service.files.version_path(task_id, 1).with_suffix(".csv")
    export.write_bytes(b"corrupt")
    assert service.get_workspace(task_id)["current_operation"]["delivery_status"] == "unavailable"


@pytest.mark.parametrize("query", ["在“总额”这一列填入数量×单价", "总额=数量*单价"])
def test_column_arithmetic_runs_and_exports_cached_values(tmp_path, monkeypatch, query):
    monkeypatch.setattr(calculation, "_recalculate_with_excel", cache_products)
    model = TaskEnvelope(goal=query, spec=FormulaTaskSpec(sheet="数据", column="总额", missing_fields=["formula"], confidence=1))
    service = make_service(tmp_path, model)
    task_id = upload(service)
    turn = service.submit_message(task_id, query)
    assert turn.plan is not None, turn
    service.approve_current_plan(task_id)
    _, data = service.get_download(task_id)
    formulas = load_workbook(io.BytesIO(data), data_only=False)
    assert formulas["数据"]["D2"].value == "=(B2*C2)"
    assert formulas["数据"]["D3"].value == "=(B3*C3)"
    formulas.close()
    values = load_workbook(io.BytesIO(data), data_only=True)
    assert values["数据"]["D2"].value == 30
    assert values["数据"]["D3"].value == 50
    values.close()
    execute(service, task_id, TaskEnvelope(goal="CSV", spec=ConversionTaskSpec(direction="xlsx_to_csv", source_sheet="数据", confidence=1)))
    rows = list(csv.reader(io.StringIO(service.get_download(task_id)[1].decode("utf-8-sig"))))
    assert float(rows[1][-1]) == 30 and float(rows[2][-1]) == 50


def test_normalize_then_export_uses_final_data(tmp_path):
    service = make_service(tmp_path)
    task_id = upload(service)
    envelope = TaskEnvelope(goal="清理后导出 CSV", spec=CompositeTaskSpec(operations=[
        NormalizeTaskSpec(sheet="数据", column="编号", operations=["nfkc", "trim"], overwrite=True, confidence=1),
        ConversionTaskSpec(direction="xlsx_to_csv", source_sheet="数据", confidence=1),
    ], confidence=1))
    execute(service, task_id, envelope)
    content = service.get_download(task_id)[1]
    assert list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))[1][0] == "001"
    envelope.spec.operations.reverse()
    _, valid = service.make_task_plan(task_id, envelope)
    assert not valid.valid


def test_unavailable_engine_pauses_without_publishing(tmp_path, monkeypatch):
    monkeypatch.setattr(calculation, "_recalculate_with_excel", lambda *a: (False, "Excel 不可用"))
    monkeypatch.setattr(calculation, "_recalculate_with_libreoffice", lambda *a: (False, "LibreOffice 不可用"))
    service = make_service(tmp_path)
    task_id = upload(service)
    envelope = TaskEnvelope(goal="计算", spec=FormulaTaskSpec(sheet="数据", column="总额", start_row=2, end_row=3, mode="set_formula", formula="=B2*C2", confidence=1))
    with pytest.raises(VerificationError):
        execute(service, task_id, envelope)
    assert service.repository.get_task(task_id)["current_version"] == 0
    run = service.repository.latest_plan_run(task_id)
    assert run["status"] == "paused"
    assert len(service.repository.list_versions(task_id)) == 1
    # Recovery re-runs the failed step from the previous verified input.
    monkeypatch.setattr(calculation, "_recalculate_with_excel", cache_products)
    service.resume_current_run(task_id)
    assert service.repository.get_task(task_id)["current_version"] == 1


def test_cache_audit_distinguishes_empty_string_missing_and_error(tmp_path):
    path = tmp_path / "cache.xlsx"
    book = Workbook()
    for row in range(1, 5):
        book.active.cell(row, 1, '=IF(1,"",1)')
    book.save(path)
    updated = path.with_suffix(".new.xlsx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(updated, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<c r="A1">', b'<c r="A1" t="str">')
                data = data.replace(b'<c r="A3">', b'<c r="A3" t="e">')
                data = data.replace(b'<c r="A4">', b'<c r="A4" t="n">')
                root = ET.fromstring(data)
                ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                cells = list(root.iter(ns + "c"))
                cells[2].find(ns + "v").text = "#DIV/0!"
                cells[3].find(ns + "v").text = "42"
                data = ET.tostring(root)
            target.writestr(info, data)
    updated.replace(path)
    audit = calculation._audit_cached_values(path)
    assert not audit.success
    assert audit.empty_string_cells == 1 and audit.calculated_cells == 2
    assert audit.blank_cached_cells == ["Sheet!A2"]
    assert audit.error_cells == {"Sheet!A3": "#DIV/0!"}


def test_backend_timeout_and_cancel_stop_owned_process():
    started = time.monotonic()
    with pytest.raises(Exception) as exc:
        calculation._run_backend([sys.executable, "-c", "import time; time.sleep(60)"], timeout=.2)
    assert exc.type is __import__("subprocess").TimeoutExpired
    token = calculation._control.set(lambda: True)
    try:
        with pytest.raises(InterruptedError):
            calculation._run_backend([sys.executable, "-c", "import time; time.sleep(60)"], timeout=90)
    finally:
        calculation._control.reset(token)
    assert time.monotonic() - started < 5


def test_backends_share_one_budget(tmp_path, monkeypatch):
    path = tmp_path / "budget.xlsx"
    path.write_bytes(workbook_bytes())
    clock = [0.0]
    budgets = []
    monkeypatch.setattr(calculation.time, "monotonic", lambda: clock[0])
    def first(path, timeout):
        budgets.append(timeout)
        clock[0] += 80
        return False, "第一个引擎失败"
    def second(path, timeout):
        budgets.append(timeout)
        return True, "第二个引擎完成"
    monkeypatch.setattr(calculation, "_recalculate_with_excel", first)
    monkeypatch.setattr(calculation, "_recalculate_with_libreoffice", second)
    assert calculation.recalculate_workbook(path, timeout=90).success
    assert budgets == [90, 10]


def test_cancellation_during_recalculation_does_not_publish(tmp_path, monkeypatch):
    import threading
    started = threading.Event()
    errors = []
    def engine(path, timeout):
        started.set()
        calculation._run_backend([sys.executable, "-c", "import time; time.sleep(60)"], timeout=timeout)
        return True, "不可到达"
    monkeypatch.setattr(calculation, "_recalculate_with_excel", engine)
    service = make_service(tmp_path)
    task_id = upload(service)
    envelope = TaskEnvelope(goal="计算", spec=FormulaTaskSpec(sheet="数据", column="总额", start_row=2, end_row=3, mode="set_formula", formula="=B2*C2", confidence=1))
    def run():
        try:
            execute(service, task_id, envelope)
        except Exception as exc:
            errors.append(exc)
    worker = threading.Thread(target=run)
    worker.start()
    assert started.wait(10)
    service.control_current_run(task_id, "cancel")
    worker.join(timeout=10)
    assert not worker.is_alive() and errors
    assert service.repository.latest_plan_run(task_id)["status"] == "cancelled"
    assert service.repository.get_task(task_id)["current_version"] == 0


def test_nonfirst_header_and_existing_total_are_preserved(tmp_path, monkeypatch):
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet.append(["说明"])
    sheet.append([None])
    sheet.append(["编号", "数量", "单价", "总额"])
    sheet.append(["001", 2, 15, 999])
    sheet.append(["002", 4, 12.5, None])
    stream = io.BytesIO()
    book.save(stream)
    model = TaskEnvelope(goal="填入总额", spec=FormulaTaskSpec(sheet="数据", column="总额", confidence=1))
    service = make_service(tmp_path, model)
    task_id = upload(service, stream.getvalue())
    assert service.load_profile(task_id).sheets[0].header_row == 3
    turn = service.submit_message(task_id, "在总额列填入数量×单价")
    assert turn.plan.steps[0].arguments["start_row"] == 4
    monkeypatch.setattr(calculation, "_recalculate_with_excel", cache_products)
    service.approve_current_plan(task_id)
    content = service.get_download(task_id)[1]
    values = load_workbook(io.BytesIO(content), data_only=True)
    assert values["数据"]["D4"].value == 999
    assert values["数据"]["D5"].value == 50
    values.close()


def test_unknown_column_expression_requires_clarification(tmp_path):
    model = TaskEnvelope(goal="计算", spec=FormulaTaskSpec(sheet="数据", column="总额", mode="set_formula", formula="=未知*单价", confidence=1))
    service = make_service(tmp_path, model)
    task_id = upload(service)
    turn = service.submit_message(task_id, "总额=未知*单价")
    assert turn.clarification and turn.clarification.field == "formula"
    assert "计算规则" in turn.clarification.question
    resumed = service.submit_message(task_id, "数量×单价")
    assert resumed.plan is not None, resumed
    assert resumed.plan.steps[0].arguments["formula"] == "=(B2*C2)"


def test_normalization_preserves_internal_spaces_null_and_numeric_type(tmp_path):
    model = TaskEnvelope(goal="标准化", spec=NormalizeTaskSpec(sheet="数据", column="编号", confidence=1))
    service = make_service(tmp_path, model)
    task_id = upload(service, workbook_bytes([[" A  B ", 2, 3, None], [None, 1, 2, None], [42, 1, 2, None]]))
    turn = service.submit_message(task_id, "编号这一列标准化")
    service.approve_current_plan(task_id, acknowledgements=turn.plan.confirmation_requirements)
    book = load_workbook(io.BytesIO(service.get_download(task_id)[1]))
    assert [book["数据"].cell(r, 1).value for r in (2, 3, 4)] == ["A  B", None, 42]
    book.close()


def test_column_arithmetic_rejects_unknown_ambiguous_or_executable_expression(tmp_path):
    source = tmp_path / "formula.xlsx"
    source.write_bytes(workbook_bytes())
    profile = profile_workbook(source, source.name, 100_000)
    assert column_formula(profile, "数据", "总额", "总额=(数量+2)*单价/2-1") == "=((((B2+2)*C2)/2)-1)"
    for text in ("总额=未知*单价", "总额=__import__('os').system('whoami')", "总额=数量**100"):
        assert column_formula(profile, "数据", "总额", text) is None
    profile.sheets[0].columns[0].display_header = "数量"
    assert column_formula(profile, "数据", "总额", "总额=数量*单价") is None
    assert column_formula(profile, "数据", "总额", "总额=B列*C列") == "=(B2*C2)"


def test_duplicate_formula_target_can_be_selected_by_letter(tmp_path, monkeypatch):
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet.append(["编号", "数量", "单价", "总额", "总额"])
    sheet.append(["001", 2, 15, None, None])
    stream = io.BytesIO()
    book.save(stream)
    model = TaskEnvelope(goal="计算", spec=FormulaTaskSpec(sheet="数据", column="总额", confidence=1))
    service = make_service(tmp_path, model)
    task_id = upload(service, stream.getvalue())
    turn = service.submit_message(task_id, "在总额列填入数量×单价")
    assert turn.clarification and turn.clarification.issue_code == "ambiguous_formula_column"
    resumed = service.submit_message(task_id, "D列")
    assert resumed.plan is not None, resumed
    monkeypatch.setattr(calculation, "_recalculate_with_excel", cache_products)
    service.approve_current_plan(task_id)
    result = load_workbook(io.BytesIO(service.get_download(task_id)[1]), data_only=True)
    assert result["数据"]["D2"].value == 30 and result["数据"]["E2"].value is None
    result.close()


@pytest.mark.parametrize("rows", [3_000, 5_000, 40_000])
def test_formula_execution_recalculation_and_verification_are_linear(tmp_path, monkeypatch, rows):
    path = tmp_path / "large.xlsx"
    book = Workbook(write_only=True)
    sheet = book.create_sheet("数据")
    sheet.append(["编号", "数量", "单价", "总额"])
    for i in range(rows):
        sheet.append([f"{i:05}", 2, 15, None])
    book.save(path)
    profile = profile_workbook(path, path.name, 100_000)
    context = ToolContext(task_id="performance", input_path=path, output_path=tmp_path / "result.xlsx", profile=profile, version_before=0, version_after=1)
    spec = FormulaTaskSpec(sheet="数据", column="总额", start_row=2, end_row=rows + 1, mode="set_formula", formula="=B2*C2", recalculate=True, confidence=1)
    monkeypatch.setattr(calculation, "_recalculate_with_excel", cache_products)
    # Catch regressions to per-cell random access in read-only worksheets.
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    monkeypatch.setattr(ReadOnlyWorksheet, "__getitem__", lambda *a: pytest.fail("random worksheet access"))
    started = time.perf_counter()
    tool = FormulaTool()
    result = tool.execute(context, spec)
    assert tool.verify(context, spec, result).valid
    assert result.counters["calculated"] == rows
    assert time.perf_counter() - started < 30

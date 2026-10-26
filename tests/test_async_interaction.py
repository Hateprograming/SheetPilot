from __future__ import annotations

import threading
from pathlib import Path

from openpyxl import Workbook, load_workbook

from sheetpilot.config import Settings
from sheetpilot.schemas import (
    ColumnRef,
    DuplicatePolicy,
    LookupTaskSpec,
    MessageKind,
    ParseDimensionsRule,
    TaskEnvelope,
    TransformOperation,
    TransformTaskSpec,
)
from sheetpilot.service import SheetPilotService


def lookup_envelope(goal: str, column: str = "客户等级") -> TaskEnvelope:
    return TaskEnvelope(
        goal=goal,
        spec=LookupTaskSpec(
            main_sheet="订单",
            source_sheet="客户",
            main_key_column="客户编号",
            source_key_column="客户编号",
            return_columns=[column],
            output_column_names=[column],
            duplicate_policy=DuplicatePolicy.BLOCK,
            confidence=1,
        ),
    )


def make_service(tmp_path: Path, llm) -> SheetPilotService:
    settings = Settings(
        project_root=tmp_path,
        data_dir=tmp_path / "data",
        database_path=tmp_path / "data" / "sheetpilot.sqlite3",
        llm_provider="fake",
        api_key="",
        base_url="",
        model="fake",
        max_file_bytes=20 * 1024 * 1024,
        max_sheet_rows=100_000,
    )
    return SheetPilotService(settings, llm)


def create_task(service: SheetPilotService, source: Path) -> str:
    profile = service.create_task(source.name, source.read_bytes())
    return service.task_id_from_profile(profile)


def test_enqueue_persists_user_message_before_model_returns(workbook_factory, tmp_path):
    started = threading.Event()
    release = threading.Event()

    class BlockingLLM:
        def parse_task(self, request, profile, **kwargs):
            started.set()
            assert release.wait(5)
            return lookup_envelope("补充客户等级")

    service = make_service(tmp_path, BlockingLLM())
    task_id = create_task(service, workbook_factory())

    run = service.enqueue_message(task_id, "按客户编号补充客户等级")
    assert started.wait(5)
    workspace = service.get_workspace(task_id)

    assert workspace["messages"][-1]["role"] == "user"
    assert workspace["messages"][-1]["content"] == "按客户编号补充客户等级"
    assert workspace["planning_run"]["status"] in {"queued", "running"}
    assert workspace["activity_events"][-1]["label"] in {
        "已收到你的要求", "正在读取工作簿结构", "正在理解你的处理要求"
    }

    release.set()
    service._background_futures[run["job_id"]].result(timeout=10)
    completed = service.get_workspace(task_id)
    assert completed["planning_run"]["status"] == "completed"
    assert completed["messages"][-1]["kind"] == MessageKind.PLAN.value


def test_new_request_supersedes_blocked_old_planning_result(workbook_factory, tmp_path):
    first_started = threading.Event()
    release_first = threading.Event()

    class SupersedingLLM:
        def parse_task(self, request, profile, **kwargs):
            if request == "按客户编号补充客户等级":
                first_started.set()
                assert release_first.wait(5)
                return lookup_envelope("旧计划", "客户等级")
            return lookup_envelope("新计划", "地区")

    service = make_service(tmp_path, SupersedingLLM())
    task_id = create_task(service, workbook_factory())
    old = service.enqueue_message(task_id, "按客户编号补充客户等级")
    assert first_started.wait(5)

    new = service.enqueue_message(task_id, "按客户编号补充地区")
    service._background_futures[new["job_id"]].result(timeout=10)
    release_first.set()
    service._background_futures[old["job_id"]].result(timeout=10)

    workspace = service.get_workspace(task_id)
    old_run = service.repository.get_planning_run(old["job_id"])
    assert old_run["status"] == "superseded"
    assert workspace["planning_run"]["job_id"] == new["job_id"]
    assert workspace["planning_run"]["status"] == "completed"
    plans = [item for item in workspace["messages"] if item["kind"] == MessageKind.PLAN.value]
    assert len(plans) == 1
    assert plans[0]["metadata"]["envelope"]["goal"] == "新计划"


def test_activity_events_do_not_store_user_prompt_or_clinical_values(workbook_factory, tmp_path):
    class ImmediateLLM:
        def parse_task(self, request, profile, **kwargs):
            return lookup_envelope("补充客户等级")

    service = make_service(tmp_path, ImmediateLLM())
    task_id = create_task(service, workbook_factory())
    secret = "患者张三的原始值-987654"
    run = service.enqueue_message(task_id, secret)
    service._background_futures[run["job_id"]].result(timeout=10)

    serialized = str(service.repository.list_activity_events(task_id))
    assert secret not in serialized
    assert "987654" not in serialized


def test_new_message_safely_cancels_execution_then_replans(workbook_factory, tmp_path):
    tool_started = threading.Event()
    release_tool = threading.Event()

    class ImmediateLLM:
        def parse_task(self, request, profile, **kwargs):
            return lookup_envelope("按最新要求补充客户等级")

    service = make_service(tmp_path, ImmediateLLM())
    task_id = create_task(service, workbook_factory())
    service.submit_message(task_id, "按客户编号补充客户等级")

    original_tool = service.tools.get("lookup")

    class BlockingLookupTool:
        name = original_tool.name
        description = original_tool.description
        spec_model = original_tool.spec_model
        risk_level = original_tool.risk_level

        def parse_spec(self, arguments):
            return original_tool.parse_spec(arguments)

        def validate(self, context, spec):
            return original_tool.validate(context, spec)

        def preview(self, context, spec):
            return original_tool.preview(context, spec)

        def execute(self, context, spec):
            tool_started.set()
            assert release_tool.wait(5)
            return original_tool.execute(context, spec)

        def verify(self, context, spec, result):
            return original_tool.verify(context, spec, result)

    service.tools._tools["lookup"] = BlockingLookupTool()
    execution = service.enqueue_current_plan_execution(task_id)
    assert tool_started.wait(5)

    planning = service.enqueue_message(task_id, "停止旧处理，仍按客户编号补充客户等级并重新规划")
    active = service.repository.latest_plan_run(task_id)
    assert active["cancel_requested"] == 1
    assert service.get_workspace(task_id)["planning_run"]["status"] == "waiting_for_execution"

    release_tool.set()
    service._background_futures[f"execution:{execution['run_id']}"].result(timeout=10)
    service._background_futures[planning["job_id"]].result(timeout=10)

    workspace = service.get_workspace(task_id)
    assert workspace["task"]["current_version"] == 0
    assert service.repository.latest_plan_run(task_id)["status"] == "cancelled"
    assert workspace["planning_run"]["status"] == "completed"
    assert workspace["task"]["status"] == "awaiting_approval"


def test_case3_row_corrections_resume_same_draft_without_second_llm_call(tmp_path):
    source = tmp_path / "case3.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tumor_Data"
    sheet.append(["病例编号", "肿瘤大小：长、宽、高（mm）"])
    for row in range(2, 17):
        value = "10*8*6"
        if row == 13:
            value = "21.*7*16.7"
        elif row == 16:
            value = "1*2*3*4"
        sheet.append([f"P{row - 1:03d}", value])
    workbook.save(source)
    workbook.close()

    envelope = TaskEnvelope(
        goal="解析肿瘤尺寸",
        spec=TransformTaskSpec(
            operation=TransformOperation.PARSE_DIMENSIONS,
            sheet="Tumor_Data",
            dimension_rule=ParseDimensionsRule(
                source=ColumnRef(
                    sheet="Tumor_Data", column_letter="B",
                    display_header="肿瘤大小：长、宽、高（mm）",
                ),
                default_unit="mm",
            ),
            confidence=1,
        ),
    )

    class CountingLLM:
        def __init__(self):
            self.calls = 0

        def parse_task(self, request, profile, **kwargs):
            self.calls += 1
            return envelope.model_copy(deep=True)

    llm = CountingLLM()
    service = make_service(tmp_path / "service", llm)
    task_id = create_task(service, source)
    first = service.enqueue_message(
        task_id,
        "解析 Tumor_Data 中的肿瘤大小列，默认单位为 mm，新增长、宽、高、最大径和解析状态列。",
    )
    service._background_futures[first["job_id"]].result(timeout=10)
    waiting = service.get_workspace(task_id)
    assert waiting["planning_run"]["status"] == "waiting_for_user"
    review = waiting["messages"][-1]["metadata"]["review"]
    assert [item["row"] for item in review["rows"]] == [13, 16]

    service.select_custom_clarification_input(task_id, "我输入具体修正规则")
    second = service.enqueue_message(
        task_id,
        "第13行数据改为：21mm7mm16.7mm 第16行数据改为：2mm3mm4mm",
    )
    service._background_futures[second["job_id"]].result(timeout=10)
    planned = service.get_workspace(task_id)

    assert llm.calls == 1
    assert planned["task"]["status"] == "awaiting_approval"
    plan = planned["messages"][-1]["metadata"]["plan"]
    assert plan["revision"] == 2
    corrections = plan["steps"][0]["arguments"]["dimension_rule"]["row_corrections"]
    assert corrections == {13: "21mm7mm16.7mm", 16: "2mm3mm4mm"} or corrections == {
        "13": "21mm7mm16.7mm", "16": "2mm3mm4mm"
    }

    executed = service.approve_current_plan(task_id)
    result_book = load_workbook(executed.output_path, data_only=True)
    result_sheet = result_book["Tumor_Data"]
    assert [result_sheet.cell(13, column).value for column in range(3, 7)] == [21, 7, 16.7, 21]
    assert result_sheet.cell(13, 7).value == "corrected_value"
    assert [result_sheet.cell(16, column).value for column in range(3, 7)] == [2, 3, 4, 4]
    assert result_sheet.cell(16, 7).value == "corrected_value"
    result_book.close()

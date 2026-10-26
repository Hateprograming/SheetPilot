from __future__ import annotations

import time
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from sheetpilot.profiler import profile_workbook
from sheetpilot.schemas import FormulaMode, FormulaTaskSpec
from sheetpilot.tools import FormulaTool, ToolContext


class _CountingCell:
    def __init__(self, row: int, column: int, value):
        self.value = value
        self.coordinate = f"{get_column_letter(column)}{row}"


class _CountingSheet:
    def __init__(self, values: dict[int, object]):
        self.values = values
        self.accesses = 0

    def cell(self, row: int, column: int):
        self.accesses += 1
        return _CountingCell(row, column, self.values.get(row))


def test_formula_inspection_accesses_cells_linearly():
    row_count = 10_000
    values = {
        row: (f"=A{row}*B{row}" if row % 2 == 0 else None)
        for row in range(2, row_count + 2)
    }
    sheet = _CountingSheet(values)
    result = FormulaTool()._inspect(sheet, 3, 2, row_count + 1)
    assert result["formula_count"] == row_count // 2
    assert sheet.accesses < row_count * 5


def _formula_workbook(path: Path, rows: int) -> Path:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("数据")
    sheet.append(["数量", "单价", "金额"])
    for row in range(2, rows + 2):
        sheet.append([row - 1, 10, None])
    workbook.save(path)
    workbook.close()
    return path


@pytest.mark.parametrize("rows", [3_000, 5_000, 40_000])
def test_set_formula_large_files_complete_under_thirty_seconds(tmp_path, rows):
    source = _formula_workbook(tmp_path / f"source-{rows}.xlsx", rows)
    output = tmp_path / f"output-{rows}.xlsx"
    profile = profile_workbook(source, source.name, 100_000)
    context = ToolContext(
        task_id=f"formula-{rows}", input_path=source, output_path=output,
        profile=profile, version_before=1, version_after=2,
    )
    spec = FormulaTaskSpec(
        sheet="数据", column="金额", start_row=2, end_row=rows + 1,
        mode=FormulaMode.SET_FORMULA, formula="=A2*B2",
        overwrite_existing=False, recalculate=False, confidence=1,
    )

    started = time.perf_counter()
    result = FormulaTool().execute(context, spec)
    elapsed = time.perf_counter() - started

    assert elapsed < 30
    assert result.counters["generated"] == rows
    assert result.details["timings_seconds"]["inspect"] < 1
    workbook = load_workbook(output, read_only=True, data_only=False)
    try:
        assert workbook["数据"]["C2"].value == "=A2*B2"
        assert workbook["数据"][f"C{rows + 1}"].value == f"=A{rows + 1}*B{rows + 1}"
    finally:
        workbook.close()

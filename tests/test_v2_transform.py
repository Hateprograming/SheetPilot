from __future__ import annotations

import csv
import sqlite3
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from sheetpilot.config import Settings
from sheetpilot.conversation import clarification_for
from sheetpilot.errors import IntentValidationError
from sheetpilot.profiler import ensure_supported_profile, profile_workbook
from sheetpilot.schemas import (
    ColumnRef,
    CompositeTaskSpec,
    ConversationRole,
    DimensionResolutionAction,
    DerivedFormulaTemplate,
    DerivedResolutionAction,
    DeriveMetricRule,
    ImputationMethod,
    InvalidNumericPolicy,
    ImputeNumericRule,
    NormalizeOperation,
    NormalizeTaskSpec,
    ParseDimensionsRule,
    MessageKind,
    TaskEnvelope,
    TransformOperation,
    TransformTaskSpec,
    ValueMapRule,
)
from sheetpilot.service import SheetPilotService
from sheetpilot.tools import ToolContext, ToolRegistry, _parse_dimensions


def make_service(tmp_path: Path) -> SheetPilotService:
    return SheetPilotService(Settings(
        project_root=tmp_path,
        data_dir=tmp_path / "data",
        database_path=tmp_path / "data" / "sheetpilot.sqlite3",
        llm_provider="fake", api_key="", base_url="", model="fake",
        max_file_bytes=20 * 1024 * 1024, max_sheet_rows=100_000,
    ))


def make_transform_book(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["BMI", "高血压", "肿瘤大小", "姓名"])
    sheet.append([20, "有", "10*20*30 mm", "A"])
    sheet.append([None, "无 ", "10×20", "B"])
    sheet.append(["无", "是", "总直径：12mm", "C"])
    sheet.append(["20..5", "不详", "10cm*20mm*30mm", "D"])
    workbook.save(path)
    workbook.close()
    return path


def create_task(service: SheetPilotService, source: Path) -> str:
    profile = service.create_task(source.name, source.read_bytes())
    return service.task_id_from_profile(profile)


def test_duplicate_and_blank_headers_receive_stable_refs(tmp_path):
    path = tmp_path / "headers.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["性别", "性别", None, "模板列"])
    sheet.append([None, "女", "有数据", None])
    workbook.save(path)
    workbook.close()

    profile = profile_workbook(path, path.name, 1000, header_overrides={"Sheet": 1})
    ensure_supported_profile(profile)
    columns = profile.sheets[0].columns
    assert [(item.letter, item.header_occurrence) for item in columns[:2]] == [("A", 1), ("B", 2)]
    assert columns[0].active is False
    assert columns[1].active is True
    assert columns[2].display_header == ""
    assert columns[2].active is True
    assert columns[3].active is False


def test_wrapped_header_matches_model_safe_header(tmp_path):
    source = tmp_path / "wrapped_header.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["合并\n高血压"])
    sheet.append(["有"])
    sheet.append(["无"])
    workbook.save(source)
    workbook.close()

    service = make_service(tmp_path / "service")
    task_id = create_task(service, source)
    envelope = TaskEnvelope(
        goal="将高血压映射为indicator",
        spec=TransformTaskSpec(
            operation=TransformOperation.MAP_VALUES,
            sheet="Data",
            map_rules=[ValueMapRule(
                source=ColumnRef(sheet="Data", display_header="合并高血压"),
                mapping={"有": 1, "无": 0},
            )],
            confidence=1,
        ),
    )
    plan, validation = service.make_task_plan(
        task_id, envelope, policy_source_text="合并高血压有为1，无为0"
    )
    assert validation.valid, validation.errors
    assert plan.steps[0].expected_changes.summary_counts["mapped"] == 2


def test_unresolved_transform_column_becomes_clarification_not_internal_error(tmp_path):
    source = tmp_path / "unresolved.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["临床分型"])
    sheet.append(["ACTH腺瘤"])
    workbook.save(source)
    workbook.close()

    service = make_service(tmp_path / "service")
    task_id = create_task(service, source)
    envelope = TaskEnvelope(
        goal="创建ACTH indicator",
        spec=TransformTaskSpec(
            operation=TransformOperation.MAP_VALUES,
            sheet="Data",
            map_rules=[ValueMapRule(
                source=ColumnRef(sheet="Data", display_header="临床分类"),
                mapping={"ACTH腺瘤": 1},
            )],
            confidence=1,
        ),
    )
    plan, validation = service.make_task_plan(
        task_id, envelope, policy_source_text="临床分类为ACTH腺瘤时映射为1"
    )
    assert not validation.valid
    assert any("不存在列“临床分类”" in error for error in validation.errors)
    assert plan.steps[0].expected_changes.summary_counts["preview_unavailable"] == 1


def test_legacy_normalize_accepts_excel_column_letter(tmp_path):
    source = tmp_path / "normalize_letter.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["年龄"])
    sheet.append(["60岁"])
    workbook.save(source)
    workbook.close()

    service = make_service(tmp_path / "service")
    task_id = create_task(service, source)
    envelope = TaskEnvelope(
        goal="年龄去掉岁并转换为数字",
        spec=NormalizeTaskSpec(
            sheet="Data",
            column="A",
            operations=[NormalizeOperation.STRIP_SUFFIX, NormalizeOperation.TO_NUMBER],
            strip_suffixes=["岁"],
            output_column="年龄_数值",
            confidence=1,
        ),
    )
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid, validation.errors
    service.approve(task_id, "approve")
    output, report, _ = service.execute_task_plan(task_id, envelope, plan)
    assert report.status == "passed"
    workbook = load_workbook(output, data_only=True)
    assert workbook["Data"].cell(2, 2).value == 60
    workbook.close()


def test_transform_rejects_payloads_for_other_operations():
    with pytest.raises(ValueError, match="cannot include payloads"):
        TransformTaskSpec(
            operation=TransformOperation.MAP_VALUES,
            sheet="Data",
            map_rules=[ValueMapRule(
                source=ColumnRef(sheet="Data", column_letter="A"),
                mapping={"有": 1},
            )],
            impute_rules=[ImputeNumericRule(
                source=ColumnRef(sheet="Data", column_letter="B")
            )],
            confidence=1,
        )


def test_model_envelope_sanitizer_repairs_one_hot_and_removes_unauthorized_defaults():
    source = ColumnRef(sheet="Data", column_letter="D", display_header="临床分型")
    dimension = TransformTaskSpec(
        operation=TransformOperation.PARSE_DIMENSIONS,
        sheet="Data",
        dimension_rule=ParseDimensionsRule(
            source=ColumnRef(sheet="Data", column_letter="C", display_header="肿瘤大小"),
            resolution_policies={
                "multiple_candidates": DimensionResolutionAction.FIRST_CANDIDATE,
                "unit_conflict": DimensionResolutionAction.CONVERT_EACH_TO_MM,
            },
        ),
        ambiguities=["默认取第一组"],
        confidence=0.8,
    )
    mapping = TransformTaskSpec(
        operation=TransformOperation.MAP_VALUES,
        sheet="Data",
        map_rules=[
            ValueMapRule(
                source=ColumnRef(sheet="Data", column_letter="B", display_header="高血压"),
                mapping={"有": 1, "无": 0}, blank_value=0, unmapped_value=0,
            ),
            ValueMapRule(
                source=source, mapping={"GH": 1, "ACTH": 1, "混合": 1, "促性腺": 1},
                blank_value=0, unmapped_value=0, output_column="GH_indicator",
            ),
            ValueMapRule(
                source=source, mapping={"ACTH": 1}, blank_value=0,
                unmapped_value=0, output_column="ACTH_indicator",
            ),
            ValueMapRule(
                source=source, mapping={"混合": 1}, blank_value=0,
                unmapped_value=0, output_column="混合_indicator",
            ),
            ValueMapRule(
                source=source, mapping={"促性腺": 1}, blank_value=0,
                unmapped_value=0, output_column="促性腺_indicator",
            ),
        ],
        ambiguities=["当前设置将未知文本映射为0"],
        confidence=0.9,
    )
    envelope = TaskEnvelope(
        goal="综合处理",
        spec=CompositeTaskSpec(
            operations=[dimension, mapping],
            ambiguities=["假设工作表已匹配"],
            confidence=0,
        ),
        assumptions=["模型生成的未授权假设"],
    )
    prepared = SheetPilotService.prepare_model_envelope(
        envelope,
        "多候选没有指定取哪组；混合单位统一换算为mm；高血压没有记录算作无；"
        "四类分别做indicator，其他分类为0",
    )
    assert isinstance(prepared.spec, CompositeTaskSpec)
    assert prepared.spec.confidence == 0.8
    assert prepared.spec.ambiguities == ["假设工作表已匹配"]
    assert prepared.assumptions == []
    parsed = prepared.spec.operations[0]
    assert isinstance(parsed, TransformTaskSpec) and parsed.dimension_rule
    assert parsed.ambiguities == ["默认取第一组"]
    assert "multiple_candidates" not in parsed.dimension_rule.resolution_policies
    assert parsed.dimension_rule.resolution_policies["unit_conflict"] == (
        DimensionResolutionAction.CONVERT_EACH_TO_MM
    )
    mapped = prepared.spec.operations[1]
    assert isinstance(mapped, TransformTaskSpec)
    assert mapped.ambiguities == ["当前设置将未知文本映射为0"]
    assert mapped.map_rules[0].blank_value == 0
    assert mapped.map_rules[0].unmapped_value is None
    assert mapped.map_rules[0].mapping == {"有": 1, "无": 0}
    assert mapped.map_rules[1].mapping == {"GH": 1}
    assert mapped.map_rules[1].blank_value is None


def test_model_envelope_splits_a_collapsed_multi_category_indicator():
    categories = ["生长激素细胞腺瘤", "ACTH腺瘤", "混合腺瘤", "促性腺激素细胞腺瘤"]
    mapping = TransformTaskSpec(
        operation=TransformOperation.MAP_VALUES,
        sheet="Data",
        map_rules=[ValueMapRule(
            source=ColumnRef(sheet="Data", column_letter="D", display_header="临床分型"),
            mapping={category: 1 for category in categories},
            unmapped_value=0,
            output_column="目标腺瘤_Flag",
        )],
        confidence=1,
    )
    envelope = TaskEnvelope(goal="四类腺瘤分别做 indicator", spec=mapping)

    prepared = SheetPilotService.prepare_model_envelope(
        envelope,
        "这四类另起四列，分别做四个indicator，其他腺瘤为0",
    )

    assert isinstance(prepared.spec, TransformTaskSpec)
    assert [rule.mapping for rule in prepared.spec.map_rules] == [
        {category: 1} for category in categories
    ]
    assert [rule.output_column for rule in prepared.spec.map_rules] == [
        f"{category}_Indicator" for category in categories
    ]
    assert all(rule.unmapped_value == 0 for rule in prepared.spec.map_rules)


def test_composite_dimension_clarification_answer_updates_same_draft(tmp_path):
    service = make_service(tmp_path / "service")
    task_id = create_task(service, make_transform_book(tmp_path / "input.xlsx"))
    dimension = TransformTaskSpec(
        operation=TransformOperation.PARSE_DIMENSIONS,
        sheet="Data",
        dimension_rule=ParseDimensionsRule(
            source=ColumnRef(sheet="Data", column_letter="C", display_header="肿瘤大小")
        ),
        confidence=1,
    )
    mapping = TransformTaskSpec(
        operation=TransformOperation.MAP_VALUES,
        sheet="Data",
        map_rules=[ValueMapRule(
            source=ColumnRef(sheet="Data", column_letter="B", display_header="高血压"),
            mapping={"有": 1, "无": 0},
        )],
        confidence=1,
    )
    previous = TaskEnvelope(
        goal="解析尺寸并映射高血压",
        spec=CompositeTaskSpec(operations=[dimension, mapping], confidence=1),
    )
    service._add_message(task_id, ConversationRole.USER, MessageKind.TEXT, "原始要求")
    service._add_message(
        task_id, ConversationRole.ASSISTANT, MessageKind.CLARIFICATION,
        "请选择多候选尺寸处理方式",
        {"review": {"issue_code": "multiple_candidates", "rows": [], "total": 2}},
    )
    service._add_message(
        task_id, ConversationRole.USER, MessageKind.TEXT, "这类记录统一取第一组尺寸"
    )
    resumed = service._resume_clarification_envelope(
        previous, "这类记录统一取第一组尺寸", service.get_conversation(task_id)
    )
    assert resumed is not None and isinstance(resumed.spec, CompositeTaskSpec)
    revised = resumed.spec.operations[0]
    assert isinstance(revised, TransformTaskSpec) and revised.dimension_rule
    assert revised.dimension_rule.resolution_policies["multiple_candidates"] == (
        DimensionResolutionAction.FIRST_CANDIDATE
    )


def test_asking_what_is_missing_repeats_actual_question_without_replanning(tmp_path):
    service = make_service(tmp_path / "service")
    task_id = create_task(service, make_transform_book(tmp_path / "input.xlsx"))
    envelope = TaskEnvelope(
        goal="解析尺寸",
        spec=TransformTaskSpec(
            operation=TransformOperation.PARSE_DIMENSIONS,
            sheet="Data",
            dimension_rule=ParseDimensionsRule(
                source=ColumnRef(sheet="Data", column_letter="C", display_header="肿瘤大小"),
                resolution_policies={
                    "multiple_candidates": DimensionResolutionAction.FIRST_CANDIDATE
                },
            ),
            confidence=1,
        ),
    )
    service.repository.update_task(
        task_id, status="needs_clarification", spec_json=envelope.model_dump_json()
    )
    service._add_message(
        task_id, ConversationRole.ASSISTANT, MessageKind.CLARIFICATION,
        "我还缺少一个会影响结果的信息。请补充你希望处理的工作表、列或规则。",
        {
            "choices": [],
            "validation_errors": [
                "步骤 1 的尺寸异常处理“first_candidate”未在用户要求中明确授权。"
            ],
            "envelope": envelope.model_dump(mode="json"),
        },
    )
    response = service.enqueue_message(task_id, "缺少什么信息呢")
    assert response["status"] == "waiting_for_user"
    assert not service._background_futures
    latest = service.get_conversation(task_id)[-1]
    assert "尺寸列中存在多组候选值" in latest.content
    assert "工作表" not in latest.content
    assert latest.metadata["choices"][0] == "这类记录统一取第一组尺寸"


def test_impute_keeps_invalid_values_out_of_mean_and_writes_bom_csv(tmp_path):
    service = make_service(tmp_path)
    task_id = create_task(service, make_transform_book(tmp_path / "input.xlsx"))
    spec = TransformTaskSpec(
        operation=TransformOperation.IMPUTE_NUMERIC,
        sheet="Data",
        impute_rules=[ImputeNumericRule(
            source=ColumnRef(sheet="Data", column_letter="A", display_header="BMI"),
            method=ImputationMethod.MEAN,
            missing_tokens=["无"],
            output_column="BMI_填补",
            indicator_column="BMI_是否填补",
        )],
        confidence=1,
    )
    envelope = TaskEnvelope(goal="用均值填补BMI，将‘无’视为缺失", spec=spec)
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid, validation.errors
    assert plan.steps[0].expected_changes.summary_counts == {
        "valid": 1, "missing": 2, "invalid": 1,
        "observed_basis": 1, "planned_fill_value": 20.0,
    }
    service.approve(task_id, "approve")
    output, report, result = service.execute_task_plan(task_id, envelope, plan)
    assert report.status == "passed_with_warnings"
    assert result.counters["imputed"] == 2
    assert result.counters["invalid"] == 1
    workbook = load_workbook(output, data_only=True)
    sheet = workbook["Data"]
    assert [sheet.cell(row, 5).value for row in range(2, 6)] == [20, 20, 20, None]
    assert [sheet.cell(row, 1).value for row in range(2, 6)] == [20, None, "无", "20..5"]
    original = load_workbook(tmp_path / "input.xlsx", data_only=True)
    for row in range(1, 6):
        for column in range(1, 5):
            assert sheet.cell(row, column).value == original["Data"].cell(row, column).value
            assert sheet.cell(row, column).style_id == original["Data"].cell(row, column).style_id
    original.close()
    workbook.close()
    artifact = Path(result.details["steps"][0]["details"]["artifact_path"])
    assert artifact.read_bytes().startswith(b"\xef\xbb\xbf")
    with artifact.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 1 and rows[0]["问题代码"] == "invalid_numeric"
    assert "姓名" not in rows[0]


def test_user_can_explicitly_impute_invalid_numeric_values(tmp_path):
    service = make_service(tmp_path)
    task_id = create_task(service, make_transform_book(tmp_path / "invalid_impute.xlsx"))
    first_spec = TransformTaskSpec(
        operation=TransformOperation.IMPUTE_NUMERIC,
        sheet="Data",
        impute_rules=[ImputeNumericRule(
            source=ColumnRef(sheet="Data", column_letter="A", display_header="BMI"),
            method=ImputationMethod.MEAN,
            missing_tokens=["无"],
            output_column="BMI_二次填补",
            indicator_column="BMI_是否填补",
        )],
        confidence=1,
    )
    first_envelope = TaskEnvelope(goal="先保守填补BMI", spec=first_spec)
    first_plan, first_validation = service.make_task_plan(task_id, first_envelope)
    assert first_validation.valid
    service.approve(task_id, "approve")
    first_output, _, _ = service.execute_task_plan(task_id, first_envelope, first_plan)
    workbook = load_workbook(first_output, data_only=True)
    assert [workbook["Data"].cell(row, 5).value for row in range(2, 6)] == [20, 20, 20, None]
    workbook.close()

    spec = first_spec.model_copy(deep=True)
    spec.impute_rules[0].invalid_policy = InvalidNumericPolicy.IMPUTE
    spec.impute_rules[0].overwrite_existing_outputs = True
    envelope = TaskEnvelope(goal="错误值也用平均值填补", spec=spec)
    plan, validation = service.make_task_plan(
        task_id,
        envelope,
        policy_source_text="现在显示为None的错误值也用平均值填补",
    )
    assert validation.valid, validation.errors
    assert plan.steps[0].expected_changes.summary_counts == {
        "valid": 1,
        "missing": 2,
        "invalid_imputed": 1,
        "observed_basis": 1,
        "planned_fill_value": 20.0,
    }
    assert plan.confirmation_requirements == ["overwrite_source_columns"]
    service.approve(
        task_id,
        "approve",
        {"acknowledgements": ["overwrite_source_columns"]},
    )
    output, report, result = service.execute_task_plan(task_id, envelope, plan)
    assert report.status == "passed"
    assert result.counters["invalid_imputed"] == 1
    assert result.counters["exceptions"] == 0
    workbook = load_workbook(output, data_only=True)
    assert [workbook["Data"].cell(row, 5).value for row in range(2, 6)] == [20, 20, 20, 20]
    workbook.close()


def test_repeated_invalid_numeric_values_require_a_reviewable_clarification(tmp_path):
    service = make_service(tmp_path)
    task_id = create_task(service, make_transform_book(tmp_path / "invalid_review.xlsx"))
    spec = TransformTaskSpec(
        operation=TransformOperation.IMPUTE_NUMERIC,
        sheet="Data",
        impute_rules=[ImputeNumericRule(
            source=ColumnRef(sheet="Data", column_letter="A", display_header="BMI"),
            method=ImputationMethod.MEAN,
            output_column="BMI_填补",
            indicator_column="BMI_是否填补",
        )],
        confidence=1,
    )
    envelope = TaskEnvelope(goal="用均值填补BMI空值", spec=spec)
    plan, validation = service.make_task_plan(task_id, envelope)

    assert not validation.valid
    assert "IMPUTE_RESOLUTION_REQUIRED:invalid_numeric:2" in validation.errors
    changes = plan.steps[0].expected_changes
    assert changes.review_totals["invalid_numeric"] == 2
    assert [item["row"] for item in changes.review_rows] == [4, 5]
    assert [item["raw_values"]["BMI"] for item in changes.review_rows] == ["无", "20..5"]
    prompt = clarification_for(validation, envelope, service.load_profile(task_id))
    assert "“BMI”列" in prompt.question
    assert "2 条非空内容无法解析为数值" in prompt.question
    assert prompt.choices == [
        "将这类错误值也用同一均值填补",
        "保持结果留空并生成异常清单",
        "我输入哪些值应视为缺失",
    ]
    review = service._clarification_review(plan, validation)
    assert review["issue_code"] == "invalid_numeric"
    assert review["total"] == 2
    assert [item["row"] for item in review["rows"]] == [4, 5]

    service._add_message(task_id, ConversationRole.USER, MessageKind.TEXT, "填补BMI空值")
    service._add_message(
        task_id, ConversationRole.ASSISTANT, MessageKind.CLARIFICATION,
        prompt.question, {"review": review},
    )
    answer = "将这类错误值也用同一均值填补"
    service._add_message(task_id, ConversationRole.USER, MessageKind.TEXT, answer)
    resumed = service._resume_clarification_envelope(
        envelope, answer, service.get_conversation(task_id)
    )
    assert resumed is not None and isinstance(resumed.spec, TransformTaskSpec)
    assert resumed.spec.impute_rules[0].invalid_policy == InvalidNumericPolicy.IMPUTE
    assert resumed.spec.impute_rules[0].invalid_policy_confirmed is True


def test_imputed_invalid_numeric_uses_excel_roundtrip_tolerance(tmp_path):
    source = tmp_path / "floating_mean.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["BMI"])
    sheet.append([0.1])
    sheet.append([0.2])
    sheet.append(["错误值"])
    workbook.save(source)
    workbook.close()

    service = make_service(tmp_path / "service")
    task_id = create_task(service, source)
    spec = TransformTaskSpec(
        operation=TransformOperation.IMPUTE_NUMERIC,
        sheet="Data",
        impute_rules=[ImputeNumericRule(
            source=ColumnRef(sheet="Data", column_letter="A", display_header="BMI"),
            method=ImputationMethod.MEAN,
            overwrite=True,
            invalid_policy=InvalidNumericPolicy.IMPUTE,
            invalid_policy_confirmed=True,
        )],
        confidence=1,
    )
    envelope = TaskEnvelope(goal="BMI错误值也用平均值填补", spec=spec)
    plan, validation = service.make_task_plan(
        task_id, envelope, policy_source_text="BMI错误值也用平均值填补"
    )
    assert validation.valid, validation.errors
    service.approve(
        task_id, "approve", {"acknowledgements": ["overwrite_source_columns"]}
    )
    output, report, result = service.execute_task_plan(task_id, envelope, plan)

    assert report.status == "passed"
    assert result.counters["invalid_imputed"] == 1
    workbook = load_workbook(output, data_only=True)
    assert workbook["Data"]["A4"].value == pytest.approx(0.15)
    workbook.close()


def test_mapping_and_dimension_parser_are_deterministic(tmp_path):
    assert _parse_dimensions("10*20*30 mm (备注)", "mm")[1] == "parsed_3d"
    assert _parse_dimensions("10×20", "mm")[1] == "parsed_2d"
    assert _parse_dimensions("总直径：12mm", "mm")[1] == "diameter_only"
    assert _parse_dimensions("10*20*30*40", "mm")[1] == "invalid_format"
    assert _parse_dimensions("21mm7mm16.7mm", "mm") == (
        [21.0, 7.0, 16.7, 21.0], "parsed_3d"
    )
    assert _parse_dimensions("10cm*20mm*30mm", "mm")[1] == "unit_conflict"
    assert _parse_dimensions("10*20 另12*14", "mm")[1] == "multiple_candidates"
    assert _parse_dimensions(
        "外院MR图像无法判断。\n\n24*22*8（病理报告显示）", "mm"
    ) == ([24.0, 22.0, 8.0, 24.0], "parsed_3d")
    assert _parse_dimensions(
        "27*21", "mm", require_context_for_partial=True
    )[1] == "invalid_format"
    assert _parse_dimensions(
        "27*21（病理报告显示）", "mm", require_context_for_partial=True
    ) == ([27.0, 21.0, None, 27.0], "parsed_2d")

    service = make_service(tmp_path)
    task_id = create_task(service, make_transform_book(tmp_path / "map.xlsx"))
    spec = TransformTaskSpec(
        operation=TransformOperation.MAP_VALUES, sheet="Data",
        map_rules=[ValueMapRule(
            source=ColumnRef(sheet="Data", column_letter="B", display_header="高血压"),
            mapping={"有": 1, "是": 1, "无": 0}, output_column="高血压_indicator",
        )], confidence=1,
    )
    envelope = TaskEnvelope(goal="将高血压有/无显式映射为1/0，是也映射为1", spec=spec)
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid
    assert plan.steps[0].expected_changes.new_columns == ["高血压_indicator"]
    service.approve(task_id, "approve")
    output, _, result = service.execute_task_plan(task_id, envelope, plan)
    assert result.counters == {"mapped": 3, "unknown": 1, "exceptions": 1}
    workbook = load_workbook(output, data_only=True)
    assert [workbook["Data"].cell(row, 5).value for row in range(2, 6)] == [1, 0, 1, None]
    assert workbook["Data"].max_column == 5
    workbook.close()


def test_age_suffix_normalization_can_overwrite_as_number(tmp_path):
    source = tmp_path / "age.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["年龄"])
    sheet.append([60])
    sheet.append(["61岁"])
    workbook.save(source)
    workbook.close()

    service = make_service(tmp_path / "service")
    task_id = create_task(service, source)
    spec = NormalizeTaskSpec(
        sheet="Data",
        column="年龄",
        operations=[NormalizeOperation.STRIP_SUFFIX, NormalizeOperation.TO_NUMBER],
        strip_suffixes=["岁"],
        overwrite=True,
        confidence=1,
    )
    envelope = TaskEnvelope(goal="年龄去掉岁并改为数字", spec=spec)
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid, validation.errors
    service.approve(
        task_id, "approve", {"acknowledgements": ["overwrite_source_columns"]}
    )
    output, report, _ = service.execute_task_plan(task_id, envelope, plan)
    assert report.status == "passed"
    workbook = load_workbook(output, data_only=True)
    assert [workbook["Data"].cell(row, 1).value for row in (2, 3)] == [60, 61]
    workbook.close()


def test_mapping_supports_explicit_blank_and_other_defaults(tmp_path):
    source = tmp_path / "map_defaults.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["高血压", "临床分类"])
    sheet.append(["有", "ACTH腺瘤"])
    sheet.append(["无", "生长激素细胞腺瘤"])
    sheet.append([None, "其他腺瘤"])
    workbook.save(source)
    workbook.close()

    service = make_service(tmp_path / "service")
    task_id = create_task(service, source)
    spec = TransformTaskSpec(
        operation=TransformOperation.MAP_VALUES,
        sheet="Data",
        map_rules=[
            ValueMapRule(
                source=ColumnRef(sheet="Data", column_letter="A", display_header="高血压"),
                mapping={"有": 1, "无": 0}, blank_value=0,
                output_column="高血压_indicator",
            ),
            ValueMapRule(
                source=ColumnRef(sheet="Data", column_letter="B", display_header="临床分类"),
                mapping={"ACTH腺瘤": 1}, unmapped_value=0,
                output_column="ACTH腺瘤_indicator",
            ),
        ],
        confidence=1,
    )
    envelope = TaskEnvelope(goal="映射分类", spec=spec)
    plan, validation = service.make_task_plan(
        task_id, envelope,
        policy_source_text="高血压有为1无为0，没有记录算作无；ACTH腺瘤为1，其他腺瘤为0",
    )
    assert validation.valid, validation.errors
    service.approve(task_id, "approve")
    output, report, result = service.execute_task_plan(task_id, envelope, plan)
    assert report.status == "passed"
    assert result.counters["mapped_blank"] == 1
    assert result.counters["mapped_default"] == 2
    workbook = load_workbook(output, data_only=True)
    sheet = workbook["Data"]
    assert [sheet.cell(row, 3).value for row in range(2, 5)] == [1, 0, 0]
    assert [sheet.cell(row, 4).value for row in range(2, 5)] == [1, 0, 0]
    workbook.close()


def test_mapping_status_column_is_only_written_when_explicitly_requested(tmp_path):
    source = tmp_path / "map_with_status.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["性别"])
    sheet.append(["男"])
    sheet.append(["女"])
    workbook.save(source)
    workbook.close()

    service = make_service(tmp_path / "service")
    task_id = create_task(service, source)
    spec = TransformTaskSpec(
        operation=TransformOperation.MAP_VALUES,
        sheet="Data",
        map_rules=[ValueMapRule(
            source=ColumnRef(sheet="Data", column_letter="A", display_header="性别"),
            mapping={"男": 1, "女": 0},
            output_column="性别_indicator",
            status_column="性别_indicator_状态",
        )],
        confidence=1,
    )
    envelope = TaskEnvelope(goal="生成性别 indicator，并明确输出映射状态列", spec=spec)
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid, validation.errors
    assert plan.steps[0].expected_changes.new_columns == [
        "性别_indicator", "性别_indicator_状态"
    ]
    service.approve(task_id, "approve")
    output, report, result = service.execute_task_plan(task_id, envelope, plan)
    assert report.status == "passed"
    assert result.counters["mapped"] == 2
    workbook = load_workbook(output, data_only=True)
    try:
        sheet = workbook["Data"]
        assert [sheet.cell(1, column).value for column in range(1, 4)] == [
            "性别", "性别_indicator", "性别_indicator_状态"
        ]
        assert [sheet.cell(row, 3).value for row in range(2, 4)] == ["mapped", "mapped"]
    finally:
        workbook.close()


def test_volume_uses_lwh_or_diameter_cubed_as_explicitly_selected(tmp_path):
    source = tmp_path / "mixed_volume.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["长_mm", "宽_mm", "高_mm", "最大径_mm"])
    sheet.append([10, 20, 30, 30])
    sheet.append([None, None, None, 10])
    sheet.append([10, 20, None, 20])
    workbook.save(source)
    workbook.close()

    service = make_service(tmp_path / "service")
    task_id = create_task(service, source)
    spec = TransformTaskSpec(
        operation=TransformOperation.DERIVE_METRIC,
        sheet="Data",
        derive_rule=DeriveMetricRule(
            sources=[
                ColumnRef(sheet="Data", column_letter="A"),
                ColumnRef(sheet="Data", column_letter="B"),
                ColumnRef(sheet="Data", column_letter="C"),
                ColumnRef(sheet="Data", column_letter="D"),
            ],
            formula_template=DerivedFormulaTemplate.RECTANGULAR_OR_DIAMETER_CUBED,
            output_column="肿瘤体积_cm3",
        ),
        confidence=1,
    )
    envelope = TaskEnvelope(goal="三维用长宽高乘积，单一直径用直径^3", spec=spec)
    plan, validation = service.make_task_plan(
        task_id, envelope,
        policy_source_text="三维按照长*宽*高，只有单一直径时按照直径^3计算",
    )
    assert validation.valid, validation.errors
    service.approve(task_id, "approve")
    output, report, result = service.execute_task_plan(task_id, envelope, plan)
    assert report.status == "passed_with_warnings"
    assert result.counters["derived"] == 2
    assert result.counters["invalid_input"] == 1
    workbook = load_workbook(output, data_only=True)
    sheet = workbook["Data"]
    assert sheet.cell(2, 5).value == 6
    assert sheet.cell(3, 5).value == 1
    assert sheet.cell(4, 5).value is None
    workbook.close()


def test_explicit_dimension_exception_policies_are_applied_and_audited(tmp_path):
    source = tmp_path / "dimension_resolutions.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["肿瘤大小"])
    sheet.append(["10*8*8；20*15*5"])
    sheet.append(["10cm*20mm*30mm"])
    sheet.append(["未测量"])
    sheet.append(["21.*7*16.7"])
    sheet.append([None])
    sheet.append(["1*2*3*4"])
    workbook.save(source)
    workbook.close()

    service = make_service(tmp_path / "service")
    task_id = create_task(service, source)
    spec = TransformTaskSpec(
        operation=TransformOperation.PARSE_DIMENSIONS,
        sheet="Data",
        dimension_rule=ParseDimensionsRule(
            source=ColumnRef(sheet="Data", column_letter="A", display_header="肿瘤大小"),
            resolution_policies={
                "multiple_candidates": DimensionResolutionAction.FIRST_CANDIDATE,
                "unit_conflict": DimensionResolutionAction.CONVERT_EACH_TO_MM,
                "invalid_format": DimensionResolutionAction.ZERO,
                "no_numeric_measurement": DimensionResolutionAction.ZERO,
                "blank": DimensionResolutionAction.ZERO,
            },
            value_corrections={"21.*7*16.7": "21*7*16.7"},
        ),
        confidence=1,
    )
    envelope = TaskEnvelope(goal="按确认规则处理尺寸异常", spec=spec)
    source_text = (
        "multiple candidate取前者；未测量、没有数值和四个数值的统一为0；"
        "单位不统一的统一为mm并逐个换算；21.*7*16.7改为21*7*16.7"
    )
    plan, validation = service.make_task_plan(
        task_id, envelope, policy_source_text=source_text
    )
    assert validation.valid, validation.errors
    service.approve(task_id, "approve")
    output, report, result = service.execute_task_plan(task_id, envelope, plan)
    assert report.status == "passed"
    assert result.counters["exceptions"] == 0
    workbook = load_workbook(output, data_only=True)
    sheet = workbook["Data"]
    assert [sheet.cell(2, column).value for column in range(2, 6)] == [10, 8, 8, 10]
    assert sheet.cell(2, 6).value == "resolved_multiple_first"
    assert [sheet.cell(3, column).value for column in range(2, 6)] == [100, 20, 30, 100]
    assert sheet.cell(3, 6).value == "resolved_unit_conflict"
    assert [sheet.cell(4, column).value for column in range(2, 6)] == [0, 0, 0, 0]
    assert sheet.cell(4, 6).value == "resolved_zero"
    assert [sheet.cell(5, column).value for column in range(2, 6)] == [21, 7, 16.7, 21]
    assert sheet.cell(5, 6).value == "corrected_value"
    assert sheet.cell(6, 6).value == "resolved_zero"
    assert sheet.cell(7, 6).value == "resolved_zero"
    workbook.close()


def test_repeated_dimension_exception_becomes_choice_clarification(tmp_path):
    source = tmp_path / "repeated_unit_conflicts.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["肿瘤大小"])
    sheet.append(["10cm*20mm*30mm"])
    sheet.append(["2cm*15mm*5mm"])
    workbook.save(source)
    workbook.close()

    service = make_service(tmp_path / "service")
    task_id = create_task(service, source)
    spec = TransformTaskSpec(
        operation=TransformOperation.PARSE_DIMENSIONS,
        sheet="Data",
        dimension_rule=ParseDimensionsRule(
            source=ColumnRef(sheet="Data", column_letter="A", display_header="肿瘤大小")
        ),
        confidence=1,
    )
    envelope = TaskEnvelope(goal="解析尺寸", spec=spec)
    plan, validation = service.make_task_plan(task_id, envelope)
    assert not validation.valid
    assert "DIMENSION_RESOLUTION_REQUIRED:unit_conflict:2" in validation.errors
    review = plan.steps[0].expected_changes.review_rows
    assert [item["row"] for item in review] == [2, 3]
    assert [item["raw_values"]["肿瘤大小"] for item in review] == [
        "10cm*20mm*30mm", "2cm*15mm*5mm"
    ]
    assert all(set(item) == {
        "sheet", "row", "source_columns", "raw_values", "issue_code", "description"
    } for item in review)
    prompt = clarification_for(validation, envelope, service.load_profile(task_id))
    assert "“肿瘤大小”列" in prompt.question
    assert "2 条记录同时使用了不同单位" in prompt.question
    assert prompt.choices == [
        "按每个数值的显式单位分别换算为 mm",
        "保持留空并生成异常清单",
    ]


def test_exception_review_is_limited_to_first_50_rows(tmp_path):
    source = tmp_path / "many_invalid_dimensions.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["肿瘤大小"])
    for index in range(55):
        sheet.append([f"{index + 1}*2*3*4"])
    workbook.save(source)
    workbook.close()

    service = make_service(tmp_path / "service")
    task_id = create_task(service, source)
    envelope = TaskEnvelope(
        goal="解析尺寸",
        spec=TransformTaskSpec(
            operation=TransformOperation.PARSE_DIMENSIONS,
            sheet="Data",
            dimension_rule=ParseDimensionsRule(
                source=ColumnRef(sheet="Data", column_letter="A", display_header="肿瘤大小")
            ),
            confidence=1,
        ),
    )
    plan, validation = service.make_task_plan(task_id, envelope)

    assert not validation.valid
    changes = plan.steps[0].expected_changes
    assert changes.review_totals["invalid_format"] == 55
    assert changes.review_truncated["invalid_format"] is True
    assert len(changes.review_rows) == 50
    assert [item["row"] for item in changes.review_rows] == list(range(2, 52))


def test_overwrite_requires_acknowledgement_and_lineage_blocks_stale_dependency(tmp_path):
    service = make_service(tmp_path)
    task_id = create_task(service, make_transform_book(tmp_path / "lineage.xlsx"))
    parse_spec = TransformTaskSpec(
        operation=TransformOperation.PARSE_DIMENSIONS, sheet="Data",
        dimension_rule=ParseDimensionsRule(
            source=ColumnRef(sheet="Data", column_letter="C", display_header="肿瘤大小")
        ), confidence=1,
    )
    parse_envelope = TaskEnvelope(goal="解析肿瘤尺寸", spec=parse_spec)
    plan, validation = service.make_task_plan(task_id, parse_envelope)
    assert validation.valid
    service.approve(task_id, "approve")
    service.execute_task_plan(task_id, parse_envelope, plan)
    assert len(service.repository.list_lineage(task_id, 1)) == 4

    overwrite = TransformTaskSpec(
        operation=TransformOperation.MAP_VALUES, sheet="Data",
        map_rules=[ValueMapRule(
            source=ColumnRef(sheet="Data", column_letter="C", display_header="肿瘤大小"),
            mapping={"10*20*30 mm": "10*20*31 mm"}, overwrite=True,
        )], confidence=1,
    )
    overwrite_envelope = TaskEnvelope(goal="覆盖修正肿瘤尺寸", spec=overwrite)
    overwrite_plan, validation = service.make_task_plan(task_id, overwrite_envelope)
    assert overwrite_plan.confirmation_requirements == ["overwrite_source_columns"]
    assert not validation.valid
    assert any("必须同时重算" in item for item in validation.errors)

    # Independent overwrite still cannot be approved without the explicit acknowledgement.
    bmi_overwrite = TransformTaskSpec(
        operation=TransformOperation.IMPUTE_NUMERIC, sheet="Data",
        impute_rules=[ImputeNumericRule(
            source=ColumnRef(sheet="Data", column_letter="A", display_header="BMI"),
            method=ImputationMethod.CONSTANT, constant_value=0, overwrite=True,
            invalid_policy_confirmed=True,
        )], confidence=1,
    )
    bmi_envelope = TaskEnvelope(goal="覆盖BMI空值为0", spec=bmi_overwrite)
    bmi_plan, validation = service.make_task_plan(task_id, bmi_envelope)
    assert validation.valid
    with pytest.raises(IntentValidationError):
        service.approve(task_id, "approve")
    service.approve(task_id, "approve", {"acknowledgements": ["overwrite_source_columns"]})


def test_derive_metric_requires_explicit_formula_in_user_text(tmp_path):
    service = make_service(tmp_path)
    task_id = create_task(service, make_transform_book(tmp_path / "derive.xlsx"))
    spec = TransformTaskSpec(
        operation=TransformOperation.DERIVE_METRIC, sheet="Data",
        derive_rule=DeriveMetricRule(
            sources=[
                ColumnRef(sheet="Data", column_letter="A", display_header="BMI"),
                ColumnRef(sheet="Data", column_letter="A", display_header="BMI"),
                ColumnRef(sheet="Data", column_letter="A", display_header="BMI"),
            ],
            formula_template=DerivedFormulaTemplate.ELLIPSOID_PI_OVER_6,
        ), confidence=1,
    )
    envelope = TaskEnvelope(goal="计算体积", spec=spec)
    _, validation = service.make_task_plan(task_id, envelope, policy_source_text="计算体积")
    assert not validation.valid
    assert any("未在用户原始要求中明确选择" in item for item in validation.errors)
    prompt = clarification_for(validation, envelope, service.load_profile(task_id))
    assert prompt.question == "计算体积需要明确公式。你希望使用哪一种？"
    assert len(prompt.choices) == 4


def test_repeated_invalid_derived_inputs_offer_choices_and_can_be_zeroed(tmp_path):
    source = tmp_path / "derive_invalids.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["长_mm", "宽_mm", "高_mm"])
    sheet.append([10, 20, 30])
    sheet.append([-10, 20, 30])
    sheet.append(["10mm", 20, 30])
    workbook.save(source)
    workbook.close()

    service = make_service(tmp_path / "service")
    task_id = create_task(service, source)
    base_rule = DeriveMetricRule(
        sources=[
            ColumnRef(sheet="Data", column_letter="A"),
            ColumnRef(sheet="Data", column_letter="B"),
            ColumnRef(sheet="Data", column_letter="C"),
        ],
        formula_template=DerivedFormulaTemplate.ELLIPSOID_PI_OVER_6,
    )
    envelope = TaskEnvelope(
        goal="按椭球公式计算体积",
        spec=TransformTaskSpec(
            operation=TransformOperation.DERIVE_METRIC,
            sheet="Data",
            derive_rule=base_rule,
            confidence=1,
        ),
    )
    formula_plan, formula_validation = service.make_task_plan(
        task_id, envelope, policy_source_text="计算体积", persist=False
    )
    assert any("体积公式未在用户原始要求中明确选择" in item for item in formula_validation.errors)
    assert service._clarification_review(formula_plan, formula_validation) == {}
    plan, validation = service.make_task_plan(
        task_id, envelope, policy_source_text="使用椭球体积 π/6 计算"
    )
    assert not validation.valid
    assert "DERIVE_RESOLUTION_REQUIRED:invalid_input:2" in validation.errors
    review = plan.steps[0].expected_changes.review_rows
    assert [item["row"] for item in review] == [3, 4]
    assert all(len(item["raw_values"]) == 3 for item in review)
    prompt = clarification_for(validation, envelope, service.load_profile(task_id))
    assert len(prompt.choices) == 3
    assert "统一填 0" in prompt.choices[1]

    resolved_rule = base_rule.model_copy(deep=True)
    resolved_rule.resolution_policies = {
        "invalid_input": DerivedResolutionAction.ZERO
    }
    resolved_envelope = TaskEnvelope(
        goal="按椭球公式计算体积，无效结果统一填0",
        spec=TransformTaskSpec(
            operation=TransformOperation.DERIVE_METRIC,
            sheet="Data",
            derive_rule=resolved_rule,
            confidence=1,
        ),
    )
    plan, validation = service.make_task_plan(
        task_id,
        resolved_envelope,
        policy_source_text="使用椭球体积 π/6 计算，将无效派生结果统一填0",
    )
    assert validation.valid, validation.errors
    service.approve(task_id, "approve")
    output, report, result = service.execute_task_plan(
        task_id, resolved_envelope, plan
    )
    assert report.status == "passed"
    assert result.counters["derived"] == 1
    assert result.counters["resolved_invalid_zero"] == 2
    workbook = load_workbook(output, data_only=True)
    sheet = workbook["Data"]
    assert sheet.cell(3, 4).value == 0
    assert sheet.cell(3, 5).value == "resolved_invalid_zero"
    assert sheet.cell(4, 4).value == 0
    workbook.close()


def test_schema_v4_planning_migration_is_backward_compatible(tmp_path):
    database = tmp_path / "legacy.sqlite3"
    with sqlite3.connect(database) as connection:
        connection.executescript(
            "CREATE TABLE schema_migrations(version INTEGER PRIMARY KEY, applied_at TEXT NOT NULL);"
            "INSERT INTO schema_migrations VALUES(2, 'legacy');"
        )
    settings = Settings(
        project_root=tmp_path, data_dir=tmp_path,
        database_path=database, llm_provider="fake", api_key="", base_url="", model="fake",
        max_file_bytes=20 * 1024 * 1024, max_sheet_rows=100_000,
    )
    SheetPilotService(settings)
    with sqlite3.connect(database) as connection:
        assert connection.execute("SELECT MAX(version) FROM schema_migrations").fetchone()[0] == 5
        assert connection.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='derived_column_lineage'"
        ).fetchone()
        assert connection.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='planning_runs'"
        ).fetchone()
        assert connection.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='activity_events'"
        ).fetchone()
    assert list((tmp_path / "backups").glob("legacy-pre-v3-*.sqlite3"))


@pytest.mark.skipif(
    not (Path(__file__).resolve().parents[2] / "Test" / "对照组(1).xlsx").is_file(),
    reason="real gold-standard workbook is not available",
)
def test_real_medical_workbook_gold_standard_profile(tmp_path):
    source = Path(__file__).resolve().parents[2] / "Test" / "对照组(1).xlsx"
    profile = profile_workbook(source, source.name, 100_000)
    assert profile.sheets[0].max_row - profile.sheets[0].header_row == 714
    assert any("重复列名" in item for item in profile.risks)
    ensure_supported_profile(profile)
    context = ToolContext(
        task_id="gold", input_path=source, output_path=tmp_path / "gold.xlsx",
        profile=profile, version_before=0, version_after=1,
        artifact_dir=tmp_path / "artifacts",
    )
    tool = ToolRegistry().get("transform_column")
    bmi = TransformTaskSpec(
        operation=TransformOperation.IMPUTE_NUMERIC, sheet="19-20",
        impute_rules=[ImputeNumericRule(
            source=ColumnRef(sheet="19-20", column_letter="E"),
            missing_tokens=["无", "空"],
        )], confidence=1,
    )
    assert tool.preview(context, bmi)["summary_counts"] | {"planned_fill_value": None} == {
        "valid": 538, "missing": 174, "invalid": 2, "observed_basis": 538,
        "planned_fill_value": None,
    }
    dimensions = TransformTaskSpec(
        operation=TransformOperation.PARSE_DIMENSIONS, sheet="19-20",
        dimension_rule=ParseDimensionsRule(
            source=ColumnRef(sheet="19-20", column_letter="V")
        ), confidence=1,
    )
    dimension_counts = tool.preview(context, dimensions)["summary_counts"]
    assert dimension_counts["parsed_3d"] == 677
    assert dimension_counts["parsed_2d"] == 3
    assert dimension_counts["diameter_only"] == 9
    assert sum(dimension_counts.values()) == 714

    mapping = TransformTaskSpec(
        operation=TransformOperation.MAP_VALUES, sheet="19-20",
        map_rules=[
            ValueMapRule(
                source=ColumnRef(sheet="19-20", column_letter="F"),
                mapping={"有": 1, "是": 1, "无": 0}, output_column="高血压_indicator",
            ),
            ValueMapRule(
                source=ColumnRef(sheet="19-20", column_letter="G"),
                mapping={"有": 1, "是": 1, "无": 0}, output_column="糖尿病_indicator",
            ),
        ], confidence=1,
    )
    result = tool.execute(context, mapping)
    workbook = load_workbook(context.output_path, data_only=True)
    sheet = workbook["19-20"]
    assert sum(sheet.cell(row, 74).value == 1 for row in range(2, 716)) == 189
    assert sum(sheet.cell(row, 75).value == 1 for row in range(2, 716)) == 74
    workbook.close()
    assert tool.verify(context, mapping, result).valid


@pytest.mark.skipif(
    not (Path(__file__).resolve().parents[2] / "Test" / "对照组(1).xlsx").is_file(),
    reason="real gold-standard workbook is not available",
)
def test_real_medical_workbook_composite_final_request(tmp_path):
    source = Path(__file__).resolve().parents[2] / "Test" / "对照组(1).xlsx"
    service = make_service(tmp_path / "service")
    task_id = create_task(service, source)
    profile = service.load_profile(task_id)
    sheet_profile = profile.sheets[0]
    sheet_name = sheet_profile.name

    def ref(index: int, *, display_header: str | None = None) -> ColumnRef:
        if display_header is not None:
            return ColumnRef(sheet=sheet_name, display_header=display_header)
        column = next(item for item in sheet_profile.columns if item.index == index)
        return ColumnRef(
            sheet=sheet_name,
            column_index=column.index,
            column_letter=column.letter,
            display_header=column.display_header,
            header_occurrence=column.header_occurrence,
        )

    clinical_categories = [
        "生长激素细胞腺瘤",
        "ACTH腺瘤",
        "混合腺瘤",
        "促性腺激素细胞腺瘤",
    ]
    operations = [
        NormalizeTaskSpec(
            sheet=sheet_name,
            column=ref(3).display_header,
            operations=[NormalizeOperation.STRIP_SUFFIX, NormalizeOperation.TO_NUMBER],
            strip_suffixes=["岁"],
            overwrite=True,
            confidence=1,
        ),
        TransformTaskSpec(
            operation=TransformOperation.MAP_VALUES,
            sheet=sheet_name,
            map_rules=[
                ValueMapRule(source=ref(2), mapping={"男": 1, "女": 0}, output_column="性别_indicator"),
                ValueMapRule(source=ref(6), mapping={"有": 1, "无": 0}, blank_value=0, output_column="高血压_indicator"),
                ValueMapRule(source=ref(7), mapping={"有": 1, "无": 0}, blank_value=0, output_column="糖尿病_indicator"),
                *[
                    ValueMapRule(
                        source=ref(46), mapping={category: 1}, unmapped_value=0,
                        output_column=f"{category}_indicator",
                    )
                    for category in clinical_categories
                ],
            ],
            confidence=1,
        ),
        TransformTaskSpec(
            operation=TransformOperation.PARSE_DIMENSIONS,
            sheet=sheet_name,
            dimension_rule=ParseDimensionsRule(
                source=ref(22),
                default_unit="mm",
                length_column="肿瘤长_mm",
                width_column="肿瘤宽_mm",
                height_column="肿瘤高_mm",
                max_dimension_column="肿瘤最大径_mm",
                status_column="肿瘤尺寸_解析状态",
                resolution_policies={
                    "multiple_candidates": DimensionResolutionAction.LEAVE_BLANK,
                    "unit_conflict": DimensionResolutionAction.CONVERT_EACH_TO_MM,
                    "invalid_format": DimensionResolutionAction.LEAVE_BLANK,
                    "no_numeric_measurement": DimensionResolutionAction.LEAVE_BLANK,
                    "blank": DimensionResolutionAction.LEAVE_BLANK,
                },
            ),
            confidence=1,
        ),
        TransformTaskSpec(
            operation=TransformOperation.DERIVE_METRIC,
            sheet=sheet_name,
            derive_rule=DeriveMetricRule(
                sources=[
                    ref(0, display_header="肿瘤长_mm"),
                    ref(0, display_header="肿瘤宽_mm"),
                    ref(0, display_header="肿瘤高_mm"),
                    ref(0, display_header="肿瘤最大径_mm"),
                ],
                formula_template=DerivedFormulaTemplate.RECTANGULAR_OR_DIAMETER_CUBED,
                output_column="肿瘤体积_cm3",
            ),
            confidence=1,
        ),
        TransformTaskSpec(
            operation=TransformOperation.IMPUTE_NUMERIC,
            sheet=sheet_name,
            impute_rules=[
                ImputeNumericRule(
                    source=ref(4), method=ImputationMethod.MEAN,
                    invalid_policy_confirmed=True,
                ),
                ImputeNumericRule(
                    source=ref(5), method=ImputationMethod.MEAN,
                    missing_tokens=["无", "空"],
                    invalid_policy_confirmed=True,
                ),
            ],
            confidence=1,
        ),
    ]
    envelope = TaskEnvelope(
        goal="完成真实医学工作簿的综合数据准备",
        spec=CompositeTaskSpec(operations=operations, confidence=1),
    )
    source_text = (
        "年龄字符串去掉岁并改为数字，覆盖年龄原列；性别男为1女为0另起一列；"
        "高血压和糖尿病有为1无为0，没有记录算作无；临床分类四类各建indicator，其他腺瘤为0；"
        "肿瘤尺寸默认mm，cm换算为mm，多候选及错误值保持留空；"
        "体积按长*宽*高或者直径^3计算；住院天数和BMI空值用mean填充。"
    )
    plan, validation = service.make_task_plan(
        task_id, envelope, policy_source_text=source_text
    )
    assert validation.valid, validation.errors
    assert len(plan.steps) == 5
    assert plan.steps[3].expected_changes.summary_counts == {
        "deferred_to_dependency": 714
    }
    service.approve(
        task_id, "approve",
        {"acknowledgements": ["overwrite_source_columns"]},
    )
    output, report, result = service.execute_task_plan(task_id, envelope, plan)
    assert output.is_file()
    assert report.status in {"passed", "passed_with_warnings"}
    assert result.rows_affected == 714 * 5
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        output_sheet = workbook[sheet_name]
        headers = {
            str(output_sheet.cell(1, column).value)
            for column in range(1, output_sheet.max_column + 1)
        }
        assert not any("indicator_状态" in header.casefold() for header in headers)
        assert {
            "性别_indicator",
            "高血压_indicator",
            "糖尿病_indicator",
            *(f"{category}_indicator" for category in clinical_categories),
        } <= headers
    finally:
        workbook.close()
    assert service.get_workspace(task_id)["task"]["current_version"] == 1

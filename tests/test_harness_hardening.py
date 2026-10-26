from __future__ import annotations

from pathlib import Path
import os
import shutil
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from sheetpilot.config import Settings
from sheetpilot.errors import IntentValidationError, VerificationError
from sheetpilot.files import sha256_file
from sheetpilot.schemas import (
    CompositeTaskSpec,
    FormulaMode,
    FormulaTaskSpec,
    NormalizeOperation,
    NormalizeTaskSpec,
    TaskEnvelope,
    ToolResult,
    VerificationReport,
)
from sheetpilot.service import SheetPilotService
from sheetpilot.storage import Repository


def make_service(tmp_path: Path, llm=None) -> SheetPilotService:
    return SheetPilotService(Settings(
        project_root=tmp_path,
        data_dir=tmp_path / "data",
        database_path=tmp_path / "data" / "sheetpilot.sqlite3",
        llm_provider="fake", api_key="", base_url="", model="fake",
        max_file_bytes=20 * 1024 * 1024, max_sheet_rows=100_000,
    ), llm=llm)


def normalize_envelope() -> TaskEnvelope:
    return TaskEnvelope(
        goal="标准化客户编号",
        spec=NormalizeTaskSpec(
            sheet="订单", column="客户编号",
            operations=[NormalizeOperation.TRIM],
            output_column="客户编号_标准化", confidence=1,
        ),
    )


def create_task(service: SheetPilotService, source: Path) -> str:
    profile = service.create_task(source.name, source.read_bytes())
    return service.task_id_from_profile(profile)


def test_original_upload_tampering_blocks_execution(workbook_factory, tmp_path):
    service = make_service(tmp_path)
    task_id = create_task(service, workbook_factory())
    envelope = normalize_envelope()
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid
    task = service.repository.get_task(task_id)
    original = Path(task["original_path"])
    original.write_bytes(original.read_bytes() + b"tampered")
    service.approve(task_id, "approve")

    with pytest.raises(IntentValidationError, match="完整性校验失败"):
        service.execute_task_plan(task_id, envelope, plan)
    assert service.repository.get_task(task_id)["current_version"] == 0


def test_unapproved_cell_change_is_blocked(workbook_factory, tmp_path, monkeypatch):
    service = make_service(tmp_path)
    task_id = create_task(service, workbook_factory())
    envelope = normalize_envelope()
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid
    tool = service.tools.get("normalize_column")
    original_execute = tool.execute

    def out_of_scope(context, spec):
        result = original_execute(context, spec)
        workbook = load_workbook(context.output_path)
        workbook["订单"]["A2"] = "unexpected"
        workbook.save(context.output_path)
        workbook.close()
        return result

    monkeypatch.setattr(tool, "execute", out_of_scope)
    service.approve(task_id, "approve")
    with pytest.raises(VerificationError, match="计划外变更"):
        service.execute_task_plan(task_id, envelope, plan)
    assert service.repository.get_task(task_id)["current_version"] == 0
    run = service.repository.latest_plan_run(task_id)
    assert run["status"] == "paused"


def test_only_one_execution_lease_can_be_acquired(workbook_factory, tmp_path):
    service = make_service(tmp_path)
    task_id = create_task(service, workbook_factory())
    envelope = normalize_envelope()
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid
    service.approve(task_id, "approve")
    task = service.repository.get_task(task_id)
    first = service.repository.acquire_plan_run(
        uuid4().hex, task_id, plan.plan_id, plan.revision,
        plan.base_version, plan.base_sha256, sha256_file(Path(task["current_path"])),
    )
    assert first["fencing_token"]
    with pytest.raises(ValueError, match="批准|占用"):
        service.repository.acquire_plan_run(
            uuid4().hex, task_id, plan.plan_id, plan.revision,
            plan.base_version, plan.base_sha256, sha256_file(Path(task["current_path"])),
        )


def test_concurrent_execution_lease_has_one_winner(workbook_factory, tmp_path):
    service = make_service(tmp_path)
    task_id = create_task(service, workbook_factory())
    envelope = normalize_envelope()
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid
    service.approve(task_id, "approve")
    task = service.repository.get_task(task_id)
    barrier = Barrier(2)

    def acquire():
        barrier.wait()
        try:
            service.repository.acquire_plan_run(
                uuid4().hex, task_id, plan.plan_id, plan.revision,
                plan.base_version, plan.base_sha256,
                sha256_file(Path(task["current_path"])),
            )
            return "won"
        except ValueError:
            return "lost"

    with ThreadPoolExecutor(max_workers=2) as pool:
        outcomes = list(pool.map(lambda _: acquire(), range(2)))
    assert sorted(outcomes) == ["lost", "won"]


def test_composite_top_level_ambiguity_is_rejected(workbook_factory, tmp_path):
    service = make_service(tmp_path)
    task_id = create_task(service, workbook_factory())
    first = normalize_envelope().spec
    second = first.model_copy(update={"output_column": "客户编号_二次"})
    envelope = TaskEnvelope(
        goal="复合标准化",
        spec=CompositeTaskSpec(
            operations=[first, second],
            confidence=0.4,
            ambiguities=["执行顺序不明确"],
        ),
    )
    _, validation = service.make_task_plan(task_id, envelope)
    assert not validation.valid
    assert any("复合任务" in item for item in validation.errors)


def test_set_formula_must_appear_in_user_source(workbook_factory, tmp_path):
    service = make_service(tmp_path)
    source = workbook_factory(main_headers=("客户编号", "订单额", "税后金额"))
    task_id = create_task(service, source)
    envelope = TaskEnvelope(
        goal="设置税后金额公式",
        spec=FormulaTaskSpec(
            sheet="订单", column="税后金额", start_row=2, end_row=4,
            mode=FormulaMode.SET_FORMULA, formula="=B2*0.9",
            overwrite_existing=True, recalculate=False, confidence=1,
        ),
    )
    _, validation = service.make_task_plan(
        task_id, envelope, policy_source_text="帮我设置税后金额公式"
    )
    assert not validation.valid
    assert any("未出现在用户原始要求" in item for item in validation.errors)


def test_file_committed_state_recovers_idempotently(workbook_factory, tmp_path):
    service = make_service(tmp_path)
    task_id = create_task(service, workbook_factory())
    envelope = normalize_envelope()
    plan, validation = service.make_task_plan(task_id, envelope)
    assert validation.valid
    service.approve(task_id, "approve")
    task = service.repository.get_task(task_id)
    run_id = uuid4().hex
    lease = service.repository.acquire_plan_run(
        run_id, task_id, plan.plan_id, plan.revision,
        plan.base_version, plan.base_sha256, sha256_file(Path(task["current_path"])),
    )
    commit_id = uuid4().hex
    candidate = service.files.version_candidate_path(task_id, 1, commit_id)
    final_path = service.files.version_path(task_id, 1)
    shutil.copy2(task["current_path"], candidate)
    workbook = load_workbook(candidate)
    sheet = workbook["订单"]
    sheet.cell(1, 3, "客户编号_标准化")
    for row_number in range(2, sheet.max_row + 1):
        value = sheet.cell(row_number, 1).value
        sheet.cell(row_number, 3, str(value).strip() if value is not None else None)
    workbook.save(candidate)
    workbook.close()
    candidate_hash = sha256_file(candidate)
    report = VerificationReport(task_id=task_id, status="passed", checks=[])
    result = ToolResult(tool_name="recovery_test")
    service.repository.create_version_commit(
        commit_id=commit_id, run_id=run_id, task_id=task_id, target_version=1,
        candidate_path=candidate, final_path=final_path,
        candidate_sha256=candidate_hash, fencing_token=lease["fencing_token"],
        profile_json=task["profile_json"], changeset_json=result.model_dump_json(),
        verification_json=report.model_dump_json(),
    )
    os.replace(candidate, final_path)
    service.repository.update_version_commit(commit_id, "file_committed")

    restarted = make_service(tmp_path)
    assert restarted.repository.get_task(task_id)["current_version"] == 1
    assert restarted.repository.get_version_commit(commit_id)["status"] == "completed"
    restarted_again = make_service(tmp_path)
    assert len(restarted_again.repository.list_versions(task_id)) == 2


def test_planning_passes_the_graph_bounded_context_once(workbook_factory, tmp_path):
    envelope = normalize_envelope()

    class CapturingLLM:
        def __init__(self):
            self.context = None

        def parse_task(self, request, profile, **kwargs):
            self.context = kwargs.get("workbook_context")
            return envelope.model_copy(deep=True)

    llm = CapturingLLM()
    service = make_service(tmp_path, llm=llm)
    task_id = create_task(service, workbook_factory())
    turn = service.submit_message(task_id, "把客户编号去掉首尾空格并做 NFKC 标准化")
    assert turn.plan is not None
    assert llm.context is not None
    assert llm.context["selected_skill"]["id"] == "normalization_skill"
    assert all("preview_rows" not in sheet for sheet in llm.context["sheets"])


def test_legacy_database_migrates_and_backs_up(workbook_factory, tmp_path):
    source = workbook_factory()
    data_dir = tmp_path / "legacy-data"
    original_dir = data_dir / "tasks" / "legacy" / "original"
    original_dir.mkdir(parents=True)
    original = original_dir / "input.xlsx"
    shutil.copy2(source, original)
    database = data_dir / "sheetpilot.sqlite3"
    connection = sqlite3.connect(database)
    connection.executescript("""
        CREATE TABLE tasks (
          task_id TEXT PRIMARY KEY,status TEXT NOT NULL,original_name TEXT NOT NULL,
          original_path TEXT NOT NULL,current_path TEXT NOT NULL,current_version INTEGER NOT NULL,
          profile_json TEXT,spec_json TEXT,plan_json TEXT,changeset_json TEXT,
          verification_json TEXT,error_json TEXT,created_at TEXT NOT NULL,updated_at TEXT NOT NULL
        );
        CREATE TABLE file_versions (
          task_id TEXT NOT NULL,version INTEGER NOT NULL,path TEXT NOT NULL,sha256 TEXT NOT NULL,
          parent_version INTEGER,profile_json TEXT,created_at TEXT NOT NULL,
          PRIMARY KEY(task_id,version)
        );
        CREATE TABLE events (
          id INTEGER PRIMARY KEY AUTOINCREMENT,task_id TEXT NOT NULL,event_type TEXT NOT NULL,
          payload_json TEXT NOT NULL,created_at TEXT NOT NULL
        );
    """)
    digest = sha256_file(original)
    connection.execute(
        "INSERT INTO tasks VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        ("legacy", "profiled", "input.xlsx", str(original), str(original), 0,
         None, None, None, None, None, None, "now", "now"),
    )
    connection.execute(
        "INSERT INTO file_versions VALUES (?,?,?,?,?,?,?)",
        ("legacy", 0, str(original), digest, None, None, "now"),
    )
    connection.commit()
    connection.close()

    repository = Repository(database)
    task = repository.get_task("legacy")
    assert task["original_sha256"] == digest
    assert task["integrity_level"] == "verified"
    assert list((data_dir / "backups").glob("sheetpilot-pre-v2-*.sqlite3"))
